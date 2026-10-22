#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Kaizen Workstation
==================
A single-file desktop app for Product Owners:

  * Connects to Redmine through its REST API (Base URL + API key)
  * Lists your open tickets and shows their details
  * Generates a PRD template from the selected ticket into an editable workspace
  * Exports the workspace as Markdown (.md) or text (.txt)
  * Management dashboard with backlog KPIs, charts, risks and a one-click
    status report for leadership

Requirements : Python 3.8+ (with Tkinter) and the 'requests' library
Run          : python redmine_workstation.py      (Windows)
               python3 redmine_workstation.py     (macOS / Linux)
"""

from __future__ import annotations

import datetime as dt
import json
import math
import os
import platform
import queue
import re
import statistics
import subprocess
import sys
import threading
import webbrowser
from collections import Counter
from pathlib import Path

import tkinter as tk
import tkinter.font as tkfont
from tkinter import filedialog, messagebox, ttk

try:  # Optional: modern Windows 11 look (python -m pip install sv-ttk)
    import sv_ttk
except ImportError:
    sv_ttk = None

try:
    import requests
except ImportError:  # Friendly message instead of a crash
    _root = tk.Tk()
    _root.withdraw()
    messagebox.showerror(
        "Missing library",
        "This app needs the 'requests' library.\n\n"
        "Open a terminal and run:\n\n"
        "    python -m pip install requests\n\n"
        "(on macOS / Linux type python3 instead of python),\n"
        "then start the app again.",
    )
    sys.exit(1)


# --------------------------------------------------------------------------
# Constants
# --------------------------------------------------------------------------
APP_NAME = "Kaizen Workstation"
APP_VERSION = "1.0"
LEGACY_CONFIG_PATH = Path.home() / ".redmine_po_workstation.json"
LEGACY_DRAFTS_PATH = Path.home() / ".redmine_po_workstation_drafts.json"


def _pick_data_dir() -> Path:
    """Keep the app's data next to this file, falling back to the home folder."""
    try:
        d = Path(__file__).resolve().parent / "PO_Workstation_Data"
        d.mkdir(exist_ok=True)
        probe = d / ".write_test"
        probe.write_text("ok", encoding="utf-8")
        probe.unlink()
        return d
    except OSError:
        d = Path.home() / "PO_Workstation_Data"
        d.mkdir(exist_ok=True)
        return d


DATA_DIR = _pick_data_dir()
CONFIG_PATH = DATA_DIR / "settings.json"
DRAFTS_PATH = DATA_DIR / "prd_drafts.json"
BACKUP_KEEP_DAYS = 14
PAGE_SIZE = 100          # Redmine's maximum page size
MAX_ISSUES = 2000        # Safety cap when downloading open tickets
STALE_DAYS = 14          # "No update in N days" threshold
IS_MAC = platform.system() == "Darwin"
IS_WINDOWS = platform.system() == "Windows"

C = {                      # warm, soft "Pinterest" palette
    "bg": "#FAF6F1",       # cream
    "panel": "#FFFFFF",
    "ink": "#3B3632",      # warm charcoal
    "muted": "#8A8078",
    "line": "#EADFD4",
    "accent": "#C2665C",   # dusty rose
    "accent_dark": "#A3534A",
    "good": "#6E9C7A",     # sage
    "warn": "#D9A05B",     # sand
    "bad": "#C9665E",      # soft red
    "info": "#8BA5C4",     # dusty blue
    "header": "#4A3F3A",   # heading text
    "topbar": "#FFFFFF",
    "soft": "#FBEDE8",     # blush tint for banners
    "soft_line": "#F0D5CC",
}
C.update({"field": "#FFFFFF", "bar": "#F1E9E1", "head_bg": "#F6EFE8", "faint": "#B3A99F", "select": "#F6DCD6",
          "odd": "#FDFAF6", "sand_soft": "#FCF1DF", "sand_line": "#EBD3A9", "sand_ink": "#A8743A",
          "btn_hover": "#F8EEE9", "btn_press": "#F1DDD5", "accent_disabled": "#E3B5AE", "track": "#F4ECE4",
          "grid": "#F1EAE2", "col": "#F4EEE8", "col_side": "#F1ECF4", "col_warn": "#FBF3E6", "ink2": "#5E554E",
          "on_accent": "#FFFFFF"})
LIGHT = dict(C)
DARK = dict(LIGHT, **{
    "bg": "#1D1A18", "panel": "#282422", "ink": "#ECE4DC", "muted": "#A89C92", "line": "#3B3531",
    "accent": "#E0887C", "accent_dark": "#C46F64", "good": "#8DBE99", "warn": "#E2B26C", "bad": "#E5807A",
    "info": "#9DB6D6", "header": "#F4EBE3", "topbar": "#282422", "soft": "#3A2B28", "soft_line": "#5C413B",
    "field": "#211E1C", "bar": "#24201E", "head_bg": "#302A27", "faint": "#6C625B", "select": "#4B3430",
    "odd": "#25211F", "sand_soft": "#352B1F", "sand_line": "#5A4A30", "sand_ink": "#E2A85C",
    "btn_hover": "#3A322E", "btn_press": "#463A36", "accent_disabled": "#6B4A45", "track": "#332D29",
    "grid": "#332D29", "col": "#24201E", "col_side": "#262330", "col_warn": "#2E2820", "ink2": "#CFC4BA"})


def apply_palette(name: str):
    C.update(DARK if name == "dark" else LIGHT)


def saved_theme(cfg: dict) -> str:
    """The Workstation's own choice, else whatever Kaizen Home uses, else light."""
    if cfg.get("theme") in ("light", "dark"):
        return cfg["theme"]
    try:
        home = load_json(Path(__file__).resolve().parent / "PO_Home_Data" / "po_home.json", {})
        return (home.get("settings") or {}).get("theme") or "light"
    except Exception:
        return "light"


CHART_COLORS = ["#D3837A", "#8FB39A", "#E3B26E", "#8BA5C4", "#B49FCC", "#7FB7BE", "#E8A0B4", "#A79E94"]
PRIORITY_COLORS = {
    "immediate": "#A8473F",
    "critical": "#A8473F",
    "blocker": "#A8473F",
    "urgent": "#C9665E",
    "high": "#C88A3C",
    "normal": "#6F8DB3",
    "medium": "#6F8DB3",
    "low": "#9C9188",
}
HIGH_PRIORITY_WORDS = ("high", "urgent", "immediate", "critical", "blocker")
AUTO_REFRESH_MINUTES = 60   # tickets (and the dashboard) refresh by themselves once an hour
TRACK_TAG = "2.0"                  # Tickets with this Redmine tag are "2.0"; all others are "Legacy"
DEFAULT_TAGS = [TRACK_TAG, "Sprint 1", "Sprint 2", "Sprint 3", "Hold", "Pre-Prod-Issue", "Re-open", "Regression",
                "DevOps", "Product-improvement"]
SPRINT_RE = re.compile(r"^\s*sprint\s*\d+", re.I)

FONTS: dict = {}  # Filled in once the Tk window exists


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------
def migrate_and_backup():
    """Move data from the old hidden location once, then keep one dated backup per day."""
    import shutil
    for new, old in ((CONFIG_PATH, LEGACY_CONFIG_PATH), (DRAFTS_PATH, LEGACY_DRAFTS_PATH)):
        try:
            if not new.exists() and old.exists() and old != new:
                shutil.copy2(old, new)
        except OSError:
            pass
    try:
        root = CONFIG_PATH.parent / "backups"
        today = root / dt.date.today().isoformat()
        if not today.exists():
            files = [f for f in (CONFIG_PATH, DRAFTS_PATH) if f.exists()]
            if files:
                today.mkdir(parents=True, exist_ok=True)
                for f in files:
                    shutil.copy2(f, today / f.name)
        days = sorted(d for d in root.iterdir() if d.is_dir()) if root.exists() else []
        for old in days[:-BACKUP_KEEP_DAYS]:
            shutil.rmtree(old, ignore_errors=True)
    except OSError:
        pass


def load_json(path: Path, default):
    try:
        with open(path, "r", encoding="utf-8") as fh:
            return json.load(fh)
    except (OSError, ValueError):
        return default


def save_json(path: Path, data) -> None:
    tmp = path.with_name(path.name + ".tmp")
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=2, ensure_ascii=False)
    tmp.replace(path)
    if not IS_WINDOWS:
        try:
            os.chmod(path, 0o600)  # Keep the API key private to your user
        except OSError:
            pass


def parse_date(value):
    if not value:
        return None
    try:
        return dt.date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def name_of(obj, default="—"):
    if isinstance(obj, dict):
        return obj.get("name") or default
    return default


def priority_color(name: str) -> str:
    low = (name or "").lower()
    for key, colour in PRIORITY_COLORS.items():
        if key in low:
            return colour
    return C["muted"]


def truncate(text: str, limit: int) -> str:
    text = text or ""
    return text if len(text) <= limit else text[: limit - 1] + "…"


def plural(n: int, word: str, plural_word: str = None) -> str:
    return f"{n} {word if n == 1 else (plural_word or word + 's')}"


def pct(part, whole) -> int:
    return 0 if not whole else round(100 * part / whole)


def md_cell(text) -> str:
    return str(text if text is not None else "—").replace("|", "\\|").replace("\n", " ")


def dedupe(items):
    seen, out = set(), []
    for item in items:
        key = item.strip().lower()
        if key and key not in seen:
            seen.add(key)
            out.append(item.strip())
    return out


def ensure_period(text: str) -> str:
    text = text.strip()
    if text and text[-1] not in ".!?:":
        text += "."
    return text


def lower_first(text: str) -> str:
    if len(text) > 1 and text[1].isupper():  # Keep acronyms like "API"
        return text
    return text[:1].lower() + text[1:]


def first_chars(text: str, limit: int) -> str:
    text = text.strip()
    if len(text) <= limit:
        return text
    cut = text[:limit]
    end = max(cut.rfind(". "), cut.rfind("! "), cut.rfind("? "))
    if end > limit * 0.5:
        return cut[: end + 1]
    return cut.rsplit(" ", 1)[0] + "…"


# --------------------------------------------------------------------------
# Redmine API client
# --------------------------------------------------------------------------
class RedmineError(Exception):
    """An error with a message that is safe to show to the user."""


def _system_ssl_context():
    """SSL check that trusts this computer's own certificate store (Windows / macOS keychain) as well as the
    usual public list. Company certificates and office firewalls that inspect HTTPS are trusted by Windows,
    so this keeps the check ON instead of having to untick 'Verify SSL'."""
    import ssl
    ctx = ssl.create_default_context()          # on Windows this loads the Windows certificate store
    try:
        import certifi
        ctx.load_verify_locations(certifi.where())
    except Exception:
        pass
    return ctx


class _SystemTrustAdapter(requests.adapters.HTTPAdapter):
    def init_poolmanager(self, *args, **kwargs):
        kwargs["ssl_context"] = _system_ssl_context()
        return super().init_poolmanager(*args, **kwargs)

    def proxy_manager_for(self, *args, **kwargs):
        kwargs["ssl_context"] = _system_ssl_context()
        return super().proxy_manager_for(*args, **kwargs)


class RedmineClient:
    def __init__(self, base_url: str, api_key: str, verify_ssl: bool = True, timeout: int = 25):
        self.base_url = (base_url or "").strip().rstrip("/")
        self.api_key = (api_key or "").strip()
        self.verify_ssl = verify_ssl
        self.timeout = timeout
        self._local = threading.local()
        self._headers = {
            "X-Redmine-API-Key": self.api_key,
            "Accept": "application/json",
            "User-Agent": f"RedminePOWorkstation/{APP_VERSION}",
        }
        if not verify_ssl:
            try:
                import urllib3
                urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
            except Exception:
                pass

    @property
    def session(self):
        s = getattr(self._local, "session", None)
        if s is None:  # one connection per worker thread
            s = requests.Session()
            s.headers.update(self._headers)
            if self.verify_ssl:
                try:
                    s.mount("https://", _SystemTrustAdapter())
                except Exception:
                    pass
            self._local.session = s
        return s

    def issue_history(self, issue_id: int) -> dict:
        return self._get(f"issues/{issue_id}.json", {"include": "journals"}).get("issue", {})

    def versions(self, project_id) -> list:
        return self._get(f"projects/{project_id}/versions.json").get("versions", [])

    def statuses(self) -> list:
        return self._get("issue_statuses.json").get("issue_statuses", [])

    def members(self, project_id) -> dict:
        items, _ = self._paginate(f"projects/{project_id}/memberships.json", "memberships", {}, max_items=2000)
        return {m["user"]["id"]: m["user"].get("name", "") for m in items if m.get("user")}

    def _get(self, path: str, params: dict = None) -> dict:
        if not self.base_url:
            raise RedmineError("The Redmine Base URL is empty. Enter it in the Configuration section.")
        if not self.base_url.startswith(("http://", "https://")):
            raise RedmineError("The Base URL must start with http:// or https://")
        if not self.api_key:
            raise RedmineError("The API Access Key is empty. Enter it in the Configuration section.")
        url = f"{self.base_url}/{path.lstrip('/')}"
        try:
            resp = self.session.get(url, params=params, timeout=self.timeout, verify=self.verify_ssl)
        except requests.exceptions.SSLError as exc:
            raw = str(exc)
            m = re.search(r"certificate verify failed: ([^(\]'\"]+)", raw)
            why = (m.group(1).strip() if m else raw[-160:])
            plain = next((txt for key, txt in (
                ("expired", "Redmine's security certificate has EXPIRED. IT needs to renew it."),
                ("local issuer", "Redmine's server doesn't send its full certificate chain. Browsers fill the gap "
                                 "themselves, apps like this one can't. IT can fix it on the server."),
                ("self-signed", "Redmine uses a self-signed / company-made certificate."),
                ("self signed", "Redmine uses a self-signed / company-made certificate."),
                ("hostname", "The certificate is for a different web address than the one typed here."),
                ("mismatch", "The certificate is for a different web address than the one typed here."),
            ) if key in why.lower()), "")
            raise RedmineError(
                "The secure (SSL) connection check failed.\n\n"
                + (f"Reason: {plain}\n\n" if plain else "")
                + "To keep working now: untick 'Verify SSL' and click Save & Connect (fine on the office "
                  "network).\n\n"
                + f"Technical reason (for IT): {why}")
        except requests.exceptions.Timeout:
            raise RedmineError(f"Redmine did not answer within {self.timeout} seconds. Try again, "
                               "or check that you are on the company network / VPN.")
        except requests.exceptions.ConnectionError:
            raise RedmineError(f"Could not reach {self.base_url}.\n\nCheck the URL spelling and "
                               "that you are connected to the company network / VPN.")
        except requests.exceptions.RequestException as exc:
            raise RedmineError(f"Network error: {exc}")

        if resp.status_code == 401:
            raise RedmineError("Redmine rejected the API key (HTTP 401).\n\nCopy the key again from "
                               "Redmine → My account → API access key → Show.")
        if resp.status_code == 403:
            raise RedmineError("Access denied (HTTP 403). Your account may lack permission, or the "
                               "REST API is switched off (Administration → Settings → API).")
        if resp.status_code == 404:
            raise RedmineError("Not found (HTTP 404). Check that the Base URL is the address of your "
                               "Redmine home page, e.g. https://redmine.company.com")
        if resp.status_code == 422:
            raise RedmineError(f"Redmine could not process the request (HTTP 422): {resp.text[:200]}")
        if not resp.ok:
            raise RedmineError(f"Redmine returned HTTP {resp.status_code}: {resp.text[:200]}")
        try:
            return resp.json()
        except ValueError:
            raise RedmineError("Redmine answered with a web page instead of data. Usually this means "
                               "the REST API is disabled (ask your admin: Administration → Settings "
                               "→ API → 'Enable REST web service') or the URL points to a login page.")

    def _paginate(self, path: str, key: str, params: dict = None, max_items: int = MAX_ISSUES):
        items, offset, total = [], 0, 0
        while True:
            page_params = dict(params or {})
            page_params.update(limit=PAGE_SIZE, offset=offset)
            data = self._get(path, page_params)
            batch = data.get(key, [])
            total = data.get("total_count", len(items) + len(batch))
            items.extend(batch)
            offset += len(batch)
            if not batch or offset >= total or len(items) >= max_items:
                break
        return items[:max_items], total

    def current_user(self) -> dict:
        return self._get("users/current.json").get("user", {})

    def projects(self):
        return self._paginate("projects.json", "projects", {}, max_items=5000)

    def priorities(self):
        return self._get("enumerations/issue_priorities.json").get("issue_priorities", [])

    def issues(self, params: dict, max_items: int = MAX_ISSUES):
        return self._paginate("issues.json", "issues", params, max_items=max_items)

    def count(self, params: dict) -> int:
        p = dict(params)
        p["limit"] = 1
        return int(self._get("issues.json", p).get("total_count", 0))

    def issues_by_ids(self, ids) -> list:
        """Fetch specific tickets (any status), e.g. parent tickets."""
        ids = sorted({int(i) for i in ids})
        out = []
        for i in range(0, len(ids), 50):
            chunk = ids[i:i + 50]
            items, _ = self._paginate("issues.json", "issues",
                                      {"issue_id": ",".join(map(str, chunk)), "status_id": "*"},
                                      max_items=len(chunk))
            out.extend(items)
        return out

    def issues_in(self, params: dict, pids, max_items: int = MAX_ISSUES):
        """Tickets from several projects (each project's own tickets only, not its subprojects)."""
        if not pids:
            return self.issues(params, max_items)
        merged = {}
        for pid in pids:
            items, _ = self.issues({**params, "project_id": pid, "subproject_id": "!*"}, max_items)
            for it in items:
                merged[it["id"]] = it
        out = list(merged.values())
        return out, len(out)

    def count_in(self, params: dict, pids) -> int:
        if not pids:
            return self.count(params)
        return sum(self.count({**params, "project_id": pid, "subproject_id": "!*"}) for pid in pids)

    def tag_map(self, project_id, mine_only: bool, tags):
        """Which open tickets carry which tag. Uses the Tags plugin's issue filter, because the
        plugin does not include tags in the API data. Returns (dict id -> [tags], supported)."""
        pids = project_id if isinstance(project_id, list) else ([project_id] if project_id else [])
        bases = [filter_params(pid, mine_only) for pid in pids] or [filter_params(None, mine_only)]
        try:
            probe = self._get("issues.json", {**with_tag(bases[0], "zz-no-such-tag-zz"), "limit": 1})
        except RedmineError:
            return {}, False
        if int(probe.get("total_count", 0)) > 0:  # Filter ignored → tags not available
            return {}, False
        result = {}
        for base in bases:
            for tag in tags:
                items, _ = self._paginate("issues.json", "issues", with_tag(base, tag))
                for it in items:
                    if tag not in result.setdefault(it["id"], []):
                        result[it["id"]].append(tag)
        return result, True

    def tag_ids_any_status(self, pids, mine_only: bool, tag: str):
        """IDs of tickets (open AND closed) carrying a tag. Returns (set, supported)."""
        bases = [filter_params(pid, mine_only) for pid in pids] or [filter_params(None, mine_only)]
        for b in bases:
            b["op[status_id]"] = "*"
        try:
            probe = self._get("issues.json", {**with_tag(bases[0], "zz-no-such-tag-zz"), "limit": 1})
        except RedmineError:
            return set(), False
        if int(probe.get("total_count", 0)) > 0:
            return set(), False
        ids = set()
        for base in bases:
            items, _ = self._paginate("issues.json", "issues", with_tag(base, tag))
            ids.update(it["id"] for it in items)
        return ids, True

    def issue(self, issue_id: int) -> dict:
        data = self._get(f"issues/{issue_id}.json", {"include": "journals,relations,children,attachments"})
        return data.get("issue", {})


def filter_params(project_id, mine_only: bool) -> dict:
    """Open tickets in scope, written as explicit Redmine query filters."""
    p = {"set_filter": "1", "f[]": ["status_id"], "op[status_id]": "o"}
    if project_id:
        p["project_id"] = project_id
        p["f[]"].append("subproject_id")  # this project's own tickets only
        p["op[subproject_id]"] = "!*"
    if mine_only:
        p["f[]"].append("assigned_to_id")
        p["op[assigned_to_id]"] = "="
        p["v[assigned_to_id][]"] = ["me"]
    return p


def with_tag(base: dict, tag: str) -> dict:
    p = {k: (list(v) if isinstance(v, list) else v) for k, v in base.items()}
    p["f[]"].append("issue_tags")
    p["op[issue_tags]"] = "="
    p["v[issue_tags][]"] = [tag]
    return p


def native_tags(issue: dict):
    """Tags if Redmine includes them in the API data (newer plugin versions), else None."""
    raw = issue.get("tags", issue.get("tag_list"))
    if raw is None:
        return None
    if isinstance(raw, str):
        return [t.strip() for t in raw.split(",") if t.strip()]
    out = []
    for t in raw:
        out.append(t.get("name", "") if isinstance(t, dict) else str(t))
    return [t for t in out if t]


def status_flags(issue: dict, today=None, hist=None, stuck_after: int = 2, sprint=None):
    """Signals that a ticket's status or dates need attention."""
    today = today or dt.date.today()
    flags = []
    status = name_of(issue.get("status"), "")
    low = status.lower()
    done = issue.get("done_ratio") or 0
    due = parse_date(issue.get("due_date"))
    upd = parse_date(issue.get("updated_on"))
    created = parse_date(issue.get("created_on"))
    if due and due < today:
        flags.append(f"Overdue {(today - due).days}d")
    if hist and hist["status_days"] > stuck_after:
        flags.append(f"{hist['status_days']} working days in {status}")
    if hist and hist["pushes"]:
        flags.append(f"Due date pushed {len(hist['pushes'])}×")
    if sprint and due and due > sprint["end"]:
        flags.append(f"Due after {sprint['name']} ends ({sprint['end'].strftime('%d-%b')})")
    if done >= 100:
        flags.append("100% done but open")
    if low == "new":
        if 0 < done < 100:
            flags.append(f"New but {done}% done")
        if due and today <= due <= today + dt.timedelta(days=3):
            flags.append("Due soon, still New")
        if created and upd and (today - created).days > 3 and (today - upd).days <= 3 and not hist:
            flags.append("New but recently active")
    elif not hist and upd and (today - upd).days > STALE_DAYS:
        flags.append(f"No update {(today - upd).days}d")
    return flags


def scope_params(project_id, mine_only: bool) -> dict:
    params = {}
    if project_id:
        params["project_id"] = project_id
    if mine_only:
        params["assigned_to_id"] = "me"
    return params


# --------------------------------------------------------------------------
# PRD template generation (reads the ticket description and structures it)
# --------------------------------------------------------------------------
SECTION_PATTERNS = [
    ("ac", re.compile(r"acceptance|criteria|definition of done|\bdod\b", re.I)),
    ("oos", re.compile(r"out[- ]of[- ]scope|not in scope|non[- ]goals?|exclusions?|excluded", re.I)),
    ("stories", re.compile(r"user stor", re.I)),
    ("steps", re.compile(r"steps to reproduce|reproduc|repro steps", re.I)),
    ("actual", re.compile(r"actual|observed|current behaviou?r|what happens|issue description", re.I)),
    ("expected", re.compile(r"expected|desired|should happen|requirement", re.I)),
    ("problem", re.compile(r"problem|background|context|summary|overview|goal|objective|purpose|why", re.I)),
]
BULLET_RE = re.compile(r"^(?:[-*•+]+|\d+[.)]|#+)\s+(.+)$")
STORY_RE = re.compile(r"\bas an?\s+[^,]{2,60},?\s+i\s+(?:want|need|would like|can)\b", re.I)
REQ_RE = re.compile(r"\b(must|should|shall|needs? to|has to|have to|able to|required?|ensure|allow)\b", re.I)
OOS_RE = re.compile(r"\b(out of scope|not in scope|won't|will not|not include|excluded|later phase|"
                    r"phase 2|future release|non-goal|not part of)\b", re.I)
BENEFIT_RE = re.compile(r"\b(?:so that|in order to|to allow|to enable|which will)\s+([^.;\n]{8,160})", re.I)
STORY_VERBS = {"add", "create", "export", "import", "show", "display", "update", "improve", "support",
               "implement", "build", "make", "remove", "delete", "integrate", "allow", "enable",
               "view", "see", "edit", "manage", "download", "upload", "send", "receive", "search",
               "filter", "sort", "track", "configure", "set", "change", "provide", "introduce", "automate"}
ROLE_WORDS = ["customer", "client", "administrator", "admin", "manager", "agent", "operator",
              "developer", "partner", "merchant", "buyer", "seller", "student", "teacher",
              "patient", "doctor", "driver", "employee", "guest", "visitor", "subscriber",
              "member", "user"]


def classify_heading(text: str):
    for key, pattern in SECTION_PATTERNS:
        if pattern.search(text):
            return key
    return None


def clean_inline(text: str) -> str:
    t = re.sub(r"</?(?:pre|code|strong|em|b|i|p|u|br)\s*/?>", "", text, flags=re.I)
    t = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", t)     # markdown links
    t = re.sub(r'"([^"]+)":\S+', r"\1", t)             # textile links
    t = re.sub(r"!\S+?!", "", t)                        # textile images
    t = re.sub(r"\*\*|__", "", t)
    t = re.sub(r"\s+", " ", t)
    return t.strip()


def analyse_description(desc: str) -> dict:
    sections = {k: {"paras": [], "bullets": []}
                for k in ("general", "problem", "ac", "oos", "stories", "steps", "actual", "expected")}
    state = {"current": "general"}
    buf = []

    def flush():
        if buf:
            sections[state["current"]]["paras"].append(" ".join(buf))
            buf.clear()

    for raw in (desc or "").replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if not line:
            flush()
            continue
        if re.fullmatch(r"</?(?:pre|code)[^>]*>", line, re.I):
            continue

        # Textile heading "h2. Title"
        hm = re.match(r"^h[1-6]\.\s+(.*)$", line)
        if hm:
            flush()
            state["current"] = classify_heading(hm.group(1)) or "general"
            continue
        # Markdown heading "## Title" (only when it names a known section;
        # otherwise "# item" is a Textile numbered list item)
        mm = re.match(r"^#{1,6}\s+(.*)$", line)
        if mm and classify_heading(mm.group(1)):
            flush()
            state["current"] = classify_heading(mm.group(1))
            continue
        # Short label lines like "Acceptance criteria:" or "*Scope*"
        stripped = clean_inline(line).strip("*_ ")
        if len(stripped) <= 60 and (line.endswith(":") or re.fullmatch(r"[*_]{1,2}[^*_]+[*_]{1,2}:?", line)):
            flush()
            state["current"] = classify_heading(stripped) or "general"
            continue
        # Bullet item
        bm = BULLET_RE.match(line)
        if bm:
            flush()
            item = clean_inline(bm.group(1))
            if item:
                sections[state["current"]]["bullets"].append(item)
            continue
        # "Problem: users cannot ..."
        km = re.match(r"^([A-Za-z][\w /'-]{1,30}):\s+(.+)$", line)
        if km and classify_heading(km.group(1)):
            flush()
            state["current"] = classify_heading(km.group(1))
            buf.append(clean_inline(km.group(2)))
            continue
        cleaned = clean_inline(line)
        if cleaned:
            buf.append(cleaned)
    flush()
    return sections


def split_sentences(paras):
    out = []
    for p in paras:
        for s in re.split(r"(?<=[.!?])\s+(?=[A-Z0-9\"'(])", p):
            s = s.strip()
            if len(s) > 3:
                out.append(s)
    return out


def detect_role(text: str) -> str:
    low = text.lower()
    counts = Counter()
    for word in ROLE_WORDS:
        n = len(re.findall(rf"\b{word}s?\b", low))
        if n:
            counts[word] = n
    if not counts:
        return "user"
    best = max(counts.values())
    for word in ROLE_WORDS:  # tie-break by list order (specific roles before "user")
        if counts.get(word) == best:
            return "administrator" if word == "admin" else word
    return "user"


# --------------------------------------------------------------------------
# Company calendar: holidays, working days and sprint planning
# --------------------------------------------------------------------------
DEFAULT_HOLIDAYS_2026 = [
    ("2026-01-01", "New Year’s Day"), ("2026-01-14", "Pongal/Makar Sankranti"),
    ("2026-01-26", "Republic Day"), ("2026-03-04", "Holi"), ("2026-03-19", "Ugadi/Gudi Padwa"),
    ("2026-09-14", "Ganesh/Vinayaka Chaturthi"), ("2026-10-02", "Gandhi Jayanti"),
    ("2026-10-20", "Dussehra"), ("2026-11-09", "Day after Diwali"), ("2026-12-25", "Christmas"),
]
DATE_FORMATS = ["%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d-%b-%y", "%d %b %Y", "%d %B %Y",
                "%d.%m.%Y", "%d-%B-%Y", "%b %d %Y", "%B %d %Y"]


def parse_user_date(text):
    t = re.sub(r"\s+", " ", (text or "").strip().replace(",", ""))
    for f in DATE_FORMATS:
        try:
            return dt.datetime.strptime(t, f).date()
        except ValueError:
            pass
    return None


def fmt_date(d, with_day=True) -> str:
    if not d:
        return "—"
    return d.strftime("%a %d-%b-%Y") if with_day else d.strftime("%d-%b-%Y")


def is_working_day(d, holidays) -> bool:
    return d.weekday() < 5 and d not in holidays


def next_working_day(d, holidays):
    for _ in range(400):
        if is_working_day(d, holidays):
            return d
        d += dt.timedelta(days=1)
    return d


def count_working_days(start, end, holidays) -> int:
    """Working days from start to end, both included."""
    if not start or not end or end < start:
        return 0
    n, d = 0, start
    while d <= end:
        n += is_working_day(d, holidays)
        d += dt.timedelta(days=1)
    return n


def working_days_since(start, today, holidays) -> int:
    """Working days that have passed after 'start' up to and including 'today'."""
    if not start or not today or today <= start:
        return 0
    if (today - start).days > 3650:
        start = today - dt.timedelta(days=3650)
    return count_working_days(start + dt.timedelta(days=1), today, holidays)


def plan_sprints(start, mode: str, length: int, count: int, holidays: dict, prefix="Sprint", first_no=1):
    """mode 'working': each sprint has `length` working days (holidays push the end date).
    mode 'calendar': each sprint lasts `length` calendar weeks (holidays reduce working days)."""
    sprints = []
    s = next_working_day(start, holidays)
    for i in range(max(1, count)):
        if mode == "working":
            end, n = s, 1
            while n < max(1, length):
                end += dt.timedelta(days=1)
                n += is_working_day(end, holidays)
        else:
            end = s + dt.timedelta(days=7 * max(1, length) - 1)
            while end > s and not is_working_day(end, holidays):
                end -= dt.timedelta(days=1)
        hols = [(d, holidays[d]) for d in sorted(holidays) if s <= d <= end and d.weekday() < 5]
        sprints.append({"name": f"{prefix} {first_no + i}".strip(), "start": s, "end": end,
                        "working": count_working_days(s, end, holidays), "holidays": hols})
        s = next_working_day(end + dt.timedelta(days=1), holidays)
    return sprints


def sprint_position(sprints, today, holidays):
    for sp in sprints:
        if sp["start"] <= today <= sp["end"]:
            day = count_working_days(sp["start"], today, holidays)
            left = count_working_days(today + dt.timedelta(days=1), sp["end"], holidays)
            return sp, day, left
    return None, 0, 0


# --------------------------------------------------------------------------
# Ticket history: due-date pushes and time in status (from Redmine journals)
# --------------------------------------------------------------------------
def _jdetails(journal, attr):
    return [d for d in journal.get("details") or []
            if d.get("property") == "attr" and d.get("name") == attr]


def history_info(issue: dict, holidays, today=None) -> dict:
    today = today or dt.date.today()
    journals = sorted(issue.get("journals") or [], key=lambda j: j.get("created_on") or "")
    status_since = parse_date(issue.get("created_on"))
    for j in journals:
        if _jdetails(j, "status_id"):
            status_since = parse_date(j.get("created_on")) or status_since

    # Who was assigned after each journal entry (walk backwards from today's assignee)
    cur = (issue.get("assigned_to") or {}).get("id")
    assignee_after = {}
    for j in reversed(journals):
        assignee_after[id(j)] = cur
        for d in _jdetails(j, "assigned_to_id"):
            old = d.get("old_value")
            cur = int(old) if old not in (None, "") and str(old).isdigit() else None

    pushes, original, seen_due = [], None, False
    for j in journals:
        for d in _jdetails(j, "due_date"):
            old, new = parse_date(d.get("old_value")), parse_date(d.get("new_value"))
            if not seen_due:
                original, seen_due = (old or new), True
            if old and new and new > old:
                pushes.append({"date": parse_date(j.get("created_on")), "old": old, "new": new,
                               "days": (new - old).days, "by": name_of(j.get("user"), "?"),
                               "assignee_id": assignee_after.get(id(j))})
    if not seen_due:
        original = parse_date(issue.get("due_date"))
    return {"status_since": status_since,
            "status_days": working_days_since(status_since, today, holidays),
            "pushes": pushes, "original_due": original}


def fetch_journals(client, ids, workers: int = 6) -> dict:
    """Download full ticket history for many tickets in parallel. Returns {id: issue}."""
    from concurrent.futures import ThreadPoolExecutor
    out = {}

    def one(iid):
        try:
            return iid, client.issue_history(iid)
        except RedmineError:
            return iid, None

    with ThreadPoolExecutor(max_workers=workers) as ex:
        for iid, issue in ex.map(one, list(ids)):
            if issue:
                out[iid] = issue
    return out


def compute_delivery(issues, open_ids, hist_by_id, names, start, end, stuck_after) -> dict:
    devs = {}

    def dev(name):
        return devs.setdefault(name, {"name": name, "tickets": set(), "pushes": 0, "days": 0,
                                      "closed_due": 0, "met": 0, "stuck": 0, "open": 0})

    changes, stuck = [], []
    for it in issues:
        h = hist_by_id.get(it["id"])
        if not h:
            continue
        cur_id = (it.get("assigned_to") or {}).get("id")
        cur_name = name_of(it.get("assigned_to"), "Unassigned")
        for p in h["pushes"]:
            if p["date"] and start <= p["date"] <= end:
                aid = p["assignee_id"]
                if not aid:
                    nm = "Unassigned"
                elif aid == cur_id:
                    nm = cur_name
                else:
                    nm = names.get(aid) or f"User #{aid}"
                d = dev(nm)
                d["tickets"].add(it["id"])
                d["pushes"] += 1
                d["days"] += p["days"]
                changes.append({**p, "issue": it, "dev": nm})
        if it["id"] in open_ids:
            d = dev(cur_name)
            d["open"] += 1
            if h["status_days"] > stuck_after:
                d["stuck"] += 1
                stuck.append({"issue": it, "dev": cur_name, "since": h["status_since"], "days": h["status_days"]})
        else:
            closed = parse_date(it.get("closed_on"))
            if closed and start <= closed <= end and h["original_due"]:
                d = dev(cur_name)
                d["closed_due"] += 1
                d["met"] += closed <= h["original_due"]
    rows = [d for d in devs.values() if d["pushes"] or d["stuck"] or d["closed_due"]]
    rows.sort(key=lambda d: (-d["days"], -d["pushes"], -d["stuck"], d["name"].lower()))
    changes.sort(key=lambda c: (c["date"], c["issue"]["id"]), reverse=True)
    stuck.sort(key=lambda s: (-s["days"], s["dev"].lower()))
    total_p = sum(d["pushes"] for d in rows)
    total_days = sum(d["days"] for d in rows)
    tickets = len({c["issue"]["id"] for c in changes})
    closed_due = sum(d["closed_due"] for d in rows)
    met = sum(d["met"] for d in rows)
    s = [f"Between {start.strftime('%d %b')} and {end.strftime('%d %b %Y')}, due dates were pushed "
         f"{plural(total_p, 'time')} on {plural(tickets, 'ticket')}, adding {plural(total_days, 'day')} in total."]
    top = [d for d in rows if d["days"] > 0][:3]
    if top:
        s.append("Most slippage: " + ", ".join(f"{d['name']} ({plural(d['days'], 'day')}, "
                                               f"{plural(d['pushes'], 'push', 'pushes')})" for d in top) + ".")
    s.append(f"{plural(len(stuck), 'open ticket')} have stayed in the same status for more than "
             f"{plural(stuck_after, 'working day')} and need follow-up.")
    if closed_due:
        s.append(f"Of {plural(closed_due, 'ticket')} closed with a due date, {pct(met, closed_due)}% "
                 f"met their original date.")
    return {"devs": rows, "changes": changes, "stuck": stuck, "start": start, "end": end,
            "stuck_after": stuck_after, "summary": " ".join(s), "total_pushes": total_p,
            "total_days": total_days, "tickets": tickets, "generated": dt.datetime.now()}


# --------------------------------------------------------------------------
# PRD in the company template
# --------------------------------------------------------------------------
LOWERABLE_FIRST = {"a", "an", "the", "when", "while", "if", "on", "in", "after", "before", "user", "users",
                   "error", "errors", "some", "sometimes", "clicking", "selecting", "saving", "opening",
                   "there", "it", "this", "data", "no", "unable", "cannot", "can't", "system", "page"}
GWT_RE = re.compile(r"^\s*(given|when|then|and|but)\b", re.I)


def soft_lower(text: str) -> str:
    first = (text.split(" ", 1)[0] if text else "").lower().strip(",.")
    return text[:1].lower() + text[1:] if first in LOWERABLE_FIRST else text


def strip_ref(subject: str) -> str:
    return re.sub(r"^\s*(\[[A-Za-z]{1,8}-?\d+\]\s*)+", "", subject or "").strip()


def build_prd(issue: dict, tags=None) -> str:
    iid = issue.get("id", "")
    subject = (issue.get("subject") or "Untitled ticket").strip()
    tracker = name_of(issue.get("tracker"), "Ticket")
    is_bug = any(w in tracker.lower() for w in ("bug", "defect", "incident"))
    desc = (issue.get("description") or "").strip()
    sec = analyse_description(desc)
    project = name_of(issue.get("project"), "")
    category = name_of(issue.get("category"), "")
    meta = [project + (f" · {category}" if category else ""), f"Priority: {name_of(issue.get('priority'))}",
            f"Status: {name_of(issue.get('status'))}",
            f"Assignee: {name_of(issue.get('assigned_to'), 'Unassigned')}"]
    if tags:
        meta.append(f"Tag: {', '.join(tags)}")
    clean_subject = ensure_period(strip_ref(subject))

    # 1. Overview
    problem = first_chars(" ".join(sec["problem"]["paras"] or sec["general"]["paras"]), 700)
    if not problem:
        bullets = sec["problem"]["bullets"] or sec["general"]["bullets"]
        problem = ensure_period("; ".join(bullets[:4])) if bullets else ""
    no_desc = not problem
    if no_desc:
        problem = (f"The ticket asks for: “{strip_ref(subject)}”. It has no further description yet; confirm the "
                   f"business need and today's pain point with {name_of(issue.get('author'), 'the requester')}.")
    expected = sec["expected"]["bullets"] + split_sentences(sec["expected"]["paras"])
    actual = sec["actual"]["bullets"] + split_sentences(sec["actual"]["paras"])
    if is_bug:
        objective = (ensure_period(expected[0]) if expected else
                     f"Resolve the defect so that “{strip_ref(subject)}” no longer occurs, reliably and every time.")
    else:
        bm = BENEFIT_RE.search(desc)
        if no_desc:
            objective = ("Deliver the requested capability so that it can be accepted against the criteria "
                         "below; confirm the expected business outcome with the requester.")
        else:
            objective = clean_subject[:-1] + (f", so that {bm.group(1).strip()}." if bm else ".")

    # 2. Current vs expected
    left_label = "Actual" if is_bug else "Current"
    pairs = []
    for i in range(max(len(actual), len(expected))):
        pairs.append((actual[i] if i < len(actual) else "", expected[i] if i < len(expected) else ""))
    if not pairs:
        if is_bug:
            pairs = [(clean_subject, "The feature works as designed, without this problem.")]
        else:
            pairs = [("This capability is not available today.", clean_subject)]

    # 3. Requirements
    ac_items = dedupe(sec["ac"]["bullets"] + split_sentences(sec["ac"]["paras"]))
    req_pool = sec["general"]["bullets"] + sec["problem"]["bullets"] + \
        [s for s in split_sentences(sec["general"]["paras"] + sec["problem"]["paras"]) if REQ_RE.search(s)]
    reqs = [x for x in dedupe(req_pool) if not OOS_RE.search(x) and not STORY_RE.search(x)][:8]
    if not reqs and ac_items:
        reqs = [a for a in ac_items if not GWT_RE.match(a)][:8]
    if not reqs:
        if is_bug:
            reqs = [f"The problem described in #{iid} must be fixed: {soft_lower(clean_subject)}",
                    "The root cause must be identified and fixed, not only the visible symptom.",
                    "The fix must not regress related functionality."]
        else:
            reqs = [clean_subject,
                    "The change must follow existing behaviour and design standards of the application.",
                    "Existing functionality must not regress."]

    # 4. Acceptance criteria
    acs = []
    for item in ac_items[:10]:
        parts = re.split(r"(?i)\b(?=given\b|when\b|then\b)", item)
        parts = [p.strip(" ,;") for p in parts if p.strip(" ,;")]
        if len(parts) >= 2 and all(GWT_RE.match(p) for p in parts):
            acs.append([p[:1].upper() + p[1:] for p in parts])
        else:
            acs.append(["Given the change for #%s is deployed," % iid,
                        "When the user performs the scenario described in the ticket,",
                        f"Then {soft_lower(ensure_period(item))}"])
    if not acs:
        has_steps = bool(sec["steps"]["bullets"] or sec["steps"]["paras"])
        if is_bug:
            acs.append([f"Given the situation described in #{iid},",
                        "When the documented steps to reproduce are followed," if has_steps else
                        "When the user repeats the action from the ticket,",
                        f"Then “{strip_ref(subject)}” no longer occurs."])
        else:
            generic = req_pool == [] and not ac_items
            for n, r in enumerate(reqs[:1] if generic else reqs[:5], 1):
                acs.append([f"Given a user working in {project or 'the application'},",
                            "When they use the new capability,",
                            f"Then REQ-{n} is met: {ensure_period(r)}"])
    acs.append([f"Given the existing functionality around this change,",
                "When it is used after the change,", "Then no regression occurs."])

    lines = [f"PRD — {tracker} #{iid}: {subject}", "  |  ".join(meta), "",
             "1. Overview", f"Problem. {problem}", f"Objective. {objective}", "",
             "2. Current vs. Expected Behavior"]
    for a, e in pairs:
        lines += [f"{left_label}: {ensure_period(a) if a else ''}".rstrip(),
                  f"Expected: {ensure_period(e) if e else ''}".rstrip(), ""]
    lines += ["3. Requirements"]
    lines += [f"REQ-{i}: {ensure_period(r)}" for i, r in enumerate(reqs, 1)]
    lines += ["", "4. Acceptance Criteria"]
    for i, ac in enumerate(acs, 1):
        lines += [f"AC-{i}:"] + ac + [""]
    return "\n".join(lines).rstrip() + "\n"


def prd_ai_prompt(issue: dict, tags, parent_label: str, base_url: str) -> str:
    tracker = name_of(issue.get("tracker"), "Ticket")
    notes = [j for j in (issue.get("journals") or []) if (j.get("notes") or "").strip()][-5:]
    comments = "\n".join("- %s (%s): %s" % (name_of(j.get("user")), (j.get("created_on") or "")[:10],
                                             truncate(" ".join(j["notes"].split()), 500))
                         for j in notes) or "(none)"
    return (
        "You are helping a Product Owner write a PRD. Write the PRD for the Redmine ticket below in EXACTLY "
        "this plain-text structure, so it can be pasted back into the app (keep the section titles and the "
        "'Actual:' / 'Expected:', 'REQ-n:' and 'AC-n:' prefixes; no Markdown symbols):\n\n"
        f"PRD — {tracker} #{issue.get('id')}: {issue.get('subject', '')}\n"
        f"{name_of(issue.get('project'), '')}  |  Priority: {name_of(issue.get('priority'))}  |  "
        f"Status: {name_of(issue.get('status'))}  |  Assignee: {name_of(issue.get('assigned_to'), 'Unassigned')}"
        f"{'  |  Tag: ' + ', '.join(tags) if tags else ''}\n\n"
        "1. Overview\nProblem. <2-4 sentences>\nObjective. <1-2 sentences>\n\n"
        "2. Current vs. Expected Behavior\nActual: <what happens today>\nExpected: <what should happen>\n"
        "(repeat Actual/Expected pairs as needed; for new features use 'Current:' instead of 'Actual:')\n\n"
        "3. Requirements\nREQ-1: <testable requirement>\nREQ-2: ...\n\n"
        "4. Acceptance Criteria\nAC-1:\nGiven <context>,\nWhen <action>,\nThen <observable result>.\n"
        "(repeat for each criterion, include one regression criterion)\n\n"
        "Use clear, testable language. Do not invent facts that are not in the ticket; if something is "
        "unclear, phrase it as something to confirm.\n\n"
        "=== TICKET ===\n"
        f"#{issue.get('id')} {issue.get('subject', '')}\nLink: {base_url.rstrip('/')}/issues/{issue.get('id')}\n"
        f"Tracker: {tracker} | Status: {name_of(issue.get('status'))} | Priority: {name_of(issue.get('priority'))}"
        f" | Release: {name_of(issue.get('fixed_version'))} | Parent: {parent_label or '—'}\n"
        f"Description:\n{issue.get('description') or '(empty)'}\n\nLatest comments:\n{comments}\n")


# --------------------------------------------------------------------------
# Simple document model shared by PRD and release notes (for Word/Markdown export)
# --------------------------------------------------------------------------
HEADING_RE = re.compile(r"^\d+\.\s+\S")
LEFT_RE = re.compile(r"^(Actual|Current|Observed)(\s*\([^)]*\))?\s*:\s*(.*)$", re.I)
RIGHT_RE = re.compile(r"^(Expected|Desired)\s*:\s*(.*)$", re.I)
REQROW_RE = re.compile(r"^(REQ-\d+)\s*[:.\-–]\s*(.*)$", re.I)
ACROW_RE = re.compile(r"^(AC-\d+)\s*[:.\-–]?\s*(.*)$", re.I)
LEAD_RE = re.compile(r"^([A-Z][\w/&-]{1,20}(?: [\w/&-]{1,20})?\.)\s+(.+)$")


def parse_doc(text: str, mode: str = "prd") -> list:
    lines = (text or "").replace("\r\n", "\n").split("\n")
    blocks, header_count, table, in_ac = [], 0, None, False

    def close_table():
        nonlocal table, in_ac
        if table:
            blocks.append(table)
        table, in_ac = None, False

    def next_nonempty(i):
        for k in range(i + 1, len(lines)):
            if lines[k].strip():
                return lines[k].strip()
        return ""

    for i, raw in enumerate(lines):
        line = raw.strip()
        if not line:
            in_ac = False if table and table["kind"] == "ac" else in_ac
            continue
        if header_count < 2 and not HEADING_RE.match(line):
            blocks.append({"type": "title" if header_count == 0 else "subtitle", "text": line})
            header_count += 1
            continue
        header_count = 2
        if HEADING_RE.match(line):
            close_table()
            blocks.append({"type": "heading", "text": line})
            continue
        if mode == "prd":
            m = LEFT_RE.match(line)
            if m:
                if not table or table["kind"] != "behavior":
                    close_table()
                    word = m.group(1).capitalize()
                    left = "Actual (Defect)" if word == "Actual" else word
                    table = {"type": "table", "kind": "behavior", "header": [left, "Expected"], "rows": []}
                table["rows"].append([m.group(3).strip(), ""])
                continue
            m = RIGHT_RE.match(line)
            if m and table and table["kind"] == "behavior":
                if table["rows"] and not table["rows"][-1][1]:
                    table["rows"][-1][1] = m.group(2).strip()
                else:
                    table["rows"].append(["", m.group(2).strip()])
                continue
            m = REQROW_RE.match(line)
            if m:
                if not table or table["kind"] != "req":
                    close_table()
                    table = {"type": "table", "kind": "req", "header": ["ID", "Requirement"], "rows": []}
                table["rows"].append([m.group(1).upper(), m.group(2).strip()])
                continue
            m = ACROW_RE.match(line)
            if m:
                if not table or table["kind"] != "ac":
                    close_table()
                    table = {"type": "table", "kind": "ac", "header": ["ID", "Acceptance Criterion"], "rows": []}
                table["rows"].append([m.group(1).upper(), [m.group(2).strip()] if m.group(2).strip() else []])
                in_ac = True
                continue
            if in_ac and table and table["kind"] == "ac":
                table["rows"][-1][1].append(line)
                continue
        close_table()
        if re.match(r"^[-•*]\s+", line):
            blocks.append({"type": "bullet", "text": re.sub(r"^[-•*]\s+", "", line)})
            continue
        nxt = next_nonempty(i)
        if len(line) <= 70 and not line.endswith((".", ":", "!", "?")) and re.match(r"^[-•*]\s+", nxt):
            blocks.append({"type": "sub", "text": line})
            continue
        m = LEAD_RE.match(line) if mode == "prd" else None
        if m:
            blocks.append({"type": "para", "lead": m.group(1), "text": m.group(2)})
        else:
            blocks.append({"type": "para", "lead": "", "text": line})
    close_table()
    return blocks


def _gwt_md(lines):
    out = []
    for ln in lines:
        m = GWT_RE.match(ln)
        out.append(f"**{ln[:m.end()].strip()}**{ln[m.end():]}" if m else ln)
    return "<br>".join(md_cell(x) for x in out)


def doc_to_markdown(blocks) -> str:
    out = []
    for b in blocks:
        t = b["type"]
        if t == "title":
            out += [f"# {b['text']}", ""]
        elif t == "subtitle":
            out += [f"_{b['text']}_", ""]
        elif t == "heading":
            out += ["", f"## {b['text']}", ""]
        elif t == "sub":
            out += ["", f"**{b['text']}**", ""]
        elif t == "bullet":
            out.append(f"- {b['text']}")
        elif t == "para":
            out += [(f"**{b['lead']}** " if b.get("lead") else "") + b["text"], ""]
        elif t == "table":
            out += ["", "| " + " | ".join(b["header"]) + " |", "|" + "---|" * len(b["header"])]
            for row in b["rows"]:
                cells = [_gwt_md(c) if isinstance(c, list) else md_cell(c) for c in row]
                out.append("| " + " | ".join(cells) + " |")
            out.append("")
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip() + "\n"


def _need_docx():
    try:
        import docx  # noqa: F401
        return True
    except ImportError:
        return False


def _shade(cell, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _table_borders(table, color="BFBFBF"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        borders.append(e)
    tblPr.append(borders)


def _para_border_bottom(paragraph, color="2F5496"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), color)
    bdr.append(b)
    pPr.append(bdr)


def prd_to_docx(blocks, path: str):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml.ns import qn
    blue, grey, white = RGBColor(0x2F, 0x54, 0x96), RGBColor(0x80, 0x80, 0x80), RGBColor(0xFF, 0xFF, 0xFF)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.75))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Calibri")

    def para(space_before=0, space_after=5.5):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(space_before), Pt(space_after), 1.15
        return p

    for b in blocks:
        t = b["type"]
        if t == "title":
            p = para(0, 2)
            r = p.add_run(b["text"])
            r.bold, r.font.size, r.font.color.rgb = True, Pt(14), blue
            _para_border_bottom(p, "BFBFBF")
        elif t == "subtitle":
            p = para(2, 8)
            r = p.add_run(b["text"])
            r.font.size, r.font.color.rgb = Pt(10), grey
        elif t == "heading":
            p = para(13, 5.5)
            r = p.add_run(b["text"])
            r.bold, r.font.size, r.font.color.rgb = True, Pt(14), blue
        elif t == "sub":
            p = para(6, 3)
            p.add_run(b["text"]).bold = True
        elif t == "bullet":
            doc.add_paragraph(b["text"], style="List Bullet")
        elif t == "para":
            p = para()
            if b.get("lead"):
                p.add_run(b["lead"] + " ").bold = True
            p.add_run(b["text"])
        elif t == "table":
            rows = b["rows"]
            tbl = doc.add_table(rows=1 + len(rows), cols=2)
            _table_borders(tbl)
            widths = (Inches(3.21), Inches(3.21)) if b["kind"] == "behavior" else (Inches(0.76), Inches(5.65))
            tbl.autofit = False
            for ci, col in enumerate(tbl.columns):
                col.width = widths[ci]
            grid = tbl._tbl.tblGrid
            if grid is not None:
                for ci, gc in enumerate(grid.findall(qn("w:gridCol"))):
                    gc.set(qn("w:w"), str(int(widths[ci].inches * 1440)))
            for ri in range(len(rows) + 1):
                for ci in range(2):
                    cell = tbl.rows[ri].cells[ci]
                    cell.width = widths[ci]
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(2), Pt(2)
                    if ri == 0:
                        _shade(cell, "2F5496")
                        r = p.add_run(b["header"][ci])
                        r.bold, r.font.size, r.font.color.rgb = True, Pt(10), white
                        continue
                    _shade(cell, "FFFFFF" if ri % 2 else "F2F5FA")
                    value = rows[ri - 1][ci]
                    if isinstance(value, list):
                        for li, ln in enumerate(value):
                            m = GWT_RE.match(ln)
                            if m:
                                r = p.add_run(ln[:m.end()])
                                r.bold, r.font.size = True, Pt(10)
                                ln = ln[m.end():]
                            p.add_run(ln).font.size = Pt(10)
                            if li < len(value) - 1:
                                p.add_run().add_break()
                    else:
                        p.add_run(value).font.size = Pt(10)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
    doc.save(path)


# --------------------------------------------------------------------------
# Release notes
# --------------------------------------------------------------------------
SECURITY_RE = re.compile(r"\b(security|xss|cross[- ]site|csrf|sql injection|injection|vulnerab\w*|cve-\d+|"
                         r"sanitiz\w*|privilege escalation)\b", re.I)
DONE_STATUS_RE = re.compile(r"release|resolved|done|verified|deploy|complete|fixed|closed", re.I)
RN_DEFAULTS = {
    "product": "My Product",
    "subtitle": "My Product | Short description",
    "website": "www.example.com",
    "components": "My Product – Version 1.0.0\n- UI: 1.0.0\n- API Service: 1.0.0",
    "known_issues": "None identified at the time of release.",
    "upgrade": "No special instructions required.",
    "template": "",
}
PAST_TENSE = {"add": "Added", "introduce": "Introduced", "improve": "Improved", "enable": "Enabled",
              "update": "Updated", "redesign": "Redesigned", "revamp": "Revamped", "refine": "Refined",
              "implement": "Implemented", "restrict": "Restricted", "provide": "Provided", "create": "Created",
              "remove": "Removed", "rename": "Renamed", "optimize": "Optimized", "optimise": "Optimised",
              "integrate": "Integrated", "migrate": "Migrated", "upgrade": "Upgraded", "replace": "Replaced",
              "extend": "Extended", "enhance": "Enhanced", "streamline": "Streamlined", "standardize": "Standardized",
              "configure": "Configured", "automate": "Automated", "simplify": "Simplified", "expand": "Expanded",
              "fix": "Fixed", "resolve": "Resolved", "correct": "Corrected", "restore": "Restored",
              "prevent": "Prevented", "validate": "Validated", "hide": "Hid", "disable": "Disabled"}
PAST_WORDS = set(PAST_TENSE.values()) | {"Added", "Fixed", "Resolved", "Corrected", "Restored", "Prevented"}


def rn_section(issue) -> str:
    subject = issue.get("subject", "")
    tracker = name_of(issue.get("tracker"), "").lower()
    if SECURITY_RE.search(subject):
        return "security"
    if any(w in tracker for w in ("bug", "defect", "incident", "error")):
        return "bugs"
    if any(w in tracker for w in ("feature", "functionality", "new")):
        return "features"
    return "enhancements"


def rn_area(issue, skip_words=()) -> str:
    cat = name_of(issue.get("category"), "")
    if cat:
        return cat
    s = strip_ref(issue.get("subject", ""))
    m = re.match(r"^([A-Za-z][\w &/().]{1,30}?)\s*(?::|\s-\s|\s–\s)\s*\S", s)
    if m:
        area = m.group(1).strip()
        if len(area.split()) <= 4 and area.lower() not in skip_words:
            return area
    return "General"


def rn_sentence(issue, section: str, area: str) -> str:
    s = strip_ref(issue.get("subject", "")).strip().rstrip(".")
    if area and area != "General":
        s = re.sub(r"^" + re.escape(area) + r"\s*(?::|\s-\s|\s–\s)\s*", "", s).strip()
    if not s:
        s = strip_ref(issue.get("subject", "")) or f"Ticket #{issue.get('id')}"
    first = s.split(" ", 1)[0]
    rest = s.split(" ", 1)[1] if " " in s else ""
    if first[:1].upper() + first[1:] in PAST_WORDS:
        return ensure_period(first[:1].upper() + first[1:] + (" " + rest if rest else ""))
    if section == "security" and first.lower() not in PAST_TENSE:
        return ensure_period("Fixed a security vulnerability: " + s)
    if section in ("bugs", "security"):
        if first.lower() in PAST_TENSE:
            return ensure_period(PAST_TENSE[first.lower()] + " " + rest)
        return ensure_period("Fixed an issue where " + soft_lower(s))
    m = re.match(r"^(?:the |a )?users? (?:should|must|shall|need to|needs to|would like to|want to|can)"
                 r"(?: be able to)? (.+)$", s, re.I) or \
        re.match(r"^as an? [^,]+,? i (?:want|need|would like) to (.+)$", s, re.I) or \
        re.match(r"^allow (?:the )?users? to (.+)$", s, re.I)
    if m:
        return ensure_period("Users can now " + m.group(1))
    m = re.match(r"^(?:ability|option) to (.+)$", s, re.I)
    if m:
        return ensure_period("Added the ability to " + m.group(1))
    m = re.match(r"^support(?: for)? (.+)$", s, re.I)
    if m:
        return ensure_period("Added support for " + m.group(1))
    if first.lower() in PAST_TENSE:
        return ensure_period(PAST_TENSE[first.lower()] + " " + rest)
    return ensure_period(s[:1].upper() + s[1:])


def join_words(items) -> str:
    items = [i for i in items if i]
    if len(items) <= 1:
        return "".join(items)
    return ", ".join(items[:-1]) + " and " + items[-1]


def build_release_notes(issues, version_label: str, settings: dict) -> str:
    product = settings.get("product") or RN_DEFAULTS["product"]
    skip = {product.lower(), f"{product.lower()} 2.0", f"{product.lower()} legacy", TRACK_TAG.lower()}
    groups = {"features": [], "enhancements": [], "bugs": {}, "security": []}
    for it in sorted(issues, key=lambda x: x["id"]):
        section = rn_section(it)
        area = rn_area(it, skip) if section == "bugs" else ""
        text = rn_sentence(it, section, area)
        target = groups["bugs"].setdefault(area, []) if section == "bugs" else groups[section]
        for entry in target:  # merge duplicates
            if entry[0].lower() == text.lower():
                entry[1].append(it["id"])
                break
        else:
            target.append([text, [it["id"]]])

    def bullet(entry):
        text, ids = entry
        ref = f"(Task #{ids[0]})" if len(ids) == 1 else "(Tasks " + ", ".join(f"#{i}" for i in ids) + ")"
        return f"- {text[:-1] if text.endswith('.') else text}. {ref}"

    n_feat, n_enh = len(groups["features"]), len(groups["enhancements"])
    n_bug = sum(len(v) for v in groups["bugs"].values())
    n_sec = len(groups["security"])
    parts = [plural(n_feat, "new feature") if n_feat else "", plural(n_enh, "enhancement") if n_enh else "",
             plural(n_bug, "bug fix", "bug fixes") if n_bug else ""]
    overview = "This release delivers " + (join_words(parts) or "maintenance updates") + "."
    if groups["features"]:
        heads = [e[0].rstrip(".") for e in groups["features"][:2]]
        overview += " Highlights: " + "; ".join(heads) + "."
    areas = [a for a, v in sorted(groups["bugs"].items(), key=lambda kv: -len(kv[1])) if a != "General"][:5]
    if areas:
        overview += f" Fixes span {join_words(areas)}."
    if n_sec:
        overview += " It also includes " + ("a security fix." if n_sec == 1 else f"{n_sec} security fixes.")

    lines = [f"Release Notes — {product} {version_label}".strip(), settings.get("subtitle") or "", "",
             "1. Overview", overview, "", "2. Enhancements / New Features"]
    if groups["features"]:
        lines += ["New Features"] + [bullet(e) for e in groups["features"]]
    if groups["enhancements"]:
        lines += ["Enhancements"] + [bullet(e) for e in groups["enhancements"]]
    if not n_feat and not n_enh:
        lines.append("No new features or enhancements in this release.")
    lines += ["", "3. Bug Fixes"]
    if n_bug:
        order = sorted(groups["bugs"], key=lambda a: (a == "General", -len(groups["bugs"][a]), a.lower()))
        for area in order:
            lines += [area] + [bullet(e) for e in groups["bugs"][area]]
    else:
        lines.append("No bug fixes in this release.")
    lines += ["", "4. Security Updates"]
    lines += [bullet(e) for e in groups["security"]] or ["None in this release."]
    lines += ["", "5. Known Issues", settings.get("known_issues") or RN_DEFAULTS["known_issues"], "",
              "6. Upgrade / Deployment Notes", settings.get("upgrade") or RN_DEFAULTS["upgrade"], "",
              "7. Component Versions"]
    lines += [ln for ln in (settings.get("components") or "").splitlines() if ln.strip()]
    return "\n".join(lines) + "\n"


def rn_ai_prompt(issues, version_label: str, settings: dict, draft: str) -> str:
    rows = "\n".join("- #%s [%s] [%s] %s :: %s" % (it["id"], name_of(it.get("tracker")),
                                                  name_of(it.get("category"), ""), it.get("subject", ""),
                                                  truncate(" ".join((it.get("description") or "").split()), 300))
                     for it in sorted(issues, key=lambda x: x["id"]))
    return ("You are helping a Product Owner write customer-facing release notes. Rewrite the DRAFT below into "
            "polished release notes. Keep EXACTLY the same plain-text structure (title line, subtitle line, the "
            "numbered section titles 1-7, sub-group lines such as 'New Features', 'Enhancements' or a bug-fix "
            "area name, and bullet lines starting with '- ' that end with the ticket reference like "
            "'(Task #1234)'). Write each bullet as one clear past-tense sentence describing the user-visible "
            "change (e.g. 'Fixed an issue where…', 'Added…', 'Improved…'). Group bug fixes by product area. "
            "Write a 3-4 sentence overview. Do not invent changes that are not in the ticket list. No Markdown "
            "symbols other than the '- ' bullets.\n\n"
            f"=== DRAFT ===\n{draft}\n=== TICKETS IN {version_label} ===\n{rows}\n")


def _protos_from_template(doc):
    from copy import deepcopy
    from docx.oxml.ns import qn
    protos, seen_heading = {}, False

    def grab(p):
        pPr = p._p.pPr
        r = p.runs[0]._r if p.runs else None
        rPr = r.rPr if r is not None else None
        return (deepcopy(pPr) if pPr is not None else None, deepcopy(rPr) if rPr is not None else None)

    texted = [p for p in doc.paragraphs if p.text.strip()]
    if texted:
        protos["title"] = grab(texted[0])
    if len(texted) > 1:
        protos["subtitle"] = grab(texted[1])
    for p in texted[2:]:
        txt = p.text.strip()
        has_num = p._p.pPr is not None and p._p.pPr.find(qn("w:numPr")) is not None
        bold = bool(p.runs and p.runs[0].bold)
        if HEADING_RE.match(txt):
            protos.setdefault("heading", grab(p))
            seen_heading = True
        elif has_num:
            protos.setdefault("bullet", grab(p))
        elif seen_heading and bold:
            protos.setdefault("sub", grab(p))
        elif seen_heading:
            protos.setdefault("para", grab(p))
    return protos


def _fresh_rn_doc(settings):
    from docx import Document
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from copy import deepcopy
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.5)
    sec.top_margin, sec.bottom_margin = Inches(1.0), Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(9)
    hp = sec.header.paragraphs[0]
    for para_ in (hp, sec.footer.paragraphs[0]):
        for pos in (3.25, 6.5):
            para_.paragraph_format.tab_stops.add_tab_stop(Inches(pos), WD_TAB_ALIGNMENT.CLEAR)
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(7.5), WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run("\t" + (settings.get("website") or ""))
    r.bold, r.font.size = True, Pt(9)
    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(7.5), WD_TAB_ALIGNMENT.RIGHT)
    r = fp.add_run("Confidential — Internal Use Only\tPage ")
    r.font.size, r.font.color.rgb = Pt(9), RGBColor(0x6B, 0x72, 0x80)
    for kind, text in (("begin", None), ("instr", "PAGE"), ("end", None)):
        run = fp.add_run()
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)

    def proto(size, bold=False, color=None, before=0, after=3, style=None):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        run = p.add_run("x")
        run.font.size, run.bold = Pt(size), bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        out = (deepcopy(p._p.pPr) if p._p.pPr is not None else None, deepcopy(run._r.rPr))
        p._p.getparent().remove(p._p)
        return out

    protos = {"title": proto(23, True, "000000", 10, 6), "subtitle": proto(9, False, "6B7280", 0, 2),
              "heading": proto(20, True, None, 6, 3), "sub": proto(8.5, True, None, 4, 2),
              "para": proto(8.5, False, None, 3, 3), "bullet": proto(8.5, False, None, 1, 1, "List Bullet")}
    return doc, protos


def rn_to_docx(blocks, path: str, settings: dict):
    from copy import deepcopy
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    template = settings.get("template") or ""
    protos = {}
    if template and os.path.exists(template):
        doc = Document(template)
        protos = _protos_from_template(doc)
        body = doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
    else:
        doc, protos = _fresh_rn_doc(settings)
    body = doc.element.body
    fallback = {"subtitle": "para", "heading": "title", "sub": "para", "bullet": "para", "para": "title"}

    def get(kind):
        k = kind
        for _ in range(3):
            if k in protos:
                return protos[k]
            k = fallback.get(k, "para")
        return (None, None)

    def set_bold(rPr, on):
        for el in rPr.findall(qn("w:b")) + rPr.findall(qn("w:bCs")):
            rPr.remove(el)
        if on:
            b = OxmlElement("w:b")
            rPr.insert(0, b)

    def add(kind, text, lead="", bold=None, prefix=""):
        pPr, rPr = get(kind)
        p = OxmlElement("w:p")
        if pPr is not None:
            p.append(deepcopy(pPr))
        for chunk, is_bold in ((lead + " " if lead else "", True), (prefix + text, bold)):
            if not chunk:
                continue
            r = OxmlElement("w:r")
            rp = deepcopy(rPr) if rPr is not None else OxmlElement("w:rPr")
            if is_bold is not None:
                set_bold(rp, is_bold)
            r.append(rp)
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = chunk
            r.append(t)
            p.append(r)
        sect = body.find(qn("w:sectPr"))
        if sect is not None:
            sect.addprevious(p)
        else:
            body.append(p)

    for b in blocks:
        t = b["type"]
        if t == "title":
            add("title", b["text"])
        elif t == "subtitle":
            add("subtitle", b["text"])
            add("para", "")
        elif t == "heading":
            add("heading", b["text"])
        elif t == "sub":
            add("sub", b["text"], bold=True)
        elif t == "bullet":
            add("bullet", b["text"], prefix="" if "bullet" in protos else "•  ")
        elif t == "para":
            add("para", b["text"], lead=b.get("lead", ""))
        elif t == "table":
            for row in b["rows"]:
                add("para", " — ".join(" ".join(c) if isinstance(c, list) else c for c in row))
    doc.save(path)


# --------------------------------------------------------------------------
# Retrospective document (standard format) and improvement tracking
# --------------------------------------------------------------------------
NO_RELEASE = "(No release)"
RETRO_STATUSES = ["Not started", "In progress", "Done"]
RETRO_THEMES = {
    "Requirements & clarity": ("requirement", "prd", "clarity", "unclear", "scope", "spec", "acceptance",
                               "ambigu", "changing", "change request"),
    "Estimation & planning": ("estimat", "planning", "plan ", "capacity", "deadline", "due date", "overcommit",
                              "spill", "carry", "timeline", "sprint goal"),
    "Testing & quality": ("test", "qa", "bug", "regression", "defect", "quality", "uat"),
    "Communication & meetings": ("communicat", "meeting", "sync", "inform", "handoff", "hand-off", "standup",
                                 "stand-up", "update", "clarif"),
    "Deployment & environments": ("deploy", "release", "environment", "build", "server", "staging", "pipeline",
                                  "downtime"),
    "Dependencies & blockers": ("depend", "block", "waiting", "access", "third party", "third-party", "vendor"),
    "Process & tracking": ("redmine", "process", "status", "ticket", "tracking", "document", "jira"),
}
RETRO_HELP = ("Anonymous: please don't add names. One point per cell, in your own words. Keep it about the work "
              "and the process; Kudos is for thanking people. Please don't rename the three column headings.")
RETRO_COLUMNS = [("well", "😊  What went well", "E2F5EA", "2E7D4F"),
                 ("better", "🔧  What could have gone better", "FFF1DC", "9A5B00"),
                 ("kudos", "🌟  Kudos", "E3ECFD", "1D4FBF")]


def build_retro_xlsx(path: str, info: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "Retro"
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = f"Retrospective — {info.get('release') or '<Release version>'}"
    ws["A1"].font = Font(size=16, bold=True, color="2F5496")
    meta = [x for x in (info.get("project"), f"Sprints: {info['sprints']}" if info.get("sprints") else "",
                        f"Retro date: {info.get('date') or '____________'}") if x]
    ws["A2"] = "   |   ".join(meta)
    ws["A2"].font = Font(size=10, color="808080")
    ws["A3"] = RETRO_HELP
    ws["A3"].font = Font(size=10, italic=True, color="808080")
    ws.merge_cells("A3:C3")
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[3].height = 30
    for ci, (_, title, fill, color) in enumerate(RETRO_COLUMNS, 1):
        c = ws.cell(row=5, column=ci, value=title)
        c.font = Font(size=12, bold=True, color=color)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
        c.border = border
        ws.column_dimensions[c.column_letter].width = 48
    ws.row_dimensions[5].height = 28
    rows = int(info.get("rows", 25))
    for r in range(6, 6 + rows):
        ws.row_dimensions[r].height = 36
        for ci in range(1, 4):
            c = ws.cell(row=r, column=ci)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = border
    ws.freeze_panes = "A6"
    ws.sheet_view.zoomScale = 110
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "5:5"
    wb.save(path)


def parse_retro_xlsx(path: str) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = {"title": "", "release": "", "well": [], "better": [], "kudos": []}
    keys = {"well": ("went well",), "better": ("better", "could have", "improve"), "kudos": ("kudos", "thank")}
    for ws in wb.worksheets:
        header_row, cols = None, {}
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
            found = {}
            for c in row:
                txt = str(c.value or "").lower()
                for k, words in keys.items():
                    if k not in found and any(w in txt for w in words) and len(txt) < 80:
                        found[k] = c.column
            if len(found) >= 2:
                header_row, cols = row[0].row, found
                break
        if not header_row:
            continue
        title = str(ws["A1"].value or "")
        out["title"] = title
        m = re.search(r"[—–-]\s*(.+)$", title)
        out["release"] = m.group(1).strip() if m and "<" not in m.group(1) else ""
        for r in range(header_row + 1, ws.max_row + 1):
            for k, col in cols.items():
                v = ws.cell(row=r, column=col).value
                text = re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""
                if text:
                    out[k].append({"text": text})
        break
    return out


def retro_themes(retros: dict):
    hits = {}
    for release, data in retros.items():
        for item in data.get("better", []):
            text = (item.get("text", "") + " " + item.get("suggestion", "")).lower()
            for theme, words in RETRO_THEMES.items():
                if any(w in text for w in words):
                    h = hits.setdefault(theme, {"count": 0, "releases": set(), "examples": []})
                    h["count"] += 1
                    h["releases"].add(release)
                    if len(h["examples"]) < 3:
                        h["examples"].append(item.get("text", ""))
    return sorted(((t, h["count"], sorted(h["releases"]), h["examples"]) for t, h in hits.items()),
                  key=lambda x: (-len(x[2]), -x[1], x[0]))


def retro_ai_prompt(release: str, data: dict, themes, open_actions) -> str:
    def lines(items, fmt):
        return "\n".join("- " + fmt(i) for i in items) or "- (none)"
    return (
        "You are an experienced agile coach helping a Product Owner act on a team retrospective. Using the "
        "retro below, write:\n1. A 3-sentence summary of the retro.\n2. The top 3 improvements to make next "
        "release, each with a concrete action, a suggested owner role and how we'll know it worked.\n"
        "3. Which previous open actions should be closed, kept or escalated.\n4. A short thank-you note "
        "for the team using the kudos.\nBe practical and specific; don't invent facts.\n\n"
        f"=== RETRO: {release} ===\nWhat went well:\n"
        + lines(data.get("well", []), lambda i: i.get("text", "")) +
        "\n\nWhat could have gone better:\n"
        + lines(data.get("better", []), lambda i: i.get("text", "")) +
        "\n\nKudos:\n" + lines(data.get("kudos", []), lambda i: i.get("text", "")) +
        "\n\nThemes that keep coming up across retros:\n"
        + lines(themes, lambda t: f"{t[0]} ({t[1]} mentions in {', '.join(t[2])})") +
        "\n\nOpen actions from earlier retros:\n"
        + lines(open_actions, lambda a: f"{a['action']} (owner {a.get('owner') or '?'}, {a['status']}, "
                                        f"from {a['release']})") + "\n")


# --------------------------------------------------------------------------
# Team, handover messages and developer suggestions
# --------------------------------------------------------------------------
DEFAULT_TEAM = []  # (name, designation, relieving date) tuples to pre-fill the Team tab on first run
DESIGNATIONS = ["Developer", "Lead / Developer", "QA", "Product Owner", "Business Analyst", "Designer", "Sales",
                "Manager", "Other"]
SESSION_OUTCOMES = ["Groomed", "Ready for sprint", "Needs PRD", "Needs clarification", "Split ticket", "Deprioritise",
                    "Not discussed"]
SUGGEST_STOP = set("the a an and or for to of in on with without is are be when while from into not no can should "
                   "user users able when then that this there their have has show shows showing issue issues "
                   "error option options".split())


def norm_name(name: str) -> str:
    return " ".join((name or "").lower().split())


def first_name(name: str) -> str:
    return (name or "").split(" ")[0].capitalize() if name else "there"


def _words(text: str) -> set:
    return {w for w in re.findall(r"[a-z][a-z0-9]{3,}", (text or "").lower()) if w not in SUGGEST_STOP}


def handover_message(person: str, last_day, tickets, base_url: str) -> str:
    lines = [f"Hi {first_name(person)},", "",
             f"Before your last day ({fmt_date(last_day)}), could we plan a smooth handover of your open tickets? "
             f"You currently have {plural(len(tickets), 'open ticket')}:", ""]
    for it in tickets:
        due = parse_date(it.get("due_date"))
        lines.append(f"• #{it['id']} {it.get('subject', '')} ({name_of(it.get('status'))}"
                     + (f", due {fmt_date(due, False)}" if due else "") + ")")
    lines += ["", "For each one, please add a short status note in Redmine (what's done, what's left, any "
                  "gotchas) and let me know who you'd suggest to take it over.", "", "Thank you!",
              f"{base_url.rstrip('/')}/issues?assigned_to_id=me" if base_url else ""]
    return "\n".join(lines).rstrip() + "\n"


def nudge_message(issue: dict, flags, base_url: str) -> str:
    person = name_of(issue.get("assigned_to"), "")
    why = "; ".join(flags) if flags else f"it is currently '{name_of(issue.get('status'))}'"
    return (f"Hi {first_name(person)},\n\nQuick check on #{issue.get('id')} {issue.get('subject', '')}: {why}.\n"
            f"Could you share an update and update the status in Redmine if it has moved on? Thanks!\n"
            f"{base_url.rstrip('/')}/issues/{issue.get('id')}\n")


def suggest_developers(issue: dict, team: list, history: list, open_issues: list, pushes_by_name: dict,
                       exclude=(), today=None, min_days_left: int = 10):
    """Rank team members for a ticket. Returns (ranked list, skipped notes)."""
    today = today or dt.date.today()
    due = parse_date(issue.get("due_date"))
    words = _words(issue.get("subject", "") + " " + (issue.get("description") or "")[:400])
    cat = name_of(issue.get("category"), "")
    parent = (issue.get("parent") or {}).get("id")
    area = rn_area(issue)
    workload = Counter(norm_name(name_of(it.get("assigned_to"), "")) for it in open_issues)
    skipped, ranked = [], []
    for m in team:
        nm, desig = norm_name(m.get("name")), (m.get("designation") or "").lower()
        if nm in exclude or "dev" not in desig:
            continue
        if m.get("status", "Active") != "Active":
            skipped.append(f"{m['name']}: inactive")
            continue
        rel = parse_date(m.get("relieving"))
        if rel and (rel < today or (due and rel < due) or (rel - today).days < min_days_left):
            skipped.append(f"{m['name']}: leaves {fmt_date(rel, False)}")
            continue
        sim, examples = 0.0, []
        for h in history:
            if h.get("id") == issue.get("id") or norm_name(name_of(h.get("assigned_to"), "")) != nm:
                continue
            sc = 0
            if cat and name_of(h.get("category"), "") == cat:
                sc += 3
            if parent and (h.get("parent") or {}).get("id") == parent:
                sc += 3
            if area != "General" and rn_area(h) == area:
                sc += 2
            overlap = len(words & _words(h.get("subject", "")))
            if overlap >= 2:
                sc += min(3, overlap - 1)
            if sc:
                sim += sc
                examples.append(h["id"])
        load = workload.get(nm, 0)
        pushes = pushes_by_name.get(nm, 0)
        score = sim - 0.6 * load - 0.5 * pushes
        reasons = []
        if examples:
            reasons.append(f"{plural(len(examples), 'similar ticket')} before (e.g. #{examples[0]})")
        else:
            reasons.append("no similar tickets found")
        reasons.append(f"{plural(load, 'open ticket')} now")
        if pushes:
            reasons.append(f"{plural(pushes, 'due-date push', 'due-date pushes')} recently")
        if rel:
            reasons.append(f"⚠ on notice, leaves {fmt_date(rel, False)}")
            score -= 2
        ranked.append({"name": m["name"], "score": score, "similar": len(examples), "lead": "lead" in desig,
                       "reasons": reasons})
    devs = sorted((r for r in ranked if not r["lead"]), key=lambda r: (-r["score"], r["name"]))
    leads = sorted((r for r in ranked if r["lead"]), key=lambda r: (-r["score"], r["name"]))
    top = devs[:3]
    if not top or all(r["similar"] == 0 for r in top):
        for r in leads[:1]:
            r["reasons"].append("lead: suggested only because nobody else fits well")
            top.append(r)
    return top, skipped


# --------------------------------------------------------------------------
# PRD status detection for the ticket list
# --------------------------------------------------------------------------
PRD_STATES = ["Needed", "Partial", "Draft", "Has PRD", "Not needed"]
PRD_ICONS = {"Needed": "✖ Needed", "Partial": "◐ Partial", "Draft": "✎ Draft", "Has PRD": "✔ Has PRD",
             "Not needed": "— Not needed"}
PRD_FILE_RE = re.compile(r"prd|spec|requirement|\bbrd\b|\bfrd\b", re.I)
PRD_LINK_RE = re.compile(r"\bPRD\b|product requirements?|docs\.google\.com|confluence|notion\.so|"
                         r"sharepoint\.com|1drv\.ms|onedrive\.live", re.I)


def prd_state(issue: dict, has_draft: bool, exported: dict = None, journals=None, override: dict = None):
    """Work out whether a ticket already has a PRD. Returns (state, reason)."""
    if override and override.get("state") in PRD_STATES:
        note = f": {override['note']}" if override.get("note") else ""
        return override["state"], f"marked by you on {override.get('date', '?')}{note}"
    for att in issue.get("attachments") or []:
        name = att.get("filename", "")
        if PRD_FILE_RE.search(name):
            return "Has PRD", f"attachment '{name}' in Redmine"
    desc = issue.get("description") or ""
    m = PRD_LINK_RE.search(desc)
    if m:
        return "Has PRD", f"description mentions or links '{m.group(0)}'"
    for j in reversed(journals or []):
        notes = j.get("notes") or ""
        if not notes.strip():
            continue
        low = notes.lower()
        m = PRD_LINK_RE.search(notes)
        written = len(notes) > 400 and "acceptance" in low and ("requirement" in low or "objective" in low)
        if m or written:
            who = name_of(j.get("user"), "someone")
            return "Has PRD", f"PRD in a comment by {who} on {(j.get('created_on') or '')[:10]}"
    if exported:
        return "Has PRD", f"PRD exported from this app on {exported.get('date', '?')}"
    if has_draft:
        return "Draft", "PRD draft in progress in this app"
    sec = analyse_description(desc)
    if (sec["ac"]["bullets"] or sec["ac"]["paras"]) and len(desc) >= 120:
        return "Partial", "ticket has acceptance criteria, but no PRD"
    return "Needed", "no PRD found"


# --------------------------------------------------------------------------
# Dashboard metrics
# --------------------------------------------------------------------------
AGE_BUCKETS = [("0–7 days", 0, 7, "#8FB39A"), ("8–30 days", 8, 30, "#8BA5C4"),
               ("31–90 days", 31, 90, "#E3B26E"), ("91–180 days", 91, 180, "#D99A6C"),
               ("180+ days", 181, 10 ** 9, "#C9665E")]


def fetch_dashboard(client: RedmineClient, project_id, mine_only: bool) -> dict:
    """Everything the dashboard needs from Redmine, downloaded once. Filters are applied afterwards
    (build_dashboard), so changing a filter is instant."""
    pids = project_id if isinstance(project_id, list) else ([project_id] if project_id else [])
    today = dt.date.today()
    scope = scope_params(None, mine_only)
    open_issues, open_total = client.issues_in({**scope, "status_id": "open", "sort": "id:desc"}, pids)
    try:
        prio_list = client.priorities()
    except RedmineError:
        prio_list = []
    since = (today - dt.timedelta(days=62)).isoformat()
    try:
        created, _ = client.issues_in({**scope, "status_id": "*", "created_on": f">={since}"}, pids)
        closed, _ = client.issues_in({**scope, "status_id": "closed", "closed_on": f">={since}"}, pids)
        flow_ok = True
    except RedmineError:
        created, closed, flow_ok = [], [], False
    native = any(native_tags(it) is not None for it in open_issues)
    if native:
        track_ids = {it["id"] for it in open_issues + created + closed
                     if any(t.lower() == TRACK_TAG.lower() for t in (native_tags(it) or []))}
        track_ok = True
    else:
        track_ids, track_ok = client.tag_ids_any_status(pids, mine_only, TRACK_TAG)
    return {"open_issues": open_issues, "open_total": open_total, "prio_list": prio_list, "created": created,
            "closed": closed, "flow_ok": flow_ok, "track_ids": track_ids, "track_ok": track_ok,
            "fetched": dt.datetime.now()}


def build_dashboard(raw: dict, scope_label: str, keep=None, filtered=False) -> dict:
    """Work out every number and chart from downloaded data. keep(issue) -> bool applies the filters."""
    keep = keep or (lambda it: True)
    today = dt.date.today()
    open_issues = [it for it in raw["open_issues"] if keep(it)]
    open_total = len(open_issues) if filtered else raw["open_total"]
    n = len(open_issues)
    prio_list = raw["prio_list"]
    prio_rank = {p.get("name"): i for i, p in enumerate(prio_list)}
    default_idx = next((i for i, p in enumerate(prio_list) if p.get("is_default")), None)
    if prio_list and default_idx is not None:
        high_names = {p.get("name") for i, p in enumerate(prio_list) if i > default_idx}
    else:
        high_names = None

    def is_high(issue):
        name = name_of(issue.get("priority"), "")
        if high_names is not None:
            return name in high_names
        return any(w in name.lower() for w in HIGH_PRIORITY_WORDS)

    ages, overdue, stale, unassigned, high = [], [], [], [], []
    described = estimated = 0
    for it in open_issues:
        created = parse_date(it.get("created_on"))
        updated = parse_date(it.get("updated_on"))
        due = parse_date(it.get("due_date"))
        if created:
            ages.append(((today - created).days, it))
        if due and due < today:
            overdue.append(it)
        if updated and (today - updated).days > STALE_DAYS:
            stale.append(it)
        if not it.get("assigned_to"):
            unassigned.append(it)
        if is_high(it):
            high.append(it)
        if len((it.get("description") or "").strip()) >= 80:
            described += 1
        if it.get("estimated_hours"):
            estimated += 1

    uses_estimates = estimated > 0
    ready = sum(1 for it in open_issues
                if len((it.get("description") or "").strip()) >= 80
                and (it.get("estimated_hours") or not uses_estimates))
    ready_rule = ("clear description and an estimate" if uses_estimates
                  else "clear description; estimates aren't used in this backlog")

    # Flow metrics, counted from the downloaded created / closed tickets (so filters apply)
    created_l = [it for it in raw["created"] if keep(it)]
    closed_l = [it for it in raw["closed"] if keep(it)]
    flow_ok = raw["flow_ok"]

    def count(items, field, start, end):
        if not flow_ok:
            return None
        return sum(1 for it in items if parse_date(it.get(field)) and start <= parse_date(it.get(field)) <= end)

    d30 = today - dt.timedelta(days=30)
    created30 = count(created_l, "created_on", d30, today)
    closed30 = count(closed_l, "closed_on", d30, today)
    weeks = []
    for k in range(7, -1, -1):
        end = today - dt.timedelta(days=7 * k)
        start = end - dt.timedelta(days=6)
        weeks.append({"label": start.strftime("%d %b"), "start": start, "end": end,
                      "created": count(created_l, "created_on", start, end),
                      "closed": count(closed_l, "closed_on", start, end)})
    closed_vals = [w["closed"] for w in weeks if w["closed"] is not None]
    throughput = (sum(closed_vals) / len(closed_vals)) if closed_vals else 0.0
    weeks_to_clear = (open_total / throughput) if throughput > 0 else None
    net30 = (created30 - closed30) if (created30 is not None and closed30 is not None) else None

    # Breakdowns
    prio_counts = Counter(name_of(it.get("priority")) for it in open_issues)
    by_priority = sorted(prio_counts.items(), key=lambda kv: -prio_rank.get(kv[0], -1))
    by_status = Counter(name_of(it.get("status")) for it in open_issues).most_common()
    by_tracker = Counter(name_of(it.get("tracker")) for it in open_issues).most_common()
    by_assignee = Counter(name_of(it.get("assigned_to"), "Unassigned") for it in open_issues).most_common()
    age_buckets = []
    for label, lo, hi, colour in AGE_BUCKETS:
        age_buckets.append((label, sum(1 for d, _ in ages if lo <= d <= hi), colour))
    age_days = [d for d, _ in ages]
    oldest = max(ages, key=lambda x: x[0]) if ages else None

    # Items needing attention
    risks = []
    for it in open_issues:
        reasons, score = [], 0.0
        due = parse_date(it.get("due_date"))
        updated = parse_date(it.get("updated_on"))
        hp = is_high(it)
        if due and due < today:
            late = (today - due).days
            reasons.append(f"Overdue by {late}d")
            score += 3 + min(late, 60) / 10
        if hp and not it.get("assigned_to"):
            reasons.append("High priority, unassigned")
            score += 3
        if hp and updated and (today - updated).days > STALE_DAYS:
            reasons.append(f"High priority, no update in {(today - updated).days}d")
            score += 2
        if reasons:
            score += 1 if hp else 0
            risks.append({"issue": it, "why": "; ".join(reasons), "score": score})
    risks.sort(key=lambda r: -r["score"])

    data = {
        "scope_label": scope_label,
        "generated": raw.get("fetched") or dt.datetime.now(),
        "open_total": open_total,
        "analysed": n,
        "high": len(high),
        "overdue": len(overdue),
        "stale": len(stale),
        "unassigned": len(unassigned),
        "described_pct": pct(described, n),
        "estimated_pct": pct(estimated, n),
        "ready_pct": pct(ready, n),
        "ready_rule": ready_rule,
        "created30": created30,
        "closed30": closed30,
        "net30": net30,
        "throughput": throughput,
        "weeks_to_clear": weeks_to_clear,
        "weeks": weeks,
        "by_priority": by_priority,
        "by_status": by_status,
        "by_tracker": by_tracker,
        "by_assignee": by_assignee,
        "age_buckets": age_buckets,
        "avg_age": statistics.mean(age_days) if age_days else 0,
        "median_age": statistics.median(age_days) if age_days else 0,
        "oldest": oldest,
        "risks": risks[:25],
        "risk_total": len(risks),
        "open_issues": open_issues,
    }
    data["summary"] = executive_summary(data)
    return data


def compute_dashboard(client: RedmineClient, project_id, mine_only: bool, scope_label: str) -> dict:
    return build_dashboard(fetch_dashboard(client, project_id, mine_only), scope_label)


def executive_summary(d: dict) -> str:
    s = [f"The backlog has {plural(d['open_total'], 'open item')}: {d['high']} high priority "
         f"and {d['overdue']} overdue."]
    if d["net30"] is not None:
        net = d["net30"]
        if net > 0:
            trend = f"the backlog grew by {net}"
        elif net < 0:
            trend = f"the backlog shrank by {abs(net)}"
        else:
            trend = "the backlog held steady"
        s.append(f"In the last 30 days the team closed {d['closed30']} and took in "
                 f"{d['created30']} new, so {trend}.")
    if d["throughput"] > 0:
        line = f"Delivery pace averages {d['throughput']:.1f} items per week"
        if d["weeks_to_clear"]:
            line += f", so today's open backlog is roughly {d['weeks_to_clear']:.0f} weeks of work"
        s.append(line + ".")
    s.append(f"{d['ready_pct']}% of open items are sprint-ready ({d['ready_rule']}).")
    if d["risk_total"]:
        s.append(f"{plural(d['risk_total'], 'item')} need attention (overdue or stalled high priority).")
    return " ".join(s)


def build_status_report(d: dict, base_url: str) -> str:
    gen = d["generated"].strftime("%d %b %Y, %H:%M")
    rows = [
        ("Open items", d["open_total"]),
        ("High priority", d["high"]),
        ("Overdue", d["overdue"]),
        ("Unassigned", d["unassigned"]),
        (f"No update in {STALE_DAYS}+ days", d["stale"]),
        ("Sprint-ready", f"{d['ready_pct']}%"),
        ("Closed in last 30 days", d["closed30"] if d["closed30"] is not None else "n/a"),
        ("Created in last 30 days", d["created30"] if d["created30"] is not None else "n/a"),
        ("Net backlog change (30 days)", f"{d['net30']:+d}" if d["net30"] is not None else "n/a"),
        ("Avg weekly throughput (8 weeks)", f"{d['throughput']:.1f}"),
        ("Average age of open items", f"{d['avg_age']:.0f} days"),
        ("Median age of open items", f"{d['median_age']:.0f} days"),
    ]
    out = [
        "# Backlog Status Report",
        "",
        f"**Scope:** {d['scope_label']}  ",
        f"**Generated:** {gen}  ",
        f"**Source:** Redmine ({base_url})",
        "",
        "## Executive Summary",
        "",
        d["summary"],
        "",
        "## Key Metrics",
        "",
        "| Metric | Value |",
        "|---|---|",
    ]
    out += [f"| {md_cell(k)} | {md_cell(v)} |" for k, v in rows]
    out += ["", "## Throughput: last 8 weeks", "", "| Week starting | Created | Closed |", "|---|---|---|"]
    for w in d["weeks"]:
        out.append(f"| {w['start'].strftime('%d %b %Y')} | "
                   f"{w['created'] if w['created'] is not None else 'n/a'} | "
                   f"{w['closed'] if w['closed'] is not None else 'n/a'} |")

    def table(title, pairs):
        total = sum(v for _, v in pairs) or 1
        rows_ = ["", f"## {title}", "", "| | Count | Share |", "|---|---|---|"]
        rows_ += [f"| {md_cell(k)} | {v} | {pct(v, total)}% |" for k, v in pairs]
        return rows_

    out += table("Open Items by Priority", d["by_priority"])
    out += table("Open Items by Status", d["by_status"])
    out += table("Open Items by Tracker", d["by_tracker"])
    out += table("Workload by Assignee", d["by_assignee"])
    out += table("Age of Open Items", [(lbl, v) for lbl, v, _ in d["age_buckets"]])
    out += ["", "## Items Needing Attention", ""]
    if d["risks"]:
        out += ["| Ticket | Subject | Priority | Assignee | Why |", "|---|---|---|---|---|"]
        for r in d["risks"][:15]:
            it = r["issue"]
            out.append(f"| [#{it['id']}]({base_url}/issues/{it['id']}) | {md_cell(truncate(it.get('subject', ''), 80))} | "
                       f"{md_cell(name_of(it.get('priority')))} | {md_cell(name_of(it.get('assigned_to'), 'Unassigned'))} | "
                       f"{md_cell(r['why'])} |")
    else:
        out.append("Nothing flagged: no overdue items and no stalled high-priority work.")
    if d["analysed"] < d["open_total"]:
        out += ["", f"_Breakdowns are based on the first {d['analysed']} of {d['open_total']} open items._"]
    out += ["", f"_Generated by {APP_NAME}._", ""]
    return "\n".join(out)


# --------------------------------------------------------------------------
# Custom widgets: scrollable frame, KPI cards and charts
# --------------------------------------------------------------------------
class ScrollableFrame(ttk.Frame):
    def __init__(self, parent, **kw):
        super().__init__(parent, style="Bg.TFrame", **kw)
        self.canvas = tk.Canvas(self, highlightthickness=0, bg=C["bg"], bd=0)
        vsb = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        self.canvas.pack(side="left", fill="both", expand=True)
        self.inner = ttk.Frame(self.canvas, style="Bg.TFrame")
        self._win = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.inner.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.bind("<Configure>", lambda e: self.canvas.itemconfigure(self._win, width=e.width))
        self.bind_all("<MouseWheel>", self._on_wheel, add="+")
        self.bind_all("<Button-4>", lambda e: self._on_wheel(e, -1), add="+")
        self.bind_all("<Button-5>", lambda e: self._on_wheel(e, 1), add="+")

    def _on_wheel(self, event, direction=None):
        try:
            widget = self.winfo_containing(event.x_root, event.y_root)
        except (KeyError, tk.TclError):
            return
        if widget is None or not str(widget).startswith(str(self)):
            return
        if isinstance(widget, (ttk.Treeview, tk.Text)):
            return
        if direction is not None:
            self.canvas.yview_scroll(direction * 3, "units")
        elif IS_MAC:
            self.canvas.yview_scroll(int(-event.delta), "units")
        else:
            self.canvas.yview_scroll(int(-event.delta / 120) * 3, "units")


class KpiCard(tk.Frame):
    def __init__(self, parent, title: str, accent: str):
        super().__init__(parent, bg=C["panel"], highlightbackground=C["line"], highlightthickness=1)
        tk.Frame(self, bg=accent, height=4).pack(fill="x", side="top")
        self.value = tk.Label(self, text="—", font=FONTS["kpi"], bg=C["panel"], fg=C["ink"])
        self.value.pack(anchor="w", padx=14, pady=(10, 0))
        tk.Label(self, text=title.upper(), font=FONTS["kpi_label"], bg=C["panel"],
                 fg=C["muted"]).pack(anchor="w", padx=14)
        self.sub = tk.Label(self, text=" ", font=FONTS["small"], bg=C["panel"], fg=C["muted"],
                            justify="left", anchor="w")
        self.sub.pack(anchor="w", fill="x", padx=14, pady=(2, 12))
        self.bind("<Configure>", lambda e: self.sub.configure(wraplength=max(120, e.width - 30)))

    def set(self, value, sub: str = "", color: str = None):
        self.value.configure(text=value, fg=color or C["ink"])
        self.sub.configure(text=sub or " ")


class ChartBase(tk.Canvas):
    def __init__(self, parent, title: str, height: int = 260):
        super().__init__(parent, bg=C["panel"], height=height, highlightthickness=1,
                         highlightbackground=C["line"], bd=0)
        self.title = title
        self.items = []
        self.empty_msg = "Click 'Refresh Dashboard' to load data"
        self.bind("<Configure>", lambda e: self.redraw())

    def set_data(self, items, empty_msg: str = "No data for this scope"):
        self.items = items
        self.empty_msg = empty_msg
        self.redraw()

    def _frame(self):
        self.delete("all")
        w, h = self.winfo_width(), self.winfo_height()
        self.create_text(16, 20, text=self.title, anchor="w", font=FONTS["h3"], fill=C["ink"])
        return w, h

    def _empty(self, w, h):
        self.create_text(w / 2, h / 2 + 10, text=self.empty_msg, fill=C["muted"], font=FONTS["body"])

    def redraw(self):
        raise NotImplementedError


class HBarChart(ChartBase):
    """items: list of (label, value, colour). Set .on_pick to make bars clickable."""

    def __init__(self, parent, title: str, height: int = 260):
        super().__init__(parent, title, height)
        self.on_pick = None
        self._rows = []
        self.bind("<Button-1>", self._click)
        self.bind("<Motion>", self._motion)

    def _row_at(self, y):
        return next((label for y1, y2, label in self._rows if y1 <= y <= y2), None)

    def _click(self, event):
        label = self._row_at(event.y)
        if label is not None and self.on_pick:
            self.on_pick(label)

    def _motion(self, event):
        self.configure(cursor="hand2" if self.on_pick and self._row_at(event.y) is not None else "")

    def redraw(self):
        w, h = self._frame()
        self._rows = []
        if w < 60:
            return
        items = [i for i in self.items if i[1] > 0][:8] if any(i[1] for i in self.items) else []
        if not items:
            self._empty(w, h)
            return
        top, bottom = 44, h - 14
        row_h = min(30, (bottom - top) / len(items))
        label_w = min(160, int(w * 0.36))
        x0 = 16 + label_w
        bar_max = max(20, w - x0 - 56)
        max_v = max(v for _, v, _ in items) or 1
        for idx, (label, value, colour) in enumerate(items):
            y = top + idx * row_h + row_h / 2
            self.create_text(x0 - 10, y, text=truncate(label, 24), anchor="e",
                             font=FONTS["small"], fill=C["ink"])
            self.create_rectangle(x0, y - row_h * 0.3, x0 + bar_max, y + row_h * 0.3,
                                  fill=C["track"], width=0)
            bw = max(3, bar_max * value / max_v)
            self.create_rectangle(x0, y - row_h * 0.3, x0 + bw, y + row_h * 0.3, fill=colour, width=0)
            self.create_text(x0 + bw + 6, y, text=str(value), anchor="w",
                             font=FONTS["small_bold"], fill=C["ink"])
            self._rows.append((y - row_h / 2, y + row_h / 2, label))
        if self.on_pick:
            self.create_text(w - 14, 20, anchor="e", font=FONTS["small"], fill=C["accent"],
                             text="👆 click a bar to see its tickets")


class DonutChart(ChartBase):
    """items: list of (label, value, colour). Set .on_pick to make slices clickable."""

    def __init__(self, parent, title: str, height: int = 260):
        super().__init__(parent, title, height)
        self.on_pick = None
        self.bind("<Button-1>", self._click)

    def redraw(self):
        w, h = self._frame()
        if w < 60:
            return
        items = [i for i in self.items if i[1] > 0]
        total = sum(v for _, v, _ in items)
        if not total:
            self._empty(w, h)
            return
        size = max(60, min(h - 64, w * 0.45))
        cx, cy = 24 + size / 2, 40 + (h - 40) / 2
        x1, y1, x2, y2 = cx - size / 2, cy - size / 2, cx + size / 2, cy + size / 2
        start = 90.0
        self._hit = {"cx": cx, "cy": cy, "r": size / 2, "hole": size * 0.31, "slices": [], "legend": []}
        for label, value, colour in items:
            extent = -360.0 * value / total
            self._hit["slices"].append((start, extent, label))
            if abs(extent) >= 359.9:
                self.create_oval(x1, y1, x2, y2, fill=colour, outline=C["panel"], width=2)
            else:
                self.create_arc(x1, y1, x2, y2, start=start, extent=extent, fill=colour,
                                outline=C["panel"], width=2, style="pieslice")
            start += extent
        hole = size * 0.31
        self.create_oval(cx - hole, cy - hole, cx + hole, cy + hole, fill=C["panel"], outline=C["panel"])
        self.create_text(cx, cy - 8, text=str(total), font=FONTS["h1"], fill=C["ink"])
        self.create_text(cx, cy + 14, text="open", font=FONTS["small"], fill=C["muted"])
        lx = x2 + 28
        ly = max(48, cy - len(items[:8]) * 13)
        for label, value, colour in items[:8]:
            self.create_rectangle(lx, ly - 6, lx + 12, ly + 6, fill=colour, width=0)
            self.create_text(lx + 20, ly, anchor="w", font=FONTS["small"], fill=C["ink"],
                             text=f"{truncate(label, 18)}  {value}  ({pct(value, total)}%)")
            self._hit["legend"].append((lx - 4, ly - 12, w - 8, ly + 12, label))
            ly += 26
        if getattr(self, "on_pick", None):
            self.create_text(w - 14, 20, anchor="e", font=FONTS["small"], fill=C["accent"],
                             text="👆 click a slice to see its tickets")

    def _click(self, event):
        hit = getattr(self, "_hit", None)
        if not hit or not getattr(self, "on_pick", None):
            return
        for x1, y1, x2, y2, label in hit["legend"]:
            if x1 <= event.x <= x2 and y1 <= event.y <= y2:
                self.on_pick(label)
                return
        dx, dy = event.x - hit["cx"], event.y - hit["cy"]
        r = math.hypot(dx, dy)
        if hit["hole"] < r <= hit["r"]:
            ang = math.degrees(math.atan2(-dy, dx)) % 360
            for start, extent, label in hit["slices"]:
                if (start - ang) % 360 <= -extent + 0.01:
                    self.on_pick(label)
                    return


class TrendChart(ChartBase):
    """items: list of week dicts with 'label', 'created', 'closed'"""

    def redraw(self):
        w, h = self._frame()
        if w < 120:
            return
        lx = w - 190
        for label, colour in (("Created", C["info"]), ("Closed", C["good"])):
            self.create_rectangle(lx, 14, lx + 12, 26, fill=colour, width=0)
            self.create_text(lx + 18, 20, text=label, anchor="w", font=FONTS["small"], fill=C["ink"])
            lx += 90
        weeks = self.items
        vals = [v for wk in weeks for v in (wk["created"], wk["closed"]) if v is not None]
        if not weeks or not vals:
            self._empty(w, h)
            return
        left, right, top, bottom = 48, w - 16, 48, h - 36
        max_v = max(vals + [1])
        step = max(1, math.ceil(max_v / 4))
        y_max = step * 4
        for i in range(5):
            y = bottom - (bottom - top) * i / 4
            self.create_line(left, y, right, y, fill=C["grid"])
            self.create_text(left - 8, y, text=str(step * i), anchor="e", font=FONTS["small"], fill=C["muted"])
        group_w = (right - left) / len(weeks)
        bar_w = max(4, min(26, group_w * 0.3))
        for idx, wk in enumerate(weeks):
            cx = left + group_w * (idx + 0.5)
            for j, (val, colour) in enumerate(((wk["created"], C["info"]), (wk["closed"], C["good"]))):
                if val is None:
                    continue
                bx = cx - bar_w - 2 + j * (bar_w + 4)
                bh = (bottom - top) * val / y_max
                self.create_rectangle(bx, bottom - bh, bx + bar_w, bottom, fill=colour, width=0)
                if val:
                    self.create_text(bx + bar_w / 2, bottom - bh - 8, text=str(val),
                                     font=FONTS["small"], fill=C["ink"])
            label = "This week" if idx == len(weeks) - 1 else wk["label"]
            self.create_text(cx, bottom + 14, text=label, font=FONTS["small"], fill=C["muted"])


# --------------------------------------------------------------------------
# Main application
# --------------------------------------------------------------------------
class MultiPick(ttk.Frame):
    """A dropdown-style button that lets you tick several options (nothing ticked = all)."""

    def __init__(self, parent, all_label: str, on_change, width: int = 22, title: str = "Choose"):
        super().__init__(parent)
        self.all_label, self.on_change, self.title = all_label, on_change, title
        self.options, self.selected = [], []
        self.btn = ttk.Button(self, text=f"{all_label}  ▾", width=width, command=self.open)
        self.btn.pack(fill="x")

    def set_options(self, options):
        self.options = list(options)
        self.selected = [o for o in self.selected if o in self.options]
        self._refresh()

    def set_selected(self, selected):
        self.selected = [o for o in selected if not self.options or o in self.options]
        self._refresh()

    def summary(self) -> str:
        if not self.selected:
            return self.all_label
        if len(self.selected) <= 2:
            return ", ".join(self.selected)
        return f"{len(self.selected)} selected"

    def _refresh(self):
        self.btn.configure(text=f"{truncate(self.summary(), 28)}  ▾")

    def open(self):
        win = tk.Toplevel(self)
        win.title(self.title)
        win.transient(self.winfo_toplevel())
        win.configure(bg=C["bg"])
        win.geometry(f"+{self.btn.winfo_rootx()}+{self.btn.winfo_rooty() + self.btn.winfo_height()}")
        frm = ttk.Frame(win, padding=10, style="Bg.TFrame")
        frm.pack(fill="both", expand=True)
        ttk.Label(frm, text="Click to tick one or more. Nothing ticked = all.", style="Muted.TLabel").pack(anchor="w")
        lf = ttk.Frame(frm)
        lf.pack(fill="both", expand=True, pady=6)
        width = max([28] + [len(o) + 4 for o in self.options])
        lb = tk.Listbox(lf, selectmode="multiple", exportselection=False, height=min(14, max(4, len(self.options))),
                        width=min(width, 60), font=FONTS["body"], activestyle="none", relief="flat",
                        highlightthickness=1, highlightbackground=C["line"], selectbackground=C["select"],
                        selectforeground=C["ink"])
        sb = ttk.Scrollbar(lf, orient="vertical", command=lb.yview)
        lb.configure(yscrollcommand=sb.set)
        lb.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        for i, o in enumerate(self.options):
            lb.insert("end", o)
            if o in self.selected:
                lb.selection_set(i)

        def apply(_e=None):
            self.selected = [self.options[i] for i in lb.curselection()]
            self._refresh()
            win.destroy()
            self.on_change()

        bb = ttk.Frame(frm, style="Bg.TFrame")
        bb.pack(fill="x")
        ttk.Button(bb, text="Clear (= all)", command=lambda: lb.selection_clear(0, "end")).pack(side="left")
        ttk.Button(bb, text="Apply", style="Accent.TButton", command=apply).pack(side="right")
        win.bind("<Return>", apply)
        win.bind("<Escape>", lambda e: win.destroy())
        lb.focus_set()



# --------------------------------------------------------------------------
# Kanban board (read-only: columns follow the tickets' Redmine status)
# --------------------------------------------------------------------------
KANBAN_ORDER = ["New", "Design Review", "In Progress", "Waiting for Review", "Merged", "Dev Testing",
                "Assigned to QA", "QA Testing", "Ready for Release", "Released", "Rework", "Re-open",
                "Closed", "Rejected"]
KANBAN_SIDE = {"rework", "reopen"}
KANBAN_ONLY_IF_USED = {"closed", "rejected"}


def status_key(name) -> str:
    return re.sub(r"[^a-z]", "", (name or "").lower())


def priority_rank(name) -> int:
    low = (name or "").lower()
    keys = list(PRIORITY_COLORS)
    return next((i for i, k in enumerate(keys) if k in low), len(keys))


class KanbanBoard(ttk.Frame):
    COL_W = 238
    GAP = 10
    MAX_CARDS = 80

    def __init__(self, parent, on_select, on_open, on_menu):
        super().__init__(parent, style="Bg.TFrame")
        self.on_select, self.on_open, self.on_menu = on_select, on_open, on_menu
        self.compact = False
        # Summary strip: every column's count at a glance (click to jump there)
        self.summary = tk.Frame(self, bg=C["bg"])
        self.summary.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 6))
        self._chips = []
        self.summary.bind("<Configure>", lambda e: self._flow())
        self.head = tk.Canvas(self, height=68, bg=C["bg"], highlightthickness=0, bd=0)
        self.body = tk.Canvas(self, bg=C["bg"], highlightthickness=0, bd=0)
        vsb = ttk.Scrollbar(self, orient="vertical", command=self.body.yview)
        self.hsb = ttk.Scrollbar(self, orient="horizontal", command=self._xview)
        self.body.configure(yscrollcommand=vsb.set, xscrollcommand=self._xset)
        self.head.grid(row=1, column=0, sticky="ew")
        self.body.grid(row=2, column=0, sticky="nsew")
        vsb.grid(row=2, column=1, sticky="ns")
        self.hsb.grid(row=3, column=0, sticky="ew")
        self.rowconfigure(2, weight=1)
        self.columnconfigure(0, weight=1)
        self.head_inner = tk.Frame(self.head, bg=C["bg"])
        self.head.create_window((0, 0), window=self.head_inner, anchor="nw")
        self.inner = tk.Frame(self.body, bg=C["bg"])
        self.body.create_window((0, 0), window=self.inner, anchor="nw")
        self.inner.bind("<Configure>", lambda e: self._region())
        self.bind_all("<MouseWheel>", self._wheel, add="+")
        self.bind_all("<Shift-MouseWheel>", lambda e: self._wheel(e, horizontal=True), add="+")
        self.bind_all("<Button-4>", lambda e: self._wheel(e, step=-1), add="+")
        self.bind_all("<Button-5>", lambda e: self._wheel(e, step=1), add="+")
        self.cards, self.sel = {}, None
        self.multi = set()          # ids picked with Ctrl/Shift-click or "select column"
        self.order = []             # (column index, id) in screen order, for Shift-click ranges
        self.on_multi = None        # callback(set_of_ids)

    def _xview(self, *args):
        self.body.xview(*args)
        self.head.xview(*args)

    def _xset(self, first, last):
        self.hsb.set(first, last)
        self.head.xview_moveto(first)

    def _region(self):
        w = max(self.inner.winfo_reqwidth(), self.head_inner.winfo_reqwidth())
        self.body.configure(scrollregion=(0, 0, w, self.inner.winfo_reqheight() + 10))
        self.head.configure(scrollregion=(0, 0, w, 68))

    def _wheel(self, event, horizontal=False, step=None):
        try:
            w = self.winfo_containing(event.x_root, event.y_root)
        except (KeyError, tk.TclError):
            return
        if w is None or not str(w).startswith(str(self)) or not self.winfo_ismapped():
            return
        n = step * 3 if step is not None else (int(-event.delta) if IS_MAC else int(-event.delta / 120) * 3)
        if horizontal:
            self._xview("scroll", n, "units")
        else:
            self.body.yview_scroll(n, "units")

    @property
    def col_w(self):
        return 196 if self.compact else self.COL_W

    def _jump(self, i, n):
        self.update_idletasks()
        total = max(1, self.inner.winfo_reqwidth())
        self._xview("moveto", max(0.0, (i * (self.col_w + self.GAP) - 10) / total))

    def _render_summary(self, columns):
        for w in self.summary.winfo_children():
            w.destroy()
        total = sum(len(c[2]) for c in columns)
        self._chips = [tk.Label(self.summary, text=f"{total} tickets:", bg=C["bg"], fg=C["header"],
                                font=FONTS["h3"])]
        for i, (title, kind, cards, _note) in enumerate(columns):
            band = {"side": C["info"], "warn": C["warn"]}.get(kind, C["accent"])
            chip = tk.Frame(self.summary, bg=C["panel"], highlightbackground=C["line"], highlightthickness=1,
                            cursor="hand2")
            a = tk.Label(chip, text=truncate(title, 22), bg=C["panel"], fg=C["ink"], font=FONTS["small"])
            a.pack(side="left", padx=(8, 4), pady=3)
            b = tk.Label(chip, text=f" {len(cards)} ", bg=band if cards else C["line"], fg="white",
                         font=FONTS["small_bold"])
            b.pack(side="left", padx=(0, 6), pady=3)
            for w in (chip, a, b):
                w.bind("<Button-1>", lambda e, ii=i: self._jump(ii, len(columns)))
            self._chips.append(chip)
        self._flow()

    def _flow(self):
        """Lay the summary chips out left to right, wrapping onto new rows when they don't fit."""
        avail = self.summary.winfo_width()
        if avail < 100:
            avail = max(600, self.winfo_width())
        x = y = row_h = 0
        for w in self._chips:
            w.update_idletasks()
            cw, ch = w.winfo_reqwidth() + 6, w.winfo_reqheight() + 4
            if x and x + cw > avail:
                x, y = 0, y + row_h
                row_h = 0
            w.place(x=x, y=y)
            x += cw
            row_h = max(row_h, ch)
        self.summary.configure(height=y + row_h)

    def render(self, columns):
        """columns: list of (title, kind, [card dicts], note). kind: main / side / warn"""
        for w in self.head_inner.winfo_children() + self.inner.winfo_children():
            w.destroy()
        self.cards = {}
        self.order = []
        self._columns = columns
        self._render_summary(columns)
        cols_w = []
        for i, (status, kind, cards, note) in enumerate(columns):
            side = kind == "side"
            band = {"side": C["info"], "warn": C["warn"]}.get(kind, C["accent"])
            colbg = {"side": C["col_side"], "warn": C["col_warn"]}.get(kind, C["col"])
            h = tk.Frame(self.head_inner, bg=C["panel"], width=self.col_w, height=62,
                         highlightbackground=C["line"], highlightthickness=1)
            h.grid(row=0, column=i, padx=(0, self.GAP), pady=(0, 4), sticky="n")
            h.pack_propagate(False)
            tk.Frame(h, bg=band, height=4).pack(fill="x", side="top")
            row = tk.Frame(h, bg=C["panel"])
            row.pack(fill="x", padx=10, pady=(6, 0))
            tk.Label(row, text=truncate(status, 22), bg=C["panel"], fg=C["header"], font=FONTS["h3"]).pack(
                side="left")
            tk.Label(row, text=f" {len(cards)} ", bg=band if cards else C["line"], fg="white",
                     font=FONTS["small_bold"]).pack(side="left", padx=8)
            if cards:
                sa = tk.Label(row, text="☑", bg=C["panel"], fg=C["accent"], font=FONTS["small_bold"],
                              cursor="hand2")
                sa.pack(side="right")
                sa.bind("<Button-1>", lambda e, cs=cards: self.select_many([c["id"] for c in cs[:self.MAX_CARDS]]))
            tk.Label(h, text=note or (" side status" if side else " "), bg=C["panel"],
                     fg=C["bad"] if kind == "warn" else C["muted"], font=FONTS["small"], anchor="w").pack(
                fill="x", padx=10)
            col = tk.Frame(self.inner, bg=colbg, width=self.col_w)
            col.grid(row=0, column=i, padx=(0, self.GAP), sticky="ns")
            cols_w.append(col)
            if not cards:
                tk.Label(col, text="Free: no tickets" if kind != "side" and note else "No tickets", bg=colbg,
                         fg=C["faint"], font=FONTS["small"]).pack(pady=16, padx=6)
            for c in cards[:self.MAX_CARDS]:
                self._card(col, c).pack(fill="x", padx=6, pady=(6, 0))
                self.order.append((i, c["id"]))
            if len(cards) > self.MAX_CARDS:
                tk.Label(col, text=f"+ {len(cards) - self.MAX_CARDS} more: narrow the filters or use ☰ List",
                         bg=colbg, fg=C["muted"], font=FONTS["small"], wraplength=self.col_w - 20).pack(pady=8)
            tk.Frame(col, bg=colbg, height=8, width=self.col_w - 12).pack()
        self.update_idletasks()
        tallest = max([c.winfo_reqheight() for c in cols_w] + [80])
        for c in cols_w:  # every column exactly COL_W wide, so it lines up with its header
            c.pack_propagate(False)
            c.configure(width=self.col_w, height=tallest)
        self.update_idletasks()
        self._region()
        self.multi &= set(self.cards)
        for cid in self.multi | ({self.sel} if self.sel in self.cards else set()):
            self._paint(cid, True)

    def _card(self, parent, c):
        wrap = self.col_w - 44
        card = tk.Frame(parent, bg=C["field"], highlightbackground=C["line"], highlightthickness=1, cursor="hand2")
        tk.Frame(card, bg=c["prio_colour"], width=5).pack(side="left", fill="y")
        b = tk.Frame(card, bg=C["field"])
        b.pack(side="left", fill="both", expand=True, padx=(8, 8), pady=6)
        labels = []
        r1 = tk.Frame(b, bg=C["field"])
        r1.pack(fill="x")
        labels.append(tk.Label(r1, text=f"#{c['id']}", bg=C["field"], fg=C["header"], font=FONTS["h3"]))
        labels[-1].pack(side="left")
        labels.append(tk.Label(r1, text=f" {c['prio']} ", bg=c["prio_colour"], fg="white",
                               font=FONTS["small_bold"]))
        labels[-1].pack(side="right")
        if c.get("groomed"):
            labels.append(tk.Label(r1, text="✅", bg=C["field"], fg=C["good"], font=FONTS["small_bold"]))
            labels[-1].pack(side="right", padx=4)
        labels.append(tk.Label(b, text=truncate(c["subject"], 70 if self.compact else 120), bg=C["field"],
                               fg=C["ink"], font=FONTS["small"] if self.compact else FONTS["body"], anchor="w",
                               justify="left", wraplength=wrap))
        labels[-1].pack(fill="x", pady=(3, 2))

        def line(text, fg=None, bold=False):
            if text:
                lb = tk.Label(b, text=text, bg=C["field"], fg=fg or C["muted"], anchor="w", justify="left",
                              wraplength=wrap, font=FONTS["small_bold"] if bold else FONTS["small"])
                lb.pack(fill="x")
                labels.append(lb)

        if self.compact:
            line(c["meta"].split("    ")[0] + (f"   · ⏰ {c['due']}" if c["overdue"] else ""),
                 fg=C["bad"] if c["overdue"] else None)
            if c["flags"]:
                line(f"⚠ {plural(len(c['flags']), 'warning')}", fg=C["bad"], bold=True)
        else:
            if c["due"]:
                line(("⏰ Overdue · due " if c["overdue"] else "Due ") + c["due"],
                     fg=C["bad"] if c["overdue"] else None, bold=c["overdue"])
            line(c["meta"])
            line(f"PRD {c['prd']}" + (f"    📅 Groom {c['groom']}" if c["groom"] else "") +
                 (f" with {first_name(c['plan_dev'])}" if c.get("plan_dev") else ""),
                 fg=C["bad"] if c["prd_state"] in ("Needed",) else None)
            if c.get("groomed"):
                line(f"✅ Groomed {c['groomed']}", fg=C["good"], bold=True)
            line("    ".join(x for x in (f"🏃 {c['sprint']}" if c["sprint"] else "",
                                         f"🗓 {c['week']}" if c.get("week") else "", f"Track {c['track']}",
                                         f"🚀 {c['release']}" if c["release"] else "") if x))
            if c["parent"]:
                line(f"↳ {truncate(c['parent'], 60)}", fg=C["faint"])
            if c["flags"]:
                line("⚠ " + "; ".join(c["flags"]), fg=C["bad"], bold=True)
        cid = c["id"]
        for w in [card, b, r1] + labels:
            w.bind("<Button-1>", lambda e: self.click(cid))
            w.bind("<Control-Button-1>" if not IS_MAC else "<Command-Button-1>", lambda e: self.click(cid, "toggle"))
            w.bind("<Shift-Button-1>", lambda e: self.click(cid, "range"))
            w.bind("<Double-1>", lambda e: self.on_open(cid))
            w.bind("<Button-3>", lambda e: self.on_menu(e, cid))
            if IS_MAC:
                w.bind("<Button-2>", lambda e: self.on_menu(e, cid))
        card._parts = [card, b, r1] + [l for l in labels if l.cget("bg") == C["field"]]
        self.cards[cid] = card
        return card

    def _paint(self, cid, on):
        card = self.cards.get(cid)
        if not card or not card.winfo_exists():
            return
        bg = C["soft"] if on else C["field"]
        for w in card._parts:
            try:
                w.configure(bg=bg)
            except tk.TclError:
                pass
        card.configure(highlightbackground=C["accent"] if on else C["line"], highlightthickness=2 if on else 1)

    def click(self, cid, mode="single"):
        """Plain click = one ticket. Ctrl+click = add/remove. Shift+click = everything between."""
        if mode == "toggle":
            base = set(self.multi) or ({self.sel} if self.sel in self.cards else set())
            base ^= {cid}
            self.select_many(base, anchor=cid)
        elif mode == "range" and self.sel in self.cards:
            ids = [i for _c, i in self.order]
            a, b = sorted((ids.index(self.sel), ids.index(cid)))
            self.select_many(set(ids[a:b + 1]) | self.multi, anchor=self.sel)
        else:
            self.select_many(set())
            self.on_select(cid)

    def select_many(self, ids, anchor=None):
        ids = set(ids)
        for old in (self.multi | {self.sel}) - ids:
            self._paint(old, False)
        self.multi = ids if len(ids) > 1 else set()
        if anchor is not None:
            self.sel = anchor
        for cid in ids:
            self._paint(cid, True)
        if self.on_multi:
            self.on_multi(ids)

    def select(self, cid, scroll=False):
        if self.sel is not None and self.sel != cid and self.sel not in self.multi:
            self._paint(self.sel, False)
        self.sel = cid
        self._paint(cid, True)
        if scroll and cid in self.cards:
            self.update_idletasks()
            card = self.cards[cid]
            x = card.winfo_rootx() - self.inner.winfo_rootx()
            y = card.winfo_rooty() - self.inner.winfo_rooty()
            W, H = max(1, self.inner.winfo_reqwidth()), max(1, self.inner.winfo_reqheight())
            self._xview("moveto", max(0.0, (x - 20) / W))
            self.body.yview_moveto(max(0.0, (y - 20) / H))


class FilterBar(tk.Frame):
    """A tidy filter panel: a top line (scope / actions) and a grid of labelled filters."""

    def __init__(self, parent, cols=6, title="FILTERS"):
        super().__init__(parent, bg=C["soft"], highlightbackground=C["soft_line"], highlightthickness=1)
        self.top = tk.Frame(self, bg=C["soft"])
        self.top.pack(fill="x", padx=12, pady=(8, 2))
        if title:
            tk.Label(self.top, text=title, bg=C["soft"], fg=C["accent_dark"], font=FONTS["small_bold"]).pack(
                side="left", padx=(0, 12))
        self.grid_f = tk.Frame(self, bg=C["soft"])
        self.grid_f.pack(fill="x", padx=12, pady=(2, 10))
        self.cols, self.r, self.c = cols, 0, 0
        for i in range(cols):
            self.grid_f.columnconfigure(i, weight=1, uniform="fb")

    def label(self, parent, text):
        tk.Label(parent, text=text, bg=C["soft"], fg=C["muted"], font=FONTS["small"]).pack(side="left",
                                                                                         padx=(10, 4))

    def field(self, label, span=1):
        if self.c + span > self.cols:
            self.r, self.c = self.r + 1, 0
        cell = tk.Frame(self.grid_f, bg=C["soft"])
        cell.grid(row=self.r, column=self.c, columnspan=span, sticky="ew", padx=(0, 12), pady=3)
        tk.Label(cell, text=label.upper(), bg=C["soft"], fg=C["muted"], font=FONTS["small_bold"]).pack(anchor="w")
        inner = tk.Frame(cell, bg=C["soft"])
        inner.pack(fill="x")
        self.c += span
        return inner


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(f"{APP_NAME}")
        self.geometry("1400x900")
        self.minsize(1050, 680)
        self.config_data = load_json(CONFIG_PATH, {})
        self.theme = saved_theme(self.config_data)
        apply_palette(self.theme)
        self.configure(bg=C["bg"])
        self._init_fonts()
        self._init_style()

        self.drafts = load_json(DRAFTS_PATH, {})
        self.issues = []
        self.issue_index = {}
        self.issue_total = 0
        self.current_issue = None
        self.draft_key = "scratch"
        self.projects = {}
        self.dashboard_data = None
        self.dashboard_dirty = True
        self._queue = queue.Queue()
        self._busy = 0
        self._detail_token = 0

        self.var_url = tk.StringVar(value=self.config_data.get("base_url", ""))
        self.var_key = tk.StringVar(value=self.config_data.get("api_key", ""))
        self.var_ssl = tk.BooleanVar(value=self.config_data.get("verify_ssl", True))
        self.var_show_key = tk.BooleanVar(value=False)
        self.var_project = tk.StringVar(value="All projects")
        self.var_scope = tk.StringVar(value=self.config_data.get("scope", "mine"))
        self.var_search = tk.StringVar()
        self.var_status_filter = tk.StringVar(value=self.config_data.get("status_filter", "All statuses"))
        self.var_prd_filter = tk.StringVar(value="Any")
        self.var_sprint_filter = tk.StringVar(value=self.config_data.get("sprint_filter", "All sprints"))
        self.var_track_filter = tk.StringVar(value=self.config_data.get("track_filter", "All tracks"))
        self.var_release_filter = tk.StringVar(value="All releases")
        self.var_check_filter = tk.StringVar(value="Any")
        self.var_assignee_filter = tk.StringVar(value=self.config_data.get("assignee_filter", "All assignees"))
        self.var_due_from = tk.StringVar()
        self.var_due_to = tk.StringVar()
        self.var_groom_filter = tk.StringVar(value="Any")
        self.groom_filter_dates = {}
        self.week_labels = {}
        self.weekly_data = None
        self.var_priority_filter = tk.StringVar(value="All priorities")
        self.team_open = []
        self.suggest_history = None
        self.user_name = ""
        self.var_group = tk.BooleanVar(value=self.config_data.get("group_by_parent", True))
        self.visible_rows = []
        self.tags_by_id = {}
        self.tags_supported = None
        self.extra_issues = {}
        self.last_loaded = None
        self._auto_job = None
        self.hist_cache = {}
        self._holiday_cache = None
        self.delivery = None
        self.rn_versions = {}
        self.rn_issues = []
        self.rn_include = {}
        self.rn_closed_status = set()
        self.var_period = tk.StringVar(value=self.config_data.get("delivery_period", "Last 30 days"))
        self.var_stuck = tk.IntVar(value=int(self.config_data.get("stuck_after", 2)))
        self.var_rn_project = tk.StringVar()
        self.var_sp_version = tk.StringVar(value=self.config_data.get("sprint_version_selected", NO_RELEASE))
        self.var_sp_project = tk.StringVar(value=self.config_data.get("sprint_project", ""))
        self.var_retro_version = tk.StringVar()
        self.sp_version_info = {}
        self.var_rn_version = tk.StringVar()
        self.var_status = tk.StringVar(value="Ready.")

        self._build_ui()
        self._load_draft("scratch")
        self.protocol("WM_DELETE_WINDOW", self._on_close)
        self.after(100, self._poll_queue)
        mod = "Command" if IS_MAC else "Control"
        self.bind_all(f"<{mod}-s>", lambda e: self.export_prd())
        self.bind_all(f"<{mod}-f>", lambda e: self._focus_search())
        self.bind_all(f"<{mod}-g>", lambda e: self.set_groom_date())
        self.bind_all("<F5>", lambda e: self.refresh_current_tab())

        if self.var_url.get() and self.var_key.get():
            self.after(300, self.connect_and_load)
        else:
            self.set_status("Welcome! Enter your Redmine URL and API key above, then click 'Save & Connect'.")

    # ----- look & feel ---------------------------------------------------
    def _init_fonts(self):
        families = set(tkfont.families(self))
        default_family = tkfont.nametofont("TkDefaultFont").actual("family")
        if IS_WINDOWS:
            fam, mono = "Segoe UI", "Consolas"
        elif IS_MAC:
            fam, mono = "Helvetica Neue", "Menlo"
        else:
            fam, mono = "DejaVu Sans", "DejaVu Sans Mono"
        if fam not in families:
            fam = default_family
        if mono not in families:
            mono = tkfont.nametofont("TkFixedFont").actual("family")
        b = 12 if IS_MAC else 10
        serif = next((f for f in ("Georgia", "Cambria", "Palatino Linotype", "DejaVu Serif") if f in families), fam)
        FONTS.update({
            "body": (fam, b), "small": (fam, b - 1), "small_bold": (fam, b - 1, "bold"),
            "h1": (serif, b + 7, "bold"), "h2": (serif, b + 4, "bold"), "h3": (fam, b + 1, "bold"),
            "kpi": (serif, b + 14, "bold"), "kpi_label": (fam, b - 2, "bold"),
            "summary": (fam, b + 2), "mono": (mono, b + 1),
        })
        for named in ("TkDefaultFont", "TkTextFont", "TkMenuFont"):
            try:
                tkfont.nametofont(named).configure(family=fam, size=b)
            except tk.TclError:
                pass

    def _init_style(self):
        st = ttk.Style(self)
        self.modern = False
        if False and sv_ttk is not None:
            try:
                sv_ttk.set_theme("light")
                self.modern = True
            except Exception:
                self.modern = False
        if self.modern:
            C["bg"] = st.lookup(".", "background") or "#fafafa"
            self.configure(bg=C["bg"])
            for named in ("TkDefaultFont", "TkTextFont", "TkMenuFont", "TkHeadingFont"):
                try:
                    tkfont.nametofont(named).configure(family=FONTS["body"][0], size=FONTS["body"][1])
                except tk.TclError:
                    pass
            st.configure("Bg.TFrame", background=C["bg"])
            st.configure("Muted.TLabel", foreground=C["muted"], font=FONTS["small"])
            st.configure("Flag.TLabel", foreground=C["bad"], font=FONTS["small_bold"])
            st.configure("TLabelframe.Label", foreground=C["header"], font=FONTS["h3"])
            st.configure("Treeview", rowheight=32 if IS_MAC else 28, font=FONTS["body"])
            st.configure("Treeview.Heading", font=FONTS["small_bold"])
            st.configure("TNotebook.Tab", font=FONTS["h3"])
            return
        try:
            st.theme_use("clam")
        except tk.TclError:
            pass
        for opt, val in (("*Text.background", C["field"]), ("*Text.foreground", C["ink"]),
                         ("*Text.insertBackground", C["ink"]), ("*Text.selectBackground", C["select"]),
                         ("*Text.selectForeground", C["ink"]), ("*Listbox.background", C["field"]),
                         ("*Listbox.foreground", C["ink"]), ("*Listbox.selectBackground", C["select"]),
                         ("*Listbox.selectForeground", C["ink"]), ("*Toplevel.background", C["bg"]),
                         ("*Menu.background", C["panel"]), ("*Menu.foreground", C["ink"]),
                         ("*Menu.activeBackground", C["select"]), ("*Menu.activeForeground", C["ink"]),
                         ("*TCombobox*Listbox.background", C["field"]),
                         ("*TCombobox*Listbox.foreground", C["ink"]),
                         ("*TCombobox*Listbox.selectBackground", C["select"]),
                         ("*TCombobox*Listbox.selectForeground", C["ink"]),
                         ("*Canvas.background", C["bg"])):
            self.option_add(opt, val)
        st.configure(".", background=C["bg"], foreground=C["ink"], font=FONTS["body"], fieldbackground=C["field"],
                     insertcolor=C["ink"], selectbackground=C["select"], selectforeground=C["ink"])
        st.configure("Bg.TFrame", background=C["bg"])
        st.configure("TFrame", background=C["bg"])
        st.configure("TLabel", background=C["bg"], foreground=C["ink"])
        st.configure("Muted.TLabel", background=C["bg"], foreground=C["muted"], font=FONTS["small"])
        st.configure("Flag.TLabel", background=C["bg"], foreground=C["bad"], font=FONTS["small_bold"])
        st.configure("TLabelframe", background=C["bg"], bordercolor=C["line"], relief="flat")
        st.configure("TLabelframe.Label", background=C["bg"], foreground=C["header"], font=FONTS["h3"])
        st.configure("TButton", padding=(12, 6), relief="flat", borderwidth=1, background=C["field"], foreground=C["ink"],
                     bordercolor=C["line"], lightcolor=C["field"], darkcolor=C["field"])
        st.map("TButton", background=[("active", C["btn_hover"]), ("pressed", C["btn_press"])])
        st.configure("Accent.TButton", background=C["accent"], foreground="white", bordercolor=C["accent"],
                     lightcolor=C["accent"], darkcolor=C["accent"], padding=(14, 6), font=FONTS["small_bold"])
        st.map("Accent.TButton",
               background=[("disabled", C["accent_disabled"]), ("active", C["accent_dark"]), ("pressed", C["accent_dark"])],
               foreground=[("disabled", "white")])
        st.configure("TEntry", fieldbackground=C["field"], foreground=C["ink"], insertcolor=C["ink"],
                     bordercolor=C["line"], lightcolor=C["line"],
                     darkcolor=C["line"], padding=4)
        st.configure("TCombobox", fieldbackground=C["field"], background=C["field"], foreground=C["ink"],
                     arrowcolor=C["muted"], bordercolor=C["line"],
                     lightcolor=C["line"], darkcolor=C["line"], arrowsize=14, padding=3)
        st.map("TCombobox", fieldbackground=[("readonly", C["field"])], foreground=[("readonly", C["ink"])],
               selectbackground=[("readonly", C["field"])], selectforeground=[("readonly", C["ink"])])
        st.configure("Treeview.Heading", foreground=C["header"])
        st.map("Treeview.Heading", background=[("active", C["btn_hover"])])
        st.map("TCheckbutton", background=[("active", C["bg"])])
        st.map("TRadiobutton", background=[("active", C["bg"])])
        st.configure("TCheckbutton", foreground=C["ink"])
        st.configure("TRadiobutton", foreground=C["ink"], indicatorbackground=C["field"])
        st.configure("TMenubutton", background=C["field"], foreground=C["ink"], bordercolor=C["line"],
                     arrowcolor=C["muted"])
        st.map("TMenubutton", background=[("active", C["btn_hover"])])
        st.configure("TCheckbutton", background=C["bg"])
        st.configure("TRadiobutton", background=C["bg"])
        st.configure("Treeview", rowheight=32 if IS_MAC else 29, background=C["field"], foreground=C["ink"],
                     fieldbackground=C["field"], bordercolor=C["line"], lightcolor=C["line"], darkcolor=C["line"])
        st.configure("TLabelframe", background=C["bg"], bordercolor=C["line"], lightcolor=C["line"],
                     darkcolor=C["line"], relief="solid", borderwidth=1)
        st.configure("TSpinbox", fieldbackground=C["field"], foreground=C["ink"], bordercolor=C["line"],
                     lightcolor=C["line"],
                     darkcolor=C["line"], arrowcolor=C["muted"])
        st.configure("TCheckbutton", indicatorcolor=C["field"], indicatorbackground=C["field"])
        st.map("TCheckbutton", indicatorcolor=[("selected", C["accent"])])
        st.map("TRadiobutton", indicatorcolor=[("selected", C["accent"])])
        st.configure("TProgressbar", background=C["accent"], troughcolor=C["line"], bordercolor=C["line"],
                     lightcolor=C["accent"], darkcolor=C["accent"])
        st.configure("Treeview.Heading", font=FONTS["small_bold"], background=C["head_bg"], relief="flat",
                     bordercolor=C["line"], lightcolor=C["head_bg"], darkcolor=C["head_bg"], padding=(6, 6))
        st.map("Treeview", background=[("selected", C["select"])], foreground=[("selected", C["ink"])])
        st.configure("TNotebook", background=C["bg"], borderwidth=0, tabmargins=(0, 0, 0, 0))
        st.configure("TNotebook.Tab", padding=(16, 9), font=FONTS["h3"], background=C["bg"],
                     bordercolor=C["line"], lightcolor=C["bg"], foreground=C["ink"])
        st.map("TNotebook.Tab", background=[("selected", C["panel"])], foreground=[("!selected", C["muted"])])
        st.configure("Sub.TNotebook", background=C["bg"], borderwidth=0)
        st.configure("Sub.TNotebook.Tab", padding=(14, 6), font=FONTS["small_bold"], background=C["bg"],
                     foreground=C["muted"], bordercolor=C["line"], lightcolor=C["bg"])
        st.map("Sub.TNotebook.Tab", background=[("selected", C["soft"])],
               foreground=[("selected", C["accent_dark"]), ("!selected", C["muted"])])
        for nm in ("Soft.TRadiobutton", "Soft.TCheckbutton"):
            st.configure(nm, background=C["soft"], foreground=C["ink"])
            st.map(nm, background=[("active", C["soft"])])
        st.configure("TPanedwindow", background=C["bg"])
        st.configure("Vertical.TScrollbar", background=C["bar"], troughcolor=C["bg"], bordercolor=C["bg"],
                     arrowcolor=C["muted"], lightcolor=C["bar"], darkcolor=C["bar"])
        st.configure("Horizontal.TScrollbar", background=C["bar"], troughcolor=C["bg"], bordercolor=C["bg"],
                     arrowcolor=C["muted"], lightcolor=C["bar"], darkcolor=C["bar"])

    # ----- layout --------------------------------------------------------
    def _build_ui(self):
        self._build_header()
        self._build_statusbar()
        self._build_config()
        self.nb = ttk.Notebook(self)
        self.nb.pack(fill="both", expand=True, padx=12, pady=(4, 6))
        self.tab_today = ttk.Frame(self.nb, style="Bg.TFrame", padding=(10, 4, 10, 8))
        self.tab_team = ttk.Frame(self.nb, style="Bg.TFrame", padding=(8, 0, 8, 8))
        self.tab_work = ttk.Frame(self.nb, style="Bg.TFrame", padding=(0, 6, 0, 0))
        self.tab_delivery = ttk.Frame(self.nb, style="Bg.TFrame", padding=(8, 0, 8, 8))  # hidden for now
        # Grooming Workspace has two inner tabs: the tickets and the weekly view
        self.work_nb = ttk.Notebook(self.tab_work, style="Sub.TNotebook")
        self.work_nb.pack(fill="both", expand=True)
        self.tab_tickets = ttk.Frame(self.work_nb, style="Bg.TFrame", padding=(0, 8, 0, 0))
        self.tab_week = ttk.Frame(self.work_nb, style="Bg.TFrame", padding=(8, 0, 8, 8))
        self.work_nb.add(self.tab_tickets, text="  🎫 Tickets  ")
        self.work_nb.add(self.tab_week, text="  🗓 Weekly view  ")
        self.tab_load = ttk.Frame(self.work_nb, style="Bg.TFrame", padding=(8, 6, 8, 8))
        self.work_nb.add(self.tab_load, text="  👩‍💻 Developer load  ")
        self.tab_sprint = ttk.Frame(self.nb, style="Bg.TFrame", padding=(8, 0, 8, 8))
        self.tab_rn = ttk.Frame(self.nb, style="Bg.TFrame", padding=(8, 0, 8, 8))
        self.tab_dash = ttk.Frame(self.nb, style="Bg.TFrame", padding=(8, 0, 0, 0))
        self.tab_retro = ttk.Frame(self.nb, style="Bg.TFrame", padding=(8, 0, 8, 8))
        self.nb.add(self.tab_today, text="  🧭 Overview  ")
        self.nb.add(self.tab_work, text="  🎫 Grooming Workspace  ")
        self.nb.add(self.tab_sprint, text="  🏃 Sprint Planner  ")
        self.nb.add(self.tab_retro, text="  🔁 Retrospective  ")
        self.nb.add(self.tab_rn, text="  📝 Release Notes  ")
        self.nb.add(self.tab_team, text="  👥 Team  ")
        self.nb.add(self.tab_dash, text="  📊 Management Dashboard  ")
        self._build_workspace(self.tab_tickets)
        self._build_today(self.tab_today)
        self._build_team(self.tab_team)
        self._build_delivery(self.tab_delivery)
        self._build_weekly(self.tab_week)
        self._build_devload(self.tab_load)
        self._build_sprint(self.tab_sprint)
        self._build_release(self.tab_rn)
        self._build_retro(self.tab_retro)
        self._build_dashboard(self.tab_dash)
        self.nb.bind("<<NotebookTabChanged>>", self._on_tab_changed)
        self.work_nb.bind("<<NotebookTabChanged>>", self._on_work_tab_changed)
        self.nb.select(self.tab_today)

    def _on_work_tab_changed(self, _event=None):
        if self.nb.select() == str(self.tab_work) and self.work_nb.select() == str(self.tab_load):
            self.render_devload()
        if self.nb.select() == str(self.tab_work) and self.work_nb.select() == str(self.tab_week):
            self._weekly_stale = False
            self._refresh_week_choices()
            if self._configured(warn=False) and self._busy == 0:
                self.refresh_weekly()

    def _show_tickets(self):
        """Go to Grooming Workspace → Tickets."""
        self.nb.select(self.tab_work)
        self.work_nb.select(self.tab_tickets)

    def _build_header(self):
        bar = tk.Frame(self, bg=C["topbar"], highlightbackground=C["line"], highlightthickness=1)
        bar.pack(fill="x", side="top")
        tk.Label(bar, text="✿  " + APP_NAME, bg=C["topbar"], fg=C["accent"],
                 font=FONTS["h1"]).pack(side="left", padx=(18, 10), pady=12)
        tk.Label(bar, text="grooming  ·  PRDs  ·  sprints  ·  team", bg=C["topbar"],
                 fg=C["muted"], font=FONTS["body"]).pack(side="left", pady=12)
        self.btn_settings = tk.Label(bar, text="⚙  Settings", bg=C["soft"], fg=C["ink"], padx=14, pady=5,
                                     font=FONTS["small_bold"], cursor="hand2")
        self.btn_settings.pack(side="right", padx=(8, 18))
        self.btn_settings.bind("<Button-1>", lambda e: self.toggle_settings())
        self.btn_theme = tk.Label(bar, text="☀  Light" if self.theme == "dark" else "🌙  Dark", bg=C["soft"],
                                  fg=C["ink"], padx=14, pady=5, font=FONTS["small_bold"], cursor="hand2")
        self.btn_theme.pack(side="right", padx=(8, 0))
        self.btn_theme.bind("<Button-1>", lambda e: self.toggle_theme())
        self.btn_help = tk.Label(bar, text="❔  Help", bg=C["soft"], fg=C["ink"], padx=14, pady=5,
                                 font=FONTS["small_bold"], cursor="hand2")
        self.btn_help.pack(side="right", padx=(8, 0))
        self.btn_help.bind("<Button-1>", lambda e: self.show_help())
        self.lbl_user = tk.Label(bar, text="○  Not connected", bg=C["topbar"], fg=C["muted"],
                                 font=FONTS["small_bold"])
        self.lbl_user.pack(side="right", padx=8)

    def _build_statusbar(self):
        bar = tk.Frame(self, bg=C["bar"])
        bar.pack(fill="x", side="bottom")
        tk.Label(bar, textvariable=self.var_status, bg=C["bar"], fg=C["muted"], anchor="w",
                 font=FONTS["small"]).pack(side="left", fill="x", expand=True, padx=12, pady=4)
        self.progress = ttk.Progressbar(bar, mode="indeterminate", length=140)
        self.progress.pack(side="right", padx=12, pady=4)

    def _build_config(self):
        box = ttk.LabelFrame(self, text="Configuration", padding=(12, 8))
        box.pack(fill="x", padx=12, pady=(10, 4))
        self.cfg_box = box
        ttk.Label(box, text="Redmine Base URL").grid(row=0, column=0, sticky="w")
        ttk.Entry(box, textvariable=self.var_url, width=40).grid(row=0, column=1, sticky="ew", padx=(6, 16))
        ttk.Label(box, text="API Access Key").grid(row=0, column=2, sticky="w")
        self.e_key = ttk.Entry(box, textvariable=self.var_key, width=34, show="•")
        self.e_key.grid(row=0, column=3, sticky="ew", padx=(6, 6))
        ttk.Checkbutton(box, text="Show", variable=self.var_show_key,
                        command=self._toggle_key).grid(row=0, column=4, padx=(0, 10))
        ttk.Checkbutton(box, text="Verify SSL", variable=self.var_ssl).grid(row=0, column=5, padx=(0, 10))
        ttk.Button(box, text="Save & Connect", style="Accent.TButton",
                   command=self.connect_and_load).grid(row=0, column=6, padx=(0, 6))
        ttk.Button(box, text="Test Connection", command=self.test_connection).grid(row=0, column=7)
        ttk.Button(box, text="?", width=3, command=self._show_key_help).grid(row=0, column=8, padx=(6, 0))
        ttk.Label(box, style="Muted.TLabel",
                  text=f"Example: https://redmine.yourcompany.com    •    Settings are saved on this "
                       f"computer in {CONFIG_PATH}").grid(row=1, column=0, columnspan=9, sticky="w", pady=(6, 0))
        box.columnconfigure(1, weight=3)
        box.columnconfigure(3, weight=2)

    def _build_workspace(self, parent):
        topf = ttk.Frame(parent, style="Bg.TFrame", padding=(0, 0, 10, 0))
        topf.pack(fill="x")
        pw = self.ws_pw = ttk.PanedWindow(parent, orient="horizontal")
        pw.pack(fill="both", expand=True)
        side = ttk.Frame(pw, style="Bg.TFrame", padding=(0, 0, 10, 0))
        right = self.ws_right = ttk.PanedWindow(pw, orient="vertical")
        pw.add(side, weight=3)
        pw.add(right, weight=2)

        # Ticket list / board header and filters (full width)
        head = ttk.Frame(topf, style="Bg.TFrame")
        head.pack(fill="x")
        ttk.Label(head, text="Tickets", font=FONTS["h2"], foreground=C["header"]).pack(side="left")
        self.var_ws_view = tk.StringVar(value=self.config_data.get("ws_view", "board"))
        self.btn_board = ttk.Button(head, text="▦ Board", command=lambda: self._set_ws_view("board"))
        self.btn_board.pack(side="left", padx=(14, 0))
        self.btn_people = ttk.Button(head, text="👥 By developer", command=lambda: self._set_ws_view("people"))
        self.btn_people.pack(side="left", padx=(4, 0))
        self.btn_groomv = ttk.Button(head, text="📅 By groom date", command=lambda: self._set_ws_view("groom"))
        self.btn_groomv.pack(side="left", padx=(4, 0))
        self.btn_list = ttk.Button(head, text="☰ List", command=lambda: self._set_ws_view("list"))
        self.btn_list.pack(side="left", padx=(4, 0))
        self.var_compact = tk.BooleanVar(value=bool(self.config_data.get("board_compact", False)))
        self.var_hide_empty = tk.BooleanVar(value=bool(self.config_data.get("board_hide_empty", False)))
        self.chk_hide_empty = ttk.Checkbutton(head, text="Hide empty columns", variable=self.var_hide_empty,
                                              command=self._board_option_changed)
        self.chk_hide_empty.pack(side="left", padx=(12, 0))
        self.chk_compact = ttk.Checkbutton(head, text="Compact cards", variable=self.var_compact,
                                           command=self._board_option_changed)
        self.chk_compact.pack(side="left", padx=(8, 0))
        ttk.Button(head, text="🔎 Why isn't a ticket here?", command=self.find_ticket).pack(side="left", padx=(12, 0))
        self.btn_details = ttk.Button(head, text="◧ Hide details", command=self._toggle_details)
        self.btn_details.pack(side="left", padx=(12, 0))
        ttk.Button(head, text="🏷 Tags…", command=self.edit_tags).pack(side="right")
        self.mb_views = ttk.Menubutton(head, text="⭐ Views ▾")
        self.menu_views = tk.Menu(self.mb_views, tearoff=0, postcommand=self._build_views_menu)
        self.mb_views["menu"] = self.menu_views
        self.mb_views.pack(side="right", padx=6)
        ttk.Button(head, text="▶ Grooming session", style="Accent.TButton",
                   command=self.start_grooming_session).pack(side="right", padx=(0, 6))
        ttk.Button(head, text="↻ Refresh", command=self.refresh_tickets).pack(side="right", padx=6)
        self.lbl_updated = ttk.Label(head, text="", style="Muted.TLabel")
        self.lbl_updated.pack(side="right", padx=6)

        fb = FilterBar(topf, cols=6)
        fb.pack(fill="x", pady=(8, 8))
        fb.label(fb.top, "Projects")
        self.mp_project = MultiPick(fb.top, "All projects", self._on_scope_changed, width=26, title="Projects")
        self.mp_project.pack(side="left")
        ttk.Radiobutton(fb.top, text="Assigned to me", value="mine", variable=self.var_scope, style="Soft.TRadiobutton",
                        command=self._on_scope_changed).pack(side="left", padx=(14, 0))
        ttk.Radiobutton(fb.top, text="All open", value="all", variable=self.var_scope, style="Soft.TRadiobutton",
                        command=self._on_scope_changed).pack(side="left", padx=(10, 0))
        ttk.Button(fb.top, text="✖ Clear all filters", command=self._clear_ws_filters).pack(side="right")
        ttk.Checkbutton(fb.top, text="Group under parent", variable=self.var_group, style="Soft.TCheckbutton",
                        command=self._on_list_filter_changed).pack(side="right", padx=12)
        fb.label(fb.top, "    🔎 Search")
        self.e_search = ttk.Entry(fb.top, textvariable=self.var_search)
        self.e_search.pack(side="left", fill="x", expand=True)
        self.var_search.trace_add("write", lambda *a: self._populate_tree())

        def combo(label, var, values):
            cb = ttk.Combobox(fb.field(label), textvariable=var, state="readonly", values=values)
            cb.pack(fill="x")
            cb.bind("<<ComboboxSelected>>", lambda e: self._on_list_filter_changed())
            return cb

        self.cb_status = combo("Status", self.var_status_filter, ["All statuses"])
        self.cb_priority = combo("Priority", self.var_priority_filter, ["All priorities"])
        self.cb_assignee = combo("Assignee", self.var_assignee_filter, ["All assignees"])
        self.mp_release = MultiPick(fb.field("Release"), "All releases", self._on_list_filter_changed, width=16,
                                    title="Releases")
        self.mp_release.pack(fill="x")
        self.mp_release.set_selected(self.config_data.get("release_selected") or [])
        self.cb_sprint = combo("Sprint", self.var_sprint_filter, ["All sprints"])
        self.cb_track = combo("Track", self.var_track_filter, ["All tracks", TRACK_TAG, "Legacy"])
        self.cb_prd = combo("PRD", self.var_prd_filter, ["Any", "Needs a PRD (Needed + Partial)"] + PRD_STATES)
        self.cb_check = combo("Status check", self.var_check_filter, ["Any", "⚠ Flagged only"])
        self.cb_groom = combo("Groom date", self.var_groom_filter, ["Any", "Not set", "Set"])
        due = fb.field("Due between", span=3)
        ttk.Entry(due, textvariable=self.var_due_from, width=11).pack(side="left")
        tk.Label(due, text="and", bg=C["soft"], fg=C["muted"], font=FONTS["small"]).pack(side="left", padx=4)
        ttk.Entry(due, textvariable=self.var_due_to, width=11).pack(side="left")
        for txt, mode in (("This week", "week"), ("This sprint", "sprint"), ("Clear", "clear")):
            ttk.Button(due, text=txt, command=lambda m=mode: self.quick_due_range(m)).pack(side="left", padx=(6, 0))
        for v in (self.var_due_from, self.var_due_to):
            v.trace_add("write", lambda *a: self._due_filter_changed())

        tf = self.ws_list = ttk.Frame(side)
        tf.pack(fill="both", expand=True)
        cols = ("status", "instatus", "assignee", "groom", "prd", "check", "sprint", "week", "track", "priority",
                "due", "release")
        self.tree = ttk.Treeview(tf, columns=cols, show="tree headings", selectmode="extended")
        self.tree.heading("#0", text="Ticket", anchor="w")
        self.tree.column("#0", width=330, minwidth=200, stretch=True, anchor="w")
        for col, text, width in (("status", "Status", 124), ("instatus", "In status", 76),
                                 ("assignee", "Assignee", 130), ("groom", "Groom date", 96), ("prd", "PRD", 112),
                                 ("check", "Status check", 190),
                                 ("sprint", "Sprint", 110), ("week", "Sprint week", 150), ("track", "Track", 62), ("priority", "Priority", 84),
                                 ("due", "Due", 88), ("release", "Release", 150)):
            self.tree.heading(col, text=text, anchor="w")
            self.tree.column(col, width=width, minwidth=50, stretch=False, anchor="w")
        vsb = ttk.Scrollbar(tf, orient="vertical", command=self.tree.yview)
        hsb = ttk.Scrollbar(tf, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        tf.rowconfigure(0, weight=1)
        tf.columnconfigure(0, weight=1)
        for key, colour in PRIORITY_COLORS.items():
            self.tree.tag_configure(f"p_{key}", foreground=colour)
        self.tree.tag_configure("ghost", foreground=C["faint"])
        self.tree.tag_configure("odd", background=C["odd"])
        self.tree.bind("<<TreeviewSelect>>", self._on_tree_select)
        self.tree.bind("<Button-3>", self._tree_menu)
        self.tree.bind("<Button-2>" if IS_MAC else "<Control-Button-1>", self._tree_menu)
        self.board = KanbanBoard(side, self._board_select, self._board_open, self._board_menu)
        self.board.on_multi = self._board_multi
        self.selbar = tk.Frame(side, bg=C["soft"], highlightbackground=C["soft_line"], highlightthickness=1)
        self.lbl_selbar = tk.Label(self.selbar, text="", bg=C["soft"], fg=C["header"], font=FONTS["h3"])
        self.lbl_selbar.pack(side="left", padx=12, pady=6)
        for txt, cmd, acc in (("✅ Mark groomed", lambda: self.mark_groomed(True), True),
                              ("↺ Not groomed", lambda: self.mark_groomed(False), False),
                              ("📅 Set groom date…", self.set_groom_date, False),
                              ("Open in Redmine ↗", lambda: self._open_ids_in_redmine(self._selected_ticket_ids()),
                               False),
                              ("✖ Clear selection", lambda: self.board.select_many(set()), False)):
            ttk.Button(self.selbar, text=txt, command=cmd, style="Accent.TButton" if acc else "TButton").pack(
                side="left", padx=(0, 6), pady=4)
        self._board_job = None
        bottom = self.ws_bottom = ttk.Frame(side, style="Bg.TFrame")
        bottom.pack(fill="x", pady=(6, 0))
        ttk.Button(bottom, text="⬇ Grooming Agenda (Excel)…", style="Accent.TButton",
                   command=self.export_agenda).pack(side="right")
        self.lbl_count = ttk.Label(bottom, text="No tickets loaded yet", style="Muted.TLabel",
                                   justify="left")
        self.lbl_count.pack(side="left", anchor="w")

        # Details panel
        det = ttk.LabelFrame(right, text="Ticket Details", padding=10)
        right.add(det, weight=2)
        top = ttk.Frame(det)
        top.pack(fill="x")
        self.lbl_prio = tk.Label(top, text="  NO TICKET  ", bg=C["muted"], fg="white",
                                 font=FONTS["small_bold"], padx=6, pady=2)
        self.lbl_prio.pack(side="left")
        self.lbl_id = ttk.Label(top, text="", font=FONTS["small_bold"], foreground=C["muted"])
        self.lbl_id.pack(side="left", padx=10)
        self.btn_open = ttk.Button(top, text="Open in Redmine ↗", command=self._open_in_browser,
                                   state="disabled")
        self.btn_open.pack(side="right")
        act = ttk.Frame(det)
        act.pack(fill="x", pady=(6, 0))
        mb = ttk.Menubutton(act, text="PRD status ▾", width=13)
        menu = tk.Menu(mb, tearoff=0)
        menu.add_command(label="✔  Has PRD (done elsewhere)", command=lambda: self.set_prd_mark("Has PRD"))
        menu.add_command(label="—  PRD not needed", command=lambda: self.set_prd_mark("Not needed"))
        menu.add_command(label="↺  Let the app decide", command=lambda: self.set_prd_mark(None))
        mb["menu"] = menu
        mb.pack(side="left")
        ttk.Button(act, text="📅 Groom date…", command=self.set_groom_date).pack(side="left", padx=6)
        ttk.Button(act, text="✅ Groomed", command=lambda: self.mark_groomed(True)).pack(side="left", padx=(0, 6))
        ttk.Button(act, text="📝 My note…", command=self.edit_ticket_note).pack(side="left")
        ttk.Button(act, text="💡 Suggest developer", command=self.suggest_for_selected).pack(side="left", padx=6)
        ttk.Button(act, text="✉ Message", command=self.copy_ticket_message).pack(side="left")
        self.lbl_subject = ttk.Label(det, text="Select a ticket from the list on the left",
                                     font=FONTS["h2"], foreground=C["header"], justify="left")
        self.lbl_subject.pack(anchor="w", fill="x", pady=(8, 2))
        det.bind("<Configure>", lambda e: self.lbl_subject.configure(wraplength=max(200, e.width - 30)))
        self.lbl_meta = ttk.Label(det, text="", style="Muted.TLabel", justify="left")
        self.lbl_meta.pack(anchor="w", fill="x", pady=(0, 2))
        self.lbl_flags = ttk.Label(det, text="", style="Flag.TLabel", justify="left")
        self.lbl_flags.pack(anchor="w", fill="x", pady=(0, 6))
        det.bind("<Configure>", lambda e: [w.configure(wraplength=max(200, e.width - 30))
                                           for w in (self.lbl_subject, self.lbl_meta, self.lbl_flags)], add="+")
        df = ttk.Frame(det)
        df.pack(fill="both", expand=True)
        self.txt_desc = tk.Text(df, wrap="word", height=8, font=FONTS["body"], relief="flat",
                                bg=C["field"], fg=C["ink"], highlightthickness=1,
                                highlightbackground=C["line"], padx=10, pady=8, state="disabled")
        dsb = ttk.Scrollbar(df, orient="vertical", command=self.txt_desc.yview)
        self.txt_desc.configure(yscrollcommand=dsb.set)
        self.txt_desc.pack(side="left", fill="both", expand=True)
        dsb.pack(side="right", fill="y")

        # PRD workspace
        prd = ttk.LabelFrame(right, text="PRD Grooming Workspace", padding=10)
        right.add(prd, weight=3)
        tb = ttk.Frame(prd)
        tb.pack(fill="x", pady=(0, 6))
        ttk.Button(tb, text="✨ Generate PRD", style="Accent.TButton",
                   command=self.generate_prd).pack(side="left")
        ttk.Button(tb, text="⬇ Export (Word)…", command=self.export_prd).pack(side="left", padx=(6, 0))
        ttk.Button(tb, text="Copy", command=self.copy_prd).pack(side="left", padx=(6, 0))
        ttk.Button(tb, text="📋 Prompt for Claude", command=self.copy_prd_prompt).pack(side="left", padx=(6, 0))
        ttk.Button(tb, text="Clear", command=self.clear_prd).pack(side="left", padx=(6, 0))
        self.lbl_wc = ttk.Label(tb, text="0 words", style="Muted.TLabel")
        self.lbl_wc.pack(side="right")
        self.lbl_draft = ttk.Label(tb, text="", style="Muted.TLabel")
        self.lbl_draft.pack(side="right", padx=14)
        pf = ttk.Frame(prd)
        pf.pack(fill="both", expand=True)
        self.txt_prd = tk.Text(pf, wrap="word", undo=True, font=FONTS["mono"], relief="flat",
                               bg=C["field"], fg=C["ink"], insertbackground=C["ink"], highlightthickness=1,
                               highlightbackground=C["line"], highlightcolor=C["accent"], padx=12, pady=10)
        psb = ttk.Scrollbar(pf, orient="vertical", command=self.txt_prd.yview)
        self.txt_prd.configure(yscrollcommand=psb.set)
        self.txt_prd.pack(side="left", fill="both", expand=True)
        psb.pack(side="right", fill="y")
        self.txt_prd.bind("<KeyRelease>", lambda e: self._update_wc())
        self.after(50, lambda: self._set_ws_view(self.var_ws_view.get(), save=False))

    def _build_dashboard(self, parent):
        tb = ttk.Frame(parent, style="Bg.TFrame", padding=(0, 10, 12, 8))
        tb.pack(fill="x")
        left = ttk.Frame(tb, style="Bg.TFrame")
        left.pack(side="left")
        ttk.Label(left, text="Backlog Health", font=FONTS["h1"], foreground=C["header"]).pack(anchor="w")
        self.lbl_dash_scope = ttk.Label(left, text="Pick your projects and filters below, then click "
                                                   "Refresh.", style="Muted.TLabel")
        self.lbl_dash_scope.pack(anchor="w")
        ttk.Button(tb, text="Copy Executive Summary", command=self.copy_summary).pack(side="right")
        ttk.Button(tb, text="Export Status Report…", command=self.export_report).pack(side="right", padx=6)
        ttk.Button(tb, text="↻ Refresh Dashboard", style="Accent.TButton",
                   command=self.refresh_dashboard).pack(side="right")
        self._build_dash_filters(parent)

        sf = self.dash_sf = ScrollableFrame(parent)
        sf.pack(fill="both", expand=True)
        body = tk.Frame(sf.inner, bg=C["bg"])
        body.pack(fill="both", expand=True, padx=(0, 10))

        banner = tk.Frame(body, bg=C["soft"], highlightbackground=C["soft_line"], highlightthickness=1)
        banner.pack(fill="x", pady=(0, 10), padx=4)
        tk.Label(banner, text="EXECUTIVE SUMMARY", bg=C["soft"], fg=C["accent_dark"],
                 font=FONTS["kpi_label"]).pack(anchor="w", padx=16, pady=(10, 0))
        self.lbl_summary = tk.Label(banner, text="Click 'Refresh Dashboard' to analyse your backlog.",
                                    bg=C["soft"], fg=C["ink"], font=FONTS["summary"], justify="left",
                                    anchor="w")
        self.lbl_summary.pack(anchor="w", fill="x", padx=16, pady=(4, 12))
        banner.bind("<Configure>", lambda e: self.lbl_summary.configure(wraplength=max(300, e.width - 40)))

        grid = tk.Frame(body, bg=C["bg"])
        grid.pack(fill="x")
        specs = [("open", "Open items", C["accent"]), ("high", "High priority", C["bad"]),
                 ("overdue", "Overdue", C["warn"]), ("ready", "Sprint-ready", C["good"]),
                 ("closed30", "Closed · last 30 days", C["good"]),
                 ("net", "Net backlog change · 30 days", C["info"]),
                 ("throughput", "Avg weekly throughput", C["accent"]),
                 ("age", "Avg age of open items", C["muted"])]
        self.kpi = {}
        for i, (key, title, colour) in enumerate(specs):
            card = KpiCard(grid, title, colour)
            card.grid(row=i // 4, column=i % 4, sticky="nsew", padx=4, pady=4)
            self.kpi[key] = card
        for c in range(4):
            grid.columnconfigure(c, weight=1, uniform="kpi")

        charts = tk.Frame(body, bg=C["bg"])
        charts.pack(fill="x", pady=(6, 0))
        self.ch_priority = DonutChart(charts, "Open items by priority")
        self.ch_status = HBarChart(charts, "Open items by status")
        self.ch_age = HBarChart(charts, "Age of open items")
        self.ch_assignee = HBarChart(charts, "Workload by assignee")
        for i, ch in enumerate((self.ch_priority, self.ch_status, self.ch_age, self.ch_assignee)):
            ch.grid(row=i // 2, column=i % 2, sticky="nsew", padx=4, pady=4)
        charts.columnconfigure(0, weight=1, uniform="ch")
        charts.columnconfigure(1, weight=1, uniform="ch")
        self.ch_priority.on_pick = self.show_priority_tickets
        self.ch_status.on_pick = lambda lbl: self.show_pick(
            f"Status: {lbl}", lambda it: name_of(it.get("status")) == lbl)
        self.ch_assignee.on_pick = lambda lbl: self.show_pick(
            f"Assignee: {lbl}", lambda it: name_of(it.get("assigned_to"), "Unassigned") == lbl)
        self.ch_age.on_pick = self._pick_age
        self.pick_box = ttk.LabelFrame(body, text="🔎 Tickets", padding=8)
        pf, self.tv_pick = self._make_table(self.pick_box, [("ticket", "Ticket", 80), ("subject", "Subject", 420),
                                                             ("assignee", "Assignee", 150), ("status", "Status", 140),
                                                             ("due", "Due", 110), ("prd", "PRD", 110),
                                                             ("groom", "Groom date", 110)], height=8)
        self.tv_pick.configure(selectmode="extended")
        self.tv_pick.bind("<Double-1>", lambda e: self._pick_open_redmine())
        pf.pack(fill="x")
        pb = ttk.Frame(self.pick_box)
        pb.pack(fill="x", pady=(6, 0))
        ttk.Button(pb, text="📅 Set groom date for selected", style="Accent.TButton",
                   command=lambda: self.set_groom_date(ids=[k[1:] for k in self.tv_pick.selection()])).pack(
            side="left")
        ttk.Button(pb, text="Open in Redmine ↗", command=self._pick_open_redmine).pack(side="left", padx=6)
        ttk.Button(pb, text="Show this ticket in Grooming Workspace",
                   command=self._open_pick_in_grooming).pack(side="left")
        ttk.Label(pb, text="Tip: double-click a ticket to open it in Redmine.", style="Muted.TLabel").pack(
            side="left", padx=10)
        ttk.Button(pb, text="Close", command=self.pick_box.pack_forget).pack(side="right")
        self.ch_trend = TrendChart(body, "Created vs. closed: last 8 weeks (7-day windows)", height=290)
        self.ch_trend.pack(fill="x", padx=4, pady=4)

        rk = ttk.LabelFrame(body, text="Items needing attention  (double-click to open in Redmine · right-click "
                                       "for more)", padding=8)
        rk.pack(fill="x", padx=4, pady=(8, 6))
        self.risk_tree = ttk.Treeview(rk, columns=("id", "priority", "subject", "assignee", "why"),
                                      show="headings", height=8)
        for col, text, width, stretch in (("id", "#", 70, False), ("priority", "Priority", 90, False),
                                          ("subject", "Subject", 380, True), ("assignee", "Assignee", 150, False),
                                          ("why", "Why it's flagged", 280, True)):
            self.risk_tree.heading(col, text=text, anchor="w")
            self.risk_tree.column(col, width=width, stretch=stretch, anchor="w")
        rsb = ttk.Scrollbar(rk, orient="vertical", command=self.risk_tree.yview)
        self.risk_tree.configure(yscrollcommand=rsb.set)
        self.risk_tree.pack(side="left", fill="x", expand=True)
        rsb.pack(side="right", fill="y")
        self.risk_tree.bind("<Double-1>", lambda e: self._risk_open_redmine())
        self.risk_tree.bind("<Button-3>", self._risk_menu)
        self.risk_tree.bind("<Button-2>" if IS_MAC else "<Control-Button-1>", self._risk_menu)
        self.lbl_asof = ttk.Label(body, text="", style="Muted.TLabel")
        self.lbl_asof.pack(anchor="w", padx=6, pady=(0, 14))

    # ----- background work ----------------------------------------------
    def run_async(self, work, on_success, busy: str = None, on_error=None):
        self._busy += 1
        self.progress.start(12)
        if busy:
            self.set_status(busy)

        def runner():
            try:
                result = work()
                self._queue.put((on_success, result))
            except Exception as exc:  # Passed back to the UI thread
                self._queue.put((on_error or self._show_error, exc))

        threading.Thread(target=runner, daemon=True).start()

    def _poll_queue(self):
        try:
            while True:
                callback, payload = self._queue.get_nowait()
                self._busy = max(0, self._busy - 1)
                if self._busy == 0:
                    self.progress.stop()
                try:
                    callback(payload)
                except Exception as exc:
                    self._show_error(exc)
        except queue.Empty:
            pass
        self.after(100, self._poll_queue)

    def _show_error(self, exc):
        msg = str(exc) if isinstance(exc, RedmineError) else f"Unexpected error: {type(exc).__name__}: {exc}"
        self.set_status("⚠ " + msg.splitlines()[0])
        messagebox.showerror(APP_NAME, msg, parent=self)

    def set_status(self, text: str):
        self.var_status.set(text)

    # ----- configuration -------------------------------------------------
    def _toggle_key(self):
        self.e_key.configure(show="" if self.var_show_key.get() else "•")

    def _show_key_help(self):
        messagebox.showinfo(
            "Where do I find these?",
            "BASE URL\nThe address of your Redmine home page, for example\n"
            "https://redmine.yourcompany.com (no /issues at the end).\n\n"
            "API ACCESS KEY\n1. Log in to Redmine in your browser.\n"
            "2. Click 'My account' (top-right corner).\n"
            "3. In the right-hand column under 'API access key', click 'Show'.\n"
            "4. Copy the long code and paste it here.\n\n"
            "If there is no 'API access key' box, ask your Redmine administrator to enable "
            "'Enable REST web service' under Administration → Settings → API.\n\n"
            "VERIFY SSL\nLeave this ticked. Untick it only if you get a certificate error "
            "with an internal company server.",
            parent=self)

    def toggle_settings(self, show=None):
        visible = bool(self.cfg_box.winfo_ismapped())
        show = (not visible) if show is None else show
        if show and not visible:
            self.cfg_box.pack(fill="x", padx=12, pady=(10, 4), before=self.nb)
        elif not show and visible:
            self.cfg_box.pack_forget()

    def _configured(self, warn: bool = True) -> bool:
        if self.var_url.get().strip() and self.var_key.get().strip():
            return True
        if warn:
            messagebox.showinfo(APP_NAME, "Please enter your Redmine Base URL and API Access Key in the "
                                          "Configuration section first.", parent=self)
        return False

    def save_config(self) -> bool:
        if not self._configured():
            return False
        url = self.var_url.get().strip().rstrip("/")
        if not url.startswith(("http://", "https://")):
            url = "https://" + url
        self.var_url.set(url)
        self.config_data.update(base_url=url, api_key=self.var_key.get().strip(),
                                verify_ssl=bool(self.var_ssl.get()), scope=self.var_scope.get())
        if self.projects:
            self.config_data["projects_selected"] = list(self.mp_project.selected)
        try:
            save_json(CONFIG_PATH, self.config_data)
        except OSError as exc:
            messagebox.showerror(APP_NAME, f"Could not save settings:\n{exc}", parent=self)
            return False
        return True

    def make_client(self) -> RedmineClient:
        return RedmineClient(self.var_url.get(), self.var_key.get(), bool(self.var_ssl.get()))

    def test_connection(self):
        if not self._configured():
            return
        client = self.make_client()

        def done(user):
            name = f"{user.get('firstname', '')} {user.get('lastname', '')}".strip() or user.get("login", "")
            self.set_status(f"Connection OK. Logged in as {name}.")
            messagebox.showinfo(APP_NAME, f"Connection successful!\n\nLogged in as: {name}", parent=self)

        self.run_async(client.current_user, done, busy="Testing connection…")

    def connect_and_load(self):
        if not self.save_config():
            return
        client = self.make_client()

        def work():
            user = client.current_user()
            projects, _ = client.projects()
            return user, projects

        def done(result):
            user, projects = result
            name = f"{user.get('firstname', '')} {user.get('lastname', '')}".strip() or user.get("login", "")
            self.lbl_user.configure(text=f"●  Connected as {name}", fg="#6E9C7A")
            self.user_name = name
            self.toggle_settings(False)
            self._set_projects(projects)
            self.after(500, self.load_sprint_versions)
            self.set_status(f"Connected to {self.var_url.get()} as {name}.")
            self.dashboard_dirty = True
            self.refresh_tickets()
            if self.nb.select() == str(self.tab_dash):
                self.refresh_dashboard()

        def fail(exc):
            self.lbl_user.configure(text="○  Not connected", fg="#C46A5F")
            self._show_error(exc)

        self.run_async(work, done, busy="Connecting to Redmine…", on_error=fail)

    def _set_projects(self, projects):
        by_id = {p["id"]: p for p in projects}
        names = Counter(p.get("name", "") for p in projects)
        self.projects = {}
        for p in projects:
            label = p.get("name", f"Project {p['id']}")
            parent = p.get("parent")
            if names[label] > 1 and parent and parent.get("id") in by_id:
                label = f"{label} ({parent.get('name')})"
            if label in self.projects:
                label = f"{label} #{p['id']}"
            self.projects[label] = p["id"]
        ordered = sorted(self.projects, key=str.lower)
        self.mp_project.set_options(ordered)
        self.mp_dash_project.set_options(ordered)
        dsaved = self.config_data.get("dash_projects")
        if dsaved is None:
            dsaved = self.config_data.get("projects_selected") or []
        self.mp_dash_project.set_selected([x for x in dsaved if x in self.projects])
        saved = self.config_data.get("projects_selected")
        if saved is None:  # settings from an earlier version
            old = self.config_data.get("project") or ""
            old = old.split(" » ")[-1]
            saved = [old] if old in self.projects else []
        self.mp_project.set_selected([x for x in saved if x in self.projects])
        self.var_project.set(self.mp_project.summary())
        self.cb_rn_project.configure(values=ordered)
        self.cb_sp_project.configure(values=ordered)
        first = (self.mp_project.selected or [""])[0]
        if self.var_rn_project.get() not in self.projects:
            self.var_rn_project.set(first if first in self.projects else (ordered or [""])[0])

    def _scope(self):
        pids = [self.projects[n] for n in self.mp_project.selected if n in self.projects]
        mine = self.var_scope.get() == "mine"
        self.var_project.set(self.mp_project.summary())
        label = f"{self.mp_project.summary()} · {'Assigned to me' if mine else 'All open items'}"
        return pids, mine, label

    def _on_scope_changed(self):
        self.var_project.set(self.mp_project.summary())
        self.config_data.update(scope=self.var_scope.get(), projects_selected=list(self.mp_project.selected))
        try:
            save_json(CONFIG_PATH, self.config_data)
        except OSError:
            pass
        if not self._configured(warn=False):
            return
        self.refresh_tickets()
        self.dashboard_dirty = True
        if self.nb.select() == str(self.tab_dash):
            self.refresh_dashboard()

    # ----- tickets -------------------------------------------------------
    def refresh_tickets(self, silent: bool = False):
        if not self._configured(warn=not silent):
            return
        client = self.make_client()
        pid, mine, label = self._scope()
        params = {**scope_params(None, mine), "status_id": "open", "sort": "priority:desc,updated_on:desc",
                  "include": "attachments"}
        tags = dedupe(list(self.config_data.get("tracked_tags") or DEFAULT_TAGS) +
                      [sp["name"] for _, sp in self.all_sprints()])

        def work():
            issues, total = client.issues_in(params, pid)
            native = any(native_tags(it) is not None for it in issues)
            if native:
                tag_by_id, supported = {it["id"]: native_tags(it) or [] for it in issues}, True
            else:
                tag_by_id, supported = client.tag_map(pid, mine, tags)
            known = {it["id"] for it in issues}
            extra = {}
            wanted = {it["parent"]["id"] for it in issues if it.get("parent")} - known
            for _ in range(4):  # parents of parents, a few levels up
                if not wanted:
                    break
                fetched = client.issues_by_ids(wanted)
                for it in fetched:
                    extra[it["id"]] = it
                wanted = ({it["parent"]["id"] for it in fetched if it.get("parent")}
                          - known - set(extra))
            return issues, total, tag_by_id, supported, extra

        def done(result):
            self.issues, self.issue_total, self.tags_by_id, self.tags_supported, self.extra_issues = result
            self.issue_index = {it["id"]: it for it in self.issues}
            self.last_loaded = dt.datetime.now()
            self._refresh_filter_values()
            self._populate_tree()
            self.lbl_updated.configure(text=f"Updated {self.last_loaded.strftime('%H:%M')}")
            note = "" if self.tags_supported else "  (Tags not available from your Redmine, see 🏷 Tags…)"
            self.set_status(f"Loaded {len(self.issues)} open ticket(s) for {label}.{note}")
            self._schedule_auto_refresh()
            self._load_histories(self.issues, then=self._populate_tree)
            if self.nb.select() == str(self.tab_work) and self.work_nb.select() == str(self.tab_load):
                self.render_devload()


        def fail(exc):
            if silent:
                self.set_status(f"⚠ Auto-refresh failed: {str(exc).splitlines()[0]}")
                self._schedule_auto_refresh()
            else:
                self._show_error(exc)

        self.run_async(work, done, busy=None if silent else "Loading tickets, tags and parent tickets…",
                       on_error=fail)

    def _schedule_auto_refresh(self):
        if self._auto_job:
            self.after_cancel(self._auto_job)
        self._auto_job = self.after(AUTO_REFRESH_MINUTES * 60 * 1000, self._auto_refresh)

    def _auto_refresh(self):
        self._auto_job = None
        if self._busy or not self._configured(warn=False):
            self._schedule_auto_refresh()
            return
        self.refresh_tickets(silent=True)
        # the dashboard follows the same hourly rhythm: refresh it now if it's on screen, else on next open
        self.dashboard_dirty = True
        try:
            if self.nb.select() == str(self.tab_dash) and self.dash_raw is not None:
                self.refresh_dashboard()
        except (tk.TclError, AttributeError):
            pass

    def _refresh_filter_values(self):
        def setvals(cb, var, first, values, extra=()):
            vals = [first] + list(extra) + values
            cb.configure(values=vals)
            if var.get() not in vals:
                var.set(first)
        statuses = sorted({name_of(it.get("status")) for it in self.issues}, key=str.lower)
        setvals(self.cb_status, self.var_status_filter, "All statuses", statuses)
        sprints = sorted({t for it in self.issues for t in self._tags(it) if self._is_sprint_tag(t)} |
                         {self.plan_info(it)["sprint"] for it in self.issues} - {""},
                         key=lambda t: [int(x) if x.isdigit() else x for x in re.split(r"(\d+)", t.lower())])
        setvals(self.cb_sprint, self.var_sprint_filter, "All sprints", sprints, ["No sprint"])
        releases = sorted({self.plan_info(it)["release"] for it in self.issues} - {""}, key=str.lower)
        self.mp_release.set_options(["No release"] + releases)
        self._refresh_groom_choices()
        prios = sorted({name_of(it.get("priority")) for it in self.issues},
                       key=lambda n: [k for k in PRIORITY_COLORS].index(next((k for k in PRIORITY_COLORS
                                                                              if k in n.lower()), "low")))
        setvals(self.cb_priority, self.var_priority_filter, "All priorities", prios)
        people = sorted({name_of(it.get("assigned_to"), "") for it in self.issues} - {""}, key=str.lower)
        setvals(self.cb_assignee, self.var_assignee_filter, "All assignees", people, ["Unassigned"])

    def _tags(self, issue):
        return self.tags_by_id.get(issue.get("id"), [])

    def _sprint(self, issue):
        return ", ".join(t for t in self._tags(issue) if self._is_sprint_tag(t))

    def _track(self, issue):
        if not self.tags_supported or issue.get("id") not in self.issue_index:
            return "?"
        return "2.0" if any(t.lower() == TRACK_TAG.lower() for t in self._tags(issue)) else "Legacy"

    def _parent_label(self, issue):
        pid = (issue.get("parent") or {}).get("id")
        if not pid:
            return ""
        par = self.issue_index.get(pid) or self.extra_issues.get(pid)
        return f"#{pid} {par.get('subject', '')}" if par else f"#{pid}"

    def edit_tags(self):
        win = tk.Toplevel(self)
        win.title("Tags to track")
        win.configure(bg=C["bg"])
        win.transient(self)
        win.geometry("460x520")
        frm = ttk.Frame(win, padding=16, style="Bg.TFrame")
        frm.pack(fill="both", expand=True)
        msg = ("Type one tag per line, exactly as it appears in Redmine. Sprint tags (Sprint 1, Sprint 2, "
               f"Sprint 3 …) feed the Sprint column. The '{TRACK_TAG}' tag decides the Track column.")
        if self.tags_supported is False:
            msg = ("⚠ Your Redmine did not answer tag questions, so tag columns show '?'. Ask your Redmine "
                   "admin whether the Tags plugin filter is enabled for the API.\n\n") + msg
        ttk.Label(frm, text=msg, wraplength=420, justify="left").pack(anchor="w")
        txt = tk.Text(frm, height=16, font=FONTS["body"], relief="flat", highlightthickness=1,
                      highlightbackground=C["line"], padx=8, pady=6)
        txt.pack(fill="both", expand=True, pady=10)
        txt.insert("1.0", "\n".join(self.config_data.get("tracked_tags") or DEFAULT_TAGS))

        def save():
            tags = [t.strip() for t in txt.get("1.0", "end").splitlines() if t.strip()]
            self.config_data["tracked_tags"] = dedupe(tags) or DEFAULT_TAGS
            try:
                save_json(CONFIG_PATH, self.config_data)
            except OSError:
                pass
            win.destroy()
            self.refresh_tickets()

        ttk.Button(frm, text="Save & Reload", style="Accent.TButton", command=save).pack(side="right")
        ttk.Button(frm, text="Cancel", command=win.destroy).pack(side="right", padx=8)

    # ----- shared helpers --------------------------------------------------
    def _ensure(self, module: str, pip_name: str) -> bool:
        """Make sure an optional add-on is installed; offer to install it if not."""
        import importlib
        try:
            importlib.import_module(module)
            return True
        except ImportError:
            pass
        if not messagebox.askyesno(APP_NAME, f"This needs a small free add-on called '{pip_name}' that isn't "
                                             "installed yet.\n\nInstall it now? (needs internet, takes under a minute)",
                                   parent=self):
            return False
        import subprocess
        self.configure(cursor="watch")
        self.set_status(f"Installing {pip_name}… please wait")
        self.update()
        detail = ""
        try:
            r = subprocess.run([sys.executable, "-m", "pip", "install", pip_name], capture_output=True, text=True,
                               timeout=600)
            detail = (r.stderr or r.stdout or "")[-500:]
        except Exception as exc:  # pip missing, no internet, blocked by policy …
            detail = str(exc)
        finally:
            self.configure(cursor="")
        importlib.invalidate_caches()
        try:
            importlib.import_module(module)
            self.set_status(f"{pip_name} installed.")
            return True
        except ImportError:
            messagebox.showerror(APP_NAME, f"Could not install {pip_name} automatically.\n\nIn the VS Code terminal "
                                           f"type:\n\n    python -m pip install {pip_name}\n\nthen restart the app."
                                           f"\n\nDetails: {detail}", parent=self)
            return False

    def _save_config_quiet(self):
        try:
            save_json(CONFIG_PATH, self.config_data)
        except OSError:
            pass

    def _offer_open(self, path):
        if messagebox.askyesno(APP_NAME, f"Saved:\n{path}\n\nOpen it now?", parent=self):
            try:
                os.startfile(path)  # Windows
            except AttributeError:
                webbrowser.open(Path(path).as_uri())
            except OSError:
                pass

    def holidays(self) -> dict:
        if self._holiday_cache is None:
            raw = self.config_data.get("holidays")
            if raw is None:
                raw = [list(h) for h in DEFAULT_HOLIDAYS_2026]
            out = {}
            for item in raw:
                d = parse_date(item[0])
                if d:
                    out[d] = item[1] if len(item) > 1 else "Holiday"
            self._holiday_cache = out
        return self._holiday_cache

    def stuck_after(self) -> int:
        try:
            return max(1, int(self.var_stuck.get()))
        except (tk.TclError, ValueError):
            return 2

    def _plans(self) -> dict:
        plans = self.config_data.get("sprint_plans")
        if plans is None:
            plans = {}
            if self.config_data.get("sprint_plan"):  # plan saved by an earlier version of the app
                plans[NO_RELEASE] = self.config_data["sprint_plan"]
            self.config_data["sprint_plans"] = plans
        return plans

    def _plan_from_cfg(self, cfg) -> list:
        if not cfg:
            return []
        out = []
        start = parse_date(cfg.get("start"))
        count = int(cfg.get("count", 3))
        if start and count > 0:
            length = cfg.get("weeks", 2) if cfg.get("mode") == "calendar" else cfg.get("wdays", 10)
            regular = plan_sprints(start, cfg.get("mode", "calendar"), int(length), count, self.holidays(),
                                   cfg.get("prefix", "Sprint"), int(cfg.get("first", 1)))
            for sp in regular:
                sp["kind"] = "Sprint"
            if regular and count >= 2 and cfg.get("regression", True):
                last = regular[-1]
                last["kind"] = "Regression"
                extra = int(cfg.get("reg_extra", 0) or 0)
                if extra > 0:
                    hols = self.holidays()
                    end, n = last["end"], 0
                    while n < extra:
                        end += dt.timedelta(days=1)
                        n += is_working_day(end, hols)
                    last["end"] = end
                    last["working"] = count_working_days(last["start"], end, hols)
                    last["holidays"] = [(d, hols[d]) for d in sorted(hols) if last["start"] <= d <= end and d.weekday() < 5]
                    last["note"] = f"+{plural(extra, 'extra working day')}"
            out.extend(regular)
        hols = self.holidays()
        for i, ex in enumerate(cfg.get("extras") or []):
            a, b = parse_date(ex.get("start")), parse_date(ex.get("end"))
            if not a:
                continue
            b = b if b and b >= a else a
            out.append({"name": ex.get("name") or f"Patch {i + 1}", "start": a, "end": b,
                        "working": count_working_days(a, b, hols), "kind": ex.get("kind", "Patch release"),
                        "holidays": [(d, hols[d]) for d in sorted(hols) if a <= d <= b and d.weekday() < 5],
                        "extra_index": i, "note": ex.get("note", "")})
        out.sort(key=lambda sp: (sp["start"], sp["end"]))
        closed = cfg.get("closed") or {}
        for sp in out:
            if sp["name"] in closed:
                sp["closed"] = closed[sp["name"]]
        return out

    def _selected_sprint_name(self):
        sel = self.tv_sprints.selection()
        return self.tv_sprints.set(sel[0], "sprint") if sel else None

    def close_sprint(self):
        key = self.var_sp_version.get() or NO_RELEASE
        name = self._selected_sprint_name()
        if not name:
            messagebox.showinfo(APP_NAME, "Click the sprint you want to close in the table first.", parent=self)
            return
        sp = next((x for x in self.sprint_plan(key) if x["name"] == name), None)
        if not sp:
            return
        prev = sp.get("closed") or {}
        open_left = [it for it in self.issues if name.lower() in {t.lower() for t in self._tags(it)}]
        win = tk.Toplevel(self)
        win.title(f"Close {name}")
        win.configure(bg=C["bg"])
        win.transient(self)
        win.geometry("520x430")
        frm = ttk.Frame(win, padding=16, style="Bg.TFrame")
        frm.pack(fill="both", expand=True)
        ttk.Label(frm, text=f"{key} · {name}", font=FONTS["h3"], foreground=C["header"]).pack(anchor="w")
        ttk.Label(frm, style="Muted.TLabel",
                  text=f"{fmt_date(sp['start'])} → {fmt_date(sp['end'])} · {sp['working']} working days").pack(anchor="w")
        if open_left:
            ttk.Label(frm, style="Flag.TLabel", wraplength=480, justify="left",
                      text=f"⚠ {plural(len(open_left), 'ticket')} tagged '{name}' are still open in Redmine "
                           f"(e.g. #{open_left[0]['id']}). Mention carry-over in your comments.").pack(anchor="w", pady=(6, 0))
        row = ttk.Frame(frm, style="Bg.TFrame")
        row.pack(fill="x", pady=(10, 4))
        ttk.Label(row, text="Closed on").pack(side="left")
        v_date = tk.StringVar(value=(parse_date(prev.get("date")) or dt.date.today()).strftime("%d-%m-%Y"))
        ttk.Entry(row, textvariable=v_date, width=12).pack(side="left", padx=8)
        ttk.Label(frm, text="Comments (what was delivered, what carried over, anything to remember)").pack(anchor="w",
                                                                                                         pady=(6, 4))
        txt = tk.Text(frm, height=8, wrap="word", font=FONTS["body"], relief="flat", highlightthickness=1,
                      highlightbackground=C["line"], padx=8, pady=6)
        txt.pack(fill="both", expand=True)
        txt.insert("1.0", prev.get("comment", ""))

        def save():
            d = parse_user_date(v_date.get()) or dt.date.today()
            cfg = self._plans().setdefault(key, {"count": 0})
            cfg.setdefault("closed", {})[name] = {"date": d.isoformat(), "comment": txt.get("1.0", "end").strip(),
                                                  "open_left": len(open_left)}
            self._save_config_quiet()
            win.destroy()
            self.render_sprints()
            self.set_status(f"{name} closed.")

        bb = ttk.Frame(frm, style="Bg.TFrame")
        bb.pack(fill="x", pady=(10, 0))
        ttk.Button(bb, text="Cancel", command=win.destroy).pack(side="right")
        ttk.Button(bb, text="✔ Close sprint" if not prev else "Save", style="Accent.TButton",
                   command=save).pack(side="right", padx=8)

    def reopen_sprint(self):
        key = self.var_sp_version.get() or NO_RELEASE
        name = self._selected_sprint_name()
        closed = (self._plans().get(key) or {}).get("closed") or {}
        if not name or name not in closed:
            messagebox.showinfo(APP_NAME, "Select a closed sprint first.", parent=self)
            return
        if messagebox.askyesno(APP_NAME, f"Reopen {name}? Its closing comments will be removed.", parent=self):
            closed.pop(name, None)
            self._save_config_quiet()
            self.render_sprints()

    def save_prod_date(self):
        key = self.var_sp_version.get() or NO_RELEASE
        txt = self.var_sp_prod.get().strip()
        d = parse_user_date(txt) if txt else None
        if txt and not d:
            messagebox.showinfo(APP_NAME, "Type the production date like 20-11-2026 (or leave it empty).", parent=self)
            return
        cfg = self._plans().setdefault(key, {"count": 0})
        if d:
            cfg["prod_date"] = d.isoformat()
        else:
            cfg.pop("prod_date", None)
        self._save_config_quiet()
        self.render_sprints()
        self.set_status(f"Production deployment date for {key} {'saved' if d else 'cleared'}.")

    def _sprint_names(self) -> set:
        return {sp["name"].lower() for _, sp in self.all_sprints()}

    def _is_sprint_tag(self, tag: str) -> bool:
        return bool(SPRINT_RE.match(tag)) or tag.lower() in self._sprint_names()

    def add_extra_sprint(self):
        key = self.var_sp_version.get() or NO_RELEASE
        win = tk.Toplevel(self)
        win.title("Add a patch / one-off sprint")
        win.configure(bg=C["bg"])
        win.transient(self)
        frm = ttk.Frame(win, padding=16, style="Bg.TFrame")
        frm.pack(fill="both", expand=True)
        existing = [sp for sp in self.sprint_plan(key) if sp.get("kind") != "Sprint"]
        today = dt.date.today()
        v_kind = tk.StringVar(value="Patch release")
        v_name = tk.StringVar(value=f"Patch {len(existing) + 1}")
        v_start = tk.StringVar(value=next_working_day(today, self.holidays()).strftime("%d-%m-%Y"))
        v_end = tk.StringVar(value=v_start.get())
        v_note = tk.StringVar()
        ttk.Label(frm, text=f"Release: {key}", font=FONTS["h3"], foreground=C["header"]).grid(
            row=0, column=0, columnspan=2, sticky="w", pady=(0, 8))
        rows = [("Type", None), ("Name", v_name), ("Start date", v_start), ("End date", v_end), ("Note (optional)", v_note)]
        for r, (label, var) in enumerate(rows, 1):
            ttk.Label(frm, text=label).grid(row=r, column=0, sticky="w", pady=4, padx=(0, 10))
            if var is None:
                cb = ttk.Combobox(frm, textvariable=v_kind, state="readonly", width=28,
                                  values=["Patch release", "One-day sprint", "Hotfix", "Other"])
                cb.grid(row=r, column=1, sticky="ew")
            else:
                ttk.Entry(frm, textvariable=var, width=30).grid(row=r, column=1, sticky="ew")
        ttk.Label(frm, style="Muted.TLabel", text="Dates like 21-10-2026. For a one-day release use the same start and "
                                                  "end date.", wraplength=360, justify="left").grid(
            row=6, column=0, columnspan=2, sticky="w", pady=(6, 0))

        def kind_changed(_e=None):
            k = v_kind.get()
            base = {"Patch release": "Patch", "One-day sprint": "One-day sprint", "Hotfix": "Hotfix"}.get(k, "Extra")
            v_name.set(f"{base} {len([e for e in existing if e.get('kind') == k]) + 1}")
            if k in ("One-day sprint", "Hotfix"):
                v_end.set(v_start.get())

        cb.bind("<<ComboboxSelected>>", kind_changed)

        def start_changed(*_):
            if v_kind.get() in ("One-day sprint", "Hotfix"):
                v_end.set(v_start.get())

        v_start.trace_add("write", start_changed)

        def ok():
            a, b = parse_user_date(v_start.get()), parse_user_date(v_end.get() or v_start.get())
            if v_kind.get() == "One-day sprint":
                b = a
            if not a or not b:
                messagebox.showinfo(APP_NAME, "Type the dates like 21-10-2026.", parent=win)
                return
            if b < a:
                messagebox.showinfo(APP_NAME, "The end date is before the start date.", parent=win)
                return
            hols = self.holidays()
            if count_working_days(a, b, hols) == 0 and not messagebox.askyesno(
                    APP_NAME, "These dates fall only on a weekend or company holiday. Add anyway?", parent=win):
                return
            plans = self._plans()
            cfg = plans.setdefault(key, {"count": 0})
            cfg.setdefault("extras", []).append({"name": v_name.get().strip() or "Patch", "kind": v_kind.get(),
                                                 "start": a.isoformat(), "end": b.isoformat(),
                                                 "note": v_note.get().strip()})
            if key in self.sp_version_info:
                cfg["release_due"] = self.sp_version_info[key].get("due_date") or cfg.get("release_due")
            self._save_config_quiet()
            win.destroy()
            self.render_sprints()
            self._populate_tree()
            self.set_status(f"{v_kind.get()} added to {key}.")

        bb = ttk.Frame(frm, style="Bg.TFrame")
        bb.grid(row=7, column=0, columnspan=2, sticky="e", pady=(14, 0))
        ttk.Button(bb, text="Cancel", command=win.destroy).pack(side="right")
        ttk.Button(bb, text="Add", style="Accent.TButton", command=ok).pack(side="right", padx=8)

    def _sprint_row_double(self):
        sel = self.tv_sprints.selection()
        if not sel:
            return
        if sel[0].startswith("w"):
            self.week_dialog(edit=True)
        elif self.tv_sprints.set(sel[0], "state").startswith("✔"):
            self.close_sprint()

    def week_dialog(self, edit=False):
        key = self.var_sp_version.get() or NO_RELEASE
        plan = self.sprint_plan(key)
        cfg = self._plans().setdefault(key, {"count": 0})
        weeks = cfg.setdefault("week_list", [])
        sel = self.tv_sprints.selection()
        idx = None
        if edit:
            if not sel or not sel[0].startswith("w"):
                messagebox.showinfo(APP_NAME, "Click a week in the table first (weeks are listed under their sprint).",
                                    parent=self)
                return
            idx = int(sel[0][1:])
        sprint_names = [sp["name"] for sp in plan]
        default_sprint = ""
        if sel:
            row = sel[0]
            parent = self.tv_sprints.parent(row)
            default_sprint = self.tv_sprints.set(parent or row, "sprint")
        if not default_sprint and plan:
            today = dt.date.today()
            default_sprint = next((sp["name"] for sp in plan if sp["start"] <= today <= sp["end"]), plan[0]["name"])
        hols = self.holidays()

        def suggest(sprint_name):
            mine = sorted((w for w in weeks if w.get("sprint") == sprint_name and parse_date(w.get("end"))),
                          key=lambda w: w["end"])
            sp = next((x for x in plan if x["name"] == sprint_name), None)
            if mine:
                a = next_working_day(parse_date(mine[-1]["end"]) + dt.timedelta(days=1), hols)
            elif sp:
                a = sp["start"]
            else:
                a = next_working_day(dt.date.today(), hols)
            b = a + dt.timedelta(days=1)
            while b.weekday() != 0:  # the following Monday (grooming day)
                b += dt.timedelta(days=1)
            return a, b

        w0 = weeks[idx] if edit else {}
        v_name = tk.StringVar(value=w0.get("name", ""))
        v_sprint = tk.StringVar(value=w0.get("sprint", default_sprint))
        if edit:
            a0, b0 = parse_date(w0.get("start")), parse_date(w0.get("end"))
        else:
            a0, b0 = suggest(v_sprint.get())
        v_from = tk.StringVar(value=a0.strftime("%d-%m-%Y") if a0 else "")
        v_to = tk.StringVar(value=b0.strftime("%d-%m-%Y") if b0 else "")
        c0 = parse_date(w0.get("cutoff"))
        v_cut = tk.StringVar(value=c0.strftime("%d-%m-%Y") if c0 else "")

        win = tk.Toplevel(self)
        win.title("Edit week" if edit else "Add a week")
        win.configure(bg=C["bg"])
        win.transient(self)
        frm = ttk.Frame(win, padding=16, style="Bg.TFrame")
        frm.pack(fill="both", expand=True)
        ttk.Label(frm, text=f"Release: {key}", font=FONTS["h3"], foreground=C["header"]).grid(
            row=0, column=0, columnspan=2, sticky="w", pady=(0, 8))
        ttk.Label(frm, text="Week name").grid(row=1, column=0, sticky="w", pady=4, padx=(0, 10))
        ttk.Entry(frm, textvariable=v_name, width=30).grid(row=1, column=1, sticky="ew")
        ttk.Label(frm, text="Sprint").grid(row=2, column=0, sticky="w", pady=4, padx=(0, 10))
        cb = ttk.Combobox(frm, textvariable=v_sprint, state="readonly", values=sprint_names + ["—"], width=28)
        cb.grid(row=2, column=1, sticky="ew")
        for r, (label, var) in enumerate((("From", v_from), ("To", v_to), ("Dev cut-off (optional)", v_cut)), 3):
            ttk.Label(frm, text=label).grid(row=r, column=0, sticky="w", pady=4, padx=(0, 10))
            ttk.Entry(frm, textvariable=var, width=30).grid(row=r, column=1, sticky="ew")
        ttk.Label(frm, style="Muted.TLabel", wraplength=380, justify="left",
                  text="Dates like 21-10-2026. From/To are only suggestions; change them as you need. Tickets due in "
                       "this week after the dev cut-off are flagged 'not enough testing time'. Leave the cut-off empty "
                       "for no check.").grid(row=6, column=0, columnspan=2, sticky="w", pady=(6, 0))

        def sprint_changed(_e=None):
            if not edit:
                a, b = suggest(v_sprint.get())
                v_from.set(a.strftime("%d-%m-%Y"))
                v_to.set(b.strftime("%d-%m-%Y"))

        cb.bind("<<ComboboxSelected>>", sprint_changed)

        def ok():
            a, b = parse_user_date(v_from.get()), parse_user_date(v_to.get())
            cut_txt = v_cut.get().strip()
            cut = parse_user_date(cut_txt) if cut_txt else None
            if not a or not b or (cut_txt and not cut):
                messagebox.showinfo(APP_NAME, "Type the dates like 21-10-2026.", parent=win)
                return
            if b < a:
                messagebox.showinfo(APP_NAME, "'To' is before 'From'.", parent=win)
                return
            rec = {"name": v_name.get().strip(), "sprint": "" if v_sprint.get() == "—" else v_sprint.get(),
                   "start": a.isoformat(), "end": b.isoformat(), "cutoff": cut.isoformat() if cut else "",
                   "goals": w0.get("goals", "")}
            if edit:
                weeks[idx] = rec
            else:
                weeks.append(rec)
            self._save_config_quiet()
            win.destroy()
            self.render_sprints()
            self._populate_tree()
            self._refresh_week_choices()
            self.set_status("Week saved.")

        bb = ttk.Frame(frm, style="Bg.TFrame")
        bb.grid(row=7, column=0, columnspan=2, sticky="e", pady=(14, 0))
        ttk.Button(bb, text="Cancel", command=win.destroy).pack(side="right")
        ttk.Button(bb, text="Save", style="Accent.TButton", command=ok).pack(side="right", padx=8)

    def remove_extra_sprint(self):
        key = self.var_sp_version.get() or NO_RELEASE
        wsel = [k for k in self.tv_sprints.selection() if k.startswith("w")]
        if wsel:
            if not messagebox.askyesno(APP_NAME, f"Delete {plural(len(wsel), 'week')}?", parent=self):
                return
            weeks = (self._plans().get(key) or {}).get("week_list") or []
            for k in sorted((int(x[1:]) for x in wsel), reverse=True):
                if k < len(weeks):
                    weeks.pop(k)
            self._save_config_quiet()
            self.render_sprints()
            self._populate_tree()
            self._refresh_week_choices()
            return
        sel = [k for k in self.tv_sprints.selection() if k.startswith("x")]
        if not sel:
            messagebox.showinfo(APP_NAME, "Select a patch / one-off sprint in the table first.\n\nRegular sprints come "
                                          "from the sprint settings: lower 'Number of sprints' and click Calculate, "
                                          "or set it to 0 to keep only patch sprints.", parent=self)
            return
        cfg = self._plans().get(key) or {}
        extras = cfg.get("extras") or []
        for k in sorted((int(x[1:]) for x in sel), reverse=True):
            if k < len(extras):
                extras.pop(k)
        self._save_config_quiet()
        self.render_sprints()
        self._populate_tree()
        self.set_status("Removed.")

    def sprint_plan(self, key=None) -> list:
        key = key or self.var_sp_version.get() or NO_RELEASE
        return self._plan_from_cfg(self._plans().get(key))

    def all_sprints(self):
        c = getattr(self, "_sprints_cache", None)
        if c and (dt.datetime.now() - c[0]).total_seconds() < 2:
            return c[1]
        out = self._all_sprints()
        self._sprints_cache = (dt.datetime.now(), out)
        return out

    def _all_sprints(self):
        out = []
        for key, cfg in self._plans().items():
            for sp in self._plan_from_cfg(cfg):
                out.append((key, sp))
        return out

    def _current_sprint(self):
        today = dt.date.today()
        for key, sp in self.all_sprints():
            if sp["start"] <= today <= sp["end"]:
                return key, sp
        return None

    def _sprint_for(self, issue):
        names = {t.lower() for t in self._tags(issue) if self._is_sprint_tag(t)}
        if not names:
            return None
        version = name_of(issue.get("fixed_version"), "")
        if version in self._plans():
            for sp in self.sprint_plan(version):
                if sp["name"].lower() in names:
                    return sp
        today = dt.date.today()
        matches = [sp for _, sp in self.all_sprints() if sp["name"].lower() in names]
        return min(matches, key=lambda sp: abs((sp["start"] - today).days)) if matches else None

    def _hist(self, issue):
        c = self.hist_cache.get(issue.get("id"))
        return history_info(c["issue"], self.holidays()) if c else None

    def _flags(self, issue):
        flags = status_flags(issue, hist=self._hist(issue), stuck_after=self.stuck_after(),
                             sprint=self._sprint_for(issue))
        due = parse_date(issue.get("due_date"))
        w = self._week_for_date(due)
        if w and w["cutoff"] and due > w["cutoff"]:
            flags.append(f"After dev cut-off ({w.get('name') or 'week'}: {w['cutoff'].strftime('%d-%b')}), "
                         "not enough testing time")
        m = self.member_for(issue)
        if m:
            rel = parse_date(m.get("relieving"))
            if m.get("status") == "Inactive":
                flags.append("Assignee is inactive")
            elif rel and rel >= dt.date.today():
                flags.append(f"Assignee leaves {fmt_date(rel, False)}" + (" (due after last day)"
                                                                          if due and due > rel else ""))
        note = (self.config_data.get("ticket_notes") or {}).get(str(issue.get("id"))) or {}
        fu = parse_date(note.get("follow_up"))
        if fu and fu <= dt.date.today():
            flags.append(f"Follow-up due {fmt_date(fu, False)}")
        return flags

    def _load_histories(self, issues, then=None, busy=None):
        need = [it["id"] for it in issues
                if (self.hist_cache.get(it["id"]) or {}).get("updated_on") != it.get("updated_on")]
        if not need:
            if then:
                then()
            return
        client = self.make_client()

        def done(res):
            for iid, issue in res.items():
                self.hist_cache[iid] = {"updated_on": issue.get("updated_on"), "issue": issue}
            if then:
                then()
            self.set_status(f"Ticket history updated ({len(res)} tickets).")

        self.run_async(lambda: fetch_journals(client, need), done,
                       busy=busy or f"Reading the history of {len(need)} tickets (time in status, due-date changes)…",
                       on_error=lambda e: self.set_status(f"⚠ Could not read ticket history: {str(e).splitlines()[0]}"))

    def _make_table(self, parent, columns, height=10):
        frame = ttk.Frame(parent)
        tree = ttk.Treeview(frame, columns=[c[0] for c in columns], show="headings", height=height)
        for key, text, width in columns:
            tree.heading(key, text=text, anchor="w", command=lambda k=key, t=tree: self._sort_tree(t, k))
            tree.column(key, width=width, minwidth=40, stretch=key in ("subject", "occasion"), anchor="w")
        vsb = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")
        tree.tag_configure("odd", background=C["odd"])
        tree.tag_configure("hot", foreground=C["bad"])
        tree.tag_configure("muted", foreground=C["faint"])
        tree.tag_configure("current", background=C["soft"])
        return frame, tree

    def _sort_tree(self, tree, col):
        rows = [(tree.set(k, col), k) for k in tree.get_children("")]

        def key(v):
            txt = v[0].replace("#", "").replace("%", "").replace(" d", "").replace("+", "").strip()
            try:
                return (0, float(txt), "")
            except ValueError:
                d = parse_user_date(txt.split(" ", 1)[-1]) or parse_user_date(txt)
                return (1, d.toordinal(), "") if d else (2, 0, txt.lower())

        desc = getattr(tree, "_sort_desc", {}).get(col, False)
        rows.sort(key=key, reverse=not desc)
        for i, (_, k) in enumerate(rows):
            tree.move(k, "", i)
        tree._sort_desc = {col: not desc}

    def _open_ticket_from(self, tree, col="ticket"):
        sel = tree.selection()
        if not sel:
            return
        m = re.search(r"\d+", tree.set(sel[0], col) or "")
        if m:
            webbrowser.open(f"{self.var_url.get().rstrip('/')}/issues/{m.group()}")

    # ----- Weekly view -------------------------------------------------------
    def _build_weekly(self, parent):
        self.var_week = tk.StringVar()
        top = ttk.Frame(parent, style="Bg.TFrame", padding=(0, 10, 0, 6))
        top.pack(fill="x")
        left = ttk.Frame(top, style="Bg.TFrame")
        left.pack(side="left")
        ttk.Label(left, text="Weekly View", font=FONTS["h1"], foreground=C["header"]).pack(anchor="w")
        ttk.Label(left, style="Muted.TLabel", text="Week goals, who is doing what, and who is being groomed on what. "
                                                   "Weeks are set up in the Sprint Planner (➕ Add week).").pack(anchor="w")
        ttk.Button(top, text="⬇ Export week (Excel)…", command=self.export_weekly).pack(side="right")
        ttk.Button(top, text="↻ Refresh", style="Accent.TButton", command=self.refresh_weekly).pack(side="right",
                                                                                                   padx=6)
        wfb = FilterBar(parent, cols=6, title="WEEK")
        wfb.pack(fill="x", pady=(0, 8))
        ctl = wfb.top
        ttk.Button(ctl, text="◀", width=3, command=lambda: self._step_week(-1)).pack(side="left")
        self.cb_week = ttk.Combobox(ctl, textvariable=self.var_week, state="readonly", width=60)
        self.cb_week.pack(side="left", padx=6)
        self.cb_week.bind("<<ComboboxSelected>>", lambda e: self.refresh_weekly())
        ttk.Button(ctl, text="▶", width=3, command=lambda: self._step_week(1)).pack(side="left")
        self.lbl_week_info = tk.Label(ctl, text="", bg=C["soft"], fg=C["muted"], font=FONTS["small"])
        self.lbl_week_info.pack(side="left", padx=12)
        wfb.grid_f.pack_forget()

        pw = ttk.PanedWindow(parent, orient="horizontal")
        pw.pack(fill="both", expand=True)
        gf = ttk.LabelFrame(pw, text="🎯 Week goals (one per line)", padding=8)
        rf = ttk.Frame(pw, style="Bg.TFrame")
        pw.add(gf, weight=1)
        pw.add(rf, weight=3)
        self.txt_goals = tk.Text(gf, wrap="word", height=10, font=FONTS["body"], relief="flat", highlightthickness=1,
                                 highlightbackground=C["line"], padx=8, pady=6, width=34)
        self.txt_goals.pack(fill="both", expand=True)
        self.txt_goals.bind("<FocusOut>", lambda e: self.save_week_goals(quiet=True))
        ttk.Button(gf, text="Save goals", command=self.save_week_goals).pack(anchor="e", pady=(6, 0))

        cols = [("status", "Status", 130), ("due", "Due", 110), ("groom", "Groom date", 110),
                ("sprint", "Sprint", 110), ("release", "Release", 140), ("prd", "PRD", 100), ("check", "Check", 300)]
        boxes = []
        for title in ("👥 Due this week: who is doing what", "🧭 To be groomed this week (by groom date)"):
            lf = ttk.LabelFrame(rf, text=title, padding=6)
            lf.pack(fill="both", expand=True, pady=(0, 6))
            f = ttk.Frame(lf)
            f.pack(fill="both", expand=True)
            tv = ttk.Treeview(f, columns=[c[0] for c in cols], show="tree headings", height=7)
            tv.heading("#0", text="Person / ticket", anchor="w")
            tv.column("#0", width=380, stretch=True)
            for key, text, width in cols:
                tv.heading(key, text=text, anchor="w")
                tv.column(key, width=width, stretch=key == "check", anchor="w")
            sb = ttk.Scrollbar(f, orient="vertical", command=tv.yview)
            tv.configure(yscrollcommand=sb.set)
            tv.pack(side="left", fill="both", expand=True)
            sb.pack(side="right", fill="y")
            tv.tag_configure("person", font=FONTS["small_bold"], background=C["head_bg"])
            tv.tag_configure("hot", foreground=C["bad"])
            tv.tag_configure("muted", foreground=C["faint"])
            tv.bind("<Double-1>", lambda e, t=tv: self._open_week_ticket(t))
            boxes.append(tv)
        self.tv_week_due, self.tv_week_groom = boxes

    def _refresh_week_choices(self):
        weeks = self.all_weeks()
        self.week_labels = {}
        for w in weeks:
            label = (f"{w.get('name') or '(unnamed week)'}  ·  {fmt_date(w['start'])} – {fmt_date(w['end'])}"
                     f"  ·  {w['release']}" + (f" / {w['sprint']}" if w.get("sprint") else ""))
            self.week_labels[label] = (w["release"], w["index"])
        labels = list(self.week_labels)
        self.cb_week.configure(values=labels)
        if self.var_week.get() not in self.week_labels:
            today = dt.date.today()
            pick = next((lb for lb, w in zip(labels, weeks) if w["start"] <= today <= w["end"]), None) or \
                next((lb for lb, w in zip(labels, weeks) if w["start"] > today), None) or (labels[-1] if labels else "")
            self.var_week.set(pick)

    def _step_week(self, step):
        labels = list(self.week_labels)
        if not labels:
            return
        cur = labels.index(self.var_week.get()) if self.var_week.get() in labels else 0
        self.var_week.set(labels[max(0, min(len(labels) - 1, cur + step))])
        self.refresh_weekly()

    def _current_week(self):
        ref = self.week_labels.get(self.var_week.get())
        if not ref:
            return None, None
        key, idx = ref
        weeks = (self._plans().get(key) or {}).get("week_list") or []
        return (weeks[idx], ref) if idx < len(weeks) else (None, None)

    def save_week_goals(self, quiet=False):
        w, _ = self._current_week()
        if not w:
            return
        w["goals"] = self.txt_goals.get("1.0", "end").strip()
        self._save_config_quiet()
        if not quiet:
            self.set_status("Week goals saved.")

    def refresh_weekly(self):
        self._refresh_week_choices()
        w, _ = self._current_week()
        for tv in (self.tv_week_due, self.tv_week_groom):
            tv.delete(*tv.get_children())
        self.txt_goals.delete("1.0", "end")
        if not w:
            self.lbl_week_info.configure(text="No weeks yet. Add weeks in the Sprint Planner (➕ Add week).")
            return
        self.txt_goals.insert("1.0", w.get("goals", ""))
        a, b = parse_date(w["start"]), parse_date(w["end"])
        cut = parse_date(w.get("cutoff"))
        self.lbl_week_info.configure(text=f"{count_working_days(a, b, self.holidays())} working days · dev cut-off "
                                          f"{fmt_date(cut) if cut else 'not set'}")
        if not self._configured(warn=False):
            return
        pids, _, _ = self._scope()
        grooms = {int(k): parse_date(v) for k, v in (self.config_data.get("groom_dates") or {}).items()}
        groom_ids = [k for k, d in grooms.items() if d and a <= d <= b]
        client = self.make_client()
        this_week = f"{(self._current_week()[1] or ('',))[0]}|{a.isoformat()}"
        picked = self.config_data.get("ticket_weeks") or {}
        moved_in = [int(k) for k, ref in picked.items() if ref == this_week]
        moved_out = {int(k) for k, ref in picked.items() if ref != this_week}

        def work():
            due, _ = client.issues_in({"status_id": "*", "due_date": f"><{a.isoformat()}|{b.isoformat()}",
                                       "sort": "due_date:asc"}, pids)
            due = [it for it in due if it["id"] not in moved_out]
            have = {it["id"] for it in due}
            extra = [i for i in moved_in if i not in have]
            if extra:
                due += client.issues_by_ids(extra)
            groom = client.issues_by_ids(groom_ids) if groom_ids else []
            return due, groom

        def done(res):
            due, groom = res
            self.weekly_data = {"week": w, "due": due, "groom": groom}
            self._fill_week_tree(self.tv_week_due, due, grooms)
            self._fill_week_tree(self.tv_week_groom, groom, grooms)
            self.set_status(f"Week loaded: {plural(len(due), 'ticket')} due, {plural(len(groom), 'ticket')} to groom.")

        self.run_async(work, done, busy="Loading the week…")

    def _fill_week_tree(self, tv, issues, grooms):
        tv.delete(*tv.get_children())
        by = {}
        for it in issues:
            by.setdefault(name_of(it.get("assigned_to"), "Unassigned"), []).append(it)
        for person in sorted(by, key=lambda n: (n == "Unassigned", n.lower())):
            items = by[person]
            pid = tv.insert("", "end", text=f"{person}  ({len(items)})", open=True, tags=("person",))
            for it in sorted(items, key=lambda x: (x.get("due_date") or "", x["id"])):
                flags = self._flags(it)
                state, _ = self._prd_info(it)
                closed = name_of(it.get("status"), "") in self.rn_closed_status or bool(it.get("closed_on"))
                tv.insert(pid, "end", iid=f"{tv}-{it['id']}", text=f"#{it['id']}   {it.get('subject', '')}",
                          values=(name_of(it.get("status")),
                                  parse_date(it["due_date"]).strftime("%a %d-%b") if it.get("due_date") else "",
                                  grooms[it["id"]].strftime("%a %d-%b") if grooms.get(it["id"]) else "",
                                  self.plan_info(it)["sprint_txt"], self.plan_info(it)["release_txt"],
                                  PRD_ICONS[state],
                                  ("📌 week picked by you   " if self.week_of(it)[1] == "manual" else "") +
                                  (("⚠ " + "; ".join(flags)) if flags else "")),
                          tags=("muted",) if closed else (("hot",) if any("cut-off" in f for f in flags) else ()))

    def _open_week_ticket(self, tv):
        sel = tv.selection()
        if sel:
            m = re.search(r"#(\d+)", tv.item(sel[0], "text"))
            if m:
                webbrowser.open(f"{self.var_url.get().rstrip('/')}/issues/{m.group(1)}")

    def export_weekly(self):
        d = self.weekly_data
        w, _ = self._current_week()
        if not d or not w:
            messagebox.showinfo(APP_NAME, "Choose a week and click ↻ Refresh first.", parent=self)
            return
        self._ensure("openpyxl", "openpyxl")
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            return
        name = re.sub(r"[^A-Za-z0-9]+", "_", w.get("name") or "Week").strip("_")
        path = filedialog.asksaveasfilename(parent=self, title="Save week", initialfile=f"{name}.xlsx",
                                            defaultextension=".xlsx", filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        grooms = {int(k): parse_date(v) for k, v in (self.config_data.get("groom_dates") or {}).items()}
        a, b, cut = parse_date(w["start"]), parse_date(w["end"]), parse_date(w.get("cutoff"))
        wb = Workbook()
        ws = wb.active
        ws.title = "Week"
        ws["A1"] = w.get("name") or "Week"
        ws["A1"].font = Font(size=16, bold=True, color="2F5496")
        ws["A2"] = (f"{fmt_date(a)} – {fmt_date(b)}  ·  {w.get('sprint') or ''}  ·  dev cut-off "
                    f"{fmt_date(cut) if cut else 'not set'}")
        ws["A2"].font = Font(size=10, color="808080")
        row = 4
        ws.cell(row=row, column=1, value="Week goals").font = Font(bold=True, size=12, color="102A43")
        for g in [x for x in (w.get("goals") or "").splitlines() if x.strip()] or ["(none)"]:
            row += 1
            ws.cell(row=row, column=1, value=f"• {g.strip()}")
        heads = ["Person", "Ticket", "Subject", "Status", "Due date", "Groom date", "Check"]
        widths = [20, 10, 60, 18, 14, 14, 50]
        for col, wd in zip("ABCDEFG", widths):
            ws.column_dimensions[col].width = wd
        for title, items in (("Due this week: who is doing what", d["due"]), ("To be groomed this week", d["groom"])):
            row += 2
            ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=12, color="102A43")
            row += 1
            for ci, h in enumerate(heads, 1):
                c = ws.cell(row=row, column=ci, value=h)
                c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="102A43")
            for it in sorted(items, key=lambda x: (name_of(x.get("assigned_to"), "~"), x.get("due_date") or "")):
                row += 1
                vals = [name_of(it.get("assigned_to"), "Unassigned"), it["id"], it.get("subject", ""),
                        name_of(it.get("status")), parse_date(it.get("due_date")), grooms.get(it["id"]),
                        "; ".join(self._flags(it))]
                for ci, v in enumerate(vals, 1):
                    c = ws.cell(row=row, column=ci, value=v)
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                    if isinstance(v, dt.date):
                        c.number_format = "DD-MMM-YYYY"
                link = ws.cell(row=row, column=2)
                link.hyperlink = f"{self.var_url.get().rstrip('/')}/issues/{it['id']}"
                link.font = Font(color="2F6FED", underline="single")
        try:
            wb.save(path)
        except OSError as exc:
            messagebox.showerror(APP_NAME, f"Could not save the file (is it open in Excel?):\n{exc}", parent=self)
            return
        self._offer_open(path)

    # ----- Delivery tracking -------------------------------------------------
    def _build_delivery(self, parent):
        top = ttk.Frame(parent, style="Bg.TFrame", padding=(0, 10, 0, 6))
        top.pack(fill="x")
        left = ttk.Frame(top, style="Bg.TFrame")
        left.pack(side="left")
        ttk.Label(left, text="Delivery Tracking", font=FONTS["h1"], foreground=C["header"]).pack(anchor="w")
        self.lbl_del_scope = ttk.Label(left, style="Muted.TLabel",
                                       text="Who pushes due dates, and which tickets are stuck in one status. "
                                            "Uses the project chosen in the Grooming Workspace (whole team).")
        self.lbl_del_scope.pack(anchor="w")
        ttk.Button(top, text="Copy summary", command=self.copy_delivery_summary).pack(side="right")
        ttk.Button(top, text="⬇ Management report (Excel)…", command=self.export_delivery).pack(side="right", padx=6)
        ttk.Button(top, text="↻ Refresh", style="Accent.TButton", command=self.refresh_delivery).pack(side="right")

        ctl = ttk.Frame(parent, style="Bg.TFrame")
        ctl.pack(fill="x", pady=(0, 8))
        ttk.Label(ctl, text="Period", style="Muted.TLabel").pack(side="left", padx=(0, 4))
        cb = ttk.Combobox(ctl, textvariable=self.var_period, state="readonly", width=15,
                          values=["Last 14 days", "Last 30 days", "Last 60 days", "Last 90 days", "Current sprint"])
        cb.pack(side="left")
        cb.bind("<<ComboboxSelected>>", lambda e: self.refresh_delivery())
        ttk.Label(ctl, text="Flag as stuck after", style="Muted.TLabel").pack(side="left", padx=(18, 4))
        sp = ttk.Spinbox(ctl, from_=1, to=30, width=4, textvariable=self.var_stuck, command=self._on_stuck_changed)
        sp.pack(side="left")
        sp.bind("<Return>", lambda e: self._on_stuck_changed())
        sp.bind("<FocusOut>", lambda e: self._on_stuck_changed())
        ttk.Label(ctl, style="Muted.TLabel",
                  text="working days in the same status  (weekends and company holidays don't count)").pack(
            side="left", padx=6)

        banner = tk.Frame(parent, bg=C["soft"], highlightbackground=C["soft_line"], highlightthickness=1)
        banner.pack(fill="x", pady=(0, 8))
        tk.Label(banner, text="SUMMARY FOR MANAGEMENT", bg=C["soft"], fg=C["accent_dark"],
                 font=FONTS["kpi_label"]).pack(anchor="w", padx=14, pady=(8, 0))
        self.lbl_del_summary = tk.Label(banner, text="Click ↻ Refresh to analyse ticket history.", bg=C["soft"],
                                        fg=C["ink"], font=FONTS["summary"], justify="left", anchor="w")
        self.lbl_del_summary.pack(anchor="w", fill="x", padx=14, pady=(2, 10))
        banner.bind("<Configure>", lambda e: self.lbl_del_summary.configure(wraplength=max(300, e.width - 30)))

        charts = tk.Frame(parent, bg=C["bg"])
        charts.pack(fill="x")
        self.ch_push = HBarChart(charts, "Days of due-date slippage, by developer", height=200)
        self.ch_stuck = HBarChart(charts, "Stuck tickets to follow up, by developer", height=200)
        self.ch_push.grid(row=0, column=0, sticky="nsew", padx=(0, 4))
        self.ch_stuck.grid(row=0, column=1, sticky="nsew", padx=(4, 0))
        charts.columnconfigure(0, weight=1, uniform="dc")
        charts.columnconfigure(1, weight=1, uniform="dc")

        nb = ttk.Notebook(parent)
        nb.pack(fill="both", expand=True, pady=(8, 0))
        f1, self.tv_devs = self._make_table(nb, [("dev", "Developer", 190), ("tickets", "Tickets pushed", 110),
                                                 ("pushes", "Pushes", 80), ("days", "Days pushed", 100),
                                                 ("avg", "Avg days / push", 120),
                                                 ("met", "Met original due date", 170),
                                                 ("stuck", "Stuck now", 90), ("open", "Open tickets", 100)])
        f2, self.tv_stuck = self._make_table(nb, [("dev", "Developer", 160), ("ticket", "Ticket", 80),
                                                  ("subject", "Subject", 380), ("status", "Status", 150),
                                                  ("since", "In this status since", 150),
                                                  ("days", "Working days", 110), ("due", "Due", 110)])
        f3, self.tv_changes = self._make_table(nb, [("ticket", "Ticket", 80), ("subject", "Subject", 340),
                                                    ("dev", "Developer", 150), ("by", "Changed by", 150),
                                                    ("when", "Changed on", 110), ("from", "From", 110),
                                                    ("to", "To", 110), ("days", "+Days", 70)])
        nb.add(f1, text="  By developer  ")
        nb.add(f2, text="  Stuck tickets: follow up  ")
        nb.add(f3, text="  Every due-date push  ")
        for tv in (self.tv_stuck, self.tv_changes):
            tv.bind("<Double-1>", lambda e, t=tv: self._open_ticket_from(t))

    def _on_stuck_changed(self):
        val = self.stuck_after()
        if self.config_data.get("stuck_after") == val:
            return
        self.config_data["stuck_after"] = val
        self._save_config_quiet()
        self._populate_tree()
        if self.delivery:
            self._recompute_delivery()

    def _delivery_period(self):
        today = dt.date.today()
        p = self.var_period.get()
        if p == "Current sprint":
            cur = self._current_sprint()
            sp = cur[1] if cur else None
            if sp:
                return sp["start"], today, f"{sp['name']} ({sp['start'].strftime('%d %b')} – today)"
            p = "Last 14 days"
        days = int(re.search(r"\d+", p).group())
        return today - dt.timedelta(days=days - 1), today, p.lower()

    def refresh_delivery(self):
        if not self._configured():
            return
        self.config_data["delivery_period"] = self.var_period.get()
        pid, _, _ = self._scope()
        start, end, _ = self._delivery_period()
        client = self.make_client()
        known = {k: v.get("updated_on") for k, v in self.hist_cache.items()}

        def work():
            scope = scope_params(None, False)
            open_issues, _ = client.issues_in({**scope, "status_id": "open"}, pid)
            recent, _ = client.issues_in({**scope, "status_id": "*", "updated_on": f">={start.isoformat()}"}, pid)
            merged = {it["id"]: it for it in recent + open_issues}
            need = [i for i, it in merged.items() if known.get(i) != it.get("updated_on")]
            fetched = fetch_journals(client, need)
            names = {}
            for one in pid:
                try:
                    names.update(client.members(one))
                except RedmineError:
                    pass
            return merged, {it["id"] for it in open_issues}, fetched, names

        def done(res):
            merged, open_ids, fetched, names = res
            for iid, issue in fetched.items():
                self.hist_cache[iid] = {"updated_on": issue.get("updated_on"), "issue": issue}
            self._del_raw = (merged, open_ids, names)
            self._recompute_delivery()
            self._populate_tree()

        self.run_async(work, done, busy="Reading ticket history for delivery tracking… (first time can take a minute)")

    def _recompute_delivery(self):
        merged, open_ids, names = self._del_raw
        names = dict(names)
        for c in self.hist_cache.values():
            for j in c["issue"].get("journals") or []:
                u = j.get("user") or {}
                if u.get("id"):
                    names.setdefault(u["id"], u.get("name", ""))
        for it in merged.values():
            for key in ("assigned_to", "author"):
                u = it.get(key) or {}
                if u.get("id"):
                    names.setdefault(u["id"], u.get("name", ""))
        hist = {i: history_info(self.hist_cache[i]["issue"], self.holidays())
                for i in merged if i in self.hist_cache}
        issues = [self.hist_cache[i]["issue"] if i in self.hist_cache else merged[i] for i in merged]
        for it in issues:  # keep list data (closed_on etc.) when the history copy lacks it
            it.setdefault("closed_on", merged[it["id"]].get("closed_on"))
        start, end, label = self._delivery_period()
        self.delivery = compute_delivery(issues, open_ids, hist, names, start, end, self.stuck_after())
        self.delivery["label"] = label
        self._render_delivery()

    def _render_delivery(self):
        d = self.delivery
        pname = self.var_project.get()
        self.lbl_del_scope.configure(text=f"Scope: {pname} · whole team · period: {d['label']}")
        self.lbl_del_summary.configure(text=d["summary"])
        self.ch_push.set_data([(r["name"], r["days"], C["warn"]) for r in d["devs"] if r["days"]][:8],
                              "No due dates were pushed in this period 🎉")
        stuck_by = Counter(x["dev"] for x in d["stuck"])
        self.ch_stuck.set_data([(n, c, C["bad"]) for n, c in stuck_by.most_common(8)],
                               "Nothing is stuck right now 🎉")
        for tv in (self.tv_devs, self.tv_stuck, self.tv_changes):
            tv.delete(*tv.get_children())
        for i, r in enumerate(d["devs"]):
            met = f"{pct(r['met'], r['closed_due'])}%  ({r['met']} of {r['closed_due']})" if r["closed_due"] else "—"
            avg = f"{r['days'] / r['pushes']:.1f}" if r["pushes"] else "—"
            self.tv_devs.insert("", "end", values=(r["name"], len(r["tickets"]), r["pushes"], r["days"], avg, met,
                                                   r["stuck"], r["open"]),
                                tags=("hot",) if r["days"] >= 7 or r["stuck"] >= 3 else ())
        for x in d["stuck"]:
            it = x["issue"]
            self.tv_stuck.insert("", "end", values=(x["dev"], f"#{it['id']}", it.get("subject", ""),
                                                    name_of(it.get("status")), fmt_date(x["since"]), x["days"],
                                                    it.get("due_date") or "—"),
                                 tags=("hot",) if x["days"] > 2 * d["stuck_after"] else ())
        for c in d["changes"]:
            it = c["issue"]
            self.tv_changes.insert("", "end", values=(f"#{it['id']}", it.get("subject", ""), c["dev"], c["by"],
                                                      fmt_date(c["date"], False), fmt_date(c["old"], False),
                                                      fmt_date(c["new"], False), f"+{c['days']}"),
                                   tags=("hot",) if c["days"] >= 7 else ())
        self.set_status("Delivery tracking updated.")

    def copy_delivery_summary(self):
        if not self.delivery:
            messagebox.showinfo(APP_NAME, "Click ↻ Refresh in Delivery Tracking first.", parent=self)
            return
        d = self.delivery
        lines = [f"Delivery update ({self.var_project.get()}, {d['label']}):", "", d["summary"], ""]
        if d["stuck"]:
            lines.append(f"Stuck for more than {d['stuck_after']} working days:")
            lines += [f"- #{x['issue']['id']} {x['issue'].get('subject', '')} ({name_of(x['issue'].get('status'))}, "
                      f"{x['days']} days, {x['dev']})" for x in d["stuck"][:15]]
        self.clipboard_clear()
        self.clipboard_append("\n".join(lines))
        self.set_status("Delivery summary copied. Paste it into email or Teams.")

    def export_delivery(self):
        if not self.delivery:
            messagebox.showinfo(APP_NAME, "Click ↻ Refresh in Delivery Tracking first.", parent=self)
            return
        self._ensure("openpyxl", "openpyxl")
        try:
            from openpyxl import Workbook
            from openpyxl.chart import BarChart, Reference
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        except ImportError:
            messagebox.showinfo(APP_NAME, "Excel export needs one more component.\n\nIn the VS Code terminal "
                                          "type:\n\n    python -m pip install openpyxl", parent=self)
            return
        d = self.delivery
        path = filedialog.asksaveasfilename(parent=self, title="Save delivery report",
                                            initialfile=f"Delivery_Report_{dt.date.today().isoformat()}.xlsx",
                                            defaultextension=".xlsx", filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        base = self.var_url.get().rstrip("/")
        navy = "102A43"
        thin = Side(style="thin", color="D9E2EC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        wb = Workbook()

        def sheet(ws, title, headers, widths, rows, link_col=None):
            ws["A1"] = title
            ws["A1"].font = Font(size=15, bold=True, color=navy)
            ws["A2"] = f"{self.var_project.get()} · {d['label']} · generated {d['generated'].strftime('%d %b %Y %H:%M')}"
            ws["A2"].font = Font(size=10, color="616E7C")
            for ci, (h, w) in enumerate(zip(headers, widths), 1):
                c = ws.cell(row=4, column=ci, value=h)
                c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=navy), border
                c.alignment = Alignment(wrap_text=True, vertical="center")
                ws.column_dimensions[c.column_letter].width = w
            for ri, row in enumerate(rows, 5):
                for ci, v in enumerate(row, 1):
                    c = ws.cell(row=ri, column=ci, value=v)
                    c.border, c.alignment = border, Alignment(wrap_text=True, vertical="top")
                if link_col:
                    c = ws.cell(row=ri, column=link_col)
                    c.hyperlink = f"{base}/issues/{c.value}"
                    c.font = Font(color="2F6FED", underline="single")
            ws.freeze_panes = "A5"
            if rows:
                ws.auto_filter.ref = f"A4:{ws.cell(row=4, column=len(headers)).column_letter}{4 + len(rows)}"
            return 4 + len(rows)

        ws = wb.active
        ws.title = "Summary"
        last = sheet(ws, "Delivery report: due dates & follow-ups",
                     ["Developer", "Tickets pushed", "Pushes", "Days pushed", "Avg days / push",
                      "Closed with due date", "Met original due date", "Met %", "Stuck now", "Open tickets"],
                     [24, 14, 10, 13, 15, 18, 20, 9, 11, 13],
                     [[r["name"], len(r["tickets"]), r["pushes"], r["days"],
                       round(r["days"] / r["pushes"], 1) if r["pushes"] else None, r["closed_due"], r["met"],
                       (r["met"] / r["closed_due"]) if r["closed_due"] else None, r["stuck"], r["open"]]
                      for r in d["devs"]])
        for ri in range(5, last + 1):
            ws.cell(row=ri, column=8).number_format = "0%"
        ws.cell(row=last + 2, column=1, value="Summary").font = Font(bold=True, color=navy)
        cell = ws.cell(row=last + 3, column=1, value=d["summary"])
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=last + 3, start_column=1, end_row=last + 3, end_column=10)
        ws.row_dimensions[last + 3].height = 60
        if d["devs"]:
            ch = BarChart()
            ch.type, ch.title, ch.style = "bar", "Days of due-date slippage by developer", 10
            ch.y_axis.title, ch.legend = "Days pushed", None
            data = Reference(ws, min_col=4, min_row=4, max_row=last)
            cats = Reference(ws, min_col=1, min_row=5, max_row=last)
            ch.add_data(data, titles_from_data=True)
            ch.set_categories(cats)
            ch.height, ch.width = 8, 18
            ws.add_chart(ch, f"A{last + 5}")
        sheet(wb.create_sheet("Stuck tickets"), f"Stuck for more than {d['stuck_after']} working days: follow up",
              ["Ticket", "Subject", "Developer", "Status", "In status since", "Working days", "Due date"],
              [10, 55, 22, 18, 16, 14, 13],
              [[x["issue"]["id"], x["issue"].get("subject", ""), x["dev"], name_of(x["issue"].get("status")),
                x["since"], x["days"], parse_date(x["issue"].get("due_date"))] for x in d["stuck"]], link_col=1)
        sheet(wb.create_sheet("Due-date pushes"), "Every due-date push in the period",
              ["Ticket", "Subject", "Developer", "Changed by", "Changed on", "From", "To", "Days added"],
              [10, 50, 22, 22, 13, 13, 13, 12],
              [[c["issue"]["id"], c["issue"].get("subject", ""), c["dev"], c["by"], c["date"], c["old"], c["new"],
                c["days"]] for c in d["changes"]], link_col=1)
        for name in ("Stuck tickets", "Due-date pushes"):
            for row in wb[name].iter_rows(min_row=5):
                for c in row:
                    if isinstance(c.value, dt.date):
                        c.number_format = "DD-MMM-YYYY"
        try:
            wb.save(path)
        except OSError as exc:
            messagebox.showerror(APP_NAME, f"Could not save the file (is it open in Excel?):\n{exc}", parent=self)
            return
        self.set_status(f"Delivery report saved to {path}")
        self._offer_open(path)

    # ----- Sprint planner ----------------------------------------------------
    def _build_sprint(self, parent):
        cfg = self._plans().get(self.config_data.get("sprint_version_selected", NO_RELEASE)) or {}
        today = dt.date.today()
        self.var_sp_prefix = tk.StringVar(value=cfg.get("prefix", "Sprint"))
        self.var_sp_first = tk.IntVar(value=int(cfg.get("first", 1)))
        default_start = parse_date(cfg.get("start")) or (today - dt.timedelta(days=today.weekday()))
        self.var_sp_start = tk.StringVar(value=default_start.strftime("%d-%m-%Y"))
        self.var_sp_mode = tk.StringVar(value=cfg.get("mode", "calendar"))
        self.var_sp_weeks = tk.IntVar(value=int(cfg.get("weeks", 2)))
        self.var_sp_wdays = tk.IntVar(value=int(cfg.get("wdays", 10)))
        self.var_sp_count = tk.IntVar(value=int(cfg.get("count", 3)))
        self.var_sp_regression = tk.BooleanVar(value=bool(cfg.get("regression", True)))
        self.var_sp_reg_extra = tk.IntVar(value=int(cfg.get("reg_extra", 0) or 0))
        self.var_sp_prod = tk.StringVar(value=(parse_date(cfg.get("prod_date")).strftime("%d-%m-%Y")
                                               if parse_date(cfg.get("prod_date")) else ""))
        self.var_hol_date = tk.StringVar()
        self.var_hol_name = tk.StringVar()

        top = ttk.Frame(parent, style="Bg.TFrame", padding=(0, 10, 0, 8))
        top.pack(fill="x")
        ttk.Label(top, text="Sprint Planner", font=FONTS["h1"], foreground=C["header"]).pack(anchor="w")
        ttk.Label(top, style="Muted.TLabel", text="Plan sprint dates around weekends and company holidays. "
                                                   "Sprint dates are also used to flag tickets due after their "
                                                   "sprint ends.").pack(anchor="w")
        body = ttk.Frame(parent, style="Bg.TFrame")
        body.pack(fill="both", expand=True)
        left = ttk.Frame(body, style="Bg.TFrame")
        left.pack(side="left", fill="y", padx=(0, 14))
        right = ttk.Frame(body, style="Bg.TFrame")
        right.pack(side="left", fill="both", expand=True)

        rl = ttk.LabelFrame(left, text="Release version (from Redmine)", padding=12)
        rl.pack(fill="x", pady=(0, 10))
        prow = ttk.Frame(rl)
        prow.pack(fill="x", pady=(0, 6))
        ttk.Label(prow, text="Project", style="Muted.TLabel").pack(side="left", padx=(0, 6))
        self.cb_sp_project = ttk.Combobox(prow, textvariable=self.var_sp_project, state="readonly", width=24)
        self.cb_sp_project.pack(side="left", fill="x", expand=True)
        self.cb_sp_project.bind("<<ComboboxSelected>>", lambda e: self.load_sprint_versions())
        ttk.Button(prow, text="↻", width=3, command=self.load_sprint_versions).pack(side="left", padx=(6, 0))
        self.cb_sp_version = ttk.Combobox(rl, textvariable=self.var_sp_version, state="readonly",
                                          values=[NO_RELEASE] + sorted(k for k in self._plans() if k != NO_RELEASE))
        self.cb_sp_version.pack(fill="x")
        self.cb_sp_version.bind("<<ComboboxSelected>>", lambda e: self._on_sp_version())
        self.lbl_sp_version = ttk.Label(rl, style="Muted.TLabel", wraplength=400, justify="left",
                                        text="Choose the release these sprints belong to. Each release keeps its "
                                             "own sprint plan.")
        self.lbl_sp_version.pack(anchor="w", pady=(6, 0))
        prod = ttk.Frame(rl)
        prod.pack(fill="x", pady=(8, 0))
        ttk.Label(prod, text="🚀 Production deployment").pack(side="left")
        ttk.Entry(prod, textvariable=self.var_sp_prod, width=12).pack(side="left", padx=6)
        ttk.Button(prod, text="Save", command=self.save_prod_date).pack(side="left")
        st = ttk.LabelFrame(left, text="Sprint settings", padding=12)
        st.pack(fill="x")
        ttk.Label(st, text="Sprint name").grid(row=0, column=0, sticky="w", pady=4)
        ttk.Entry(st, textvariable=self.var_sp_prefix, width=12).grid(row=0, column=1, sticky="w")
        ttk.Label(st, text="first number").grid(row=0, column=2, sticky="e", padx=(10, 4))
        ttk.Spinbox(st, from_=1, to=999, width=5, textvariable=self.var_sp_first).grid(row=0, column=3, sticky="w")
        ttk.Label(st, text="First sprint starts").grid(row=1, column=0, sticky="w", pady=4)
        ttk.Entry(st, textvariable=self.var_sp_start, width=12).grid(row=1, column=1, sticky="w")
        ttk.Label(st, text="DD-MM-YYYY", style="Muted.TLabel").grid(row=1, column=2, columnspan=2, sticky="w",
                                                                    padx=(10, 0))
        ttk.Label(st, text="Sprint length").grid(row=2, column=0, sticky="nw", pady=(8, 2))
        r1 = ttk.Frame(st)
        r1.grid(row=2, column=1, columnspan=3, sticky="w", pady=(8, 2))
        ttk.Radiobutton(r1, text="", value="calendar", variable=self.var_sp_mode).pack(side="left")
        ttk.Spinbox(r1, from_=1, to=8, width=3, textvariable=self.var_sp_weeks).pack(side="left")
        ttk.Label(r1, text=" calendar weeks").pack(side="left")
        ttk.Label(st, style="Muted.TLabel", text="Fixed cadence; holidays reduce the working days.").grid(
            row=3, column=1, columnspan=3, sticky="w")
        r2 = ttk.Frame(st)
        r2.grid(row=4, column=1, columnspan=3, sticky="w", pady=(8, 2))
        ttk.Radiobutton(r2, text="", value="working", variable=self.var_sp_mode).pack(side="left")
        ttk.Spinbox(r2, from_=1, to=40, width=3, textvariable=self.var_sp_wdays).pack(side="left")
        ttk.Label(r2, text=" working days").pack(side="left")
        ttk.Label(st, style="Muted.TLabel", text="Fixed effort; holidays push the end date out.").grid(
            row=5, column=1, columnspan=3, sticky="w")
        ttk.Label(st, text="Number of sprints").grid(row=6, column=0, sticky="w", pady=(10, 4))
        ttk.Spinbox(st, from_=0, to=52, width=5, textvariable=self.var_sp_count).grid(row=6, column=1, sticky="w",
                                                                                     pady=(10, 4))
        rg = ttk.Frame(st)
        rg.grid(row=7, column=0, columnspan=4, sticky="w", pady=(8, 0))
        ttk.Checkbutton(rg, text="Last sprint = regression, +", variable=self.var_sp_regression).pack(side="left")
        ttk.Spinbox(rg, from_=0, to=20, width=3, textvariable=self.var_sp_reg_extra).pack(side="left", padx=4)
        ttk.Label(rg, text="extra days").pack(side="left")
        ttk.Button(st, text="Calculate sprint dates", style="Accent.TButton",
                   command=self.calculate_sprints).grid(row=8, column=0, columnspan=4, sticky="ew", pady=(10, 0))

        hl = ttk.LabelFrame(left, text="Company holidays", padding=12)
        hl.pack(fill="both", expand=True, pady=(12, 0))
        hf, self.tv_hol = self._make_table(hl, [("date", "Date", 100), ("day", "Day", 90),
                                                ("occasion", "Occasion", 190)], height=9)
        hf.pack(fill="both", expand=True)
        add = ttk.Frame(hl)
        add.pack(fill="x", pady=(8, 0))
        ttk.Entry(add, textvariable=self.var_hol_date, width=12).pack(side="left")
        ttk.Entry(add, textvariable=self.var_hol_name).pack(side="left", fill="x", expand=True, padx=6)
        ttk.Button(add, text="Add", command=self.add_holiday).pack(side="left")
        ttk.Label(hl, text="Date (DD-MM-YYYY) and occasion, then Add.", style="Muted.TLabel").pack(anchor="w")
        btns = ttk.Frame(hl)
        btns.pack(fill="x", pady=(6, 0))
        ttk.Button(btns, text="Remove selected", command=self.remove_holiday).pack(side="left")
        ttk.Button(btns, text="Reset to 2026 defaults", command=self.reset_holidays).pack(side="right")

        self.sp_banner = tk.Frame(right, bg=C["soft"], highlightbackground=C["soft_line"], highlightthickness=1)
        self.sp_banner.pack(fill="x")
        self.lbl_sp_now = tk.Label(self.sp_banner, text="Set the first sprint's start date and click Calculate.",
                                   bg=C["soft"], fg=C["ink"], font=FONTS["h3"], justify="left", anchor="w")
        self.lbl_sp_now.pack(fill="x", padx=14, pady=12)
        self.sp_banner.bind("<Configure>", lambda e: self.lbl_sp_now.configure(wraplength=max(300, e.width - 30)))
        rf, self.tv_sprints = self._make_table(right, [("sprint", "Sprint", 110), ("kind", "Type", 100),
                                                       ("start", "Starts", 138), ("end", "Ends", 138),
                                                       ("working", "Days", 50),
                                                       ("state", "Status", 150),
                                                       ("holidays", "Holidays / notes / comments", 300)], height=8)
        rf.pack(fill="both", expand=True, pady=(10, 0))
        self.tv_sprints.configure(show="tree headings")
        self.tv_sprints.column("#0", width=28, minwidth=28, stretch=False)
        self.tv_sprints.tag_configure("week", foreground=C["ink2"])
        self.tv_sprints.bind("<Double-1>", lambda e: self._sprint_row_double())
        xb = ttk.Frame(right, style="Bg.TFrame")
        xb.pack(fill="x", pady=(6, 0))
        ttk.Button(xb, text="➕ Add patch / one-off sprint…", style="Accent.TButton",
                   command=self.add_extra_sprint).pack(side="left")
        ttk.Button(xb, text="➕ Add week…", command=self.week_dialog).pack(side="left", padx=(6, 0))
        ttk.Button(xb, text="✏ Edit week…", command=lambda: self.week_dialog(edit=True)).pack(side="left", padx=(6, 0))
        ttk.Button(xb, text="🗑 Remove selected", command=self.remove_extra_sprint).pack(side="left", padx=6)
        ttk.Button(xb, text="↺ Reopen", command=self.reopen_sprint).pack(side="right")
        ttk.Button(xb, text="✔ Close selected sprint…", style="Accent.TButton",
                   command=self.close_sprint).pack(side="right", padx=6)
        ttk.Label(right, text="Saved release plans  (double-click to open)", font=FONTS["h3"],
                  foreground=C["header"]).pack(anchor="w", pady=(12, 4))
        pf, self.tv_plans = self._make_table(right, [("release", "Release", 200), ("sprints", "Sprints", 200),
                                                     ("from", "From", 130), ("to", "To", 130),
                                                     ("due", "Release due", 120), ("golive", "🚀 Production", 120),
                                                     ("check", "Check", 220)], height=5)
        pf.pack(fill="x")
        self.tv_plans.bind("<Double-1>", lambda e: self._open_saved_plan())
        pb = ttk.Frame(right, style="Bg.TFrame")
        pb.pack(fill="x", pady=(6, 0))
        ttk.Button(pb, text="🔗 Link selected plan to a release…", command=self.link_plan_to_release).pack(side="left")
        ttk.Button(pb, text="🗑 Delete selected plan", command=self.delete_sprint_plan).pack(side="left", padx=6)
        rb = ttk.Frame(right, style="Bg.TFrame")
        rb.pack(fill="x", pady=(8, 0))
        ttk.Button(rb, text="⬇ Export to Excel…", command=self.export_sprints).pack(side="right")
        ttk.Button(rb, text="Copy as text", command=self.copy_sprints).pack(side="right", padx=6)
        self._render_holidays()
        self.render_sprints()

    def _render_holidays(self):
        self.tv_hol.delete(*self.tv_hol.get_children())
        for i, (d, name) in enumerate(sorted(self.holidays().items())):
            self.tv_hol.insert("", "end", iid=d.isoformat(), values=(d.strftime("%d-%b-%Y"), d.strftime("%A"), name),
                               tags=("muted",) if d < dt.date.today() else ())

    def _save_holidays(self, hols: dict):
        self.config_data["holidays"] = [[d.isoformat(), n] for d, n in sorted(hols.items())]
        self._holiday_cache = None
        self._save_config_quiet()
        self._render_holidays()
        self.render_sprints()
        self._populate_tree()

    def add_holiday(self):
        d = parse_user_date(self.var_hol_date.get())
        if not d:
            messagebox.showinfo(APP_NAME, "Type the date like 25-12-2026, then the occasion, then click Add.",
                                parent=self)
            return
        hols = dict(self.holidays())
        hols[d] = self.var_hol_name.get().strip() or "Holiday"
        self.var_hol_date.set("")
        self.var_hol_name.set("")
        self._save_holidays(hols)

    def remove_holiday(self):
        sel = self.tv_hol.selection()
        if not sel:
            return
        hols = dict(self.holidays())
        for k in sel:
            hols.pop(parse_date(k), None)
        self._save_holidays(hols)

    def reset_holidays(self):
        if messagebox.askyesno(APP_NAME, "Replace the list with the default 2026 holiday calendar (India)?", parent=self):
            self._save_holidays({parse_date(d): n for d, n in DEFAULT_HOLIDAYS_2026})

    def calculate_sprints(self):
        start = parse_user_date(self.var_sp_start.get())
        try:
            zero = int(self.var_sp_count.get()) == 0
        except (tk.TclError, ValueError):
            zero = False
        if not start and not zero:
            messagebox.showinfo(APP_NAME, "Type the first sprint's start date like 19-10-2026.", parent=self)
            return
        start = start or dt.date.today()
        key = self.var_sp_version.get() or NO_RELEASE
        try:
            cfg = {"prefix": self.var_sp_prefix.get().strip() or "Sprint", "first": int(self.var_sp_first.get()),
                   "start": start.isoformat(), "mode": self.var_sp_mode.get(), "weeks": int(self.var_sp_weeks.get()),
                   "wdays": int(self.var_sp_wdays.get()), "count": int(self.var_sp_count.get()),
                   "regression": bool(self.var_sp_regression.get()), "reg_extra": int(self.var_sp_reg_extra.get())}
        except (tk.TclError, ValueError):
            messagebox.showinfo(APP_NAME, "Please use whole numbers for the sprint length and count.", parent=self)
            return
        info = self.sp_version_info.get(key) or {}
        old = self._plans().get(key) or {}
        cfg["release_due"] = info.get("due_date") or old.get("release_due")
        for k in ("extras", "closed", "prod_date"):
            if k in old:
                cfg[k] = old[k]
        self._plans()[key] = cfg
        self.config_data["sprint_version_selected"] = key
        self._save_config_quiet()
        self.render_sprints()
        self._populate_tree()
        self.set_status(f"Sprint dates calculated and saved under {key}.")

    def _on_sp_version(self):
        key = self.var_sp_version.get() or NO_RELEASE
        self.config_data["sprint_version_selected"] = key
        cfg = self._plans().get(key)
        if cfg:
            self.var_sp_prefix.set(cfg.get("prefix", "Sprint"))
            self.var_sp_first.set(int(cfg.get("first", 1)))
            if parse_date(cfg.get("start")):
                self.var_sp_start.set(parse_date(cfg["start"]).strftime("%d-%m-%Y"))
            self.var_sp_mode.set(cfg.get("mode", "calendar"))
            self.var_sp_weeks.set(int(cfg.get("weeks", 2)))
            self.var_sp_wdays.set(int(cfg.get("wdays", 10)))
            self.var_sp_count.set(int(cfg.get("count", 3)))
            self.var_sp_regression.set(bool(cfg.get("regression", True)))
            self.var_sp_reg_extra.set(int(cfg.get("reg_extra", 0) or 0))
        pd_ = parse_date((cfg or {}).get("prod_date"))
        self.var_sp_prod.set(pd_.strftime("%d-%m-%Y") if pd_ else "")
        if cfg:
            pass
        else:  # suggest: start the next working day after the latest saved plan
            ends = [sp["end"] for k, sp in self.all_sprints() if k != NO_RELEASE] or \
                [sp["end"] for _, sp in self.all_sprints()]
            if ends:
                nxt = next_working_day(max(ends) + dt.timedelta(days=1), self.holidays())
                self.var_sp_start.set(nxt.strftime("%d-%m-%Y"))
            self.var_sp_first.set(1)
            self.var_sp_count.set(3)
        self._save_config_quiet()
        self.render_sprints()

    def _selected_plan_key(self):
        sel = self.tv_plans.selection()
        if sel:
            return self.tv_plans.set(sel[0], "release")
        return self.var_sp_version.get() or NO_RELEASE

    def link_plan_to_release(self):
        key = self._selected_plan_key()
        if key not in self._plans():
            messagebox.showinfo(APP_NAME, "Select a saved plan in the 'Saved release plans' table first.", parent=self)
            return
        win = tk.Toplevel(self)
        win.title("Link sprint plan to a release")
        win.configure(bg=C["bg"])
        win.transient(self)
        frm = ttk.Frame(win, padding=16, style="Bg.TFrame")
        frm.pack(fill="both", expand=True)
        ttk.Label(frm, text=f"Move the sprints saved under “{key}” to this release:").pack(anchor="w")
        var = tk.StringVar()
        names = [n for n in self.sp_version_info if n != key]
        cb = ttk.Combobox(frm, textvariable=var, values=names, width=40)
        cb.pack(fill="x", pady=8)
        ttk.Label(frm, style="Muted.TLabel", wraplength=380, justify="left",
                  text="Pick a release from Redmine, or type its name if it isn't in the list.").pack(anchor="w")

        def ok():
            new = var.get().strip()
            if not new or new == NO_RELEASE:
                return
            plans = self._plans()
            if new in plans and not messagebox.askyesno(APP_NAME, f"“{new}” already has a sprint plan. Replace it?",
                                                        parent=win):
                return
            cfg = plans.pop(key)
            info = self.sp_version_info.get(new) or {}
            cfg["release_due"] = info.get("due_date") or cfg.get("release_due")
            plans[new] = cfg
            self.var_sp_version.set(new)
            self.config_data["sprint_version_selected"] = new
            vals = list(self.cb_sp_version.cget("values") or [])
            if new not in vals:
                self.cb_sp_version.configure(values=vals + [new])
            self._save_config_quiet()
            win.destroy()
            self._on_sp_version()
            self._populate_tree()
            self.set_status(f"Sprint plan linked to {new}.")

        bb = ttk.Frame(frm, style="Bg.TFrame")
        bb.pack(fill="x", pady=(12, 0))
        ttk.Button(bb, text="Cancel", command=win.destroy).pack(side="right")
        ttk.Button(bb, text="Link", style="Accent.TButton", command=ok).pack(side="right", padx=8)

    def delete_sprint_plan(self):
        key = self._selected_plan_key()
        if key not in self._plans():
            messagebox.showinfo(APP_NAME, "Select a saved plan in the 'Saved release plans' table first.", parent=self)
            return
        if not messagebox.askyesno(APP_NAME, f"Delete the sprint plan saved under “{key}”?\n\nThis only removes it "
                                             "from this app; nothing changes in Redmine.", parent=self):
            return
        self._plans().pop(key, None)
        self.config_data.pop("sprint_plan", None)
        self._save_config_quiet()
        self.render_sprints()
        self._populate_tree()
        self.set_status(f"Sprint plan for {key} deleted.")

    def _open_saved_plan(self):
        sel = self.tv_plans.selection()
        if sel:
            self.var_sp_version.set(self.tv_plans.set(sel[0], "release"))
            self._on_sp_version()

    def load_sprint_versions(self):
        if not self._configured(warn=False):
            return
        if self.var_sp_project.get() not in self.projects:
            first = (self.mp_project.selected or [""])[0]
            pick = first if first in self.projects else self.var_rn_project.get()
            self.var_sp_project.set(pick if pick in self.projects else "")
        pid = self.projects.get(self.var_sp_project.get())
        if not pid:
            self.lbl_sp_version.configure(text="Choose the project above to load its release versions.")
            return
        self.config_data["sprint_project"] = self.var_sp_project.get()
        client = self.make_client()

        def done(versions):
            versions.sort(key=lambda v: (v.get("status") != "open", v.get("due_date") or "9999", v.get("name", "")))
            self.sp_version_info = {v.get("name", ""): v for v in versions}
            names = [NO_RELEASE] + list(self.sp_version_info) + \
                [k for k in self._plans() if k not in self.sp_version_info and k != NO_RELEASE]
            self.cb_sp_version.configure(values=names)
            self.cb_retro_version.configure(values=[n for n in names if n != NO_RELEASE] +
                                            [k for k in (self.config_data.get("retros") or {}) if k not in names])
            self.render_sprints()

        self.run_async(lambda: client.versions(pid), done, busy="Loading release versions…",
                       on_error=lambda e: self.set_status("⚠ Could not load release versions."))

    def render_sprints(self):
        plan = self.sprint_plan()
        key = self.var_sp_version.get() or NO_RELEASE
        self.tv_sprints.delete(*self.tv_sprints.get_children())
        today = dt.date.today()
        for n, sp in enumerate(plan):
            if sp.get("closed"):
                state = "✔ Closed"
            else:
                state = "Ended, not closed" if sp["end"] < today else ("Current" if sp["start"] <= today else "Upcoming")
            hols = ", ".join(f"{d.strftime('%d-%b')} {nm}" for d, nm in sp["holidays"])
            comment = f"💬 {sp['closed']['comment']}" if sp.get("closed", {}).get("comment") else ""
            hols = "; ".join(x for x in (hols, sp.get("note", ""), comment) if x) or "—"
            iid = f"x{sp['extra_index']}" if "extra_index" in sp else f"r{n}"
            self.tv_sprints.insert("", "end", iid=iid, values=(sp["name"], sp.get("kind", "Sprint"),
                                                               fmt_date(sp["start"]), fmt_date(sp["end"]),
                                                               sp["working"], state, hols),
                                   open=True,
                                   tags=("current",) if state == "Current" else (("muted",) if state == "✔ Closed" else
                                                                                 (("hot",) if state.startswith("Ended")
                                                                                  else ())))
        hols_all = self.holidays()
        for i, w in enumerate((self._plans().get(key) or {}).get("week_list") or []):
            a, b = parse_date(w.get("start")), parse_date(w.get("end"))
            if not a or not b:
                continue
            parent_iid = next((k for k in self.tv_sprints.get_children("")
                               if self.tv_sprints.set(k, "sprint") == w.get("sprint")), "")
            wstate = "Done" if b < today else ("Current" if a <= today else "Upcoming")
            cut = parse_date(w.get("cutoff"))
            notes = [f"Dev cut-off {fmt_date(cut)}" if cut else "No dev cut-off"]
            if (w.get("goals") or "").strip():
                notes.append(f"🎯 {len([g for g in w['goals'].splitlines() if g.strip()])} goals")
            self.tv_sprints.insert(parent_iid, "end", iid=f"w{i}",
                                   values=(w.get("name") or "(unnamed week)", "Week", fmt_date(a), fmt_date(b),
                                           count_working_days(a, b, hols_all), wstate, "; ".join(notes)),
                                   tags=("current",) if wstate == "Current" else ("week",))
        info = self.sp_version_info.get(key) or {}
        due = parse_date(info.get("due_date") or (self._plans().get(key) or {}).get("release_due"))
        vtxt = f"Release due {fmt_date(due)} (from Redmine)." if due else ""
        if key == NO_RELEASE:
            vtxt = "Not linked to a release. Choose a release version to save these sprints under it."
        self.lbl_sp_version.configure(text=vtxt or "This release has no due date in Redmine.")
        # saved plans overview
        self.tv_plans.delete(*self.tv_plans.get_children())
        for k, cfg in sorted(self._plans().items(), key=lambda kv: (kv[1] or {}).get("start", "")):
            sprints = self._plan_from_cfg(cfg)
            if not sprints:
                continue
            kinfo = self.sp_version_info.get(k) or {}
            kdue = parse_date(kinfo.get("due_date") or (cfg or {}).get("release_due"))
            check = "—"
            if kdue:
                check = "⚠ Ends after due date" if sprints[-1]["end"] > kdue else "✔ Fits before due date"
            n_reg = sum(1 for x in sprints if "extra_index" not in x)
            n_x = len(sprints) - n_reg
            desc = " + ".join(t for t in (plural(n_reg, "sprint") if n_reg else "",
                                          plural(n_x, "one-off") if n_x else "") if t)
            kprod = parse_date((cfg or {}).get("prod_date"))
            if kprod and kprod < sprints[-1]["end"]:
                check = "⚠ Production before last sprint ends"
            self.tv_plans.insert("", "end", values=(k, desc,
                                                    fmt_date(sprints[0]["start"], False),
                                                    fmt_date(sprints[-1]["end"], False),
                                                    fmt_date(kdue, False) if kdue else "—",
                                                    fmt_date(kprod, False) if kprod else "—", check),
                                 tags=("current",) if k == key else (("hot",) if check.startswith("⚠") else ()))
        if not plan:
            prod = parse_date((self._plans().get(key) or {}).get("prod_date"))
            self.lbl_sp_now.configure(text=f"No sprints saved for {key} yet. Set the first sprint's start date and "
                                           "click Calculate." + (f"\n🚀 Production deployment: {fmt_date(prod)}"
                                                                 if prod else ""))
            return
        sp, day, left = sprint_position(plan, today, self.holidays())
        warn = ""
        if due and plan[-1]["end"] > due:
            warn = f"  ⚠ The last sprint ends after the release due date ({fmt_date(due)})."
        prod = parse_date((self._plans().get(key) or {}).get("prod_date"))
        if prod:
            hols = self.holidays()
            gap = count_working_days(today + dt.timedelta(days=1), prod, hols) if prod > today else 0
            warn += (f"\n🚀 Production deployment: {fmt_date(prod)}"
                     + (f" ({plural(gap, 'working day')} from today)" if prod > today else
                        (" (today)" if prod == today else " (done)")))
            if prod < plan[-1]["end"]:
                warn += f"  ⚠ This is before the last sprint ends ({fmt_date(plan[-1]['end'])})."
            if not is_working_day(prod, hols):
                warn += "  ⚠ It falls on a weekend or company holiday."
        if sp:
            extra = f" · holidays: {', '.join(n for _, n in sp['holidays'])}" if sp["holidays"] else ""
            self.lbl_sp_now.configure(text=f"{key}: today is working day {day} of {sp['working']} in {sp['name']} · "
                                           f"{plural(left, 'working day')} left · ends {fmt_date(sp['end'])}"
                                           f"{extra}{warn}")
        elif today < plan[0]["start"]:
            self.lbl_sp_now.configure(text=f"{key}: {plan[0]['name']} starts {fmt_date(plan[0]['start'])}; "
                                           f"last sprint ends {fmt_date(plan[-1]['end'])}.{warn}")
        else:
            self.lbl_sp_now.configure(text=f"{key}: all sprints are finished (ended {fmt_date(plan[-1]['end'])}).{warn}")

    def copy_sprints(self):
        plan = self.sprint_plan()
        if not plan:
            return
        key = self.var_sp_version.get() or NO_RELEASE
        lines = [f"{key}:"] + [f"{sp['name']} ({sp.get('kind', 'Sprint')}): {fmt_date(sp['start'])} → "
                               f"{fmt_date(sp['end'])} ({sp['working']} working days"
                               + (f"; holidays: {', '.join(n for _, n in sp['holidays'])}" if sp["holidays"] else "")
                               + ")" + (f" ✔ closed {sp['closed'].get('date', '')}: {sp['closed'].get('comment', '')}"
                                        if sp.get("closed") else "")
                               for sp in plan]
        prod = parse_date((self._plans().get(key) or {}).get("prod_date"))
        if prod:
            lines.append(f"Production deployment: {fmt_date(prod)}")
        self.clipboard_clear()
        self.clipboard_append("\n".join(lines))
        self.set_status("Sprint dates copied.")

    def export_sprints(self):
        plan = self.sprint_plan()
        if not plan:
            messagebox.showinfo(APP_NAME, "Calculate the sprint dates first.", parent=self)
            return
        self._ensure("openpyxl", "openpyxl")
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            messagebox.showinfo(APP_NAME, "Excel export needs: python -m pip install openpyxl", parent=self)
            return
        path = filedialog.asksaveasfilename(parent=self, title="Save sprint plan", initialfile="Sprint_Plan.xlsx",
                                            defaultextension=".xlsx", filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "Sprint plan"
        ws["A1"] = f"Sprint plan: {self.var_sp_version.get() or NO_RELEASE}"
        ws["A1"].font = Font(size=15, bold=True, color="102A43")
        prod = parse_date((self._plans().get(self.var_sp_version.get() or NO_RELEASE) or {}).get("prod_date"))
        ws["A2"] = f"Production deployment: {fmt_date(prod)}" if prod else "Production deployment: not set"
        ws["A2"].font = Font(size=11, bold=True, color="2F6FED")
        heads = ["Sprint", "Start", "End", "Working days", "Holidays / note", "Type", "Closed on", "Closing comments"]
        for ci, (h, w) in enumerate(zip(heads, [16, 16, 16, 14, 40, 16, 14, 60]), 1):
            c = ws.cell(row=3, column=ci, value=h)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="102A43")
            ws.column_dimensions[c.column_letter].width = w
        for ri, sp in enumerate(plan, 4):
            vals = [sp["name"], sp["start"], sp["end"], sp["working"],
                    "; ".join(x for x in (", ".join(f"{d.strftime('%d-%b')} {n}" for d, n in sp["holidays"]),
                                          sp.get("note", "")) if x), sp.get("kind", "Sprint"),
                    parse_date((sp.get("closed") or {}).get("date")), (sp.get("closed") or {}).get("comment", "")]
            for ci, v in enumerate(vals, 1):
                c = ws.cell(row=ri, column=ci, value=v)
                if isinstance(v, dt.date):
                    c.number_format = "DDD DD-MMM-YYYY"
        al = wb.create_sheet("All releases")
        al.append(["Release", "Sprint", "Start", "End", "Working days", "Type", "Closed", "Production"])
        for key, sp in self.all_sprints():
            al.append([key, sp["name"], sp["start"], sp["end"], sp["working"], sp.get("kind", "Sprint"),
                       "Yes" if sp.get("closed") else "", parse_date((self._plans().get(key) or {}).get("prod_date"))])
            al.cell(row=al.max_row, column=8).number_format = "DD-MMM-YYYY"
            for col in (3, 4):
                al.cell(row=al.max_row, column=col).number_format = "DD-MMM-YYYY"
        for col, w in zip("ABCDE", (22, 12, 14, 14, 13)):
            al.column_dimensions[col].width = w
        hs = wb.create_sheet("Holidays")
        hs.append(["Date", "Day", "Occasion"])
        for d, n in sorted(self.holidays().items()):
            hs.append([d, d.strftime("%A"), n])
            hs.cell(row=hs.max_row, column=1).number_format = "DD-MMM-YYYY"
        hs.column_dimensions["A"].width, hs.column_dimensions["B"].width, hs.column_dimensions["C"].width = 14, 12, 32
        try:
            wb.save(path)
        except OSError as exc:
            messagebox.showerror(APP_NAME, f"Could not save the file:\n{exc}", parent=self)
            return
        self._offer_open(path)

    # ----- Release notes -----------------------------------------------------
    def rn_settings(self) -> dict:
        out = dict(RN_DEFAULTS)
        out.update(self.config_data.get("rn_settings") or {})
        return out

    def _build_release(self, parent):
        top = ttk.Frame(parent, style="Bg.TFrame", padding=(0, 10, 0, 6))
        top.pack(fill="x")
        left = ttk.Frame(top, style="Bg.TFrame")
        left.pack(side="left")
        ttk.Label(left, text="Release Notes", font=FONTS["h1"], foreground=C["header"]).pack(anchor="w")
        ttk.Label(left, style="Muted.TLabel", text="Pick a target version, check which tickets go in, and generate "
                                                   "release notes in your own format.").pack(anchor="w")
        ttk.Button(top, text="⚙ Document settings…", command=self.rn_settings_dialog).pack(side="right")

        rfb = FilterBar(parent, cols=6, title="RELEASE")
        rfb.pack(fill="x", pady=(0, 8))
        rfb.grid_f.pack_forget()
        ctl = rfb.top
        rfb.label(ctl, "Project")
        self.cb_rn_project = ttk.Combobox(ctl, textvariable=self.var_rn_project, state="readonly", width=26)
        self.cb_rn_project.pack(side="left")
        self.cb_rn_project.bind("<<ComboboxSelected>>", lambda e: self.load_versions())
        rfb.label(ctl, "   Target version")
        self.cb_rn_version = ttk.Combobox(ctl, textvariable=self.var_rn_version, state="readonly", width=34)
        self.cb_rn_version.pack(side="left")
        self.cb_rn_version.bind("<<ComboboxSelected>>", lambda e: self.load_version_tickets())
        ttk.Button(ctl, text="↻ Reload", command=self.load_versions).pack(side="left", padx=8)
        ttk.Button(ctl, text="✅ Mark this release as done", style="Accent.TButton",
                   command=self.rn_mark_done).pack(side="right")

        self.rn_nb = ttk.Notebook(parent, style="Sub.TNotebook")
        self.rn_nb.pack(fill="both", expand=True)
        write = ttk.Frame(self.rn_nb, style="Bg.TFrame", padding=(0, 8, 0, 0))
        donef = ttk.Frame(self.rn_nb, style="Bg.TFrame", padding=(0, 8, 0, 0))
        self.rn_nb.add(write, text="  ✍ Write release notes  ")
        self.rn_nb.add(donef, text="  ✅ Done releases  ")
        self._build_rn_done(donef)
        parent = write

        pw = ttk.PanedWindow(parent, orient="horizontal")
        pw.pack(fill="both", expand=True)
        lf = ttk.Frame(pw, style="Bg.TFrame", padding=(0, 0, 10, 0))
        rf = ttk.LabelFrame(pw, text="Release notes (editable)", padding=10)
        pw.add(lf, weight=1)
        pw.add(rf, weight=1)
        ttk.Label(lf, text="Tickets in this version  (double-click to include / exclude)",
                  font=FONTS["h3"], foreground=C["header"]).pack(anchor="w", pady=(0, 6))
        tf, self.tv_rn = self._make_table(lf, [("inc", "In notes", 80), ("ticket", "Ticket", 70),
                                               ("tracker", "Tracker", 100), ("status", "Status", 130),
                                               ("section", "Section", 110), ("subject", "Subject", 300)], height=16)
        tf.pack(fill="both", expand=True)
        self.tv_rn.bind("<Double-1>", lambda e: self._toggle_rn())
        self.tv_rn.bind("<space>", lambda e: self._toggle_rn())
        self.lbl_rn_count = ttk.Label(lf, text="Choose a target version.", style="Muted.TLabel")
        self.lbl_rn_count.pack(anchor="w", pady=(6, 0))
        b = ttk.Frame(lf, style="Bg.TFrame")
        b.pack(fill="x", pady=(6, 0))
        ttk.Button(b, text="Only finished", command=lambda: self._rn_select("done")).pack(side="left")
        ttk.Button(b, text="All", command=lambda: self._rn_select("all")).pack(side="left", padx=6)
        ttk.Button(b, text="None", command=lambda: self._rn_select("none")).pack(side="left")
        ttk.Button(b, text="✨ Generate Release Notes", style="Accent.TButton",
                   command=self.generate_release_notes).pack(side="right")

        tb = ttk.Frame(rf)
        tb.pack(fill="x", pady=(0, 6))
        ttk.Button(tb, text="⬇ Export (Word)…", style="Accent.TButton", command=self.export_rn).pack(side="left")
        ttk.Button(tb, text="Markdown…", command=lambda: self.export_rn(".md")).pack(side="left", padx=6)
        ttk.Button(tb, text="Copy", command=self.copy_rn).pack(side="left")
        ttk.Button(tb, text="📋 Prompt for Claude", command=self.copy_rn_prompt).pack(side="left", padx=6)
        xf = ttk.Frame(rf)
        xf.pack(fill="both", expand=True)
        self.txt_rn = tk.Text(xf, wrap="word", undo=True, font=FONTS["mono"], relief="flat", bg=C["field"],
                              fg=C["ink"], insertbackground=C["ink"], highlightthickness=1,
                              highlightbackground=C["line"], highlightcolor=C["accent"], padx=12, pady=10)
        xsb = ttk.Scrollbar(xf, orient="vertical", command=self.txt_rn.yview)
        self.txt_rn.configure(yscrollcommand=xsb.set)
        self.txt_rn.pack(side="left", fill="both", expand=True)
        xsb.pack(side="right", fill="y")
        draft = self.config_data.get("rn_draft")
        if draft:
            self.txt_rn.insert("1.0", draft)

    def load_versions(self):
        if not self._configured(warn=False):
            return
        pid = self.projects.get(self.var_rn_project.get())
        if not pid:
            self.lbl_rn_count.configure(text="Choose a project first.")
            return
        client = self.make_client()

        def work():
            versions = client.versions(pid)
            statuses = client.statuses()
            return versions, statuses

        def done(res):
            versions, statuses = res
            self.rn_closed_status = {s.get("name") for s in statuses if s.get("is_closed")}
            versions.sort(key=lambda v: (v.get("status") != "open", v.get("due_date") or "9999", v.get("name", "")))
            self.rn_versions = {}
            for v in versions:
                label = v.get("name", "")
                extra = [v.get("status", "")] + ([f"due {v['due_date']}"] if v.get("due_date") else [])
                label = f"{label}  ({', '.join(x for x in extra if x)})"
                self.rn_versions[label] = v
            self.cb_rn_version.configure(values=list(self.rn_versions))
            if self.var_rn_version.get() not in self.rn_versions:
                self.var_rn_version.set("")
            self.lbl_rn_count.configure(text=f"{plural(len(versions), 'version')} found. Choose a target version.")
            self.render_rn_done()

        self.run_async(work, done, busy="Loading versions…")

    def load_version_tickets(self):
        v = self.rn_versions.get(self.var_rn_version.get())
        if not v:
            return
        client = self.make_client()

        def done(res):
            self.rn_issues, _ = res
            self.rn_include = {it["id"]: self._rn_default(it) for it in self.rn_issues}
            self._render_rn_list()

        self.run_async(lambda: client.issues({"fixed_version_id": v["id"], "status_id": "*", "sort": "id:asc"}),
                       done, busy=f"Loading tickets for {v.get('name')}…")

    def _rn_default(self, it):
        st = name_of(it.get("status"), "")
        return st in self.rn_closed_status or bool(DONE_STATUS_RE.search(st))

    def _render_rn_list(self):
        names = {"features": "New feature", "enhancements": "Enhancement", "bugs": "Bug fix", "security": "Security"}
        self.tv_rn.delete(*self.tv_rn.get_children())
        for it in self.rn_issues:
            inc = self.rn_include.get(it["id"])
            self.tv_rn.insert("", "end", iid=str(it["id"]),
                              values=("✔ yes" if inc else "✖ no", f"#{it['id']}", name_of(it.get("tracker")),
                                      name_of(it.get("status")), names[rn_section(it)], it.get("subject", "")),
                              tags=() if inc else ("muted",))
        n_in = sum(1 for v in self.rn_include.values() if v)
        self.lbl_rn_count.configure(text=f"{plural(len(self.rn_issues), 'ticket')} in this version · {n_in} in the "
                                         f"notes · {len(self.rn_issues) - n_in} left out (not finished)")

    def _toggle_rn(self):
        for k in self.tv_rn.selection():
            self.rn_include[int(k)] = not self.rn_include.get(int(k))
        sel = self.tv_rn.selection()
        self._render_rn_list()
        if sel:
            self.tv_rn.selection_set(sel)

    def _rn_select(self, mode):
        for it in self.rn_issues:
            self.rn_include[it["id"]] = True if mode == "all" else (False if mode == "none" else self._rn_default(it))
        self._render_rn_list()

    def _rn_version_label(self):
        v = self.rn_versions.get(self.var_rn_version.get()) or {}
        name = v.get("name", "")
        m = re.search(r"\d+(?:\.\d+)+", name)
        return m.group() if m else name

    def generate_release_notes(self):
        chosen = [it for it in self.rn_issues if self.rn_include.get(it["id"])]
        if not chosen:
            messagebox.showinfo(APP_NAME, "Choose a target version and include at least one ticket.", parent=self)
            return
        if self.txt_rn.get("1.0", "end-1c").strip() and not messagebox.askyesno(
                APP_NAME, "Replace the current release notes text with a freshly generated version?", parent=self):
            return
        text = build_release_notes(chosen, self._rn_version_label(), self.rn_settings())
        self.txt_rn.delete("1.0", "end")
        self.txt_rn.insert("1.0", text)
        self._stash_rn()
        self.set_status(f"Release notes generated from {plural(len(chosen), 'ticket')}. Edit freely before exporting.")

    def _stash_rn(self):
        self.config_data["rn_draft"] = self.txt_rn.get("1.0", "end-1c")
        self._save_config_quiet()

    def export_rn(self, ext=".docx"):
        text = self.txt_rn.get("1.0", "end-1c")
        if not text.strip():
            messagebox.showinfo(APP_NAME, "Generate the release notes first.", parent=self)
            return
        self._stash_rn()
        settings = self.rn_settings()
        name = re.sub(r"[^A-Za-z0-9.]+", "_", f"{settings.get('product', '')}_{self._rn_version_label()}").strip("_")
        types = [("Word document", "*.docx")] if ext == ".docx" else [("Markdown", "*.md"), ("Text file", "*.txt")]
        path = filedialog.asksaveasfilename(parent=self, title="Save release notes",
                                            initialfile=f"{name or 'Release'}_Release_Notes{ext}",
                                            defaultextension=ext, filetypes=types)
        if not path:
            return
        blocks = parse_doc(text, "rn")
        try:
            if path.lower().endswith(".docx"):
                if not self._ensure("docx", "python-docx"):
                    messagebox.showinfo(APP_NAME, "Word export needs one more component.\n\nIn the VS Code "
                                                  "terminal type:\n\n    python -m pip install python-docx",
                                        parent=self)
                    return
                rn_to_docx(blocks, path, settings)
            else:
                with open(path, "w", encoding="utf-8") as fh:
                    fh.write(doc_to_markdown(blocks) if path.lower().endswith(".md") else text)
        except Exception as exc:
            messagebox.showerror(APP_NAME, f"Could not save the file (is it open in Word?):\n{exc}", parent=self)
            return
        self._rn_record(how=f"Exported ({Path(path).suffix.lstrip('.').upper()})", path=path)
        self.set_status(f"Release notes saved to {path}")
        self._offer_open(path)

    # ----- done releases ------------------------------------------------------------
    def _build_rn_done(self, parent):
        head = ttk.Frame(parent, style="Bg.TFrame")
        head.pack(fill="x", pady=(0, 6))
        ttk.Label(head, text="Releases whose notes are done", font=FONTS["h2"], foreground=C["header"]).pack(
            side="left")
        ttk.Label(head, text="   Added when you export the notes or click '✅ Mark this release as done'. Releases "
                             "closed in Redmine are listed too.", style="Muted.TLabel").pack(side="left")
        tf, self.tv_rn_done = self._make_table(parent, [("release", "Release", 200), ("project", "Project", 180),
                                                        ("done", "Done on", 130), ("how", "How", 160),
                                                        ("tickets", "Tickets", 80), ("file", "File", 420)], height=14)
        tf.pack(fill="both", expand=True)
        self.tv_rn_done.bind("<Double-1>", lambda e: self._rn_done_open())
        bb = ttk.Frame(parent, style="Bg.TFrame")
        bb.pack(fill="x", pady=(6, 0))
        ttk.Button(bb, text="📂 Open the notes file", style="Accent.TButton", command=self._rn_done_open).pack(
            side="left")
        ttk.Button(bb, text="↺ Not done (remove from list)", command=self._rn_done_remove).pack(side="left", padx=6)
        self.render_rn_done()

    def _rn_record(self, how, path=""):
        v = self.rn_versions.get(self.var_rn_version.get()) or {}
        name = v.get("name") or self.var_rn_version.get().split("  (")[0]
        if not name:
            return False
        hist = self.config_data.setdefault("rn_history", {})
        key = f"{self.var_rn_project.get()}|{name}"
        rec = hist.get(key) or {}
        rec.update({"release": name, "project": self.var_rn_project.get(), "done": dt.date.today().isoformat(),
                    "how": how, "tickets": sum(1 for k, inc in (getattr(self, "rn_include", {}) or {}).items() if inc)
                    or rec.get("tickets", "")})
        if path:
            rec["file"] = path
        hist[key] = rec
        self._save_config_quiet()
        self.render_rn_done()
        return True

    def rn_mark_done(self):
        if not self.var_rn_version.get():
            messagebox.showinfo(APP_NAME, "Choose the target version first.", parent=self)
            return
        if self._rn_record(how="Marked as done"):
            self.rn_nb.select(1)
            self.set_status(f"{self.var_rn_version.get().split('  (')[0]} marked as done ✅")

    def render_rn_done(self):
        if not hasattr(self, "tv_rn_done"):
            return
        tv = self.tv_rn_done
        tv.delete(*tv.get_children())
        hist = self.config_data.get("rn_history") or {}
        rows = sorted(hist.items(), key=lambda kv: kv[1].get("done", ""), reverse=True)
        names = {r.get("release") for _k, r in rows}
        for k, r in rows:
            tv.insert("", "end", iid=f"h::{k}", values=(r.get("release", ""), r.get("project", ""),
                                                        fmt_date(parse_date(r.get("done")), False),
                                                        r.get("how", ""), r.get("tickets", ""), r.get("file", "")))
        for v in self.rn_versions.values():
            if v.get("status") == "closed" and v.get("name") not in names:
                tv.insert("", "end", iid=f"v::{v.get('id')}", values=(
                    v.get("name", ""), self.var_rn_project.get(), fmt_date(parse_date(v.get("due_date")), False)
                    if v.get("due_date") else "—", "Closed in Redmine", "", ""), tags=("muted",))
        if not tv.get_children():
            tv.insert("", "end", values=("Nothing yet", "", "", "Export notes or click ✅ Mark this release as done",
                                         "", ""), tags=("muted",))

    def _rn_done_open(self):
        sel = self.tv_rn_done.selection()
        if not sel:
            return
        path = self.tv_rn_done.set(sel[0], "file")
        if not path:
            messagebox.showinfo(APP_NAME, "There's no saved file for this release (it was only marked as done).",
                                parent=self)
            return
        if not Path(path).exists():
            messagebox.showinfo(APP_NAME, f"The file isn't there any more:\n{path}", parent=self)
            return
        try:
            if IS_WINDOWS:
                os.startfile(path)  # noqa
            else:
                webbrowser.open(Path(path).as_uri())
        except OSError as exc:
            messagebox.showerror(APP_NAME, str(exc), parent=self)

    def _rn_done_remove(self):
        sel = [k for k in self.tv_rn_done.selection() if k.startswith("h::")]
        if not sel:
            return
        hist = self.config_data.get("rn_history") or {}
        for k in sel:
            hist.pop(k[3:], None)
        self._save_config_quiet()
        self.render_rn_done()

    def copy_rn(self):
        self.clipboard_clear()
        self.clipboard_append(self.txt_rn.get("1.0", "end-1c"))
        self.set_status("Release notes copied.")

    def copy_rn_prompt(self):
        chosen = [it for it in self.rn_issues if self.rn_include.get(it["id"])]
        draft = self.txt_rn.get("1.0", "end-1c")
        if not chosen or not draft.strip():
            messagebox.showinfo(APP_NAME, "Generate the release notes first.", parent=self)
            return
        self.clipboard_clear()
        self.clipboard_append(rn_ai_prompt(chosen, self._rn_version_label(), self.rn_settings(), draft))
        messagebox.showinfo(APP_NAME, "A request with your draft and the ticket list is on your clipboard.\n\n"
                                      "Paste it into Claude (claude.ai), copy the polished answer back into this "
                                      "box, then Export to Word.", parent=self)

    def rn_settings_dialog(self):
        cfg = self.rn_settings()
        win = tk.Toplevel(self)
        win.title("Release notes: document settings")
        win.configure(bg=C["bg"])
        win.transient(self)
        win.geometry("640x640")
        frm = ttk.Frame(win, padding=16, style="Bg.TFrame")
        frm.pack(fill="both", expand=True)
        vars_ = {k: tk.StringVar(value=cfg.get(k, "")) for k in ("product", "subtitle", "website", "template",
                                                                  "known_issues", "upgrade")}
        rows = [("Product name (title)", "product"), ("Subtitle line", "subtitle"), ("Website (header)", "website"),
                ("Known issues text", "known_issues"), ("Upgrade / deployment notes", "upgrade")]
        for r, (label, key) in enumerate(rows):
            ttk.Label(frm, text=label).grid(row=r, column=0, sticky="w", pady=4)
            ttk.Entry(frm, textvariable=vars_[key]).grid(row=r, column=1, columnspan=2, sticky="ew", pady=4)
        r = len(rows)
        ttk.Label(frm, text="Component versions").grid(row=r, column=0, sticky="nw", pady=4)
        comp = tk.Text(frm, height=9, font=FONTS["body"], relief="flat", highlightthickness=1,
                       highlightbackground=C["line"], padx=6, pady=4)
        comp.grid(row=r, column=1, columnspan=2, sticky="nsew", pady=4)
        comp.insert("1.0", cfg.get("components", ""))
        ttk.Label(frm, style="Muted.TLabel", text="First line is the bold product line; lines starting with '- ' "
                                                  "become bullets.").grid(row=r + 1, column=1, columnspan=2, sticky="w")
        ttk.Label(frm, text="Word template (optional)").grid(row=r + 2, column=0, sticky="w", pady=(14, 4))
        ttk.Entry(frm, textvariable=vars_["template"]).grid(row=r + 2, column=1, sticky="ew", pady=(14, 4))

        def browse():
            path = filedialog.askopenfilename(parent=win, title="Choose your last release notes Word file",
                                              filetypes=[("Word document", "*.docx")])
            if path:
                vars_["template"].set(path)

        ttk.Button(frm, text="Browse…", command=browse).grid(row=r + 2, column=2, padx=(6, 0), pady=(14, 4))
        ttk.Label(frm, style="Muted.TLabel", wraplength=440, justify="left",
                  text="Choose your previous release notes (e.g. Product_7.2.0_Release_Notes.docx). The app reuses "
                       "its logo, header, footer, fonts and bullet style, and replaces the content.").grid(
            row=r + 3, column=1, columnspan=2, sticky="w")
        frm.columnconfigure(1, weight=1)
        frm.rowconfigure(r, weight=1)

        def save():
            new = {k: v.get().strip() for k, v in vars_.items()}
            new["components"] = comp.get("1.0", "end").strip()
            self.config_data["rn_settings"] = new
            self._save_config_quiet()
            win.destroy()
            self.set_status("Release notes settings saved.")

        bb = ttk.Frame(frm, style="Bg.TFrame")
        bb.grid(row=r + 4, column=0, columnspan=3, sticky="e", pady=(16, 0))
        ttk.Button(bb, text="Cancel", command=win.destroy).pack(side="right")
        ttk.Button(bb, text="Save", style="Accent.TButton", command=save).pack(side="right", padx=8)

    # ----- Retrospective ---------------------------------------------------------
    def _build_retro(self, parent):
        top = ttk.Frame(parent, style="Bg.TFrame", padding=(0, 10, 0, 6))
        top.pack(fill="x")
        left = ttk.Frame(top, style="Bg.TFrame")
        left.pack(side="left")
        ttk.Label(left, text="Retrospective", font=FONTS["h1"], foreground=C["header"]).pack(anchor="w")
        ttk.Label(left, style="Muted.TLabel",
                  text="1) Create the retro sheet for a release  2) share it with the team (anonymous)  3) import the "
                       "filled sheet  4) turn points into improvement actions and track them here.").pack(anchor="w")
        rfb = FilterBar(parent, cols=6, title="RELEASE")
        rfb.pack(fill="x", pady=(0, 8))
        rfb.grid_f.pack_forget()
        ctl = rfb.top
        known = sorted((self.config_data.get("retros") or {}).keys())
        self.cb_retro_version = ttk.Combobox(ctl, textvariable=self.var_retro_version, state="readonly", width=30,
                                             values=known)
        self.cb_retro_version.pack(side="left")
        self.cb_retro_version.bind("<<ComboboxSelected>>", lambda e: self.render_retro())
        ttk.Button(ctl, text="⬇ Create retro sheet (Excel)…", style="Accent.TButton",
                   command=self.create_retro_doc).pack(side="left", padx=(12, 6))
        ttk.Button(ctl, text="⬆ Import filled retro…", command=self.import_retro).pack(side="left")
        ttk.Button(ctl, text="🗑 Delete this retro", command=self.delete_retro).pack(side="right")
        ttk.Button(ctl, text="Blank template…", command=lambda: self.create_retro_doc(blank=True)).pack(side="left",
                                                                                                    padx=6)

        cols = tk.Frame(parent, bg=C["bg"])
        cols.pack(fill="x")
        boxes = []
        for title, spec in (("😊 What went well", [("text", "Point", 380)]),
                            ("🔧 What could have gone better", [("text", "Point", 380)]),
                            ("🌟 Kudos", [("text", "Kudos", 380)])):
            lf = ttk.LabelFrame(cols, text=title, padding=6)
            f, tv = self._make_table(lf, spec, height=7)
            f.pack(fill="both", expand=True)
            boxes.append((lf, tv))
        for i, (lf, _) in enumerate(boxes):
            lf.grid(row=0, column=i, sticky="nsew", padx=(0 if i == 0 else 4, 0 if i == 2 else 4))
            cols.columnconfigure(i, weight=1, uniform="rc")
        self.tv_well, self.tv_better, self.tv_kudos = (b[1] for b in boxes)

        banner = tk.Frame(parent, bg=C["sand_soft"], highlightbackground=C["sand_line"], highlightthickness=1)
        banner.pack(fill="x", pady=(8, 8))
        tk.Label(banner, text="KEEPS COMING UP (across all imported retros)", bg=C["sand_soft"], fg=C["sand_ink"],
                 font=FONTS["kpi_label"]).pack(anchor="w", padx=14, pady=(8, 0))
        self.lbl_themes = tk.Label(banner, text="Import a filled retro document to see recurring themes.",
                                   bg=C["sand_soft"], fg=C["ink"], font=FONTS["body"], justify="left", anchor="w")
        self.lbl_themes.pack(anchor="w", fill="x", padx=14, pady=(2, 10))
        banner.bind("<Configure>", lambda e: self.lbl_themes.configure(wraplength=max(300, e.width - 30)))

        head = ttk.Frame(parent, style="Bg.TFrame")
        head.pack(fill="x")
        ttk.Label(head, text="Improvement actions (all releases)  ·  double-click to change status",
                  font=FONTS["h3"], foreground=C["header"]).pack(side="left")
        ttk.Button(head, text="📋 Prompt for Claude", command=self.copy_retro_prompt).pack(side="right")
        ttk.Button(head, text="⬇ Improvement tracker (Excel)…", command=self.export_retro_tracker).pack(side="right",
                                                                                                      padx=6)
        ttk.Button(head, text="Delete selected", command=self.delete_retro_action).pack(side="right")
        ttk.Button(head, text="➕ Add action", style="Accent.TButton",
                   command=self.add_retro_action).pack(side="right", padx=6)
        af, self.tv_actions = self._make_table(parent, [("release", "Release", 150), ("action", "Action", 420),
                                                        ("owner", "Owner", 140), ("due", "Due", 110),
                                                        ("status", "Status", 110)], height=7)
        af.pack(fill="both", expand=True, pady=(6, 0))
        self.tv_actions.bind("<Double-1>", lambda e: self.cycle_retro_action())
        self.render_retro()

    def _retro_actions(self) -> list:
        return self.config_data.setdefault("retro_actions", [])

    def render_retro(self):
        retros = self.config_data.get("retros") or {}
        rel = self.var_retro_version.get()
        data = retros.get(rel) or {}
        for tv, key, fields in ((self.tv_well, "well", ("text",)), (self.tv_better, "better", ("text",)),
                                (self.tv_kudos, "kudos", ("text",))):
            tv.delete(*tv.get_children())
            for item in data.get(key, []):
                tv.insert("", "end", values=[item.get(f, "") for f in fields])
        themes = retro_themes(retros)
        if themes:
            self.lbl_themes.configure(text="   ".join(
                f"• {t} ({n}× in {', '.join(r)})" for t, n, r, _ in themes[:5]))
        else:
            self.lbl_themes.configure(text="Import a filled retro document to see recurring themes.")
        self.tv_actions.delete(*self.tv_actions.get_children())
        order = {s: i for i, s in enumerate(RETRO_STATUSES)}
        for i, a in sorted(enumerate(self._retro_actions()), key=lambda x: (order.get(x[1]["status"], 0), x[1]["release"])):
            self.tv_actions.insert("", "end", iid=str(i), values=(a["release"], a["action"], a.get("owner", ""),
                                                                  a.get("due", ""), a["status"]),
                                   tags=("muted",) if a["status"] == "Done" else (("hot",) if a["status"] == "Not started"
                                                                                  else ()))

    def cycle_retro_action(self):
        sel = self.tv_actions.selection()
        if not sel:
            return
        acts = self._retro_actions()
        for k in sel:
            a = acts[int(k)]
            a["status"] = RETRO_STATUSES[(RETRO_STATUSES.index(a["status"]) + 1) % len(RETRO_STATUSES)]
        self._save_config_quiet()
        self.render_retro()

    def delete_retro_action(self):
        sel = self.tv_actions.selection()
        if sel and messagebox.askyesno(APP_NAME, f"Delete {plural(len(sel), 'action')}?", parent=self):
            acts = self._retro_actions()
            for k in sorted((int(x) for x in sel), reverse=True):
                acts.pop(k)
            self._save_config_quiet()
            self.render_retro()

    def create_retro_doc(self, blank=False):
        self._ensure("openpyxl", "openpyxl")
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            messagebox.showinfo(APP_NAME, "Excel export needs: python -m pip install openpyxl", parent=self)
            return
        rel = "" if blank else self.var_retro_version.get()
        if not blank and not rel:
            messagebox.showinfo(APP_NAME, "Choose the release version first (or use 'Blank template').", parent=self)
            return
        name = re.sub(r"[^A-Za-z0-9.]+", "_", rel).strip("_")
        path = filedialog.asksaveasfilename(parent=self, title="Save retro sheet", defaultextension=".xlsx",
                                            initialfile=f"Retro_{name}.xlsx" if name else "Retro_Template.xlsx",
                                            filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        info = {"release": rel, "project": "" if blank or not self.mp_project.selected
                else ", ".join(self.mp_project.selected)}
        plan = self.sprint_plan(rel) if rel else []
        if plan:
            info["sprints"] = (f"{plan[0]['name']} – {plan[-1]['name']} ({fmt_date(plan[0]['start'], False)} – "
                               f"{fmt_date(plan[-1]['end'], False)})")
        try:
            build_retro_xlsx(path, info)
        except OSError as exc:
            messagebox.showerror(APP_NAME, f"Could not save the file (is it open in Excel?):\n{exc}", parent=self)
            return
        self.set_status(f"Retro sheet saved to {path}")
        self._offer_open(path)
    def rn_versions_by_name(self):
        return {v.get("name", ""): v for v in self.rn_versions.values()}

    def import_retro(self):
        path = filedialog.askopenfilename(parent=self, title="Choose the filled retro sheet",
                                          filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        try:
            data = parse_retro_xlsx(path)
        except Exception as exc:
            messagebox.showerror(APP_NAME, f"Could not read this file (is it still open in Excel?):\n{exc}",
                                 parent=self)
            return
        if not (data["well"] or data["better"] or data["kudos"]):
            messagebox.showinfo(APP_NAME, "No retro points found. Make sure the three column headings (What went "
                                          "well / What could have gone better / Kudos) are still there.", parent=self)
            return
        rel = self.var_retro_version.get() or data.get("release") or Path(path).stem
        if data.get("release") and self.var_retro_version.get() and data["release"] != self.var_retro_version.get():
            if messagebox.askyesno(APP_NAME, f"This sheet is for '{data['release']}', but '{rel}' is selected.\n\n"
                                             f"Import it under '{data['release']}' instead?", parent=self):
                rel = data["release"]
        retros = self.config_data.setdefault("retros", {})
        retros[rel] = {"file": path, "imported": dt.date.today().isoformat(),
                       "well": data["well"], "better": data["better"], "kudos": data["kudos"]}
        self._save_config_quiet()
        vals = list(self.cb_retro_version.cget("values") or [])
        if rel not in vals:
            self.cb_retro_version.configure(values=vals + [rel])
        self.var_retro_version.set(rel)
        self.render_retro()
        messagebox.showinfo(APP_NAME, f"Imported the retro for {rel}:\n\n"
                                      f"• {plural(len(data['well']), 'point')} that went well\n"
                                      f"• {plural(len(data['better']), 'point')} that could have gone better\n"
                                      f"• {plural(len(data['kudos']), 'kudos', 'kudos')}\n\n"
                                      "Select a 'could have gone better' point and click ➕ Add action to turn it "
                                      "into an improvement.", parent=self)

    def delete_retro(self):
        rel = self.var_retro_version.get()
        retros = self.config_data.get("retros") or {}
        if rel not in retros:
            messagebox.showinfo(APP_NAME, "Choose a release whose retro you imported.", parent=self)
            return
        n_act = sum(1 for a in self._retro_actions() if a["release"] == rel)
        also = messagebox.askyesnocancel(
            APP_NAME, f"Delete the imported retro for “{rel}”?\n\nYes = delete the retro and its {plural(n_act, 'action')}"
                      "\nNo = delete the retro but keep its actions\nCancel = keep everything", parent=self)
        if also is None:
            return
        retros.pop(rel, None)
        if also:
            self.config_data["retro_actions"] = [a for a in self._retro_actions() if a["release"] != rel]
        self._save_config_quiet()
        self.render_retro()
        self.set_status(f"Retro for {rel} deleted.")

    def add_retro_action(self):
        from tkinter import simpledialog
        rel = self.var_retro_version.get() or NO_RELEASE
        sel = self.tv_better.selection()
        start = self.tv_better.set(sel[0], "text") if sel else ""
        action = simpledialog.askstring("Improvement action", "What will we do differently?", initialvalue=start,
                                        parent=self)
        if not action or not action.strip():
            return
        owner = simpledialog.askstring("Improvement action", "Owner (optional):", parent=self) or ""
        due = simpledialog.askstring("Improvement action", "Due date (optional, e.g. 30-11-2026):", parent=self) or ""
        self._retro_actions().append({"release": rel, "action": action.strip(), "owner": owner.strip(),
                                      "due": due.strip(), "status": "Not started"})
        self._save_config_quiet()
        self.render_retro()
    def export_retro_tracker(self):
        acts = self._retro_actions()
        retros = self.config_data.get("retros") or {}
        if not acts and not retros:
            messagebox.showinfo(APP_NAME, "Import a filled retro document first.", parent=self)
            return
        self._ensure("openpyxl", "openpyxl")
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            messagebox.showinfo(APP_NAME, "Excel export needs: python -m pip install openpyxl", parent=self)
            return
        path = filedialog.asksaveasfilename(parent=self, title="Save improvement tracker",
                                            initialfile=f"Retro_Improvements_{dt.date.today().isoformat()}.xlsx",
                                            defaultextension=".xlsx", filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        wb = Workbook()

        def sheet(ws, title, headers, widths, rows):
            ws["A1"] = title
            ws["A1"].font = Font(size=15, bold=True, color="102A43")
            for ci, (h, w) in enumerate(zip(headers, widths), 1):
                c = ws.cell(row=3, column=ci, value=h)
                c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="102A43")
                ws.column_dimensions[c.column_letter].width = w
            for ri, row in enumerate(rows, 4):
                for ci, v in enumerate(row, 1):
                    ws.cell(row=ri, column=ci, value=v).alignment = Alignment(wrap_text=True, vertical="top")
            ws.freeze_panes = "A4"

        ws = wb.active
        ws.title = "Actions"
        sheet(ws, "Retro improvement actions", ["Release", "Action", "Owner", "Due", "Status"], [20, 60, 18, 14, 14],
              [[a["release"], a["action"], a.get("owner", ""), a.get("due", ""), a["status"]] for a in acts])
        fills = {"Done": "DDF3E7", "In progress": "FCEFD8", "Not started": "FDE2E1"}
        for r in range(4, 4 + len(acts)):
            c = ws.cell(row=r, column=5)
            c.fill = PatternFill("solid", fgColor=fills.get(c.value, "FFFFFF"))
        sheet(wb.create_sheet("Themes"), "Themes that keep coming up", ["Theme", "Mentions", "Releases", "Examples"],
              [28, 11, 30, 80], [[t, n, ", ".join(r), " | ".join(e)] for t, n, r, e in retro_themes(retros)])
        rows = []
        for rel, d in retros.items():
            rows += [[rel, "Went well", i.get("text", "")] for i in d.get("well", [])]
            rows += [[rel, "Could have gone better", i.get("text", "")] for i in d.get("better", [])]
            rows += [[rel, "Kudos", i.get("text", "")] for i in d.get("kudos", [])]
        sheet(wb.create_sheet("All retro points"), "Everything the team said (anonymous)",
              ["Release", "Type", "Point"], [18, 24, 90], rows)
        try:
            wb.save(path)
        except OSError as exc:
            messagebox.showerror(APP_NAME, f"Could not save the file:\n{exc}", parent=self)
            return
        self._offer_open(path)

    def copy_retro_prompt(self):
        retros = self.config_data.get("retros") or {}
        rel = self.var_retro_version.get()
        if rel not in retros:
            messagebox.showinfo(APP_NAME, "Choose a release whose filled retro you have imported.", parent=self)
            return
        open_actions = [a for a in self._retro_actions() if a["status"] != "Done"]
        self.clipboard_clear()
        self.clipboard_append(retro_ai_prompt(rel, retros[rel], retro_themes(retros), open_actions))
        messagebox.showinfo(APP_NAME, "A request with this retro is on your clipboard. Paste it into Claude "
                                      "(claude.ai) to get a summary, the top improvements and a thank-you note.",
                            parent=self)

    # ----- groom dates, due-date range, weeks ------------------------------------
    def _selected_ticket_ids(self):
        ids = [i for i in self.tree.selection() if i.isdigit() and int(i) in self.issue_index]
        if not ids and self.current_issue:
            ids = [str(self.current_issue["id"])]
        return ids

    def set_groom_date(self, clear=False, ids=None):
        ids = ids or self._selected_ticket_ids()
        if not ids:
            messagebox.showinfo(APP_NAME, "Select one or more tickets in the list first.", parent=self)
            return
        grooms = self.config_data.setdefault("groom_dates", {})
        if clear:
            for i in ids:
                grooms.pop(i, None)
        else:
            from tkinter import simpledialog
            current = parse_date(grooms.get(ids[0]))
            if not current:  # suggest the next grooming Monday
                current = dt.date.today()
                while current.weekday() != 0:
                    current += dt.timedelta(days=1)
            txt = simpledialog.askstring("Groom date", f"Groom date for {plural(len(ids), 'ticket')} (DD-MM-YYYY):",
                                         initialvalue=current.strftime("%d-%m-%Y"), parent=self)
            if txt is None:
                return
            d = parse_user_date(txt)
            if not d:
                messagebox.showinfo(APP_NAME, "Type the date like 26-10-2026.", parent=self)
                return
            for i in ids:
                grooms[i] = d.isoformat()
        self._save_config_quiet()
        self._ticket_meta_changed(f"Groom date {'cleared' if clear else 'set'} for {plural(len(ids), 'ticket')}.")

    def _refresh_groom_choices(self):
        dates = sorted({d for d in (parse_date(v) for v in (self.config_data.get("groom_dates") or {}).values()) if d})
        self.groom_filter_dates = {fmt_date(d): d for d in dates}
        values = ["Any", "Not set", "Set", "✅ Groomed", "Not groomed yet"] + list(self.groom_filter_dates)
        self.cb_groom.configure(values=values)
        if self.var_groom_filter.get() not in values:
            self.var_groom_filter.set("Any")

    def _due_filter_changed(self):
        for v in (self.var_due_from, self.var_due_to):
            txt = v.get().strip()
            if txt and not parse_user_date(txt):
                return  # wait until the date is complete
        self._populate_tree()

    def quick_due_range(self, mode):
        today = dt.date.today()
        a = b = None
        if mode == "week":
            w = self._week_for_date(today)
            if w:
                a, b = w["start"], w["end"]
            else:  # Tuesday to Monday around today
                a = today - dt.timedelta(days=(today.weekday() - 1) % 7)
                b = a + dt.timedelta(days=6)
        elif mode == "sprint":
            cur = self._current_sprint()
            if not cur:
                messagebox.showinfo(APP_NAME, "No sprint in the Sprint Planner covers today.", parent=self)
                return
            a, b = cur[1]["start"], cur[1]["end"]
        self.var_due_from.set(a.strftime("%d-%m-%Y") if a else "")
        self.var_due_to.set(b.strftime("%d-%m-%Y") if b else "")
        self._populate_tree()

    def all_weeks(self) -> list:
        c = getattr(self, "_weeks_cache", None)
        if c and (dt.datetime.now() - c[0]).total_seconds() < 2:
            return c[1]
        out = self._all_weeks()
        self._weeks_cache = (dt.datetime.now(), out)
        return out

    def _all_weeks(self) -> list:
        out = []
        for key, cfg in self._plans().items():
            for i, w in enumerate((cfg or {}).get("week_list") or []):
                a, b = parse_date(w.get("start")), parse_date(w.get("end"))
                if a and b:
                    out.append({**w, "start": a, "end": b, "cutoff": parse_date(w.get("cutoff")),
                                "release": key, "index": i})
        out.sort(key=lambda w: (w["start"], w["end"]))
        return out

    def _week_for_date(self, d):
        if not d:
            return None
        for w in self.all_weeks():
            if w["start"] <= d <= w["end"]:
                return w
        return None

    # ----- sprint week per ticket (from the due date, or picked by you) -----------------
    @staticmethod
    def _week_id(w):
        start = w["start"] if isinstance(w["start"], dt.date) else parse_date(w["start"])
        return f"{w.get('release', '')}|{start.isoformat()}"

    @staticmethod
    def _week_short(w):
        return f"{w.get('name') or 'Week'} ({w['start'].strftime('%d %b')}–{w['end'].strftime('%d %b')})"

    def week_of(self, issue):
        """(week, how): how is 'manual' when you picked the week, 'due' when it comes from the due date."""
        ref = (self.config_data.get("ticket_weeks") or {}).get(str(issue.get("id")))
        if ref:
            w = next((w for w in self.all_weeks() if self._week_id(w) == ref), None)
            if w:
                return w, "manual"
        w = self._week_for_date(parse_date(issue.get("due_date")))
        return (w, "due") if w else (None, "")

    def week_text(self, issue, short=False):
        w, how = self.week_of(issue)
        if not w:
            return ""
        return (w.get("name") or "Week") if short else self._week_short(w) + (" 📌" if how == "manual" else "")

    def plan_info(self, issue):
        """Sprint and release for a ticket. Redmine's own values win (sprint tag, target version); when they're
        empty they are worked out from the sprint week / due date and the Sprint Planner (marked auto)."""
        tag_sprint = self._sprint(issue)
        red_rel = name_of(issue.get("fixed_version"), "")
        w, _ = self.week_of(issue)
        sprint, s_auto, rel_from_plan = tag_sprint, False, ""
        if w:
            rel_from_plan = w["release"] if w["release"] != NO_RELEASE else ""
        if not sprint:
            if w and w.get("sprint"):
                sprint, s_auto = w["sprint"], True
            else:
                d = (w["start"] if w else None) or parse_date(issue.get("due_date"))
                if d:
                    hits = [(k, sp) for k, sp in self.all_sprints() if sp["start"] <= d <= sp["end"]]
                    hits.sort(key=lambda h: h[0] != red_rel)
                    if hits:
                        sprint, s_auto = hits[0][1]["name"], True
                        rel_from_plan = rel_from_plan or (hits[0][0] if hits[0][0] != NO_RELEASE else "")
        release, r_auto = (red_rel, False) if red_rel else (rel_from_plan, bool(rel_from_plan))
        return {"sprint": sprint, "sprint_auto": s_auto, "release": release, "release_auto": r_auto,
                "sprint_txt": (sprint + (" · auto" if s_auto else "")) if sprint else "",
                "release_txt": (release + (" · auto" if r_auto else "")) if release else ""}

    def set_ticket_week(self, ids, week=None):
        """Pick a sprint week for tickets; week=None goes back to 'from the due date'."""
        store = self.config_data.setdefault("ticket_weeks", {})
        for i in ids:
            if week:
                store[str(i)] = self._week_id(week)
            else:
                store.pop(str(i), None)
        self._save_config_quiet()
        self._ticket_meta_changed(f"Sprint week {'set to ' + self._week_short(week) if week else 'set from due date'}"
                                  f" for {plural(len(ids), 'ticket')}.")

    def set_groom_dates(self, ids, d=None):
        grooms = self.config_data.setdefault("groom_dates", {})
        for i in ids:
            if d:
                grooms[str(i)] = d.isoformat()
            else:
                grooms.pop(str(i), None)
        self._save_config_quiet()
        self._ticket_meta_changed(f"Groom date {fmt_date(d) if d else 'cleared'} for {plural(len(ids), 'ticket')}.")

    def _ticket_meta_changed(self, msg):
        """Something you set locally (groom date / sprint week) changed: refresh every view that shows it."""
        self._weeks_cache = self._sprints_cache = None
        self._refresh_groom_choices()
        self._populate_tree()
        if self.current_issue:
            self._show_issue(self.current_issue)
        try:
            if self.nb.select() == str(self.tab_work) and self.work_nb.select() == str(self.tab_load):
                self.render_devload()

        except (tk.TclError, AttributeError):
            pass
        self._weekly_stale = True
        try:
            self.render_today()
        except Exception:
            pass
        self.set_status(msg)

    # ----- shortcuts & help ------------------------------------------------------
    def _focus_search(self):
        self._show_tickets()
        self.e_search.focus_set()
        self.e_search.select_range(0, "end")

    def refresh_current_tab(self):
        cur = self.nb.select()
        if cur == str(self.tab_dash):
            self.refresh_dashboard()
        elif cur == str(self.tab_delivery):
            self.refresh_delivery()
        elif cur == str(self.tab_work) and self.work_nb.select() == str(self.tab_week):
            self.refresh_weekly()
        elif cur == str(self.tab_team):
            self.refresh_team()
        elif cur == str(self.tab_rn):
            self.load_versions()
        else:
            self.refresh_tickets()

    HELP = {
        "load": ("Developer load", "One row per developer for the chosen period. Each #ticket chip sits on its due "
                                   "date (✅ = groomed). Red days = busy (up to their latest due date), green = free, "
                                   "grey = leave or holiday (leave comes from Kaizen Home's Team availability). The left "
                                   "card says until when they're busy, from when they're free and how many free days "
                                   "they have, so you can decide who to give the next ticket to. Tickets without a "
                                   "due date are listed at the end and are not counted. Click a person to list all "
                                   "their tickets below; double-click a ticket to open it.\n\nGrooming plan: pick "
                                   "the session (next Monday) in the GROOMING row at the top. Each person shows 🧭 "
                                   "n/3. Click a person, then ➕ Pick tickets… on the bottom right. That sets the "
                                   "ticket's groom date (shown everywhere); Redmine isn't changed. 📋 Copy plan "
                                   "gives a message for Teams / email; ▶ Start session runs it."),
        "today": ("Overview", "Your morning check in one place: what's due or overdue, stuck tickets, follow-ups, who is "
                           "being groomed next and who is leaving soon. Double-click a ticket to open it in the "
                           "Grooming Workspace. ▶ Start grooming session walks through the next session's tickets."),
        "work": ("Grooming Workspace", "Filter tickets, see PRD status and warnings, set groom dates (right-click), add "
                                       "your private notes and follow-ups (📝), get 💡 developer suggestions, and "
                                       "generate a PRD. ⭐ Views saves filter combinations. Export the grooming "
                                       "agenda to Excel."),
        "week": ("Weekly View", "Pick a week to see its goals, who is doing what (tickets due that week) and who is "
                                "being groomed. Weeks are created in the Sprint Planner (➕ Add week)."),
        "delivery": ("Delivery Tracking", "Due-date pushes per developer and tickets stuck in one status. Export a "
                                          "management report to Excel or copy a summary."),
        "sprint": ("Sprint Planner", "Choose the project and release, set the sprint dates, regression extra days, "
                                     "weeks, patches and the production date. Close sprints with comments."),
        "rn": ("Release Notes", "Pick a target version, tick which tickets go in, generate notes and export to Word "
                                "using your template (⚙ Document settings)."),
        "retro": ("Retrospective", "Create the anonymous retro sheet, share it, import it back, then turn points into "
                           "improvement actions."),
        "team": ("Team", "Your team list: designation, Active/Inactive, last working day and days left. 🧳 Handover "
                         "lists the open tickets of people who are leaving; copy a handover message or get 💡 "
                         "suggestions for new owners."),
        "dash": ("Management Dashboard", "Backlog health for management. Click a priority in the donut to list its "
                                         "tickets and set groom dates."),
    }

    def show_help(self):
        cur = self.nb.select()
        if cur == str(self.tab_work) and self.work_nb.select() in (str(self.tab_week), str(self.tab_load)):
            cur = self.work_nb.select()
        key = next((k for k, tab in (("today", self.tab_today), ("work", self.tab_work), ("week", self.tab_week),
                                     ("load", self.tab_load),
                                     ("delivery", self.tab_delivery), ("sprint", self.tab_sprint),
                                     ("rn", self.tab_rn), ("retro", self.tab_retro), ("team", self.tab_team),
                                     ("dash", self.tab_dash)) if str(tab) == cur), "work")
        title, text = self.HELP[key]
        mod = "Cmd" if IS_MAC else "Ctrl"
        messagebox.showinfo(f"Help: {title}", f"{text}\n\nShortcuts\n  {mod}+F   search tickets\n  F5        refresh "
                                              f"this tab\n  {mod}+G   set groom date for selected tickets\n  "
                                              f"{mod}+S   export the PRD\n\nYour data is saved in:\n{DATA_DIR}\n"
                                              "(with daily backups in the 'backups' folder)", parent=self)

    # ----- saved views -------------------------------------------------------------
    def _current_view(self) -> dict:
        return {"projects": list(self.mp_project.selected), "scope": self.var_scope.get(),
                "status": self.var_status_filter.get(), "sprint": self.var_sprint_filter.get(),
                "track": self.var_track_filter.get(), "releases": list(self.mp_release.selected),
                "assignee": self.var_assignee_filter.get(), "prd": self.var_prd_filter.get(),
                "check": self.var_check_filter.get(), "search": self.var_search.get(),
                "due_from": self.var_due_from.get(), "due_to": self.var_due_to.get(),
                "groom": self.var_groom_filter.get(), "priority": self.var_priority_filter.get(),
                "group": bool(self.var_group.get())}

    def _build_views_menu(self):
        m = self.menu_views
        m.delete(0, "end")
        m.add_command(label="💾  Save current view…", command=self.save_view)
        views = self.config_data.get("saved_views") or {}
        if views:
            m.add_separator()
            for name in sorted(views, key=str.lower):
                m.add_command(label=f"   {name}", command=lambda n=name: self.apply_view(n))
            m.add_separator()
            m.add_command(label="🗑  Delete a view…", command=self.delete_view)

    def save_view(self):
        from tkinter import simpledialog
        name = simpledialog.askstring("Save view", "Name for this view (e.g. '7.3.0 – New – needs PRD'):", parent=self)
        if name and name.strip():
            self.config_data.setdefault("saved_views", {})[name.strip()] = self._current_view()
            self._save_config_quiet()
            self.set_status(f"View '{name.strip()}' saved. Open it from ⭐ Views.")

    def apply_view(self, name):
        v = (self.config_data.get("saved_views") or {}).get(name)
        if not v:
            return
        projects_changed = list(v.get("projects", [])) != list(self.mp_project.selected) or \
            v.get("scope", "all") != self.var_scope.get()
        self.mp_project.set_selected(v.get("projects", []))
        self.var_scope.set(v.get("scope", "all"))
        for var, key, default in ((self.var_status_filter, "status", "All statuses"),
                                  (self.var_sprint_filter, "sprint", "All sprints"),
                                  (self.var_track_filter, "track", "All tracks"),
                                  (self.var_assignee_filter, "assignee", "All assignees"),
                                  (self.var_prd_filter, "prd", "Any"), (self.var_check_filter, "check", "Any"),
                                  (self.var_search, "search", ""), (self.var_due_from, "due_from", ""),
                                  (self.var_due_to, "due_to", ""), (self.var_groom_filter, "groom", "Any"),
                                  (self.var_priority_filter, "priority", "All priorities")):
            var.set(v.get(key, default))
        self.var_group.set(v.get("group", True))
        self.mp_release.set_selected(v.get("releases", []))
        self._show_tickets()
        if projects_changed:
            self._on_scope_changed()
        else:
            self._populate_tree()
        self.set_status(f"View '{name}' applied.")

    def delete_view(self):
        from tkinter import simpledialog
        views = self.config_data.get("saved_views") or {}
        name = simpledialog.askstring("Delete view", "Type the name of the view to delete:\n" + "\n".join(sorted(views)),
                                      parent=self)
        if name and name.strip() in views:
            views.pop(name.strip())
            self._save_config_quiet()
            self.set_status(f"View '{name.strip()}' deleted.")

    # ----- private notes & follow-ups ----------------------------------------------
    def edit_ticket_note(self):
        ids = self._selected_ticket_ids()
        if not ids:
            messagebox.showinfo(APP_NAME, "Select a ticket first.", parent=self)
            return
        key = ids[0]
        notes = self.config_data.setdefault("ticket_notes", {})
        cur = notes.get(key) or {}
        it = self.issue_index.get(int(key)) or {}
        win = tk.Toplevel(self)
        win.title(f"My note: #{key}")
        win.configure(bg=C["bg"])
        win.transient(self)
        win.geometry("520x380")
        frm = ttk.Frame(win, padding=16, style="Bg.TFrame")
        frm.pack(fill="both", expand=True)
        ttk.Label(frm, text=f"#{key} {truncate(it.get('subject', ''), 60)}", font=FONTS["h3"],
                  foreground=C["header"]).pack(anchor="w")
        ttk.Label(frm, text="Private note (only you can see it)", style="Muted.TLabel").pack(anchor="w", pady=(6, 2))
        txt = tk.Text(frm, height=8, wrap="word", font=FONTS["body"], relief="flat", highlightthickness=1,
                      highlightbackground=C["line"], padx=8, pady=6)
        txt.pack(fill="both", expand=True)
        txt.insert("1.0", cur.get("note", ""))
        row = ttk.Frame(frm, style="Bg.TFrame")
        row.pack(fill="x", pady=(8, 0))
        ttk.Label(row, text="Follow up on (optional)").pack(side="left")
        fu = parse_date(cur.get("follow_up"))
        v_fu = tk.StringVar(value=fu.strftime("%d-%m-%Y") if fu else "")
        ttk.Entry(row, textvariable=v_fu, width=12).pack(side="left", padx=8)
        ttk.Label(row, text="DD-MM-YYYY", style="Muted.TLabel").pack(side="left")

        def save(clear=False):
            if clear:
                notes.pop(key, None)
            else:
                d_txt = v_fu.get().strip()
                d = parse_user_date(d_txt) if d_txt else None
                if d_txt and not d:
                    messagebox.showinfo(APP_NAME, "Type the date like 23-10-2026.", parent=win)
                    return
                notes[key] = {"note": txt.get("1.0", "end").strip(), "follow_up": d.isoformat() if d else ""}
            self._save_config_quiet()
            win.destroy()
            self._populate_tree()
            if self.current_issue:
                self._show_issue(self.current_issue)
            self.set_status("Note saved." if not clear else "Note removed.")

        bb = ttk.Frame(frm, style="Bg.TFrame")
        bb.pack(fill="x", pady=(12, 0))
        ttk.Button(bb, text="Delete note", command=lambda: save(clear=True)).pack(side="left")
        ttk.Button(bb, text="Cancel", command=win.destroy).pack(side="right")
        ttk.Button(bb, text="Save", style="Accent.TButton", command=save).pack(side="right", padx=8)

    # ----- messages --------------------------------------------------------------
    def copy_ticket_message(self):
        ids = self._selected_ticket_ids()
        if not ids:
            messagebox.showinfo(APP_NAME, "Select a ticket first.", parent=self)
            return
        texts = []
        for i in ids:
            it = self.issue_index.get(int(i))
            if it:
                flags = [f for f in self._flags(it) if not f.startswith("Follow-up")]
                texts.append(nudge_message(it, flags, self.var_url.get()))
        self.clipboard_clear()
        self.clipboard_append("\n---\n\n".join(texts))
        self.set_status(f"{plural(len(texts), 'message')} copied. Paste into Teams, Slack or email (Ctrl+V).")

    # ----- team -------------------------------------------------------------------
    def team(self) -> list:
        t = self.config_data.get("team")
        if t is None:
            t = [{"name": n, "designation": d, "status": "Active", "relieving": r, "notes": ""}
                 for n, d, r in DEFAULT_TEAM]
            self.config_data["team"] = t
        return t

    def member_for(self, issue):
        nm = norm_name(name_of(issue.get("assigned_to"), ""))
        if not nm:
            return None
        for m in self.team():
            if norm_name(m.get("name")) == nm:
                return m
        return None

    def _open_for(self, name):
        pool = {it["id"]: it for it in self.issues}
        pool.update({it["id"]: it for it in self.team_open})
        nm = norm_name(name)
        return sorted((it for it in pool.values() if norm_name(name_of(it.get("assigned_to"), "")) == nm),
                      key=lambda it: (it.get("due_date") or "9999", it["id"]))

    def _build_team(self, parent):
        top = ttk.Frame(parent, style="Bg.TFrame", padding=(0, 10, 0, 6))
        top.pack(fill="x")
        left = ttk.Frame(top, style="Bg.TFrame")
        left.pack(side="left")
        ttk.Label(left, text="Team", font=FONTS["h1"], foreground=C["header"]).pack(anchor="w")
        ttk.Label(left, style="Muted.TLabel", text="Who is who, who is leaving and when. Double-click a person to "
                                                   "edit. Open-ticket counts come from Redmine (all projects).").pack(
            anchor="w")
        ttk.Button(top, text="↻ Refresh open tickets", command=self.refresh_team).pack(side="right")
        bar = ttk.Frame(parent, style="Bg.TFrame")
        bar.pack(fill="x", pady=(0, 6))
        ttk.Button(bar, text="➕ Add person", style="Accent.TButton", command=lambda: self.person_dialog()).pack(
            side="left")
        ttk.Button(bar, text="✏ Edit", command=lambda: self.person_dialog(edit=True)).pack(side="left", padx=6)
        ttk.Button(bar, text="⇄ Active / Inactive", command=self.toggle_person).pack(side="left")
        ttk.Button(bar, text="🗑 Delete", command=self.delete_person).pack(side="left", padx=6)
        tf, self.tv_team = self._make_table(parent, [("name", "Name", 180), ("designation", "Designation", 150),
                                                     ("status", "Status", 90), ("relieving", "Last working day", 150),
                                                     ("days", "Days left", 90), ("wdays", "Working days left", 130),
                                                     ("open", "Open tickets", 100), ("notes", "Notes", 300)],
                                            height=10)
        tf.pack(fill="x")
        self.tv_team.tag_configure("amber", foreground="#B07A2A")
        self.tv_team.tag_configure("red", foreground=C["bad"])
        self.tv_team.bind("<Double-1>", lambda e: self.person_dialog(edit=True))

        hl = ttk.LabelFrame(parent, text="🧳 Leaving soon: open tickets to hand over", padding=8)
        hl.pack(fill="both", expand=True, pady=(10, 0))
        hb = ttk.Frame(hl)
        hb.pack(fill="x", pady=(0, 6))
        ttk.Button(hb, text="✉ Copy handover message", style="Accent.TButton",
                   command=self.copy_handover_message).pack(side="left")
        ttk.Button(hb, text="💡 Suggest new owners", command=self.suggest_new_owners).pack(side="left", padx=6)
        ttk.Button(hb, text="⬇ Export handover (Excel)…", command=self.export_handover).pack(side="left")
        ttk.Label(hb, text="Click a person (or one of their tickets) first.", style="Muted.TLabel").pack(side="left",
                                                                                                      padx=10)
        f = ttk.Frame(hl)
        f.pack(fill="both", expand=True)
        cols = [("status", "Status", 140), ("due", "Due", 120), ("priority", "Priority", 100),
                ("release", "Release", 130), ("note", "Note", 300)]
        self.tv_hand = ttk.Treeview(f, columns=[c[0] for c in cols], show="tree headings", height=8)
        self.tv_hand.heading("#0", text="Person / ticket", anchor="w")
        self.tv_hand.column("#0", width=420, stretch=True)
        for k, t, wd in cols:
            self.tv_hand.heading(k, text=t, anchor="w")
            self.tv_hand.column(k, width=wd, stretch=k == "note", anchor="w")
        sb = ttk.Scrollbar(f, orient="vertical", command=self.tv_hand.yview)
        self.tv_hand.configure(yscrollcommand=sb.set)
        self.tv_hand.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        self.tv_hand.tag_configure("person", font=FONTS["small_bold"], background=C["head_bg"])
        self.tv_hand.tag_configure("hot", foreground=C["bad"])
        self.render_team()

    def render_team(self):
        today = dt.date.today()
        hols = self.holidays()
        self.tv_team.delete(*self.tv_team.get_children())
        for i, m in enumerate(self.team()):
            rel = parse_date(m.get("relieving"))
            days = wdays = ""
            tag = ()
            if rel:
                left = (rel - today).days
                days = f"{left} days" if left >= 0 else "left"
                wdays = str(count_working_days(today + dt.timedelta(days=1), rel, hols)) if left >= 0 else "—"
                tag = ("red",) if 0 <= left < 14 else (("amber",) if 0 <= left < 30 else ())
            if m.get("status") == "Inactive":
                tag = ("muted",)
            n_open = len(self._open_for(m["name"]))
            self.tv_team.insert("", "end", iid=f"p{i}", values=(m["name"], m.get("designation", ""),
                                                                m.get("status", "Active"),
                                                                fmt_date(rel) if rel else "—", days or "—",
                                                                wdays or "—", n_open, m.get("notes", "")), tags=tag)
        self.tv_hand.delete(*self.tv_hand.get_children())
        leaving = sorted((m for m in self.team() if m.get("status") == "Active" and parse_date(m.get("relieving"))
                          and parse_date(m.get("relieving")) >= today), key=lambda m: m["relieving"])
        for m in leaving:
            rel = parse_date(m["relieving"])
            tickets = self._open_for(m["name"])
            pid = self.tv_hand.insert("", "end", iid=f"h{norm_name(m['name'])}", open=True, tags=("person",),
                                      text=f"{m['name']}  ·  last day {fmt_date(rel)}  ·  {(rel - today).days} days "
                                           f"left  ·  {plural(len(tickets), 'open ticket')}")
            for it in tickets:
                due = parse_date(it.get("due_date"))
                late = bool(due and due > rel)
                self.tv_hand.insert(pid, "end", iid=f"t{m['name']}|{it['id']}",
                                    text=f"#{it['id']}   {it.get('subject', '')}",
                                    values=(name_of(it.get("status")), fmt_date(due, False) if due else "—",
                                            name_of(it.get("priority")), name_of(it.get("fixed_version"), "—"),
                                            "⚠ due after last day" if late else ""),
                                    tags=("hot",) if late else ())

    def refresh_team(self):
        if not self._configured(warn=False):
            return
        client = self.make_client()

        def done(res):
            self.team_open, _ = res
            self.render_team()
            self.render_today()
            self.set_status("Team open tickets updated.")

        self.run_async(lambda: client.issues({"status_id": "open", "sort": "due_date:asc"}), done,
                       busy="Loading everyone's open tickets…")

    def person_dialog(self, edit=False):
        team = self.team()
        idx = None
        if edit:
            sel = self.tv_team.selection()
            if not sel:
                messagebox.showinfo(APP_NAME, "Click a person in the list first.", parent=self)
                return
            idx = int(sel[0][1:])
        m = team[idx] if edit else {}
        win = tk.Toplevel(self)
        win.title("Edit person" if edit else "Add person")
        win.configure(bg=C["bg"])
        win.transient(self)
        frm = ttk.Frame(win, padding=16, style="Bg.TFrame")
        frm.pack(fill="both", expand=True)
        rel = parse_date(m.get("relieving"))
        v = {"name": tk.StringVar(value=m.get("name", "")), "designation": tk.StringVar(value=m.get("designation",
                                                                                                    "Developer")),
             "status": tk.StringVar(value=m.get("status", "Active")),
             "relieving": tk.StringVar(value=rel.strftime("%d-%m-%Y") if rel else ""),
             "notes": tk.StringVar(value=m.get("notes", ""))}
        rows = [("Name", "name", None), ("Designation", "designation", DESIGNATIONS),
                ("Status", "status", ["Active", "Inactive"]), ("Last working day (if on notice)", "relieving", None),
                ("Notes", "notes", None)]
        for r, (label, key, values) in enumerate(rows):
            ttk.Label(frm, text=label).grid(row=r, column=0, sticky="w", pady=4, padx=(0, 10))
            if values:
                ttk.Combobox(frm, textvariable=v[key], values=values, width=30,
                             state="readonly" if key == "status" else "normal").grid(row=r, column=1, sticky="ew")
            else:
                ttk.Entry(frm, textvariable=v[key], width=32).grid(row=r, column=1, sticky="ew")
        ttk.Label(frm, text="Dates like 28-10-2026 (with the year). Leave empty if not leaving.",
                  style="Muted.TLabel").grid(row=5, column=0, columnspan=2, sticky="w", pady=(6, 0))

        def ok():
            name = v["name"].get().strip()
            if not name:
                messagebox.showinfo(APP_NAME, "Please type a name.", parent=win)
                return
            rtxt = v["relieving"].get().strip()
            d = parse_user_date(rtxt) if rtxt else None
            if rtxt and not d:
                messagebox.showinfo(APP_NAME, "Type the date like 28-10-2026.", parent=win)
                return
            rec = {"name": name, "designation": v["designation"].get().strip(), "status": v["status"].get(),
                   "relieving": d.isoformat() if d else "", "notes": v["notes"].get().strip()}
            if edit:
                team[idx] = rec
            else:
                team.append(rec)
            self._save_config_quiet()
            win.destroy()
            self.render_team()
            self._populate_tree()

        bb = ttk.Frame(frm, style="Bg.TFrame")
        bb.grid(row=6, column=0, columnspan=2, sticky="e", pady=(14, 0))
        ttk.Button(bb, text="Cancel", command=win.destroy).pack(side="right")
        ttk.Button(bb, text="Save", style="Accent.TButton", command=ok).pack(side="right", padx=8)

    def toggle_person(self):
        team = self.team()
        for k in self.tv_team.selection():
            m = team[int(k[1:])]
            m["status"] = "Inactive" if m.get("status", "Active") == "Active" else "Active"
        self._save_config_quiet()
        self.render_team()
        self._populate_tree()

    def delete_person(self):
        sel = self.tv_team.selection()
        if not sel or not messagebox.askyesno(APP_NAME, f"Delete {plural(len(sel), 'person', 'people')} from the "
                                                        "list?", parent=self):
            return
        team = self.team()
        for k in sorted((int(x[1:]) for x in sel), reverse=True):
            team.pop(k)
        self._save_config_quiet()
        self.render_team()

    def _selected_leaver(self):
        sel = self.tv_hand.selection()
        if not sel:
            return None
        row = sel[0]
        if row.startswith("t"):
            row = self.tv_hand.parent(row)
        nm = row[1:]
        return next((m for m in self.team() if norm_name(m["name"]) == nm), None)

    def copy_handover_message(self):
        m = self._selected_leaver()
        if not m:
            messagebox.showinfo(APP_NAME, "Click a person in the handover list first.", parent=self)
            return
        text = handover_message(m["name"], parse_date(m["relieving"]), self._open_for(m["name"]), self.var_url.get())
        self.clipboard_clear()
        self.clipboard_append(text)
        self.set_status(f"Handover message for {m['name']} copied. Paste it into Teams or email.")

    def export_handover(self):
        self._ensure("openpyxl", "openpyxl")
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            return
        path = filedialog.asksaveasfilename(parent=self, title="Save handover list",
                                            initialfile=f"Handover_{dt.date.today().isoformat()}.xlsx",
                                            defaultextension=".xlsx", filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "Handover"
        ws["A1"] = "Handover: open tickets of people who are leaving"
        ws["A1"].font = Font(size=15, bold=True, color="4A3F3A")
        heads = ["Person", "Last working day", "Days left", "Ticket", "Subject", "Status", "Due date", "Priority",
                 "Release", "Due after last day?", "New owner", "Handover done?"]
        for ci, (h, wd) in enumerate(zip(heads, [20, 16, 10, 10, 55, 16, 13, 12, 16, 16, 18, 14]), 1):
            c = ws.cell(row=3, column=ci, value=h)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="C2665C")
            ws.column_dimensions[c.column_letter].width = wd
        r = 3
        today = dt.date.today()
        for m in sorted((m for m in self.team() if m.get("status") == "Active" and parse_date(m.get("relieving"))
                         and parse_date(m["relieving"]) >= today), key=lambda m: m["relieving"]):
            rel = parse_date(m["relieving"])
            for it in self._open_for(m["name"]) or [None]:
                r += 1
                due = parse_date(it.get("due_date")) if it else None
                vals = [m["name"], rel, (rel - today).days, it["id"] if it else "", it.get("subject", "") if it else
                        "(no open tickets)", name_of(it.get("status")) if it else "", due,
                        name_of(it.get("priority")) if it else "", name_of(it.get("fixed_version"), "") if it else "",
                        "Yes" if due and due > rel else "", "", ""]
                for ci, val in enumerate(vals, 1):
                    c = ws.cell(row=r, column=ci, value=val)
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                    if isinstance(val, dt.date):
                        c.number_format = "DD-MMM-YYYY"
                if it:
                    c = ws.cell(row=r, column=4)
                    c.hyperlink = f"{self.var_url.get().rstrip('/')}/issues/{it['id']}"
                    c.font = Font(color="C2665C", underline="single")
        ws.freeze_panes = "A4"
        try:
            wb.save(path)
        except OSError as exc:
            messagebox.showerror(APP_NAME, f"Could not save the file (is it open in Excel?):\n{exc}", parent=self)
            return
        self._offer_open(path)

    # ----- developer suggestions ---------------------------------------------------
    def _with_history(self, then):
        fresh = self.suggest_history and (dt.datetime.now() - self.suggest_history[0]).seconds < 1800
        if fresh or not self._configured(warn=False):
            then(self.suggest_history[1] if self.suggest_history else [])
            return
        pids, _, _ = self._scope()
        since = (dt.date.today() - dt.timedelta(days=180)).isoformat()
        client = self.make_client()

        def done(res):
            items, _ = res
            self.suggest_history = (dt.datetime.now(), items)
            then(items)

        self.run_async(lambda: client.issues_in({"status_id": "*", "updated_on": f">={since}"}, pids, 1500), done,
                       busy="Looking at who did similar work in the last 6 months…")

    def _show_suggestions(self, issues, exclude=(), title="💡 Suggested developers", min_days_left=10):
        def go(history):
            pool = {it["id"]: it for it in history}
            pool.update({it["id"]: it for it in self.issues})
            pushes = {norm_name(d["name"]): d["pushes"] for d in (self.delivery or {}).get("devs", [])}
            open_now = list({it["id"]: it for it in self.issues + self.team_open}.values())
            lines = []
            for it in issues:
                top, skipped = suggest_developers(it, self.team(), list(pool.values()), open_now, pushes, exclude,
                                                  min_days_left=min_days_left)
                cur = name_of(it.get("assigned_to"), "Unassigned")
                lines.append(f"#{it['id']} {it.get('subject', '')}\n   (currently: {cur})")
                if not top:
                    lines.append("   No available developer found.")
                for n, r in enumerate(top, 1):
                    lead = "  (lead, fallback)" if r["lead"] else ""
                    lines.append(f"   {n}. {r['name']}{lead}: " + "; ".join(r["reasons"]))
                if skipped:
                    lines.append("   Not suggested: " + ", ".join(skipped))
                lines.append("")
            win = tk.Toplevel(self)
            win.title(title)
            win.configure(bg=C["bg"])
            win.transient(self)
            win.geometry("760x480")
            frm = ttk.Frame(win, padding=14, style="Bg.TFrame")
            frm.pack(fill="both", expand=True)
            ttk.Label(frm, style="Muted.TLabel", wraplength=720, justify="left",
                      text="Suggestions only: based on similar past tickets, current workload and recent due-date "
                           "pushes. Inactive people and anyone leaving before the due date are left out. Nothing "
                           "is changed in Redmine.").pack(anchor="w")
            txt = tk.Text(frm, wrap="word", font=FONTS["body"], relief="flat", highlightthickness=1,
                          highlightbackground=C["line"], padx=10, pady=8)
            txt.pack(fill="both", expand=True, pady=8)
            txt.insert("1.0", "\n".join(lines))
            txt.configure(state="disabled")
            bb = ttk.Frame(frm, style="Bg.TFrame")
            bb.pack(fill="x")

            def copy():
                self.clipboard_clear()
                self.clipboard_append("\n".join(lines))
                self.set_status("Suggestions copied.")

            ttk.Button(bb, text="Copy", command=copy).pack(side="left")
            ttk.Button(bb, text="Close", style="Accent.TButton", command=win.destroy).pack(side="right")

        self._with_history(go)

    def suggest_for_selected(self):
        ids = self._selected_ticket_ids()
        issues = [self.issue_index[int(i)] for i in ids if int(i) in self.issue_index]
        if not issues:
            messagebox.showinfo(APP_NAME, "Select one or more tickets first.", parent=self)
            return
        self._show_suggestions(issues)

    def suggest_new_owners(self):
        m = self._selected_leaver()
        if not m:
            messagebox.showinfo(APP_NAME, "Click a person in the handover list first.", parent=self)
            return
        tickets = self._open_for(m["name"])
        if not tickets:
            messagebox.showinfo(APP_NAME, f"{m['name']} has no open tickets.", parent=self)
            return
        self._show_suggestions(tickets, exclude={norm_name(m["name"])},
                               title=f"💡 New owners for {m['name']}'s tickets", min_days_left=45)

    # ----- dashboard: filters ---------------------------------------------------
    def _build_dash_filters(self, parent):
        fb = FilterBar(parent, cols=8)
        fb.pack(fill="x", padx=(0, 12), pady=(0, 8))
        fb.label(fb.top, "Projects")
        self.mp_dash_project = MultiPick(fb.top, "All projects", self._dash_scope_changed, width=24, title="Projects")
        self.mp_dash_project.pack(side="left")
        self.var_dash_scope = tk.StringVar(value=self.config_data.get("dash_scope", "all"))
        ttk.Radiobutton(fb.top, text="Assigned to me", value="mine", variable=self.var_dash_scope,
                        style="Soft.TRadiobutton", command=self._dash_scope_changed).pack(side="left", padx=(12, 0))
        ttk.Radiobutton(fb.top, text="All open", value="all", variable=self.var_dash_scope,
                        style="Soft.TRadiobutton", command=self._dash_scope_changed).pack(side="left", padx=(8, 0))
        ttk.Button(fb.top, text="✖ Clear filters", command=self._dash_clear).pack(side="right")
        self.mp_dash_priority = MultiPick(fb.field("Priority"), "All priorities", self._dash_apply, width=14,
                                          title="Priorities")
        self.mp_dash_priority.pack(fill="x")
        self.mp_dash_status = MultiPick(fb.field("Status"), "All statuses", self._dash_apply, width=14,
                                        title="Statuses")
        self.mp_dash_status.pack(fill="x")
        self.mp_dash_release = MultiPick(fb.field("Release"), "All releases", self._dash_apply, width=16,
                                         title="Releases")
        self.mp_dash_release.pack(fill="x")
        self.mp_dash_assignee = MultiPick(fb.field("Assignee"), "All assignees", self._dash_apply, width=16,
                                          title="Assignees")
        self.mp_dash_assignee.pack(fill="x")
        self.var_dash_track = tk.StringVar(value="All tracks")
        cb = ttk.Combobox(fb.field("Track"), textvariable=self.var_dash_track, state="readonly",
                          values=["All tracks", TRACK_TAG, "Legacy"])
        cb.pack(fill="x")
        cb.bind("<<ComboboxSelected>>", lambda e: self._dash_apply())
        cr = fb.field("Created between", span=3)
        self.var_dash_from, self.var_dash_to = tk.StringVar(), tk.StringVar()
        ttk.Entry(cr, textvariable=self.var_dash_from, width=11).pack(side="left")
        tk.Label(cr, text="and", bg=C["soft"], fg=C["muted"], font=FONTS["small"]).pack(side="left", padx=4)
        ttk.Entry(cr, textvariable=self.var_dash_to, width=11).pack(side="left")
        for txt, days in (("30 d", 30), ("90 d", 90)):
            ttk.Button(cr, text=txt, width=5, command=lambda d=days: self._dash_created_range(d)).pack(side="left",
                                                                                                   padx=(6, 0))
        for v in (self.var_dash_from, self.var_dash_to):
            v.trace_add("write", lambda *a: self._dash_date_typed())
        self.lbl_dash_filters = tk.Label(fb, text="", bg=C["soft"], fg=C["accent_dark"], font=FONTS["small_bold"],
                                         anchor="w")
        self.lbl_dash_filters.pack(fill="x", padx=12, pady=(0, 8))
        self.dash_raw = None
        self._dash_date_job = None

    def _dash_created_range(self, days):
        self.var_dash_from.set((dt.date.today() - dt.timedelta(days=days)).strftime("%d-%m-%Y"))
        self.var_dash_to.set(dt.date.today().strftime("%d-%m-%Y"))

    def _dash_date_typed(self):
        if self._dash_date_job:
            self.after_cancel(self._dash_date_job)
        self._dash_date_job = self.after(500, self._dash_apply)

    def _dash_clear(self):
        for mp in (self.mp_dash_release, self.mp_dash_assignee, self.mp_dash_priority, self.mp_dash_status):
            mp.set_selected([])
        self.var_dash_track.set("All tracks")
        self.var_dash_from.set("")
        self.var_dash_to.set("")
        self._dash_apply()

    def _dash_scope(self):
        pids = [self.projects[n] for n in self.mp_dash_project.selected if n in self.projects]
        mine = self.var_dash_scope.get() == "mine"
        return pids, mine, f"{self.mp_dash_project.summary()} · {'Assigned to me' if mine else 'All open items'}"

    def _dash_scope_changed(self):
        self.config_data["dash_projects"] = list(self.mp_dash_project.selected)
        self.config_data["dash_scope"] = self.var_dash_scope.get()
        self._save_config_quiet()
        self.refresh_dashboard()

    def _dash_keep(self):
        """Returns (keep(issue), active, description) for the non-project filters."""
        rel = set(self.mp_dash_release.selected)
        asg = set(self.mp_dash_assignee.selected)
        pri = set(self.mp_dash_priority.selected)
        sts = set(self.mp_dash_status.selected)
        closed_ids = {it["id"] for it in (self.dash_raw or {}).get("closed", [])}
        track = self.var_dash_track.get()
        d_from, d_to = parse_user_date(self.var_dash_from.get()), parse_user_date(self.var_dash_to.get())
        raw = self.dash_raw or {}
        track_ids, track_ok = raw.get("track_ids", set()), raw.get("track_ok", False)
        bits = []
        if rel:
            bits.append("Release: " + ", ".join(sorted(rel)))
        if asg:
            bits.append("Assignee: " + ", ".join(sorted(asg)))
        if pri:
            bits.append("Priority: " + ", ".join(sorted(pri)))
        if sts:
            bits.append("Status: " + ", ".join(sorted(sts, key=self._status_rank)))
        if track != "All tracks" and track_ok:
            bits.append(f"Track: {track}")
        if d_from or d_to:
            bits.append("Created " + (fmt_date(d_from, False) if d_from else "…") + " → " +
                        (fmt_date(d_to, False) if d_to else "today"))

        def keep(it):
            if rel and (name_of(it.get("fixed_version"), "") or "No release") not in rel:
                return False
            if asg and name_of(it.get("assigned_to"), "Unassigned") not in asg:
                return False
            if pri and name_of(it.get("priority")) not in pri:
                return False
            # status filter narrows the open items; closed-ticket stats (throughput) stay as they are
            if sts and it["id"] not in closed_ids and name_of(it.get("status")) not in sts:
                return False
            if track != "All tracks" and track_ok:
                is2 = it["id"] in track_ids
                if (track == TRACK_TAG) != is2:
                    return False
            if d_from or d_to:
                c = parse_date(it.get("created_on"))
                if not c or (d_from and c < d_from) or (d_to and c > d_to):
                    return False
            return True

        return keep, bool(bits), " · ".join(bits)

    def _dash_options(self, raw):
        items = raw["open_issues"] + raw["created"] + raw["closed"]
        rels = sorted({name_of(it.get("fixed_version"), "") for it in items} - {""}, key=str.lower)
        self.mp_dash_release.set_options(["No release"] + rels)
        people = sorted({name_of(it.get("assigned_to"), "") for it in items} - {""}, key=str.lower)
        self.mp_dash_assignee.set_options(["Unassigned"] + people)
        order = [p.get("name") for p in reversed(raw["prio_list"])]
        seen = {name_of(it.get("priority")) for it in items}
        self.mp_dash_priority.set_options([p for p in order if p in seen] +
                                          sorted(seen - set(order), key=str.lower))
        stats = {name_of(it.get("status")) for it in raw["open_issues"]} - {""}
        self.mp_dash_status.set_options(sorted(stats, key=self._status_rank))

    @staticmethod
    def _status_rank(name):
        return (KANBAN_ORDER.index(name) if name in KANBAN_ORDER else len(KANBAN_ORDER), name.lower())

    def _dash_apply(self):
        if not self.dash_raw:
            return
        keep, active, desc = self._dash_keep()
        self.lbl_dash_filters.configure(text=("Showing: " + desc) if active else "No filters: showing everything")
        label = self._dash_scope_label + (f"  ·  Filtered by {desc}" if active else "")
        self._render_dashboard(build_dashboard(self.dash_raw, label, keep, active))

    # ----- dashboard: click a chart to see its tickets ----------------------------------
    def show_priority_tickets(self, priority):
        self.show_pick(f"Priority: {priority}", lambda it: name_of(it.get("priority")) == priority)

    def _pick_age(self, label):
        rng = next(((lo, hi) for lbl, lo, hi, _c in AGE_BUCKETS if lbl == label), None)
        if not rng:
            return
        today = dt.date.today()

        def age_ok(it):
            c = parse_date(it.get("created_on"))
            return bool(c) and rng[0] <= (today - c).days <= rng[1]

        self.show_pick(f"Age: {label}", age_ok)

    def show_pick(self, title, test):
        d = self.dashboard_data or {}
        items = [it for it in d.get("open_issues", []) if test(it)]
        grooms = self.config_data.get("groom_dates") or {}
        self.tv_pick.delete(*self.tv_pick.get_children())
        for it in sorted(items, key=lambda x: (x.get("due_date") or "9999", x["id"])):
            g = parse_date(grooms.get(str(it["id"])))
            state, _ = self._prd_info(it)
            self.tv_pick.insert("", "end", iid=f"i{it['id']}",
                                values=(f"#{it['id']}", it.get("subject", ""), name_of(it.get("assigned_to"),
                                                                                         "Unassigned"),
                                        name_of(it.get("status")), it.get("due_date") or "—", PRD_ICONS[state],
                                        fmt_date(g, False) if g else ""))
        kids = self.tv_pick.get_children()
        if kids:
            self.tv_pick.selection_set(kids[0])
        self.pick_box.configure(text=f"🔎 {title}: {plural(len(items), 'open ticket')}  "
                                     f"(double-click to open in Redmine · Ctrl+click to pick several)")
        self.pick_box.pack(fill="x", padx=4, pady=(4, 8), before=self.ch_trend)
        self.update_idletasks()
        try:
            inner = self.dash_sf.inner
            y = self.pick_box.winfo_rooty() - inner.winfo_rooty()
            self.dash_sf.canvas.yview_moveto(max(0.0, (y - 40) / max(1, inner.winfo_height())))
        except tk.TclError:
            pass

    def _open_ids_in_redmine(self, ids):
        ids = list(ids)
        if not ids:
            messagebox.showinfo(APP_NAME, "Click a ticket in the list first.", parent=self)
            return
        if len(ids) > 10 and not messagebox.askyesno(APP_NAME, f"Open {len(ids)} tickets in your browser?",
                                                     parent=self):
            return
        base = self.var_url.get().rstrip("/")
        for i in ids:
            webbrowser.open(f"{base}/issues/{i}")

    def _pick_open_redmine(self):
        self._open_ids_in_redmine(k[1:] for k in self.tv_pick.selection())

    def _open_pick_in_grooming(self):
        sel = self.tv_pick.selection()
        if not sel:
            messagebox.showinfo(APP_NAME, "Click a ticket in the list first.", parent=self)
            return
        self.show_in_workspace(int(sel[0][1:]))

    # ----- Grooming Workspace → Developer load -----------------------------------------
    LOAD_WINDOWS = ["Rest of this sprint", "This week", "Next 2 weeks", "Next 4 weeks"]
    LOAD_SORTS = ["Most free first", "Busiest first", "Name (A–Z)"]

    def _load_colours(self):
        dark = self.theme == "dark"
        return {"free": "#22352A" if dark else "#E4F1E6", "free_ink": "#8FCB9F" if dark else "#3E7A4E",
                "busy": "#3B2A27" if dark else "#F7E1DC", "busy_ink": "#E39A8F" if dark else "#A8473F",
                "off": C["head_bg"], "today": C["accent"]}

    def _build_devload(self, parent):
        fb = FilterBar(parent, cols=6, title="SHOW")
        fb.pack(fill="x", pady=(0, 6))
        fb.grid_f.pack_forget()
        self.var_load_win = tk.StringVar(value=self.config_data.get("load_window", self.LOAD_WINDOWS[0]))
        self.var_load_who = tk.StringVar(value=self.config_data.get("load_who", "Developers"))
        self.var_load_sort = tk.StringVar(value=self.config_data.get("load_sort", self.LOAD_SORTS[0]))
        self.var_load_all = tk.BooleanVar(value=bool(self.config_data.get("load_all_projects", False)))
        for label, var, values, width in (("Period", self.var_load_win, self.LOAD_WINDOWS, 20),
                                          ("People", self.var_load_who, ["Developers", "Everyone with tickets"], 20),
                                          ("Sort", self.var_load_sort, self.LOAD_SORTS, 16)):
            fb.label(fb.top, label)
            cb = ttk.Combobox(fb.top, textvariable=var, values=values, state="readonly", width=width)
            cb.pack(side="left")
            cb.bind("<<ComboboxSelected>>", lambda e: self._devload_changed())
        ttk.Checkbutton(fb.top, text="Count their tickets in all projects", variable=self.var_load_all,
                        style="Soft.TCheckbutton", command=self._devload_changed).pack(side="left", padx=(14, 0))
        ttk.Button(fb.top, text="↻ Refresh from Redmine", command=self._devload_refresh).pack(side="right")
        g = tk.Frame(fb, bg=C["soft"])
        g.pack(fill="x", padx=12, pady=(4, 10))
        tk.Label(g, text="🧭 GROOMING", bg=C["soft"], fg=C["accent_dark"], font=FONTS["small_bold"]).pack(
            side="left", padx=(0, 12))
        fb.label(g, "Session")
        self._gp_labels = {}
        self.var_gp_day = tk.StringVar()
        self.cb_gp_day = ttk.Combobox(g, textvariable=self.var_gp_day, state="readonly", width=22)
        self.cb_gp_day.pack(side="left")
        self.cb_gp_day.bind("<<ComboboxSelected>>", lambda e: self.render_devload())
        fb.label(g, "Tickets per developer")
        self.var_gp_target = tk.IntVar(value=self._gp_target())
        spn = ttk.Spinbox(g, from_=1, to=10, width=4, textvariable=self.var_gp_target,
                          command=self._gp_target_changed)
        spn.pack(side="left")
        spn.bind("<FocusOut>", lambda e: self._gp_target_changed())
        self.lbl_gp_sum = tk.Label(g, text="", bg=C["soft"], fg=C["accent_dark"], font=FONTS["small_bold"])
        self.lbl_gp_sum.pack(side="left", padx=14)
        ttk.Button(g, text="▶ Start session", command=self._gp_start).pack(side="right")
        ttk.Button(g, text="📋 Copy plan (Teams / email)", command=self._gp_copy).pack(side="right", padx=6)
        self.lbl_load_info = tk.Label(parent, text="", bg=C["bg"], fg=C["muted"], font=FONTS["small"], anchor="w",
                                      justify="left")
        self.lbl_load_info.pack(fill="x", pady=(0, 4))
        lg = tk.Frame(parent, bg=C["bg"])
        lg.pack(fill="x", pady=(0, 6))
        col = self._load_colours()
        for txt, bg, fg in (("  Busy  ", col["busy"], col["busy_ink"]),
                            ("  Free  ", col["free"], col["free_ink"]),
                            ("  Leave / holiday  ", col["off"], C["muted"])):
            tk.Label(lg, text=txt, bg=bg, fg=fg, font=FONTS["small_bold"], pady=2).pack(side="left", padx=(0, 6))
        tk.Label(lg, text="   Busy = up to their latest due date.  Chips sit on their due date.  Click a person to "
                          "see their tickets and plan their grooming below.", bg=C["bg"], fg=C["muted"],
                 font=FONTS["small"]).pack(side="left")
        pw = ttk.PanedWindow(parent, orient="vertical")
        pw.pack(fill="both", expand=True)
        top = tk.Frame(pw, bg=C["panel"], highlightbackground=C["line"], highlightthickness=1)
        self.load_canvas = tk.Canvas(top, bg=C["panel"], highlightthickness=0)
        sb = ttk.Scrollbar(top, orient="vertical", command=self.load_canvas.yview)
        sbx = ttk.Scrollbar(top, orient="horizontal", command=self.load_canvas.xview)
        self.load_canvas.configure(yscrollcommand=sb.set, xscrollcommand=sbx.set)
        sbx.pack(side="bottom", fill="x")
        sb.pack(side="right", fill="y")
        self.load_canvas.pack(side="left", fill="both", expand=True)
        self.load_grid = tk.Frame(self.load_canvas, bg=C["line"])
        self._load_win = self.load_canvas.create_window((0, 0), window=self.load_grid, anchor="nw")
        self._load_minw = 0
        self.load_canvas.bind("<Configure>", lambda e: self.load_canvas.itemconfigure(
            self._load_win, width=max(e.width, self._load_minw)))
        self.load_grid.bind("<Configure>",
                            lambda e: self.load_canvas.configure(scrollregion=self.load_canvas.bbox("all")))
        for w in (self.load_canvas, self.load_grid):
            w.bind("<MouseWheel>", lambda e: self.load_canvas.yview_scroll(int(-e.delta / 120), "units"))
            w.bind("<Button-4>", lambda e: self.load_canvas.yview_scroll(-2, "units"))
            w.bind("<Button-5>", lambda e: self.load_canvas.yview_scroll(2, "units"))
        pw.add(top, weight=3)
        bottom_all = ttk.PanedWindow(pw, orient="horizontal")
        pw.add(bottom_all, weight=2)
        bottom = tk.Frame(bottom_all, bg=C["bg"])
        bottom_all.add(bottom, weight=3)
        self.gp_panel = tk.Frame(bottom_all, bg=C["panel"], highlightbackground=C["line"], highlightthickness=1,
                                 width=380)
        bottom_all.add(self.gp_panel, weight=1)
        self.load_hpw, self._load_hsash_set = bottom_all, False
        self.load_pw, self._load_sash_set = pw, False
        bh = tk.Frame(bottom, bg=C["bg"])
        bh.pack(fill="x", pady=(8, 4))
        self.lbl_load_person = tk.Label(bh, text="Click a person above to see all their tickets here.", bg=C["bg"],
                                        fg=C["header"], font=FONTS["h3"], anchor="w")
        self.lbl_load_person.pack(side="left")
        ttk.Button(bh, text="🌐 Redmine", command=self._load_open_redmine).pack(side="right")
        ttk.Button(bh, text="🎫 Tickets", command=self._load_open_ws).pack(side="right", padx=6)
        tk.Label(bottom, text="✏ Click 📅 Groom date or 🗓 Sprint week to change it  ·  Ctrl/Shift for several  ·  "
                              "auto = from due date  ·  📌 = you picked",
                 bg=C["bg"], fg=C["accent_dark"], font=FONTS["small_bold"], anchor="w").pack(fill="x", pady=(0, 4))
        tf, self.tv_load = self._make_table(bottom, [("ticket", "Ticket", 360), ("status", "Status", 120),
                                                     ("priority", "Priority", 80), ("due", "Due", 130),
                                                     ("groom", "📅 Groom date ▾", 130),
                                                     ("week", "🗓 Sprint week ▾", 210), ("sprint", "Sprint", 110),
                                                     ("release", "Release", 150), ("groomed", "Groomed on", 100),
                                                     ("note", "Note", 220)], height=7)
        tf.pack(fill="both", expand=True)
        self.tv_load.configure(selectmode="extended")
        self.tv_load.tag_configure("over", foreground=C["bad"])
        self.tv_load.tag_configure("nodate", foreground=C["warn"])
        self.tv_load.bind("<Double-1>", self._load_double)
        self.tv_load.bind("<ButtonRelease-1>", self._load_cell_click)
        self.tv_load.bind("<Button-3>", self._load_cell_click)
        self.tv_load.bind("<Button-2>", self._load_cell_click)
        self._load_person = None
        self._load_rows = {}

    def _devload_changed(self):
        for k, v in (("load_window", self.var_load_win), ("load_who", self.var_load_who),
                     ("load_sort", self.var_load_sort), ("load_all_projects", self.var_load_all)):
            self.config_data[k] = v.get()
        self._save_config_quiet()
        if self.var_load_all.get() and not self.team_open and self._configured(warn=False):
            self._devload_fetch_all()
        else:
            self.render_devload()

    def _devload_refresh(self):
        self.refresh_tickets()
        if self.var_load_all.get():
            self._devload_fetch_all()

    def _devload_fetch_all(self):
        client = self.make_client()

        def done(res):
            self.team_open, _ = res
            self.render_devload()

        self.run_async(lambda: client.issues({"status_id": "open", "sort": "due_date:asc"}), done,
                       busy="Loading everyone's open tickets in all projects…")

    def _load_window(self):
        t = dt.date.today()
        w = self.var_load_win.get()
        note = ""
        if w == "Rest of this sprint":
            cur = self._current_sprint()
            if cur:
                key, sp = cur
                rel_txt = "" if key == NO_RELEASE else f" ({key})"
                return t, max(sp["end"], t), f"{sp['name']}{rel_txt}, ends {fmt_date(sp['end'])}"
            note = "No sprint is running today (set one up in the Sprint Planner), so showing the next 2 weeks. "
            w = "Next 2 weeks"
        if w == "This week":
            return t, t + dt.timedelta(days=max(0, 4 - t.weekday())), note + "this week"
        if w == "Next 4 weeks":
            return t, t + dt.timedelta(days=27), note + "the next 4 weeks"
        return t, t + dt.timedelta(days=13), note + "the next 2 weeks"

    def _home_leave(self):
        try:
            home = load_json(Path(__file__).resolve().parent / "PO_Home_Data" / "po_home.json", {})
            return home.get("leave") or []
        except Exception:
            return []

    def _devload_people(self, pool):
        devs = [m for m in self.team() if m.get("status", "Active") == "Active" and
                re.search(r"develop|engineer", m.get("designation") or "", re.I)]
        people = {norm_name(m["name"]): m["name"] for m in devs}
        if self.var_load_who.get() != "Developers":
            for it in pool:
                nm = name_of(it.get("assigned_to"), "")
                if nm and norm_name(nm) not in people:
                    people[norm_name(nm)] = nm
        return people

    def _devload_stats(self, key, name, tickets, days, leave, hols):
        """Works out busy/free days for one person in the period."""
        t = dt.date.today()
        member = next((m for m in self.team() if norm_name(m["name"]) == key), {})
        rel = parse_date(member.get("relieving"))
        dues = [parse_date(it.get("due_date")) for it in tickets]
        overdue = [it for it, d in zip(tickets, dues) if d and d < t]
        nodate = [it for it, d in zip(tickets, dues) if not d]
        future = [d for d in dues if d and d >= t]
        busy_until = max(future) if future else (t if overdue else None)
        my_leave = [lv for lv in leave if norm_name(lv.get("person")) == key]

        def off_on(d):
            if rel and d > rel:
                return 1.0, "Left"
            if d in hols:
                return 1.0, "Holiday"
            for lv in my_leave:
                a, b = parse_date(lv.get("from")), parse_date(lv.get("to"))
                if a and b and a <= d <= b and lv.get("type", "Leave") != "WFH":
                    part = lv.get("part", "Full day")
                    half = part.startswith(("First", "Second"))
                    return (0.5 if half else 1.0), lv.get("type", "Leave") + (" ½" if half else "")
            return 0.0, ""

        cells, free, busy, off, gone = {}, 0.0, 0.0, 0.0, 0
        free_from = None
        for d in days:
            o, why = off_on(d)
            avail = 1 - o
            is_busy = bool(busy_until and d <= busy_until)
            cells[d] = (o, why, is_busy)
            off += o if why != "Left" else 0
            gone += why == "Left"
            if is_busy:
                busy += avail
            else:
                free += avail
                if avail > 0 and free_from is None:
                    free_from = d
        if free_from is None and busy_until and busy_until >= (days[-1] if days else t):
            d = busy_until + dt.timedelta(days=1)
            while not is_working_day(d, hols):
                d += dt.timedelta(days=1)
            free_from = d
        return {"name": name, "member": member, "rel": rel, "tickets": tickets, "overdue": overdue,
                "nodate": nodate, "busy_until": busy_until, "free": free, "busy": busy, "off": off,
                "free_from": free_from, "cells": cells, "gone": gone}

    def render_devload(self):
        if not hasattr(self, "load_grid"):
            return
        for w in self.load_grid.winfo_children():
            w.destroy()
        self._gp_refresh_days()
        self._gp_summary()
        col = self._load_colours()
        t = dt.date.today()
        hols = self.holidays()
        start, end, wlabel = self._load_window()
        days = [start + dt.timedelta(days=i) for i in range((end - start).days + 1)
                if (start + dt.timedelta(days=i)).weekday() < 5]
        pool = {it["id"]: it for it in self.issues}
        if self.var_load_all.get():
            pool.update({it["id"]: it for it in self.team_open})
        pool = list(pool.values())
        people = self._devload_people(pool)
        by_person = {k: [] for k in people}
        for it in pool:
            k = norm_name(name_of(it.get("assigned_to"), ""))
            if k in by_person:
                by_person[k].append(it)
        leave = self._home_leave()
        stats = [self._devload_stats(k, people[k], sorted(v, key=lambda it: (it.get("due_date") or "9999", it["id"])),
                                     days, leave, hols) for k, v in by_person.items()]
        sort = self.var_load_sort.get()
        if sort == "Most free first":
            stats.sort(key=lambda s: (-s["free"], len(s["overdue"]), s["name"].lower()))
        elif sort == "Busiest first":
            stats.sort(key=lambda s: (s["free"], -len(s["tickets"]), s["name"].lower()))
        else:
            stats.sort(key=lambda s: s["name"].lower())
        self._load_stats = {norm_name(s["name"]): s for s in stats}
        _, mine, _ = self._scope()
        warn = ("  ⚠ Your ticket list is set to 'Assigned to me', so other people's tickets aren't loaded. Switch the "
                "Tickets tab to 'All open'." if mine else "")
        src = "all projects" if self.var_load_all.get() else f"the projects picked in Tickets ({self.mp_project.summary()})"
        self.lbl_load_info.configure(
            text=f"Period: {wlabel}  ·  {plural(len(days), 'working day')}  ·  tickets from {src}  ·  leave from Kaizen Home."
                 f"{warn}")
        g = self.load_grid
        ncols = len(days) + 3
        g.columnconfigure(0, weight=0, minsize=290)
        self._load_minw = 290 + (ncols - 1) * 78
        if not self._load_sash_set and self.load_pw.winfo_height() > 300:
            self.load_pw.sashpos(0, int(self.load_pw.winfo_height() * 0.56))
            self._load_sash_set = True
        if not self._load_hsash_set and self.load_hpw.winfo_width() > 800:
            self.load_hpw.sashpos(0, self.load_hpw.winfo_width() - 400)
            self._load_hsash_set = True
        self.load_canvas.itemconfigure(self._load_win, width=max(self.load_canvas.winfo_width(), self._load_minw))
        for j in range(1, ncols):
            g.columnconfigure(j, weight=1, uniform="ld", minsize=74)

        def head(j, text, fg=None, bg=None):
            tk.Label(g, text=text, bg=bg or C["head_bg"], fg=fg or C["muted"], font=FONTS["small_bold"],
                     pady=4).grid(row=0, column=j, sticky="nsew", padx=(0, 1), pady=(0, 1))

        head(0, "Person  ·  capacity in this period")
        head(1, "Overdue", fg=C["bad"])
        for j, d in enumerate(days, 2):
            head(j, d.strftime("%a\n%d %b"), fg="white" if d == t else None, bg=col["today"] if d == t else None)
        head(len(days) + 2, "No due\ndate", fg=C["warn"])
        if not stats:
            tk.Label(g, text="No developers found. Add them in the Team tab (designation 'Developer'), or switch "
                             "People to 'Everyone with tickets'.", bg=C["panel"], fg=C["muted"], pady=20).grid(
                row=1, column=0, columnspan=ncols, sticky="nsew")
        for i, s in enumerate(stats, 1):
            self._devload_row(g, i, s, days, col, t)
        if self._load_person in self._load_stats:
            self._load_show(self._load_person)
        self._gp_render_panel()

    def _devload_row(self, g, i, s, days, col, t):
        key = norm_name(s["name"])
        sel = key == self._load_person
        bg = C["select"] if sel else C["panel"]
        card = tk.Frame(g, bg=bg, cursor="hand2", padx=10, pady=6)
        card.grid(row=i, column=0, sticky="nsew", padx=(0, 1), pady=(0, 1))
        widgets = [card]
        top = tk.Frame(card, bg=bg)
        top.pack(fill="x")
        widgets.append(top)
        widgets.append(tk.Label(top, text=s["name"], bg=bg, fg=C["header"], font=FONTS["h3"], anchor="w"))
        widgets[-1].pack(side="left")
        n = len(s["tickets"])
        widgets.append(tk.Label(top, text=f"{n} open", bg=C["accent"] if n else C["muted"], fg="white",
                                font=FONTS["small_bold"], padx=6))
        widgets[-1].pack(side="right")
        widgets.append(self._gp_chip(top, key, bg))
        fmt = lambda d: d.strftime("%a %d %b")  # noqa: E731
        if not s["busy_until"]:
            if s["free"] > 0:
                verdict, vcol = f"🟢 Free now · {s['free']:g} free day(s)", col["free_ink"]
            else:
                verdict, vcol = "⚪ Not available in this period", C["muted"]
        elif s["free"] <= 0:
            verdict = f"🔴 Fully booked until {fmt(s['busy_until'])}"
            if s["free_from"]:
                verdict += f" · free from {fmt(s['free_from'])}"
            vcol = col["busy_ink"]
        else:
            verdict = f"{'🟢' if s['free'] >= 2 else '🟠'} Busy until {fmt(s['busy_until'])} · " \
                      f"{s['free']:g} free day(s) from {fmt(s['free_from'])}"
            vcol = col["free_ink"] if s["free"] >= 2 else C["warn"]
        if s["rel"] and days and s["rel"] <= days[-1]:
            verdict += f"\n👋 Leaves on {fmt(s['rel'])}"
        widgets.append(tk.Label(card, text=verdict, bg=bg, fg=vcol, font=FONTS["small_bold"], anchor="w",
                                justify="left", wraplength=270))
        widgets[-1].pack(fill="x", pady=(2, 3))
        bar = tk.Canvas(card, height=10, bg=C["track"], highlightthickness=0, width=268)
        bar.pack(anchor="w")
        widgets.append(bar)
        total = max(len(days), 1)
        x = 0
        for part, colour in ((s["busy"], col["busy_ink"]), (s["free"], col["free_ink"]),
                             (s["off"] + s["gone"], C["muted"])):
            w = 268 * part / total
            if w > 0:
                bar.create_rectangle(x, 0, x + w, 10, fill=colour, width=0)
            x += w
        extra = []
        if s["overdue"]:
            extra.append(f"⚠ {len(s['overdue'])} overdue")
        if s["nodate"]:
            extra.append(f"{len(s['nodate'])} without due date (not counted)")
        if s["off"]:
            extra.append(f"{s['off']:g} day(s) off")
        if extra:
            widgets.append(tk.Label(card, text="  ·  ".join(extra), bg=bg, fg=C["muted"], font=FONTS["small"],
                                    anchor="w", justify="left", wraplength=270))
            widgets[-1].pack(fill="x", pady=(3, 0))
        for w in widgets:
            w.bind("<Button-1>", lambda e, k=key: self._load_pick(k))
        # overdue column
        cell = tk.Frame(g, bg=col["busy"] if s["overdue"] else C["panel"])
        cell.grid(row=i, column=1, sticky="nsew", padx=(0, 1), pady=(0, 1))
        for it in s["overdue"]:
            self._load_chip(cell, it, key)
        due_map = {}
        for it in s["tickets"]:
            d = parse_date(it.get("due_date"))
            if d and d >= t:
                if d.weekday() >= 5:   # weekend → the Friday before (or the first day shown)
                    d = max(d - dt.timedelta(days=d.weekday() - 4), days[0] if days else d)
                due_map.setdefault(d, []).append(it)
        for j, d in enumerate(days, 2):
            o, why, busy = s["cells"][d]
            if o >= 1:
                cbg = col["off"]
            else:
                cbg = col["busy"] if busy else col["free"]
            cell = tk.Frame(g, bg=cbg, highlightthickness=2 if d == t else 0, highlightbackground=col["today"])
            cell.grid(row=i, column=j, sticky="nsew", padx=(0, 1), pady=(0, 1))
            if why:
                tk.Label(cell, text=truncate(why, 10), bg=cbg, fg=C["muted"], font=FONTS["small"]).pack(pady=(4, 0))
            for it in due_map.get(d, []):
                self._load_chip(cell, it, key)
            if not why and not due_map.get(d) and not busy:
                tk.Label(cell, text="free", bg=cbg, fg=col["free_ink"], font=FONTS["small"]).pack(pady=(6, 0))
            cell.bind("<Button-1>", lambda e, k=key: self._load_pick(k))
        cell = tk.Frame(g, bg=C["panel"])
        cell.grid(row=i, column=len(days) + 2, sticky="nsew", padx=(0, 1), pady=(0, 1))
        for it in s["nodate"]:
            self._load_chip(cell, it, key, faded=True)
        # tickets due after the period: one summary chip in the last day column
        later = [it for it in s["tickets"] if parse_date(it.get("due_date")) and days and
                 parse_date(it["due_date"]) > days[-1]]
        if later:
            last = g.grid_slaves(row=i, column=len(days) + 1)
            if last:
                tk.Label(last[0], text=f"+{len(later)} later →", bg=last[0].cget("bg"), fg=col["busy_ink"],
                         font=FONTS["small_bold"]).pack(side="bottom", pady=(0, 3))

    def _load_chip(self, parent, it, key, faded=False):
        groomed = self.groomed_on(it["id"])
        d = parse_date(it.get("due_date"))
        wk = f" {d.strftime('%a')}" if d and d.weekday() >= 5 and d >= dt.date.today() else ""
        lb = tk.Label(parent, text=f"#{it['id']}" + wk + (" ✅" if groomed else ""), font=FONTS["small_bold"],
                      bg=C["muted"] if faded else priority_color(name_of(it.get("priority"))), fg="white",
                      cursor="hand2", padx=3)
        lb.pack(fill="x", padx=3, pady=(3, 0))
        lb.bind("<Button-1>", lambda e: self._load_pick(key, it["id"]))
        lb.bind("<Double-1>", lambda e: self.show_in_workspace(it["id"]))

    def _load_pick(self, key, issue_id=None):
        changed = key != self._load_person
        self._load_person = key
        if changed:
            self.render_devload()
        else:
            self._load_show(key)
            self._gp_render_panel()
        if issue_id and self.tv_load.exists(f"l{issue_id}"):
            self.tv_load.selection_set(f"l{issue_id}")
            self.tv_load.see(f"l{issue_id}")

    def _load_show(self, key):
        s = self._load_stats.get(key)
        if not s:
            return
        t = dt.date.today()
        bits = [s["name"], plural(len(s["tickets"]), "open ticket")]
        if s["busy_until"]:
            bits.append(f"busy until {fmt_date(s['busy_until'])}")
        if s["free_from"]:
            bits.append(f"free from {fmt_date(s['free_from'])}")
        self.lbl_load_person.configure(text="  ·  ".join(bits))
        tv = self.tv_load
        keep = list(tv.selection())
        tv.delete(*tv.get_children())
        order = sorted(s["tickets"], key=lambda it: (0 if it in s["overdue"] else (2 if it in s["nodate"] else 1),
                                                      it.get("due_date") or "", it["id"]))
        for it in order:
            d = parse_date(it.get("due_date"))
            g = self.groomed_on(it["id"])
            if it in s["overdue"]:
                note, tag = f"⚠ overdue by {(t - d).days} day(s)", ("over",)
            elif not d:
                note, tag = "No due date: ask for one before planning", ("nodate",)
            else:
                note, tag = (f"due in {(d - t).days} day(s)" if d > t else "due today"), ()
            gd = parse_date((self.config_data.get("groom_dates") or {}).get(str(it["id"])))
            w, how = self.week_of(it)
            pi = self.plan_info(it)
            wtxt = (self._week_short(w) + (" 📌 picked" if how == "manual" else " · auto")) if w else \
                ("—  pick a week" if not d else "— (no week covers this date)")
            tv.insert("", "end", iid=f"l{it['id']}", values=(
                f"#{it['id']}  {it.get('subject', '')}", name_of(it.get("status")), name_of(it.get("priority")),
                fmt_date(d) if d else "—", fmt_date(gd) if gd else "—  set", wtxt, pi["sprint_txt"] or "—",
                pi["release_txt"] or "—", fmt_date(g, False) if g else "—", note), tags=tag)
        keep = [k for k in keep if tv.exists(k)]
        if keep:
            tv.selection_set(keep)

    def _load_ids_for(self, row):
        sel = list(self.tv_load.selection())
        if row not in sel:
            self.tv_load.selection_set(row)
            sel = [row]
        return [int(r[1:]) for r in sel]

    def _load_cell_click(self, event):
        tv = self.tv_load
        row, col = tv.identify_row(event.y), tv.identify_column(event.x)
        if not row:
            return
        name = tv["columns"][int(col[1:]) - 1] if col and col != "#0" else ""
        right = event.num in (2, 3)
        if name not in ("groom", "week") and not right:
            return
        if event.state & 0x0005 and not right:   # Ctrl/Shift click: just selecting rows
            return
        ids = self._load_ids_for(row)
        menu = tk.Menu(self, tearoff=0)
        if name == "groom" or right:
            self._groom_menu(menu, ids)
        if right:
            menu.add_separator()
        if name == "week" or right:
            self._week_menu(menu, ids)
        if right:
            menu.add_separator()
            menu.add_command(label="🎫 Open in Tickets", command=lambda: self.show_in_workspace(ids[0]))
            menu.add_command(label="🌐 Open in Redmine", command=lambda: webbrowser.open(
                f"{self.var_url.get().rstrip('/')}/issues/{ids[0]}"))
        menu.tk_popup(event.x_root, event.y_root)

    def _groom_menu(self, menu, ids):
        menu.add_command(label=f"📅 Groom date for {plural(len(ids), 'ticket')}", state="disabled")
        t = dt.date.today()
        nxt = t + dt.timedelta(days=(7 - t.weekday()) % 7 or 7)   # next Monday
        session_days = sorted({parse_date(k) for k in (self.config_data.get("groom_sessions") or {})
                               if parse_date(k) and parse_date(k) >= t})
        choices = []
        for d in [t] + session_days + [nxt + dt.timedelta(days=7 * k) for k in range(4)]:
            if d not in choices:
                choices.append(d)
        for d in choices:
            menu.add_command(label=("   Today · " if d == t else "   ") + fmt_date(d),
                             command=lambda dd=d: self.set_groom_dates(ids, dd))
        menu.add_command(label="   Another date…", command=lambda: self.set_groom_date(ids=[str(i) for i in ids]))
        menu.add_command(label="   ✖ Clear groom date", command=lambda: self.set_groom_dates(ids, None))

    def _week_menu(self, menu, ids):
        menu.add_command(label=f"🗓 Sprint week for {plural(len(ids), 'ticket')}", state="disabled")
        t = dt.date.today()
        weeks = [w for w in self.all_weeks() if w["end"] >= t - dt.timedelta(days=7)]
        if not weeks:
            menu.add_command(label="   No weeks yet: add them in the Sprint Planner (➕ Add week)",
                             command=lambda: self.nb.select(self.tab_sprint))
            return
        issues = {it["id"]: it for it in self._load_stats_tickets()}
        cur = {self._week_id(w) for i in ids for w in [self.week_of(issues.get(i, {"id": i}))[0]] if w}
        for w in weeks:
            mark = "✔ " if self._week_id(w) in cur else "   "
            now = "  ← this week" if w["start"] <= t <= w["end"] else ""
            label = f"{mark}{self._week_short(w)}" + (f"  · {w['sprint']}" if w.get("sprint") else "") + \
                    (f"  · {w['release']}" if w["release"] != NO_RELEASE else "") + now
            menu.add_command(label=label, command=lambda ww=w: self.set_ticket_week(ids, ww))
        menu.add_command(label="   ↺ Use the due date's week (undo my pick)",
                         command=lambda: self.set_ticket_week(ids, None))

    def _load_stats_tickets(self):
        s = (getattr(self, "_load_stats", {}) or {}).get(self._load_person) or {}
        return s.get("tickets", [])

    def _load_double(self, event):
        col = self.tv_load.identify_column(event.x)
        name = self.tv_load["columns"][int(col[1:]) - 1] if col and col != "#0" else ""
        if name not in ("groom", "week"):
            self._load_open_ws()

    def _load_sel_id(self):
        sel = self.tv_load.selection()
        if not sel:
            messagebox.showinfo(APP_NAME, "Click a ticket in the list first.", parent=self)
            return None
        return int(sel[0][1:])

    def _load_open_ws(self):
        iid = self._load_sel_id()
        if iid:
            self.show_in_workspace(iid)

    def _load_open_redmine(self):
        iid = self._load_sel_id()
        if iid:
            webbrowser.open(f"{self.var_url.get().rstrip('/')}/issues/{iid}")

    # ----- Developer load → grooming plan (about 3 tickets per developer for Monday) ----------
    def planned_dev(self, issue_id):
        """Developer you planned for this ticket's grooming session (only while it still has that groom date)."""
        rec = (self.config_data.get("groom_plan") or {}).get(str(issue_id))
        gd = (self.config_data.get("groom_dates") or {}).get(str(issue_id))
        return rec["dev"] if rec and gd and rec.get("date") == gd else ""

    def _gp_sessions(self):
        t = dt.date.today()
        nxt = t + dt.timedelta(days=(7 - t.weekday()) % 7 or 7)
        days = {nxt + dt.timedelta(days=7 * k) for k in range(4)}
        if t.weekday() == 0:
            days.add(t)
        days |= {d for d in (parse_date(v) for v in (self.config_data.get("groom_dates") or {}).values())
                 if d and d >= t}
        return sorted(days)

    def _gp_day(self):
        return self._gp_labels.get(self.var_gp_day.get()) or self._gp_sessions()[0]

    def _gp_target(self):
        return int(self.config_data.get("gp_target", 3))

    def _gp_refresh_days(self):
        days = self._gp_sessions()
        self._gp_labels = {(fmt_date(d) + ("  (today)" if d == dt.date.today() else "")): d for d in days}
        self.cb_gp_day.configure(values=list(self._gp_labels))
        if self.var_gp_day.get() not in self._gp_labels:
            self.var_gp_day.set(next(iter(self._gp_labels)))

    def _gp_target_changed(self):
        try:
            n = max(1, min(10, int(self.var_gp_target.get())))
        except (tk.TclError, ValueError):
            return
        if n != self._gp_target():
            self.config_data["gp_target"] = n
            self._save_config_quiet()
            self.render_devload()

    def _gp_plan_for(self, day):
        plan = self.config_data.get("groom_plan") or {}
        grooms = self.config_data.get("groom_dates") or {}
        out = {}
        for k, rec in plan.items():
            if rec.get("date") == day.isoformat() and grooms.get(k) == rec["date"]:
                out.setdefault(norm_name(rec["dev"]), []).append(int(k))
        return out

    def _gp_devs(self, day):
        out = []
        for m in self.team():
            if m.get("status", "Active") != "Active" or not re.search(r"develop|engineer", m.get("designation") or "",
                                                                         re.I):
                continue
            rel = parse_date(m.get("relieving"))
            if rel and rel < day:
                continue
            out.append(m)
        return sorted(out, key=lambda m: m["name"].lower())

    def _gp_summary(self):
        day = self._gp_day()
        plan = self._gp_plan_for(day)
        devs = self._gp_devs(day)
        target = self._gp_target()
        done = sum(min(len(plan.get(norm_name(m["name"]), [])), target) for m in devs)
        self.lbl_gp_sum.configure(text=f"{done} of {len(devs) * target} picked",
                                  fg=C["good"] if devs and done >= len(devs) * target else C["accent_dark"])

    def _gp_chip(self, parent, key, bg):
        """Small 🧭 n/3 chip on a developer's card."""
        n = len(self._gp_plan_for(self._gp_day()).get(key, []))
        target = self._gp_target()
        colour = C["good"] if n >= target else (C["warn"] if n else C["muted"])
        lb = tk.Label(parent, text=f"🧭 {n}/{target}", bg=colour, fg="white", font=FONTS["small_bold"], padx=5)
        lb.pack(side="right", padx=(0, 4))
        return lb

    def _gp_render_panel(self):
        """Right side of the bottom area: the selected developer's tickets for the grooming session."""
        p = self.gp_panel
        for w in p.winfo_children():
            w.destroy()
        day = self._gp_day()
        key = self._load_person
        s = (getattr(self, "_load_stats", {}) or {}).get(key)
        if not s:
            tk.Label(p, text="🧭 Grooming plan", bg=C["panel"], fg=C["header"], font=FONTS["h3"], anchor="w").pack(
                fill="x", padx=12, pady=(10, 4))
            tk.Label(p, text="Click a developer above to pick their tickets for the next grooming session.",
                     bg=C["panel"], fg=C["muted"], font=FONTS["small"], anchor="w", justify="left",
                     wraplength=330).pack(fill="x", padx=12)
            return
        ids = self._gp_plan_for(day).get(key, [])
        target = self._gp_target()
        head = tk.Frame(p, bg=C["panel"])
        head.pack(fill="x", padx=12, pady=(10, 2))
        tk.Label(head, text=f"🧭 Grooming · {day.strftime('%a %d %b')}", bg=C["panel"], fg=C["header"],
                 font=FONTS["h3"]).pack(side="left")
        n = len(ids)
        tk.Label(head, text=f"{n}/{target}", bg=C["good"] if n >= target else (C["warn"] if n else C["muted"]),
                 fg="white", font=FONTS["small_bold"], padx=6).pack(side="right")
        tk.Label(p, text=f"Tickets for {s['name']} to go through:", bg=C["panel"], fg=C["muted"],
                 font=FONTS["small"], anchor="w").pack(fill="x", padx=12)
        warn = []
        for lv in self._home_leave():
            a, b = parse_date(lv.get("from")), parse_date(lv.get("to"))
            if norm_name(lv.get("person")) == key and a and b and a <= day <= b and lv.get("type") != "WFH":
                warn.append(f"🌴 {lv.get('type', 'Leave')} that day")
        if day in self.holidays():
            warn.append("🎉 that day is a holiday")
        if s["rel"] and s["rel"] < day + dt.timedelta(days=14):
            warn.append(f"👋 leaves {s['rel'].strftime('%d %b')}")
        if warn:
            tk.Label(p, text="   ".join(warn), bg=C["panel"], fg=C["bad"], font=FONTS["small_bold"],
                     anchor="w").pack(fill="x", padx=12)
        for tid in ids:
            it = self.issue_index.get(tid) or {"id": tid, "subject": "(not in the current ticket list)"}
            r = tk.Frame(p, bg=C["soft"])
            r.pack(fill="x", padx=12, pady=(4, 0))
            tk.Label(r, text=f"#{tid}  {truncate(it.get('subject', ''), 42)}", bg=C["soft"], fg=C["ink"],
                     font=FONTS["small"], anchor="w").pack(side="left", padx=6, pady=3)
            x = tk.Label(r, text="✖", bg=C["soft"], fg=C["bad"], cursor="hand2", font=FONTS["small_bold"])
            x.pack(side="right", padx=6)
            x.bind("<Button-1>", lambda e, i_=tid: self._gp_remove(i_))
        for _ in range(max(0, target - n)):
            tk.Label(p, text="   ·  empty", bg=C["panel"], fg=C["faint"], font=FONTS["small"], anchor="w").pack(
                fill="x", padx=12, pady=(4, 0))
        ttk.Button(p, text=f"➕ Pick tickets for {first_name(s['name'])}…", style="Accent.TButton",
                   command=lambda: self._gp_picker(key, s["name"])).pack(anchor="w", padx=12, pady=(10, 4))

    def _gp_picker(self, key, dev):
        """A small window to choose tickets for this developer's grooming."""
        day = self._gp_day()
        win = tk.Toplevel(self)
        win.title(f"Pick tickets for {dev} · grooming {fmt_date(day)}")
        win.configure(bg=C["bg"])
        win.geometry("1060x600")
        win.transient(self)
        frm = tk.Frame(win, bg=C["bg"], padx=14, pady=12)
        frm.pack(fill="both", expand=True)
        tk.Label(frm, text=f"Pick tickets for {dev}", bg=C["bg"], fg=C["header"], font=FONTS["h2"],
                 anchor="w").pack(fill="x")
        tk.Label(frm, text=f"For grooming on {fmt_date(day)}. Select one or more (Ctrl/Shift) and click Add. "
                           "This only sets the groom date in this app; Redmine isn't changed.", bg=C["bg"],
                 fg=C["muted"], font=FONTS["small"], anchor="w").pack(fill="x", pady=(0, 8))
        row = tk.Frame(frm, bg=C["bg"])
        row.pack(fill="x", pady=(0, 6))
        v_q, v_p, v_hide = tk.StringVar(), tk.StringVar(value="All priorities"), tk.BooleanVar(value=True)
        tk.Label(row, text="🔎", bg=C["bg"], fg=C["muted"]).pack(side="left")
        e = ttk.Entry(row, textvariable=v_q, width=28)
        e.pack(side="left", padx=(2, 10))
        prios = sorted({name_of(it.get("priority")) for it in self.issues}, key=priority_rank)
        cb = ttk.Combobox(row, textvariable=v_p, values=["All priorities"] + prios, state="readonly", width=14)
        cb.pack(side="left")
        ttk.Checkbutton(row, text="Hide groomed / already planned", variable=v_hide).pack(side="left", padx=10)
        tf, tv = self._make_table(frm, [("ticket", "Ticket", 330), ("priority", "Priority", 80),
                                        ("status", "Status", 110), ("assignee", "Assignee", 110),
                                        ("sprint", "Sprint", 90), ("release", "Release", 110),
                                        ("prd", "PRD", 80), ("groom", "Groom date", 110)], height=12)
        tf.pack(fill="both", expand=True)
        tv.configure(selectmode="extended")
        grooms = self.config_data.get("groom_dates") or {}

        def fill(*_):
            tv.delete(*tv.get_children())
            q = v_q.get().strip().lower()
            rows = []
            for it in self.issues:
                pdev = self.planned_dev(it["id"])
                if v_hide.get() and (self.groomed_on(it["id"]) or pdev):
                    continue
                if v_p.get() != "All priorities" and name_of(it.get("priority")) != v_p.get():
                    continue
                if q and q not in it.get("subject", "").lower() and q not in str(it["id"]):
                    continue
                rows.append(it)
            rows.sort(key=lambda it: (priority_rank(name_of(it.get("priority"))), it["id"]))
            for it in rows:
                pi = self.plan_info(it)
                gd = parse_date(grooms.get(str(it["id"])))
                pdev = self.planned_dev(it["id"])
                tv.insert("", "end", iid=f"p{it['id']}", values=(
                    f"#{it['id']}  {it.get('subject', '')}", name_of(it.get("priority")), name_of(it.get("status")),
                    name_of(it.get("assigned_to"), "—"), pi["sprint_txt"] or "—", pi["release_txt"] or "—",
                    PRD_ICONS[self._prd_info(it)[0]],
                    (fmt_date(gd, False) + (f" · {first_name(pdev)}" if pdev else "")) if gd else "—"))
            lbl_n.configure(text=plural(len(rows), "ticket"))

        def add(*_):
            ids = [int(r[1:]) for r in tv.selection()]
            if not ids:
                messagebox.showinfo(APP_NAME, "Select one or more tickets first.", parent=win)
                return
            self._gp_assign(ids, dev, day)
            win.destroy()

        bb = tk.Frame(frm, bg=C["bg"])
        bb.pack(fill="x", pady=(8, 0))
        lbl_n = tk.Label(bb, text="", bg=C["bg"], fg=C["muted"], font=FONTS["small"])
        lbl_n.pack(side="left")
        ttk.Button(bb, text="Close", command=win.destroy).pack(side="right")
        ttk.Button(bb, text=f"➕ Add to {first_name(dev)}", style="Accent.TButton", command=add).pack(
            side="right", padx=8)
        v_q.trace_add("write", fill)
        cb.bind("<<ComboboxSelected>>", fill)
        v_hide.trace_add("write", fill)
        tv.bind("<Double-1>", add)
        fill()
        e.focus_set()
        self._gp_win = win
        return win

    def _gp_assign(self, ids, dev, day):
        plan = self.config_data.setdefault("groom_plan", {})
        grooms = self.config_data.setdefault("groom_dates", {})
        for i in ids:
            plan[str(i)] = {"dev": dev, "date": day.isoformat()}
            grooms[str(i)] = day.isoformat()
        self._save_config_quiet()
        self._ticket_meta_changed(f"{plural(len(ids), 'ticket')} planned with {dev} for grooming on {fmt_date(day)}.")

    def _gp_remove(self, tid):
        plan = self.config_data.setdefault("groom_plan", {})
        rec = plan.pop(str(tid), None)
        grooms = self.config_data.setdefault("groom_dates", {})
        if rec and grooms.get(str(tid)) == rec.get("date"):
            grooms.pop(str(tid), None)
        self._save_config_quiet()
        self._ticket_meta_changed(f"#{tid} taken out of the grooming plan.")

    def _gp_copy(self):
        day = self._gp_day()
        plan = self._gp_plan_for(day)
        lines = [f"Grooming – {fmt_date(day)}", ""]
        for m in self._gp_devs(day):
            ids = plan.get(norm_name(m["name"]), [])
            if not ids:
                continue
            lines.append(m["name"])
            for i in ids:
                it = self.issue_index.get(i) or {"subject": ""}
                lines.append(f"  • #{i} {it.get('subject', '')} ({name_of(it.get('priority'), '')})")
            lines.append("")
        if len(lines) == 2:
            messagebox.showinfo(APP_NAME, "Nothing planned for this session yet.", parent=self)
            return
        lines.append("Please go through your tickets before the session. Thanks!")
        self.clipboard_clear()
        self.clipboard_append("\n".join(lines))
        self.set_status("Grooming plan copied. Paste it into Teams or Outlook (Ctrl+V).")

    def _gp_start(self):
        day = self._gp_day()
        ids = sorted(int(k) for k, v in (self.config_data.get("groom_dates") or {}).items() if parse_date(v) == day)
        if not ids:
            messagebox.showinfo(APP_NAME, "Nothing planned for this session yet.", parent=self)
            return
        self._open_session(day, [self.issue_index.get(i) or {"id": i, "subject": "?"} for i in ids])

    def show_in_workspace(self, issue_id: int):
        """Jump to the Grooming Workspace with this exact ticket opened."""
        self._show_tickets()
        key = str(issue_id)
        if self.tree.exists(key):
            self.tree.selection_set(key)
            self.tree.see(key)
        else:
            self.load_issue(issue_id)
            if issue_id in self.issue_index:
                self.set_status(f"#{issue_id} is hidden by your filters, so it isn't on the board/list, but its "
                                f"details are open on the right.")
        if hasattr(self, "board"):
            self.board.select(issue_id, scroll=True)

    def _risk_ids(self):
        return [k[1:] for k in self.risk_tree.selection() if k.startswith("r")]

    def _risk_open_redmine(self):
        self._open_ids_in_redmine(self._risk_ids())

    def _risk_menu(self, event):
        row = self.risk_tree.identify_row(event.y)
        if row:
            self.risk_tree.selection_set(row)
        ids = self._risk_ids()
        if not ids:
            return
        menu = tk.Menu(self, tearoff=0)
        menu.add_command(label="Open in Redmine ↗", command=self._risk_open_redmine)
        menu.add_command(label="Show in Grooming Workspace", command=lambda: self.show_in_workspace(int(ids[0])))
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    # ----- Today ---------------------------------------------------------------
    def _build_today(self, parent):
        head = ttk.Frame(parent, style="Bg.TFrame", padding=(0, 12, 0, 6))
        head.pack(fill="x")
        left = ttk.Frame(head, style="Bg.TFrame")
        left.pack(side="left", fill="x", expand=True)
        self.lbl_hello = ttk.Label(left, text="Hello!", font=FONTS["h1"], foreground=C["header"])
        self.lbl_hello.pack(anchor="w")
        self.lbl_today_sub = ttk.Label(left, text="", style="Muted.TLabel", justify="left")
        self.lbl_today_sub.pack(anchor="w", pady=(2, 0))
        ttk.Button(head, text="▶ Start grooming session", style="Accent.TButton",
                   command=self.start_grooming_session).pack(side="right")
        ttk.Button(head, text="↻ Refresh", command=self.refresh_tickets).pack(side="right", padx=6)

        grid = tk.Frame(parent, bg=C["bg"])
        grid.pack(fill="x", pady=(8, 8))
        specs = [("due", "Due today", C["accent"]), ("overdue", "Overdue", C["bad"]),
                 ("stuck", "Stuck in one status", C["warn"]), ("groom", "Next grooming", C["info"]),
                 ("follow", "Follow-ups due", "#B49FCC"), ("leaving", "Leaving in 30 days", C["good"])]
        self.today_cards = {}
        for i, (key, title, colour) in enumerate(specs):
            card = KpiCard(grid, title, colour)
            card.grid(row=0, column=i, sticky="nsew", padx=5)
            grid.columnconfigure(i, weight=1, uniform="tc")
            self.today_cards[key] = card

        pw = ttk.PanedWindow(parent, orient="horizontal")
        pw.pack(fill="both", expand=True)
        boxes = []
        for title, cols in (("⚡ Needs your attention", [("why", "Why", 260), ("who", "Assignee", 130)]),
                            ("🧭 Next grooming session", [("who", "Assignee", 130), ("prd", "PRD", 100)]),
                            ("🧳 Leaving soon", [("left", "Days left", 80), ("open", "Open", 60)])):
            lf = ttk.LabelFrame(pw, text=title, padding=6)
            pw.add(lf, weight=3 if title.startswith("⚡") else 2)
            f, tv = self._make_table(lf, cols, height=14)
            tv.configure(show="tree headings")
            tv.column("#0", width=300 if not title.startswith("🧳") else 200, stretch=True)
            tv.heading("#0", text="Ticket" if not title.startswith("🧳") else "Person", anchor="w")
            f.pack(fill="both", expand=True)
            tv.bind("<Double-1>", lambda e, t=tv: self._open_today_item(t))
            boxes.append((lf, tv))
        (self.lf_attention, self.tv_attention), (self.lf_nextgroom, self.tv_nextgroom), (_, self.tv_leaving) = boxes
        self.render_today()

    def _next_groom_date(self):
        today = dt.date.today()
        dates = sorted({d for d in (parse_date(v) for v in (self.config_data.get("groom_dates") or {}).values())
                        if d and d >= today})
        return dates[0] if dates else None

    def render_today(self):
        if not hasattr(self, "tv_attention"):
            return
        today = dt.date.today()
        hour = dt.datetime.now().hour
        hello = "Good morning" if hour < 12 else ("Good afternoon" if hour < 17 else "Good evening")
        who = first_name(self.user_name) if self.user_name else ""
        self.lbl_hello.configure(text=f"{hello}{', ' + who if who else ''} ✿")
        sub = [today.strftime("%A, %d %B %Y")]
        cur = self._current_sprint()
        if cur:
            sp = cur[1]
            _, day, left = sprint_position([sp], today, self.holidays())
            sub.append(f"{cur[0]} · {sp['name']}: day {day} of {sp['working']} ({plural(left, 'working day')} left)")
        prods = sorted((parse_date(c.get("prod_date")), k) for k, c in self._plans().items()
                       if parse_date((c or {}).get("prod_date")) and parse_date(c.get("prod_date")) >= today)
        if prods:
            d, k = prods[0]
            sub.append(f"🚀 {k} goes to production {fmt_date(d)} "
                       f"({plural(count_working_days(today + dt.timedelta(days=1), d, self.holidays()), 'working day')})")
        if not self.issues:
            sub.append("Connect to Redmine (⚙ Settings) to see your tickets here.")
        self.lbl_today_sub.configure(text="     ·     ".join(sub))

        notes = self.config_data.get("ticket_notes") or {}
        att = []
        for it in self.issues:
            due = parse_date(it.get("due_date"))
            reasons = []
            if due and due < today:
                reasons.append(f"overdue {(today - due).days}d")
            elif due == today:
                reasons.append("due today")
            h = self._hist(it)
            if h and h["status_days"] > self.stuck_after():
                reasons.append(f"{h['status_days']} working days in {name_of(it.get('status'))}")
            fu = parse_date((notes.get(str(it["id"])) or {}).get("follow_up"))
            if fu and fu <= today:
                reasons.append("follow-up due")
            m = self.member_for(it)
            if m and parse_date(m.get("relieving")) and 0 <= (parse_date(m["relieving"]) - today).days <= 14:
                reasons.append(f"assignee leaves {fmt_date(parse_date(m['relieving']), False)}")
            if reasons:
                att.append((it, reasons))
        order = lambda r: (0 if any("overdue" in x or "today" in x for x in r[1]) else 1, r[0].get("due_date") or "")
        self.tv_attention.delete(*self.tv_attention.get_children())
        for it, reasons in sorted(att, key=order)[:200]:
            self.tv_attention.insert("", "end", iid=f"a{it['id']}", text=f"#{it['id']}  {it.get('subject', '')}",
                                     values=("; ".join(reasons), name_of(it.get("assigned_to"), "Unassigned")),
                                     tags=("hot",) if any("overdue" in x for x in reasons) else ())
        nxt = self._next_groom_date()
        grooms = self.config_data.get("groom_dates") or {}
        ids = [int(k) for k, v in grooms.items() if parse_date(v) == nxt] if nxt else []
        self.tv_nextgroom.delete(*self.tv_nextgroom.get_children())
        for i in ids:
            it = self.issue_index.get(i) or {"id": i, "subject": "(not in the current ticket list)"}
            state = self._prd_info(it)[0] if i in self.issue_index else "Needed"
            self.tv_nextgroom.insert("", "end", iid=f"g{i}", text=f"#{i}  {it.get('subject', '')}",
                                     values=((self.planned_dev(i) + " (planned)") if self.planned_dev(i) else
                                             name_of(it.get("assigned_to"), "—"), PRD_ICONS[state]))
        self.lf_nextgroom.configure(text=f"🧭 Next grooming session: {fmt_date(nxt)}" if nxt else
                                    "🧭 Next grooming session (set groom dates on tickets)")
        self.tv_leaving.delete(*self.tv_leaving.get_children())
        n_leave = 0
        for m in sorted((m for m in self.team() if m.get("status") == "Active" and parse_date(m.get("relieving"))
                         and parse_date(m["relieving"]) >= today), key=lambda m: m["relieving"]):
            rel = parse_date(m["relieving"])
            left = (rel - today).days
            n_leave += left <= 30
            self.tv_leaving.insert("", "end", iid=f"l{norm_name(m['name'])}", text=f"{m['name']} · {fmt_date(rel, False)}",
                                   values=(left, len(self._open_for(m["name"]))),
                                   tags=("hot",) if left < 14 else ())
        cnt = lambda pred: sum(1 for it, r in att if any(pred(x) for x in r))
        c = self.today_cards
        c["due"].set(str(cnt(lambda x: x == "due today")), "tickets due today")
        c["overdue"].set(str(cnt(lambda x: "overdue" in x)), "past their due date", C["bad"] if cnt(
            lambda x: "overdue" in x) else C["good"])
        c["stuck"].set(str(cnt(lambda x: "working days in" in x)), f"over {plural(self.stuck_after(), 'working day')}")
        c["groom"].set(str(len(ids)), fmt_date(nxt) if nxt else "no session planned")
        c["follow"].set(str(cnt(lambda x: x == "follow-up due")), "your own reminders")
        c["leaving"].set(str(n_leave), "see the Team tab for handover")

    def _open_today_item(self, tv):
        sel = tv.selection()
        if not sel:
            return
        row = sel[0]
        if row.startswith("l"):
            self.nb.select(self.tab_team)
            return
        iid = row[1:]
        self._show_tickets()
        if self.tree.exists(iid):
            self.tree.selection_set(iid)
            self.tree.see(iid)
        else:
            self.load_issue(int(iid))

    # ----- grooming session ------------------------------------------------------
    def start_grooming_session(self):
        from tkinter import simpledialog
        nxt = self._next_groom_date() or dt.date.today()
        txt = simpledialog.askstring("Grooming session", "Which grooming date? (DD-MM-YYYY)",
                                     initialvalue=nxt.strftime("%d-%m-%Y"), parent=self)
        if not txt:
            return
        day = parse_user_date(txt)
        if not day:
            messagebox.showinfo(APP_NAME, "Type the date like 26-10-2026.", parent=self)
            return
        ids = sorted(int(k) for k, v in (self.config_data.get("groom_dates") or {}).items() if parse_date(v) == day)
        if not ids:
            messagebox.showinfo(APP_NAME, f"No tickets have the groom date {fmt_date(day)}.\n\nSelect tickets in the "
                                          "Grooming Workspace and use 📅 Groom date… first.", parent=self)
            return
        missing = [i for i in ids if i not in self.issue_index]
        if missing and self._configured(warn=False):
            client = self.make_client()

            def done(found):
                extra = {it["id"]: it for it in found}
                self._open_session(day, [self.issue_index.get(i) or extra.get(i) or {"id": i, "subject": "?"}
                                         for i in ids])

            self.run_async(lambda: client.issues_by_ids(missing), done, busy="Loading the session's tickets…")
        else:
            self._open_session(day, [self.issue_index.get(i) or {"id": i, "subject": "?"} for i in ids])

    def _open_session(self, day, tickets):
        store = self.config_data.setdefault("groom_sessions", {}).setdefault(day.isoformat(), {})
        win = tk.Toplevel(self)
        win.title(f"Grooming session: {fmt_date(day)}")
        win.configure(bg=C["bg"])
        win.geometry("1200x760")
        state = {"i": 0}
        v_out = tk.StringVar()
        top = ttk.Frame(win, padding=(16, 12), style="Bg.TFrame")
        top.pack(fill="x")
        ttk.Label(top, text=f"Grooming session · {fmt_date(day)}", font=FONTS["h1"],
                  foreground=C["header"]).pack(side="left")
        lbl_prog = ttk.Label(top, text="", style="Muted.TLabel")
        lbl_prog.pack(side="left", padx=14)
        body = ttk.PanedWindow(win, orient="horizontal")
        body.pack(fill="both", expand=True, padx=16, pady=(0, 12))
        lf = ttk.Frame(body, style="Bg.TFrame")
        rf = ttk.Frame(body, style="Bg.TFrame", padding=(12, 0, 0, 0))
        body.add(lf, weight=1)
        body.add(rf, weight=3)
        lb = tk.Listbox(lf, font=FONTS["body"], activestyle="none", relief="flat", highlightthickness=1,
                        highlightbackground=C["line"], selectbackground=C["select"], selectforeground=C["ink"],
                        exportselection=False, selectmode="extended")
        lb.pack(fill="both", expand=True)
        lbb = ttk.Frame(lf, style="Bg.TFrame")
        lbb.pack(fill="x", pady=(6, 0))
        ttk.Label(lf, text="Ctrl/Shift+click to pick several tickets.", style="Muted.TLabel").pack(anchor="w")
        lbl_title = ttk.Label(rf, text="", font=FONTS["h2"], foreground=C["header"], wraplength=760, justify="left")
        lbl_title.pack(anchor="w")
        lbl_meta = ttk.Label(rf, text="", style="Muted.TLabel", wraplength=760, justify="left")
        lbl_meta.pack(anchor="w", pady=(2, 6))
        desc = tk.Text(rf, height=11, wrap="word", font=FONTS["body"], relief="flat", highlightthickness=1,
                       highlightbackground=C["line"], padx=10, pady=8, bg=C["field"])
        desc.pack(fill="both", expand=True)
        ttk.Label(rf, text="Discussion notes", font=FONTS["h3"], foreground=C["header"]).pack(anchor="w", pady=(10, 2))
        notes = tk.Text(rf, height=6, wrap="word", font=FONTS["body"], relief="flat", highlightthickness=1,
                        highlightbackground=C["line"], padx=10, pady=8)
        notes.pack(fill="x")
        orow = ttk.Frame(rf, style="Bg.TFrame")
        orow.pack(fill="x", pady=(8, 0))
        ttk.Label(orow, text="Outcome:").pack(side="left", padx=(0, 6))
        for o in SESSION_OUTCOMES:
            ttk.Radiobutton(orow, text=o, value=o, variable=v_out).pack(side="left", padx=(0, 8))
        nav = ttk.Frame(rf, style="Bg.TFrame")
        nav.pack(fill="x", pady=(12, 0))

        def label(i):
            it = tickets[i]
            done = "✔ " if (store.get(str(it["id"])) or {}).get("outcome") else "   "
            return f"{done}#{it['id']}  {truncate(it.get('subject', ''), 46)}"

        def save_current():
            it = tickets[state["i"]]
            store[str(it["id"])] = {"outcome": v_out.get(), "notes": notes.get("1.0", "end").strip()}
            lb.delete(state["i"])
            lb.insert(state["i"], label(state["i"]))
            self._save_config_quiet()

        def show(i):
            state["i"] = i
            it = tickets[i]
            rec = store.get(str(it["id"])) or {}
            lbl_title.configure(text=f"#{it['id']}  {it.get('subject', '')}")
            prd, _ = self._prd_info(it) if it.get("subject") != "?" else ("Needed", "")
            lbl_meta.configure(text=f"{name_of(it.get('tracker'))} · {name_of(it.get('status'))} · Priority "
                                    f"{name_of(it.get('priority'))} · Assignee {name_of(it.get('assigned_to'), 'Unassigned')}"
                                    f" · Due {it.get('due_date') or '—'} · PRD: {prd}" +
                                    (f"\n👩‍💻 Planned developer for this grooming: {self.planned_dev(it['id'])}"
                                     if self.planned_dev(it.get("id")) else ""))
            desc.configure(state="normal")
            desc.delete("1.0", "end")
            desc.insert("1.0", it.get("description") or "(no description)")
            desc.configure(state="disabled")
            notes.delete("1.0", "end")
            notes.insert("1.0", rec.get("notes", ""))
            v_out.set(rec.get("outcome", ""))
            lb.selection_clear(0, "end")
            lb.selection_set(i)
            lb.see(i)
            n_done = sum(1 for t in tickets if (store.get(str(t["id"])) or {}).get("outcome"))
            lbl_prog.configure(text=f"Ticket {i + 1} of {len(tickets)} · {n_done} decided")

        def go(step):
            save_current()
            show(max(0, min(len(tickets) - 1, state["i"] + step)))

        def pick(_e=None):
            sel = lb.curselection()
            if len(sel) == 1 and sel[0] != state["i"]:
                save_current()
                show(sel[0])

        def mark_selected():
            """Bulk: every ticket picked in the list (or the current one) becomes 'Groomed'."""
            idx = list(lb.curselection()) or [state["i"]]
            save_current()
            for i in idx:
                rec = store.setdefault(str(tickets[i]["id"]), {})
                if not rec.get("outcome"):
                    rec["outcome"] = "Groomed"
                self.config_data.setdefault("groomed", {})[str(tickets[i]["id"])] = {"date": day.isoformat()}
                lb.delete(i)
                lb.insert(i, label(i))
            self._save_config_quiet()
            show(state["i"])
            lb.selection_clear(0, "end")
            for i in idx:
                lb.selection_set(i)
            self._populate_tree()
            self.set_status(f"{plural(len(idx), 'ticket')} marked as groomed ✅.")

        def add_ticket():
            from tkinter import simpledialog
            txt = simpledialog.askstring("Add tickets to this session",
                                         "Ticket number(s) to add, e.g. 10240, 10255:", parent=win)
            ids = [int(x) for x in re.findall(r"\d+", txt or "")]
            ids = [i for i in ids if i not in {t["id"] for t in tickets}]
            if not ids:
                return

            def add(found):
                byid = {it["id"]: it for it in found}
                grooms = self.config_data.setdefault("groom_dates", {})
                for i in ids:
                    it = self.issue_index.get(i) or byid.get(i)
                    if not it:
                        messagebox.showinfo(APP_NAME, f"Couldn't find #{i} in Redmine.", parent=win)
                        continue
                    tickets.append(it)
                    grooms[str(i)] = day.isoformat()
                    lb.insert("end", label(len(tickets) - 1))
                self._save_config_quiet()
                self._refresh_groom_choices()
                self._populate_tree()
                show(len(tickets) - 1)

            missing = [i for i in ids if i not in self.issue_index]
            if missing and self._configured(warn=False):
                client = self.make_client()
                self.run_async(lambda: client.issues_by_ids(missing), add, busy="Fetching tickets…")
            else:
                add([])

        def export():
            save_current()
            self._export_session(day, tickets, store)

        def close():
            save_current()
            win.destroy()

        for i in range(len(tickets)):
            lb.insert("end", label(i))
        lb.bind("<<ListboxSelect>>", pick)
        ttk.Button(lbb, text="✅ Mark selected groomed", style="Accent.TButton", command=mark_selected).pack(
            fill="x")
        ttk.Button(lbb, text="➕ Add a ticket to this session…", command=add_ticket).pack(fill="x", pady=(4, 0))
        ttk.Button(nav, text="◀ Previous", command=lambda: go(-1)).pack(side="left")
        ttk.Button(nav, text="Next ▶", style="Accent.TButton", command=lambda: go(1)).pack(side="left", padx=6)
        ttk.Button(nav, text="Open in Redmine ↗", command=lambda: webbrowser.open(
            f"{self.var_url.get().rstrip('/')}/issues/{tickets[state['i']]['id']}")).pack(side="left", padx=6)
        ttk.Button(nav, text="💡 Suggest developer", command=lambda: self._show_suggestions(
            [tickets[state["i"]]])).pack(side="left")
        ttk.Button(nav, text="Close", command=close).pack(side="right")
        ttk.Button(nav, text="⬇ Export minutes (Excel)…", command=export).pack(side="right", padx=6)
        win.protocol("WM_DELETE_WINDOW", close)
        show(0)

    def _export_session(self, day, tickets, store):
        self._ensure("openpyxl", "openpyxl")
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            return
        path = filedialog.asksaveasfilename(parent=self, title="Save grooming minutes",
                                            initialfile=f"Grooming_Minutes_{day.isoformat()}.xlsx",
                                            defaultextension=".xlsx", filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "Minutes"
        ws["A1"] = f"Grooming session: {fmt_date(day)}"
        ws["A1"].font = Font(size=15, bold=True, color="4A3F3A")
        counts = Counter((store.get(str(t["id"])) or {}).get("outcome") or "Not decided" for t in tickets)
        ws["A2"] = " · ".join(f"{k}: {v}" for k, v in counts.most_common())
        ws["A2"].font = Font(size=10, color="8A8078")
        heads = ["Ticket", "Subject", "Assignee", "Status", "Priority", "Due date", "Outcome", "Discussion notes"]
        for ci, (h, wd) in enumerate(zip(heads, [10, 50, 18, 16, 12, 13, 20, 70]), 1):
            c = ws.cell(row=4, column=ci, value=h)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="C2665C")
            ws.column_dimensions[c.column_letter].width = wd
        for r, it in enumerate(tickets, 5):
            rec = store.get(str(it["id"])) or {}
            vals = [it["id"], it.get("subject", ""), name_of(it.get("assigned_to"), "Unassigned"),
                    name_of(it.get("status")), name_of(it.get("priority")), parse_date(it.get("due_date")),
                    rec.get("outcome", ""), rec.get("notes", "")]
            for ci, val in enumerate(vals, 1):
                c = ws.cell(row=r, column=ci, value=val)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(val, dt.date):
                    c.number_format = "DD-MMM-YYYY"
            c = ws.cell(row=r, column=1)
            c.hyperlink = f"{self.var_url.get().rstrip('/')}/issues/{it['id']}"
            c.font = Font(color="C2665C", underline="single")
        ws.freeze_panes = "A5"
        try:
            wb.save(path)
        except OSError as exc:
            messagebox.showerror(APP_NAME, f"Could not save the file (is it open in Excel?):\n{exc}", parent=self)
            return
        self._offer_open(path)

    def _prd_info(self, issue: dict):
        key = str(issue.get("id"))
        exported = (self.config_data.get("exported_prds") or {}).get(key)
        cached = self.hist_cache.get(issue.get("id"))
        journals = issue.get("journals") or (cached["issue"].get("journals") if cached else None)
        override = (self.config_data.get("prd_overrides") or {}).get(key)
        return prd_state(issue, bool(self.drafts.get(key, "").strip()), exported, journals, override)

    def set_prd_mark(self, state):
        """state: 'Has PRD', 'Not needed', or None to let the app decide again."""
        ids = [i for i in self.tree.selection() if i.isdigit() and int(i) in self.issue_index]
        if not ids and self.current_issue:
            ids = [str(self.current_issue["id"])]
        if not ids:
            messagebox.showinfo(APP_NAME, "Select one or more tickets in the list first.", parent=self)
            return
        marks = self.config_data.setdefault("prd_overrides", {})
        note = ""
        if state == "Has PRD":
            from tkinter import simpledialog
            note = simpledialog.askstring("PRD location", "Where is the PRD? (optional: a link, file name or "
                                                          "'in comments')", parent=self) or ""
        for i in ids:
            if state:
                marks[i] = {"state": state, "note": note.strip(), "date": dt.date.today().isoformat()}
            else:
                marks.pop(i, None)
                (self.config_data.get("exported_prds") or {}).pop(i, None)
        self._save_config_quiet()
        self._populate_tree()
        if self.current_issue:
            self._show_issue(self.current_issue)
        label = state or "decided by the app"
        self.set_status(f"PRD status for {plural(len(ids), 'ticket')} set to: {label}.")

    def _tree_menu(self, event):
        row = self.tree.identify_row(event.y)
        if row and row not in self.tree.selection():
            self.tree.selection_set(row)
        self._ticket_menu(event.x_root, event.y_root)

    def _ticket_menu(self, x_root, y_root):
        menu = tk.Menu(self, tearoff=0)
        menu.add_command(label="✔  Mark: Has PRD", command=lambda: self.set_prd_mark("Has PRD"))
        menu.add_command(label="—  Mark: PRD not needed", command=lambda: self.set_prd_mark("Not needed"))
        menu.add_command(label="↺  Let the app decide", command=lambda: self.set_prd_mark(None))
        menu.add_separator()
        menu.add_command(label="📅  Set groom date…", command=self.set_groom_date)
        menu.add_command(label="✖  Clear groom date", command=lambda: self.set_groom_date(clear=True))
        menu.add_command(label="✅  Mark as groomed", command=lambda: self.mark_groomed(True))
        menu.add_command(label="↺  Not groomed yet", command=lambda: self.mark_groomed(False))
        menu.add_command(label="📝  My note / follow-up…", command=self.edit_ticket_note)
        menu.add_command(label="💡  Suggest developer", command=self.suggest_for_selected)
        menu.add_command(label="✉  Copy follow-up message", command=self.copy_ticket_message)
        menu.add_separator()
        menu.add_command(label="Open in Redmine ↗", command=self._open_in_browser)
        try:
            menu.tk_popup(x_root, y_root)
        finally:
            menu.grab_release()

    # ----- board / list switch -----------------------------------------------------
    def _set_ws_view(self, view, save=True):
        self.var_ws_view.set(view)
        if view in ("board", "people", "groom"):
            self.ws_list.pack_forget()
            self.board.pack(fill="both", expand=True, before=self.ws_bottom)
            self.chk_hide_empty.pack(side="left", padx=(12, 0), after=self.btn_list)
            self.chk_compact.pack(side="left", padx=(8, 0), after=self.chk_hide_empty)
            self._render_board()
            self.after(120, lambda: self._ws_sash(0.70))
        else:
            self.board.pack_forget()
            self.ws_list.pack(fill="both", expand=True, before=self.ws_bottom)
            self.chk_hide_empty.pack_forget()
            self.chk_compact.pack_forget()
            self.after(120, lambda: self._ws_sash(0.60))
        self.btn_board.configure(style="Accent.TButton" if view == "board" else "TButton")
        self.btn_people.configure(style="Accent.TButton" if view == "people" else "TButton")
        self.btn_groomv.configure(style="Accent.TButton" if view == "groom" else "TButton")
        self.btn_list.configure(style="Accent.TButton" if view == "list" else "TButton")
        if save:
            self.config_data["ws_view"] = view
            self._save_config_quiet()

    def _ws_sash(self, share, tries=0):
        """Give the board (or list) its share of the width once the tab is actually on screen."""
        try:
            if len(self.ws_pw.panes()) < 2:
                return
            self.update_idletasks()
            w = self.ws_pw.winfo_width()
            if w < 400 or not self.ws_pw.winfo_ismapped():
                if tries < 40:
                    self.after(250, lambda: self._ws_sash(share, tries + 1))
                return
            self.ws_pw.sashpos(0, int(w * share))
        except tk.TclError:
            pass

    def _toggle_details(self):
        panes = [str(p) for p in self.ws_pw.panes()]
        if str(self.ws_right) in panes:
            self.ws_pw.forget(self.ws_right)
            self.btn_details.configure(text="◨ Show details")
        else:
            self.ws_pw.add(self.ws_right, weight=2)
            self.btn_details.configure(text="◧ Hide details")

    def _board_option_changed(self):
        self.config_data["board_hide_empty"] = bool(self.var_hide_empty.get())
        self.config_data["board_compact"] = bool(self.var_compact.get())
        self._save_config_quiet()
        self._render_board()

    def _schedule_board(self):
        if self.var_ws_view.get() not in ("board", "people", "groom"):
            return
        if self._board_job:
            self.after_cancel(self._board_job)
        self._board_job = self.after(150, self._render_board)

    def _render_board(self):
        self._board_job = None
        view = self.var_ws_view.get() if hasattr(self, "var_ws_view") else "list"
        if not hasattr(self, "board") or view not in ("board", "people", "groom"):
            return
        self.board.compact = bool(self.var_compact.get())
        grooms = self.config_data.get("groom_dates") or {}
        today = dt.date.today()
        by_status = {}
        names = {}
        for it, state, _reason, flags in self.visible_rows:
            st = name_of(it.get("status"))
            k = status_key(st)
            names.setdefault(k, st)
            hist = self._hist(it)
            due = parse_date(it.get("due_date"))
            groom = parse_date(grooms.get(str(it["id"])))
            prio = name_of(it.get("priority"))
            who = name_of(it.get("assigned_to"), "Unassigned")
            days = f"{hist['status_days']} d" if hist else "…"
            by_status.setdefault(k, []).append({
                "status": st, "skey": k,
                "meta": (f"📌 {st}    ⏱ {days} in status" if view == "people" else
                         (f"👤 {who}  ·  📌 {st}" if view == "groom" else f"👤 {who}    ⏱ {days} in status")),
                "groomed": fmt_date(self.groomed_on(it["id"]), False) if self.groomed_on(it["id"]) else "",
                "groom_d": groom, "plan_dev": self.planned_dev(it["id"]),
                "id": it["id"], "subject": it.get("subject", ""), "prio": prio,
                "prio_colour": priority_color(prio) if priority_color(prio) != C["muted"] else "#9C9188",
                "rank": priority_rank(prio), "assignee": name_of(it.get("assigned_to"), "Unassigned"),
                "due": fmt_date(due, False) if due else "", "due_d": due or dt.date(9999, 1, 1),
                "overdue": bool(due and due < today), "in_status": f"{hist['status_days']} d" if hist else "…",
                "prd": PRD_ICONS[state], "prd_state": state, "groom": fmt_date(groom) if groom else "",
                "sprint": self.plan_info(it)["sprint"], "track": self._track(it),
                "week": self.week_text(it, short=True), "release": self.plan_info(it)["release"], "flags": flags,
                "parent": self._parent_label(it)})
        for cards in by_status.values():
            cards.sort(key=lambda c: (c["rank"], c["due_d"], c["id"]))
        if view in ("people", "groom"):
            allc = [c for cs in by_status.values() for c in cs]
            self.board.render(self._people_columns(allc) if view == "people" else self._groom_columns(allc))
            if self.current_issue:
                self.board.select(self.current_issue.get("id"))
            return
        f_status = self.var_status_filter.get()
        order = [status_key(s) for s in KANBAN_ORDER]
        label = {status_key(s): s for s in KANBAN_ORDER}
        label.update(names)
        known = [k for k in order if k not in KANBAN_SIDE and k not in KANBAN_ONLY_IF_USED]
        extra = sorted((k for k in by_status if k not in order), key=lambda k: names[k].lower())
        seq = known + extra + [k for k in order if k in KANBAN_SIDE] + [k for k in order if k in KANBAN_ONLY_IF_USED]
        cols = []
        for k in seq:
            cards = by_status.get(k, [])
            if f_status != "All statuses" and status_key(f_status) != k:
                continue
            if not cards and (k in KANBAN_ONLY_IF_USED or self.var_hide_empty.get()):
                continue
            cols.append((label.get(k, k), "side" if k in KANBAN_SIDE else "main", cards, ""))
        self.board.render(cols)
        if self.current_issue:
            self.board.select(self.current_issue.get("id"))

    # ----- grooming done / by groom date ------------------------------------------------
    def groomed_on(self, issue_id):
        """Date a ticket was groomed (marked by you, or given an outcome in a grooming session), else None."""
        rec = (self.config_data.get("groomed") or {}).get(str(issue_id))
        if rec:
            return parse_date(rec.get("date")) or dt.date.today()
        for day, store in (self.config_data.get("groom_sessions") or {}).items():
            if (store.get(str(issue_id)) or {}).get("outcome") not in (None, "", "Not discussed"):
                return parse_date(day)
        return None

    def mark_groomed(self, done=True, ids=None):
        ids = ids or self._selected_ticket_ids()
        if not ids:
            messagebox.showinfo(APP_NAME, "Select one or more tickets first.", parent=self)
            return
        store = self.config_data.setdefault("groomed", {})
        grooms = self.config_data.get("groom_dates") or {}
        for i in ids:
            if done:
                d = parse_date(grooms.get(str(i)))
                store[str(i)] = {"date": (d if d and d <= dt.date.today() else dt.date.today()).isoformat()}
            else:
                store.pop(str(i), None)
                for sess in (self.config_data.get("groom_sessions") or {}).values():
                    if str(i) in sess:
                        sess[str(i)]["outcome"] = ""
        self._save_config_quiet()
        self._populate_tree()
        if self.current_issue:
            self._show_issue(self.current_issue)
        if hasattr(self, "board"):
            self.board.select_many(set())
        self.set_status(f"{plural(len(ids), 'ticket')} marked as {'groomed ✅' if done else 'not groomed'}.")

    def _groom_columns(self, cards):
        """One column per grooming date (oldest first), then tickets with no groom date."""
        today = dt.date.today()
        by_day, none = {}, []
        for c in cards:
            (by_day.setdefault(c["groom_d"], []) if c["groom_d"] else none).append(c)
        cols = []
        for d in sorted(by_day):
            cs = sorted(by_day[d], key=lambda c: (bool(c["groomed"]), c["rank"], c["id"]))
            done = sum(1 for c in cs if c["groomed"])
            when = "today" if d == today else ("past" if d < today else f"in {(d - today).days} d")
            kind = "warn" if d < today and done < len(cs) else ("main" if d >= today else "side")
            cols.append((fmt_date(d), kind, cs, f"{when} · {done} of {len(cs)} groomed ✅"))
        if none and not self.var_hide_empty.get():
            cols.append(("No groom date yet", "side", sorted(none, key=lambda c: (c["rank"], c["id"])),
                         "right-click a card → 📅 Set groom date"))
        if not cols:
            cols.append(("No groom dates yet", "side", [], "select tickets → 📅 Groom date…"))
        return cols

    # ----- why isn't my ticket here? ------------------------------------------------------
    def find_ticket(self):
        from tkinter import simpledialog
        txt = simpledialog.askstring("Find a ticket", "Ticket number (e.g. 10240):", parent=self)
        if not txt:
            return
        m = re.search(r"\d+", txt)
        if not m:
            return
        iid = int(m.group())
        if iid in self.issue_index:
            self._explain_hidden(self.issue_index[iid])
            return
        if not self._configured():
            return
        client = self.make_client()
        self.run_async(lambda: client.issue(iid), lambda it: self._explain_not_loaded(iid, it),
                       busy=f"Looking up #{iid} in Redmine…",
                       on_error=lambda exc: messagebox.showinfo(
                           APP_NAME, f"Redmine couldn't find ticket #{iid}.\n\nIt may not exist, or your Redmine "
                                     f"account may not be allowed to see it.\n\n({str(exc).splitlines()[0]})",
                           parent=self))

    def _explain_hidden(self, it):
        """The ticket IS loaded: say which filter hides it, and offer to clear that filter."""
        iid = it["id"]
        if self.tree.exists(str(iid)):
            self.show_in_workspace(iid)
            messagebox.showinfo(APP_NAME, f"#{iid} is here. I've highlighted it and opened its details.",
                                parent=self)
            return
        reasons, fixes = [], []
        st, prio = name_of(it.get("status")), name_of(it.get("priority"))
        who = name_of(it.get("assigned_to"), "Unassigned")
        rel = self.plan_info(it)["release"] or "No release"
        grooms = self.config_data.get("groom_dates") or {}
        groom = parse_date(grooms.get(str(iid)))
        checks = [
            (self.var_status_filter.get() not in ("All statuses", st), f"Status filter is '{self.var_status_filter.get()}'"
             f" but the ticket is '{st}'", lambda: self.var_status_filter.set("All statuses")),
            (self.var_assignee_filter.get() not in ("All assignees", who), f"Assignee filter is "
             f"'{self.var_assignee_filter.get()}' but it's assigned to {who}",
             lambda: self.var_assignee_filter.set("All assignees")),
            (self.var_priority_filter.get() not in ("All priorities", prio), f"Priority filter is "
             f"'{self.var_priority_filter.get()}' but it's {prio}", lambda: self.var_priority_filter.set("All priorities")),
            (bool(self.mp_release.selected) and rel not in self.mp_release.selected, f"Release filter doesn't include "
             f"'{rel}'", lambda: self.mp_release.set_selected([])),
            (self.var_sprint_filter.get() != "All sprints", f"Sprint filter is '{self.var_sprint_filter.get()}'",
             lambda: self.var_sprint_filter.set("All sprints")),
            (self.var_track_filter.get() != "All tracks", f"Track filter is '{self.var_track_filter.get()}' and this "
             f"ticket is '{self._track(it)}'", lambda: self.var_track_filter.set("All tracks")),
            (self.var_prd_filter.get() != "Any", f"PRD filter is '{self.var_prd_filter.get()}'",
             lambda: self.var_prd_filter.set("Any")),
            (self.var_check_filter.get() != "Any", "Status check filter shows flagged tickets only",
             lambda: self.var_check_filter.set("Any")),
            (self.var_groom_filter.get() != "Any", f"Groom date filter is '{self.var_groom_filter.get()}'"
             + (f" (this ticket: {fmt_date(groom)})" if groom else " (this ticket has no groom date)"),
             lambda: self.var_groom_filter.set("Any")),
            (bool(self.var_due_from.get().strip() or self.var_due_to.get().strip()), "Due-date range filter is on"
             f" (this ticket is due {it.get('due_date') or 'never'})",
             lambda: (self.var_due_from.set(""), self.var_due_to.set(""))),
            (bool(self.var_search.get().strip()), f"Search box contains '{self.var_search.get().strip()}'",
             lambda: self.var_search.set("")),
        ]
        for hit, why, fix in checks:
            if hit:
                reasons.append(why)
                fixes.append(fix)
        if not reasons:
            reasons.append("It's grouped under a parent ticket that is collapsed. Untick 'Group under parent'.")
            fixes.append(lambda: self.var_group.set(False))
        msg = (f"#{iid} '{truncate(it.get('subject', ''), 70)}' is loaded, but hidden by:\n\n• " +
               "\n• ".join(reasons) + "\n\nClear these filters and show it?")
        if messagebox.askyesno(APP_NAME, msg, parent=self):
            for f in fixes:
                f()
            self._on_list_filter_changed()
            self.after(300, lambda: self.show_in_workspace(iid))

    def _explain_not_loaded(self, iid, it):
        """The ticket is NOT in the loaded list: work out why from Redmine's own data."""
        if not it:
            messagebox.showinfo(APP_NAME, f"Redmine returned nothing for #{iid}.", parent=self)
            return
        reasons = []
        status = it.get("status") or {}
        if status.get("is_closed") or name_of(status).lower() in ("closed", "rejected", "resolved", "done"):
            reasons.append(f"It is closed (status '{name_of(status)}'). This app only loads OPEN tickets.")
        pid = (it.get("project") or {}).get("id")
        pname = name_of(it.get("project"))
        pids, mine, _label = self._scope()
        if pids and pid not in pids:
            reasons.append(f"It belongs to project '{pname}', which isn't ticked in the Projects list at the top. "
                           f"(Sub-projects have to be ticked separately.)")
        if mine:
            who = name_of(it.get("assigned_to"), "nobody")
            if who != getattr(self, "user_name", who + "?"):
                reasons.append(f"'Assigned to me' is selected, but it's assigned to {who}.")
        if len(self.issues) < self.issue_total:
            reasons.append(f"Only the first {len(self.issues)} of {self.issue_total} open tickets were loaded.")
        if not reasons:
            reasons.append("It may have been created or reopened after the last refresh. Click ↻ Refresh.")
        msg = (f"#{iid} '{truncate(it.get('subject', ''), 70)}' ({pname}, {name_of(status)}, "
               f"{name_of(it.get('assigned_to'), 'Unassigned')}) isn't in this app because:\n\n• " +
               "\n• ".join(reasons) + "\n\nOpen it in Redmine?")
        if messagebox.askyesno(APP_NAME, msg, parent=self):
            webbrowser.open(f"{self.var_url.get().rstrip('/')}/issues/{iid}")

    def _people_columns(self, cards):
        """One column per person: Unassigned first, then your active developers (even with no tickets,
        so you can see who is free), then anyone else holding tickets."""
        order = {status_key(s): i for i, s in enumerate(KANBAN_ORDER)}
        by_person = {}
        for c in cards:
            by_person.setdefault(c["assignee"], []).append(c)
        for cs in by_person.values():
            cs.sort(key=lambda c: (order.get(c["skey"], 50), c["rank"], c["due_d"], c["id"]))
        team = [m for m in (self.config_data.get("team") or []) if m.get("name")]
        info = {norm_name(m["name"]): m for m in team}
        today = dt.date.today()
        hide_empty = self.var_hide_empty.get()
        f_assignee = self.var_assignee_filter.get()
        filtered = f_assignee != "All assignees" or self.var_search.get().strip()

        def is_dev(m):
            d = (m.get("designation") or "").lower()
            return m.get("status", "Active") == "Active" and ("develop" in d or "lead" in d)

        names = []
        if "Unassigned" in by_person:
            names.append("Unassigned")
        holders = {n for n in by_person if n != "Unassigned"}
        for m in team:
            hit = next((n for n in holders if norm_name(n) == norm_name(m["name"])), None)
            if hit:
                names.append(hit)
            elif is_dev(m) and not hide_empty and not filtered:
                names.append(m["name"])
        names += sorted((n for n in holders if n not in names), key=str.lower)
        cols = []
        for n in names:
            cs = by_person.get(n, [])
            m = info.get(norm_name(n), {})
            note, kind = m.get("designation") or "", "main"
            if n == "Unassigned":
                note, kind = "needs an owner", "side"
            rel = parse_date(m.get("relieving"))
            if rel and rel >= today:
                note, kind = f"leaves {fmt_date(rel, False)} ({(rel - today).days} d)", "warn"
            if cs:
                hi = sum(1 for c in cs if c["rank"] <= 4)
                late = sum(1 for c in cs if c["overdue"])
                extra = [x for x in (f"{hi} high+" if hi else "", f"{late} overdue" if late else "") if x]
                if extra:
                    note = (note + " · " if note else "") + ", ".join(extra)
            cols.append((n, kind, cs, note))
        return cols

    def _board_select(self, issue_id):
        self.selbar.pack_forget()
        self.board.select(issue_id)
        key = str(issue_id)
        if self.tree.exists(key):
            if self.tree.selection() != (key,):
                self.tree.selection_set(key)
        else:
            self.load_issue(issue_id)

    def _board_multi(self, ids):
        """Several cards picked: mirror them in the list selection and show the bulk-action bar."""
        ids = [str(i) for i in ids if self.tree.exists(str(i))]
        if len(ids) > 1:
            self._multi_sync = True
            self.tree.selection_set(ids)
            self._multi_sync = False
            self.lbl_selbar.configure(text=f"{len(ids)} tickets selected")
            if not self.selbar.winfo_ismapped():
                self.selbar.pack(fill="x", pady=(0, 6), before=self.board)
        else:
            self.selbar.pack_forget()

    def _board_open(self, issue_id):
        webbrowser.open(f"{self.var_url.get().rstrip('/')}/issues/{issue_id}")

    def _board_menu(self, event, issue_id):
        self._board_select(issue_id)
        self._ticket_menu(event.x_root, event.y_root)

    def _clear_ws_filters(self):
        for var, val in ((self.var_status_filter, "All statuses"), (self.var_priority_filter, "All priorities"),
                         (self.var_assignee_filter, "All assignees"), (self.var_sprint_filter, "All sprints"),
                         (self.var_track_filter, "All tracks"), (self.var_prd_filter, "Any"),
                         (self.var_check_filter, "Any"), (self.var_groom_filter, "Any"), (self.var_due_from, ""),
                         (self.var_due_to, ""), (self.var_search, "")):
            var.set(val)
        self.mp_release.set_selected([])
        self._on_list_filter_changed()
        self.set_status("All filters cleared.")

    def _on_list_filter_changed(self):
        self.config_data["status_filter"] = self.var_status_filter.get()
        self.config_data["release_selected"] = list(self.mp_release.selected)
        self._populate_tree()

    def _populate_tree(self):
        query = self.var_search.get().strip().lower()
        f_status = self.var_status_filter.get()
        f_prd = self.var_prd_filter.get()
        f_sprint = self.var_sprint_filter.get()
        f_track = self.var_track_filter.get()
        f_releases = set(self.mp_release.selected)
        d_from, d_to = parse_user_date(self.var_due_from.get()), parse_user_date(self.var_due_to.get())
        f_groom = self.var_groom_filter.get()
        grooms = self.config_data.get("groom_dates") or {}
        f_check = self.var_check_filter.get()
        f_assignee = self.var_assignee_filter.get()
        f_priority = self.var_priority_filter.get()
        today = dt.date.today()
        self.tree.delete(*self.tree.get_children())
        self.visible_rows = []
        counts = Counter()
        flagged = 0
        for it in self.issues:
            subject = it.get("subject", "")
            prio = name_of(it.get("priority"))
            status = name_of(it.get("status"))
            pi = self.plan_info(it)
            sprint = pi["sprint"]
            track = self._track(it)
            release = pi["release"]
            if f_status != "All statuses" and status != f_status:
                continue
            if f_sprint == "No sprint" and sprint:
                continue
            if f_sprint not in ("All sprints", "No sprint") and f_sprint not in self._tags(it) \
                    and f_sprint.lower() != sprint.lower():
                continue
            if f_track == TRACK_TAG and track != "2.0":
                continue
            if f_track == "Legacy" and track != "Legacy":
                continue
            if f_releases and (release or "No release") not in f_releases:
                continue
            due_d = parse_date(it.get("due_date"))
            if (d_from or d_to) and (not due_d or (d_from and due_d < d_from) or (d_to and due_d > d_to)):
                continue
            groom = parse_date(grooms.get(str(it["id"])))
            if f_groom == "Not set" and groom:
                continue
            if f_groom == "Set" and not groom:
                continue
            if f_groom == "✅ Groomed" and not self.groomed_on(it["id"]):
                continue
            if f_groom == "Not groomed yet" and (not groom or self.groomed_on(it["id"])):
                continue
            if f_groom in self.groom_filter_dates and groom != self.groom_filter_dates[f_groom]:
                continue
            if query and query not in subject.lower() and query not in str(it["id"]) \
                    and query not in prio.lower() and query not in " ".join(self._tags(it)).lower():
                continue
            state, reason = self._prd_info(it)
            if f_prd.startswith("Needs") and state not in ("Needed", "Partial"):
                continue
            if f_prd in PRD_STATES and state != f_prd:
                continue
            flags = self._flags(it)
            hist = self._hist(it)
            if f_check.startswith("⚠") and not flags:
                continue
            assignee = name_of(it.get("assigned_to"), "Unassigned")
            if f_assignee != "All assignees" and assignee != f_assignee:
                continue
            if f_priority != "All priorities" and prio != f_priority:
                continue
            tag = next((f"p_{k}" for k in PRIORITY_COLORS if k in prio.lower()), "")
            self.tree.insert("", "end", iid=str(it["id"]), text=f"#{it['id']}   {subject}", open=True,
                             values=(status, f"{hist['status_days']} d" if hist else "…", assignee,
                                     (groom.strftime("%a %d-%b") if groom else "") +
                                     (f" · {first_name(self.planned_dev(it['id']))}" if self.planned_dev(it["id"])
                                      else "") +
                                     (" ✅" if self.groomed_on(it["id"]) else ""), PRD_ICONS[state],
                                     ("⚠ " + "; ".join(flags)) if flags else "",
                                     pi["sprint_txt"] or "—", self.week_text(it) or "—", track, prio,
                                     it.get("due_date") or "—", pi["release_txt"] or "—"),
                             tags=(tag,) if tag else ())
            self.visible_rows.append((it, state, reason, flags))
            counts[state] += 1
            flagged += bool(flags)

        if self.var_group.get():
            self._nest_under_parents()

        shown = len(self.visible_rows)
        text = f"{shown} shown · {self.issue_total} open"
        if len(self.issues) < self.issue_total:
            text += f" (first {len(self.issues)} loaded)"
        if shown:
            text += (f" · ⚠ {flagged} status check\nPRD: {counts['Has PRD']} ready · "
                     f"{plural(counts['Draft'], 'draft')} · {counts['Partial']} partial · {counts['Needed']} needed · "
                     f"{counts['Not needed']} not needed")
        self.lbl_count.configure(text=text)
        if self.draft_key != "scratch" and self.tree.exists(self.draft_key):
            self.tree.selection_set(self.draft_key)
            self.tree.see(self.draft_key)
        if hasattr(self, "tv_attention"):
            self.after_idle(self.render_today)
        self._schedule_board()

    def _nest_under_parents(self):
        visible = {str(it["id"]) for it, *_ in self.visible_rows}

        def ensure(pid, depth=0):
            """Return the tree row for ticket pid, adding a grey placeholder if it isn't listed."""
            key = str(pid)
            if self.tree.exists(key):
                return key
            par = self.issue_index.get(pid) or self.extra_issues.get(pid) or {}
            grand = (par.get("parent") or {}).get("id")
            parent_row = ensure(grand, depth + 1) if grand and depth < 6 else ""
            status = name_of(par.get("status"), "?")
            self.tree.insert(parent_row, "end", iid=key, open=True, tags=("ghost",),
                             text=f"#{pid}   {par.get('subject', '(parent ticket)')}",
                             values=(status, "", name_of(par.get("assigned_to"), ""), "", "",
                                     "parent (not in this list)", "", "", "",
                                     name_of(par.get("priority"), ""), par.get("due_date") or "",
                                     name_of(par.get("fixed_version"), "")))
            return key

        for it, *_ in list(self.visible_rows):
            pid = (it.get("parent") or {}).get("id")
            if not pid:
                continue
            row = str(it["id"])
            target = str(pid) if str(pid) in visible else ensure(pid)
            # avoid moving a row under its own descendant
            node, ok = target, True
            while node:
                if node == row:
                    ok = False
                    break
                node = self.tree.parent(node)
            if ok:
                self.tree.move(row, target, "end")

    def _on_tree_select(self, _event=None):
        sel = self.tree.selection()
        if not sel:
            return
        if getattr(self, "_multi_sync", False):
            return
        if len(sel) > 1:
            self.set_status(f"{len(sel)} tickets selected. Right-click for ✅ Mark as groomed, 📅 groom date and more.")
        issue_id = int(sel[0])
        if hasattr(self, "board"):
            self.board.select(issue_id)
        if (sel[0] == self.draft_key and self.current_issue
                and self.current_issue.get("id") == issue_id):
            return
        self.load_issue(issue_id)

    def load_issue(self, issue_id: int):
        self._stash_draft()
        self._load_draft(str(issue_id))
        self._detail_token += 1
        token = self._detail_token
        summary = self.issue_index.get(issue_id)
        if summary:
            self.current_issue = summary
            self._show_issue(summary)
        else:
            self.current_issue = None
            self.lbl_subject.configure(text=f"Loading ticket #{issue_id}…")
        client = self.make_client()

        def done(issue):
            if token != self._detail_token or not issue:
                return
            self.current_issue = issue
            if "journals" in issue:
                self.hist_cache[issue_id] = {"updated_on": issue.get("updated_on"), "issue": issue}
            self._show_issue(issue)
            self.set_status(f"Ticket #{issue_id} loaded.")

        self.run_async(lambda: client.issue(issue_id), done, busy=f"Loading ticket #{issue_id}…")

    def _show_issue(self, issue: dict):
        prio = name_of(issue.get("priority"))
        self.lbl_prio.configure(text=f"  {prio.upper()}  ", bg=priority_color(prio))
        self.lbl_id.configure(text=f"#{issue.get('id')}  ·  {name_of(issue.get('tracker'))}")
        self.lbl_subject.configure(text=issue.get("subject") or "(no subject)")
        meta = [f"Status: {name_of(issue.get('status'))}",
                f"Project: {name_of(issue.get('project'))}",
                f"Assignee: {name_of(issue.get('assigned_to'), 'Unassigned')}",
                f"Updated: {(issue.get('updated_on') or '')[:10]}"]
        if issue.get("due_date"):
            meta.append(f"Due: {issue['due_date']}")
        if issue.get("estimated_hours"):
            meta.append(f"Estimate: {issue['estimated_hours']} h")
        state, reason = self._prd_info(issue)
        line2 = [f"Release: {name_of(issue.get('fixed_version'))}",
                 f"Sprint: {self._sprint(issue) or '—'}",
                 f"Track: {self._track(issue)}",
                 f"Tags: {', '.join(self._tags(issue)) or '—'}"]
        if issue.get("parent"):
            line2.append(f"Parent: {self._parent_label(issue)}")
        line2.append(f"PRD: {state} ({reason})")
        groom = parse_date((self.config_data.get("groom_dates") or {}).get(str(issue.get("id"))))
        if groom:
            line2.append(f"Groom date: {fmt_date(groom)}")
        note = (self.config_data.get("ticket_notes") or {}).get(str(issue.get("id"))) or {}
        if note.get("note"):
            line2.append(f"📝 {truncate(' '.join(note['note'].split()), 70)}")
        if parse_date(note.get("follow_up")):
            line2.append(f"Follow up: {fmt_date(parse_date(note['follow_up']))}")
        self.lbl_meta.configure(text="     •     ".join(meta) + "\n" + "     •     ".join(line2))
        flags = self._flags(issue)
        self.lbl_flags.configure(text=("⚠ Status check: " + "; ".join(flags)) if flags else "")
        desc = issue.get("description") or "(This ticket has no description.)"
        self.txt_desc.configure(state="normal")
        self.txt_desc.delete("1.0", "end")
        self.txt_desc.insert("1.0", desc)
        self.txt_desc.configure(state="disabled")
        self.btn_open.configure(state="normal")

    def _open_in_browser(self):
        if self.current_issue:
            webbrowser.open(f"{self.var_url.get().rstrip('/')}/issues/{self.current_issue['id']}")

    # ----- PRD workspace -------------------------------------------------
    def _prd_text(self) -> str:
        return self.txt_prd.get("1.0", "end-1c")

    def _set_prd(self, text: str):
        self.txt_prd.delete("1.0", "end")
        self.txt_prd.insert("1.0", text)
        self.txt_prd.edit_reset()
        self._update_wc()

    def _update_wc(self):
        words = len(self._prd_text().split())
        self.lbl_wc.configure(text=plural(words, "word"))

    def _stash_draft(self):
        text = self._prd_text()
        if text.strip():
            self.drafts[self.draft_key] = text
        else:
            self.drafts.pop(self.draft_key, None)
        try:
            save_json(DRAFTS_PATH, self.drafts)
        except OSError:
            pass

    def _load_draft(self, key: str):
        self.draft_key = key
        self._set_prd(self.drafts.get(key, ""))
        if key == "scratch":
            self.lbl_draft.configure(text="Scratch pad · auto-saved")
        else:
            self.lbl_draft.configure(text=f"#{key} · auto-saved")

    def generate_prd(self):
        if not self.current_issue:
            messagebox.showinfo(APP_NAME, "Select a ticket in the list first, then generate the PRD.", parent=self)
            return
        if self._prd_text().strip() and not messagebox.askyesno(
                APP_NAME, "The workspace already contains text.\n\nReplace it with a freshly generated PRD?",
                parent=self):
            return
        self._set_prd(build_prd(self.current_issue, tags=self._tags(self.current_issue)))
        self._stash_draft()
        self._populate_tree()
        self._show_issue(self.current_issue)
        self.set_status(f"PRD generated from #{self.current_issue.get('id')} in the company template. "
                        "Edit freely; it auto-saves.")

    def copy_prd_prompt(self):
        if not self.current_issue:
            messagebox.showinfo(APP_NAME, "Select a ticket first.", parent=self)
            return
        text = prd_ai_prompt(self.current_issue, self._tags(self.current_issue),
                             self._parent_label(self.current_issue), self.var_url.get())
        self.clipboard_clear()
        self.clipboard_append(text)
        messagebox.showinfo(APP_NAME, "A ready-made request with this ticket's details is on your clipboard.\n\n"
                                      "1. Open Claude (claude.ai) and paste it (Ctrl + V).\n"
                                      "2. Copy Claude's answer.\n"
                                      "3. Paste it into the PRD workspace here, then Export to Word.\n\n"
                                      "Only this one ticket is included.", parent=self)

    def export_prd(self):
        text = self._prd_text()
        if not text.strip():
            messagebox.showinfo(APP_NAME, "The PRD workspace is empty. Nothing to export yet.", parent=self)
            return
        base = "PRD"
        if self.current_issue and self.draft_key != "scratch":
            slug = re.sub(r"[^A-Za-z0-9]+", "_", self.current_issue.get("subject", ""))[:40].strip("_")
            base = f"PRD_{self.current_issue.get('id')}_{slug}" if slug else f"PRD_{self.current_issue.get('id')}"
        path = filedialog.asksaveasfilename(
            parent=self, title="Export PRD", initialfile=f"{base}.docx", defaultextension=".docx",
            filetypes=[("Word document", "*.docx"), ("Markdown", "*.md"), ("Text file", "*.txt")])
        if not path:
            return
        ext = Path(path).suffix.lower()
        try:
            if ext == ".docx":
                if not self._ensure("docx", "python-docx"):
                    messagebox.showinfo(APP_NAME, "Word export needs one more component.\n\nIn the VS Code "
                                                  "terminal type:\n\n    python -m pip install python-docx\n\n"
                                                  "then start the app again.", parent=self)
                    return
                prd_to_docx(parse_doc(text, "prd"), path)
            else:
                with open(path, "w", encoding="utf-8") as fh:
                    fh.write(doc_to_markdown(parse_doc(text, "prd")) if ext == ".md" else text)
        except OSError as exc:
            messagebox.showerror(APP_NAME, f"Could not save the file (is it open in Word?):\n{exc}", parent=self)
            return
        self.set_status(f"PRD exported to {path}")
        if self.draft_key != "scratch":
            self._stash_draft()
            self.config_data.setdefault("exported_prds", {})[self.draft_key] = {
                "path": path, "date": dt.date.today().isoformat()}
            self._save_config_quiet()
            self._populate_tree()
            if self.current_issue:
                self._show_issue(self.current_issue)
        self._offer_open(path)

    def export_agenda(self):
        if not self.visible_rows:
            messagebox.showinfo(APP_NAME, "No tickets are shown in the list. Load tickets or change the "
                                          "filters first.", parent=self)
            return
        self._ensure("openpyxl", "openpyxl")
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.worksheet.datavalidation import DataValidation
        except ImportError:
            messagebox.showinfo(APP_NAME, "Excel export needs one more component.\n\nIn the VS Code terminal "
                                          "type:\n\n    python -m pip install openpyxl\n\nthen start the app again.",
                                parent=self)
            return
        today = dt.date.today()
        path = filedialog.asksaveasfilename(
            parent=self, title="Save Grooming Agenda", initialfile=f"Grooming_Agenda_{today.isoformat()}.xlsx",
            defaultextension=".xlsx", filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        pid, mine, scope_label = self._scope()
        base = self.var_url.get().rstrip("/")
        order = {"Needed": 0, "Partial": 1, "Draft": 2, "Has PRD": 3, "Not needed": 4}
        rows = sorted(self.visible_rows, key=lambda r: order[r[1]])
        filters = (f"Status: {self.var_status_filter.get()} · Sprint: {self.var_sprint_filter.get()} · "
                   f"Track: {self.var_track_filter.get()} · Release: {self.mp_release.summary()} · "
                   f"Due: {self.var_due_from.get() or '…'} to {self.var_due_to.get() or '…'} · "
                   f"Groom date: {self.var_groom_filter.get()} · "
                   f"PRD: {self.var_prd_filter.get()}")

        navy, white = "102A43", "FFFFFF"
        fills = {"Needed": "FDE2E1", "Partial": "FCEFD8", "Draft": "DCE8FF", "Has PRD": "DDF3E7",
                 "Not needed": "EEF0F3"}
        thin = Side(style="thin", color="D9E2EC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        wrap = Alignment(wrap_text=True, vertical="top")

        wb = Workbook()
        ws = wb.active
        ws.title = "Grooming Agenda"
        headers = ["Ticket", "Subject", "Parent", "Track", "Sprint", "Status", "Priority", "Assignee", "Groom date",
                   "Due date", "Release", "PRD", "PRD note", "Status check", "Discussion notes", "Decision"]
        widths = [10, 48, 30, 9, 11, 16, 11, 16, 13, 12, 16, 11, 30, 26, 40, 20]
        grooms = self.config_data.get("groom_dates") or {}
        ws["A1"] = f"Backlog Grooming Agenda: {today.strftime('%d %b %Y')}"
        ws["A1"].font = Font(size=16, bold=True, color=navy)
        ws["A2"] = f"Scope: {scope_label}   |   {filters}"
        ws["A2"].font = Font(size=10, color="616E7C")
        hdr_row = 4
        for col, (h, w) in enumerate(zip(headers, widths), start=1):
            c = ws.cell(row=hdr_row, column=col, value=h)
            c.font = Font(bold=True, color=white)
            c.fill = PatternFill("solid", fgColor=navy)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.border = border
            ws.column_dimensions[c.column_letter].width = w
        ws.row_dimensions[hdr_row].height = 24
        r = hdr_row
        for it, state, reason, flags in rows:
            r += 1
            values = [it["id"], it.get("subject", ""), self._parent_label(it), self._track(it),
                      self._sprint(it), name_of(it.get("status")), name_of(it.get("priority")),
                      name_of(it.get("assigned_to"), "Unassigned"), parse_date(grooms.get(str(it["id"]))),
                      parse_date(it.get("due_date")),
                      name_of(it.get("fixed_version"), ""), state, reason, "; ".join(flags), "", ""]
            for col, v in enumerate(values, start=1):
                c = ws.cell(row=r, column=col, value=v)
                c.alignment = wrap
                c.border = border
            link = ws.cell(row=r, column=1)
            link.hyperlink = f"{base}/issues/{it['id']}"
            link.font = Font(color="2F6FED", underline="single")
            ws.cell(row=r, column=9).number_format = "DDD DD-MMM-YYYY"
            ws.cell(row=r, column=10).number_format = "DD-MMM-YYYY"
            ws.cell(row=r, column=12).fill = PatternFill("solid", fgColor=fills[state])
            ws.cell(row=r, column=12).font = Font(bold=True)
            if flags:
                ws.cell(row=r, column=14).fill = PatternFill("solid", fgColor="FFF4DE")
                ws.cell(row=r, column=14).font = Font(color="B54708")
        ws.freeze_panes = ws.cell(row=hdr_row + 1, column=3)
        ws.auto_filter.ref = f"A{hdr_row}:P{max(r, hdr_row + 1)}"
        dv = DataValidation(type="list", allow_blank=True, formula1='"Ready for sprint,Needs PRD,'
                                                                 'Needs clarification,Split ticket,Deprioritise,Close"')
        ws.add_data_validation(dv)
        dv.add(f"P{hdr_row + 1}:P{max(r, hdr_row + 1) + 50}")
        ws.sheet_view.zoomScale = 100

        sm = wb.create_sheet("Summary")
        sm["A1"] = "Grooming summary"
        sm["A1"].font = Font(size=16, bold=True, color=navy)
        sm["A2"] = f"{today.strftime('%d %b %Y')} · {scope_label}"
        sm["A2"].font = Font(size=10, color="616E7C")
        sm.column_dimensions["A"].width = 34
        sm.column_dimensions["B"].width = 12
        row = 4

        def block(title, pairs):
            nonlocal row
            c = sm.cell(row=row, column=1, value=title)
            c.font = Font(bold=True, color=white)
            c.fill = PatternFill("solid", fgColor=navy)
            sm.cell(row=row, column=2).fill = PatternFill("solid", fgColor=navy)
            row += 1
            for k, v in pairs:
                sm.cell(row=row, column=1, value=k).border = border
                sm.cell(row=row, column=2, value=v).border = border
                row += 1
            row += 1

        n = len(rows)
        by_state = Counter(st for _, st, _, _ in rows)
        block("Overview", [("Tickets to groom", n),
                           ("Need a PRD (Needed + Partial)", by_state["Needed"] + by_state["Partial"]),
                           ("PRD draft in progress", by_state["Draft"]),
                           ("PRD ready", by_state["Has PRD"]),
                           ("PRD not needed", by_state["Not needed"]),
                           ("Status looks wrong (⚠)", sum(1 for *_, f in rows if f))])
        block("By sprint", Counter(self._sprint(it) or "No sprint" for it, *_ in rows).most_common())
        block("By track", Counter(self._track(it) for it, *_ in rows).most_common())
        block("By status", Counter(name_of(it.get("status")) for it, *_ in rows).most_common())
        block("By release", Counter(name_of(it.get("fixed_version"), "No release") for it, *_ in rows).most_common())
        block("By assignee", Counter(name_of(it.get("assigned_to"), "Unassigned") for it, *_ in rows).most_common())
        try:
            wb.save(path)
        except OSError as exc:
            messagebox.showerror(APP_NAME, f"Could not save the file (is it open in Excel?):\n{exc}", parent=self)
            return
        self.set_status(f"Grooming agenda saved to {path}")
        if messagebox.askyesno(APP_NAME, f"Saved:\n{path}\n\nOpen it now?", parent=self):
            try:
                os.startfile(path)  # Windows
            except AttributeError:
                webbrowser.open(Path(path).as_uri())
            except OSError:
                pass

    def copy_prd(self):
        self.clipboard_clear()
        self.clipboard_append(self._prd_text())
        self.set_status("PRD copied to the clipboard.")

    def clear_prd(self):
        if self._prd_text().strip() and messagebox.askyesno(APP_NAME, "Clear the workspace?", parent=self):
            self._set_prd("")
            self._stash_draft()
            self._populate_tree()

    # ----- dashboard -----------------------------------------------------
    def _on_tab_changed(self, _event=None):
        cur = self.nb.select()
        ready = self._configured(warn=False) and self._busy == 0
        if cur == str(self.tab_today):
            self.render_today()
        elif cur == str(self.tab_team):
            self.render_team()
            if ready and not self.team_open:
                self.refresh_team()
        elif cur == str(self.tab_dash) and self.dashboard_dirty and ready:
            self.refresh_dashboard()
        elif cur == str(self.tab_week):
            self._refresh_week_choices()
            if ready:
                self.refresh_weekly()
        elif cur == str(self.tab_delivery) and self.delivery is None and ready:
            self.refresh_delivery()
        elif cur == str(self.tab_rn) and not self.rn_versions and ready:
            self.load_versions()
        elif cur in (str(self.tab_sprint), str(self.tab_retro)):
            if not self.sp_version_info and ready:
                self.load_sprint_versions()
            self.render_sprints()

    def refresh_dashboard(self):
        if not self._configured():
            return
        client = self.make_client()
        pid, mine, label = self._dash_scope()
        self.dashboard_dirty = False

        def done(raw):
            self.dash_raw = raw
            self._dash_scope_label = label
            self._dash_options(raw)
            self._dash_apply()

        self.run_async(lambda: fetch_dashboard(client, pid, mine), done,
                       busy="Crunching backlog metrics… (this can take a few seconds)")

    def _render_dashboard(self, d: dict):
        self.dashboard_data = d
        n = d["analysed"]
        scope_text = f"Scope: {d['scope_label']}"
        if n < d["open_total"]:
            scope_text += f"   •   breakdowns based on the first {n} of {d['open_total']} open items"
        self.lbl_dash_scope.configure(text=scope_text)
        self.lbl_summary.configure(text=d["summary"])

        k = self.kpi
        k["open"].set(str(d["open_total"]), f"{d['unassigned']} unassigned · {d['stale']} untouched for "
                                            f"{STALE_DAYS}+ days")
        k["high"].set(str(d["high"]), f"{pct(d['high'], n)}% of open items",
                      C["bad"] if d["high"] else C["good"])
        k["overdue"].set(str(d["overdue"]), "past their due date" if d["overdue"] else "nothing past due",
                         C["bad"] if d["overdue"] else C["good"])
        rp = d["ready_pct"]
        k["ready"].set(f"{rp}%", d["ready_rule"],
                       C["good"] if rp >= 60 else (C["warn"] if rp >= 30 else C["bad"]))
        if d["closed30"] is not None:
            k["closed30"].set(str(d["closed30"]), f"{d['created30']} new items came in")
        else:
            k["closed30"].set("n/a", "your Redmine did not return closure data")
        if d["net30"] is not None:
            net = d["net30"]
            k["net"].set(f"{net:+d}", "backlog shrinking" if net < 0 else ("backlog growing" if net > 0 else
                                                                         "backlog steady"),
                         C["good"] if net <= 0 else C["warn"])
        else:
            k["net"].set("n/a", "")
        wtc = f" · ≈{d['weeks_to_clear']:.0f} wks to clear backlog" if d["weeks_to_clear"] else ""
        k["throughput"].set(f"{d['throughput']:.1f}", f"items closed per week (8-week avg){wtc}")
        if d["oldest"]:
            days, oldest = d["oldest"]
            k["age"].set(f"{d['avg_age']:.0f} d", f"median {d['median_age']:.0f} d · oldest #{oldest['id']} "
                                                  f"({days} d)")
        else:
            k["age"].set("—", "")

        prio_items = []
        for i, (name, count) in enumerate(d["by_priority"]):
            colour = priority_color(name)
            if colour == C["muted"]:
                colour = CHART_COLORS[i % len(CHART_COLORS)]
            prio_items.append((name, count, colour))
        self.ch_priority.set_data(prio_items)
        self.ch_status.set_data([(s, c, CHART_COLORS[i % len(CHART_COLORS)])
                                 for i, (s, c) in enumerate(d["by_status"])])
        self.ch_age.set_data(d["age_buckets"])
        self.ch_assignee.set_data([(a, c, C["warn"] if a == "Unassigned" else C["info"])
                                   for a, c in d["by_assignee"]])
        self.ch_trend.set_data(d["weeks"], "Closure history not available")

        self.risk_tree.delete(*self.risk_tree.get_children())
        for r in d["risks"]:
            it = r["issue"]
            self.risk_tree.insert("", "end", iid=f"r{it['id']}", values=(
                f"#{it['id']}", name_of(it.get("priority")), it.get("subject", ""),
                name_of(it.get("assigned_to"), "Unassigned"), r["why"]))
        if not d["risks"]:
            self.risk_tree.insert("", "end", values=("", "", "Nothing flagged: no overdue or stalled "
                                                             "high-priority work.", "", ""))
        self.lbl_asof.configure(text=f"Data as of {d['generated'].strftime('%d %b %Y, %H:%M')} · "
                                     f"source {self.var_url.get()}")
        self.set_status("Dashboard updated.")

    def _open_risk(self, _event=None):
        ids = self._risk_ids()
        if ids:
            self.show_in_workspace(int(ids[0]))

    def copy_summary(self):
        if not self.dashboard_data:
            messagebox.showinfo(APP_NAME, "Refresh the dashboard first.", parent=self)
            return
        d = self.dashboard_data
        text = (f"Backlog status ({d['scope_label']}, {d['generated'].strftime('%d %b %Y')}):\n\n"
                f"{d['summary']}")
        self.clipboard_clear()
        self.clipboard_append(text)
        self.set_status("Executive summary copied. Paste it into email, Slack or Teams.")

    def export_report(self):
        if not self.dashboard_data:
            messagebox.showinfo(APP_NAME, "Refresh the dashboard first.", parent=self)
            return
        stamp = dt.date.today().strftime("%Y-%m-%d")
        path = filedialog.asksaveasfilename(
            parent=self, title="Export Status Report", initialfile=f"Backlog_Status_{stamp}.md",
            defaultextension=".md", filetypes=[("Markdown", "*.md"), ("Text file", "*.txt"), ("All files", "*.*")])
        if not path:
            return
        try:
            with open(path, "w", encoding="utf-8") as fh:
                fh.write(build_status_report(self.dashboard_data, self.var_url.get().rstrip("/")))
        except OSError as exc:
            messagebox.showerror(APP_NAME, f"Could not save the file:\n{exc}", parent=self)
            return
        self.set_status(f"Status report exported to {path}")
        messagebox.showinfo(APP_NAME, f"Saved:\n{path}", parent=self)

    # ----- shutdown ------------------------------------------------------
    def toggle_theme(self):
        """Switch light/dark. The Workstation reopens itself in the new colours (your work is saved first)."""
        new = "light" if self.theme == "dark" else "dark"
        if not messagebox.askyesno(APP_NAME, f"Switch to {new} mode?\n\nThe Workstation will close and reopen "
                                             f"itself in a few seconds. Your drafts and settings are saved first.",
                                   parent=self):
            return
        self.config_data["theme"] = new
        try:
            save_json(CONFIG_PATH, self.config_data)
        except OSError:
            pass
        exe = Path(sys.executable)
        pyw = exe.with_name("pythonw.exe")
        runner = str(pyw) if IS_WINDOWS and pyw.exists() else str(exe)
        try:
            subprocess.Popen([runner, str(Path(__file__).resolve())] + sys.argv[1:],
                             cwd=str(Path(__file__).resolve().parent))
        except OSError as exc:
            messagebox.showinfo(APP_NAME, f"Saved. Please reopen the Workstation to see {new} mode.\n({exc})",
                                parent=self)
        self._on_close()

    def _on_close(self):
        self._stash_draft()
        self.config_data["rn_draft"] = self.txt_rn.get("1.0", "end-1c")
        self.config_data["delivery_period"] = self.var_period.get()
        if self.projects:
            self.config_data["projects_selected"] = list(self.mp_project.selected)
        self.config_data.update(scope=self.var_scope.get(),
                                status_filter=self.var_status_filter.get(),
                                sprint_filter=self.var_sprint_filter.get(),
                                assignee_filter=self.var_assignee_filter.get(),
                                track_filter=self.var_track_filter.get(),
                                group_by_parent=bool(self.var_group.get()))
        if self.config_data.get("base_url"):
            try:
                save_json(CONFIG_PATH, self.config_data)
            except OSError:
                pass
        self.destroy()


def main():
    if IS_WINDOWS:
        try:  # Crisp text on high-resolution Windows screens
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
    migrate_and_backup()
    app = App()
    app.mainloop()


if __name__ == "__main__":
    main()
